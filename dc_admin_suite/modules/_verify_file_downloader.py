#!/usr/bin/env python3
"""Verification suite for Batch File Downloader (native files).

The version tested is whatever `MANIFEST["version"]` says in the module
beside this file, and the banner prints it. This file carried "v1.2" in its
name and its first line while the module had moved to v1.3 — a version in a
filename is a second home for something MANIFEST already owns, and it goes
stale the first time nobody renames it. Nothing here cites a version.

Leading underscore keeps it out of the shell's module scan.

Run:  python3 modules/_verify_file_downloader.py
"""

import inspect
import io
import json
import os
import re
import sys
import tempfile
import time
import textwrap
import threading
import urllib.error
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dc_file_downloader as M          # noqa: E402

# v1.34.1: runs are also claimed in a lock file named for the Chrome they
# drive. The suite must never share that folder with a real run on the same
# machine — a test would be refused by it, or leave a lock a real Start then
# has to judge — so every test that starts a run claims in a folder of its
# own.
import atexit as _atexit                # noqa: E402
import shutil as _shutil                # noqa: E402
M.RUN_LOCK_DIR = tempfile.mkdtemp(prefix="dcfd-runlock-")
_atexit.register(_shutil.rmtree, M.RUN_LOCK_DIR, True)

PASS, FAIL = [], []


def _identity_markers():
    """The markers `make_distribution.py --generic` enforces, from its own
    reader, so this suite cannot check a list the build does not.

    Imported by path rather than by name: `docs/` is not a package and the
    suite must stay runnable from anywhere in the tree. Returns () when
    the build script or the profiles are absent, which a generic copy is
    entitled to be — the caller says what that means.
    """
    import importlib.util
    here = os.path.dirname(os.path.abspath(__file__))
    script = os.path.join(os.path.dirname(here), "docs", "make_distribution.py")
    if not os.path.exists(script):
        return ()
    try:
        spec = importlib.util.spec_from_file_location("_mkdist", script)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return tuple(mod.identity_markers())
    except Exception:
        return ()


def _private_terms():
    """The build script, for its vendor-and-licence reader.

    `make_distribution.private_terms()` reads the patterns from the
    institution profile(s), so neither this file nor the build spells
    them out (2026-10-03: the list used to sit in the build script, which
    ships in the public copy). Returns the module; raises when the build
    is missing, because the reader is not optional in any copy. What it
    returns may be () in a copy with no named profile; the caller says
    so."""
    import importlib.util
    here = os.path.dirname(os.path.abspath(__file__))
    script = os.path.join(os.path.dirname(here), "docs", "make_distribution.py")
    spec = importlib.util.spec_from_file_location("_mkdist_terms", script)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def check(name, got, want):
    if got == want:
        PASS.append(name)
    else:
        FAIL.append("{}\n     got:  {!r}\n     want: {!r}".format(
            name, got, want))


def check_true(name, cond, detail=""):
    if cond:
        PASS.append(name)
    else:
        FAIL.append("{}{}".format(name, "\n     " + detail if detail else ""))


XFAIL, XPASS = [], []


def check_xfail(name, cond, hypothesis):
    """A hazard found by reading the code and not yet confirmed against real
    HTML. `cond` is what SHOULD hold, so a False lands in XFAIL and does not
    fail the run — the defect is suspected, not proven, and CI must not go
    red on a hypothesis. A True lands in XPASS, which is the interesting
    case: the hypothesis was wrong, or someone fixed it and left the marker
    behind. Either way an XPASS has to be resolved rather than ignored, so
    both lists are printed on every run.
    """
    (XPASS if cond else XFAIL).append(
        "{}\n     {}".format(name, hypothesis))



BASE = "https://digitalcommons.example.edu"


# ---------------------------------------------------------------------------
# 1 · content_link_type — one case per URL shape.
# ---------------------------------------------------------------------------
def test_content_link_type():
    cases = [
        ("query native",
         "/cgi/viewcontent.cgi?article=1000&context=jesuit-gallery161"
         "&type=native", "native"),
        ("query pdf with extras",
         "/cgi/viewcontent.cgi?article=12&context=series&type=pdf"
         "&unstamped=yes&date=1699999999&preview_mode=1", "pdf"),
        ("path form native",
         "/context/jesuit-gallery161/article/1000/type/native/viewcontent",
         "native"),
        ("path form, type last segment",
         "/context/c/article/9/type/native", "native"),
        ("bare link, no type at all",
         "/cgi/viewcontent.cgi?article=1000&context=jesuit-gallery161", ""),
        ("additional file, query",
         "/cgi/viewcontent.cgi?filename=2&type=additional"
         "&preview_mode=1", "additional"),
        ("absolute URL, path form",
         BASE + "/context/c/article/1/type/native/viewcontent", "native"),
    ]
    for label, href, want in cases:
        check("content_link_type · " + label, M.content_link_type(href), want)

    # The bare form is the regression: v1.1 dropped it. Prove it is now
    # classified as primary rather than discarded.
    html = _pickvers_html([
        ['<a href="/cgi/viewcontent.cgi?article=1000&context=g">'
         'Download</a>']])
    revs = M.parse_pickvers(html)
    check("bare link is kept as primary, not dropped",
          [f["kind"] for f in revs[0]["files"]], ["primary"])


# ---------------------------------------------------------------------------
# 2 · native_variant — derive the fallback from a parsed href.
# ---------------------------------------------------------------------------
def test_native_variant():
    # A query-form link is rewritten into the PATH form, not given
    # "&type=native". The CGI ignores that parameter and answers with the
    # "No PDF has been provided" page — verified against jesuit-gallery25
    # article 1000 on 2026-09-07. This assertion previously pinned the
    # broken shape, which is how the defect survived 95 green tests.
    check("native_variant · query form is rewritten to the path form",
          M.native_variant(BASE + "/cgi/viewcontent.cgi?article=1&context=c"),
          BASE + "/context/c/article/1/type/native/viewcontent?preview_mode=1")
    check("native_variant · query form keeps extra parameters out of it",
          M.native_variant(BASE + "/cgi/viewcontent.cgi?article=1000"
                                  "&context=jesuit-gallery25&preview_mode=1"),
          BASE + "/context/jesuit-gallery25/article/1000/type/native"
                 "/viewcontent?preview_mode=1")
    check("native_variant · no fallback without both article and context",
          M.native_variant(BASE + "/cgi/viewcontent.cgi?article=1"), "")
    check("native_variant · path form",
          M.native_variant(BASE + "/context/c/article/1/viewcontent"),
          BASE + "/context/c/article/1/type/native/viewcontent?preview_mode=1")
    check("native_variant · already typed returns empty",
          M.native_variant(BASE + "/cgi/viewcontent.cgi?article=1&context=c"
                                  "&type=native"), "")
    check("native_variant · not a content URL returns empty",
          M.native_variant(BASE + "/cgi/editor.cgi?window=abstract"), "")


# ---------------------------------------------------------------------------
# 3 · parse_pickvers — a multi-revision PDF record and the single-revision
#     native-only record this release exists for.
# ---------------------------------------------------------------------------
def _pickvers_html(link_sets, checked_first=True):
    rows = []
    for i, links in enumerate(link_sets):
        radio = ' checked' if (i == 0 and checked_first) else ''
        rows.append(
            "<tr><td>editor{i}</td><td>comment {i}</td>"
            "<td>2026-01-0{i} 10:00</td>"
            "<td><input type='radio' name='sel'{r}></td>"
            "<td>{links}</td></tr>".format(i=i + 1, r=radio,
                                           links=" ".join(links)))
    return ("<html><body><table id='revision'>"
            "<tr><td>User</td><td>Comment</td><td>Date</td>"
            "<td>Select</td><td>Files</td></tr>"
            + "".join(rows) + "</table></body></html>")


def test_parse_pickvers():
    # (a) A conventional PDF record with two revisions, query-form links.
    html = _pickvers_html([
        ['<a href="/cgi/viewcontent.cgi?article=5&context=s&type=pdf'
         '&unstamped=yes&date=17&preview_mode=1">PDF</a>',
         '<a href="/cgi/viewcontent.cgi?article=5&context=s&type=native'
         '&date=17">native</a>'],
        ['<a href="/cgi/viewcontent.cgi?article=5&context=s&type=pdf'
         '&unstamped=yes&date=16&preview_mode=1">PDF</a>'],
    ])
    revs = M.parse_pickvers(html)
    check("pickvers · two revisions parsed", len(revs), 2)
    check("pickvers · newest row kinds",
          sorted(f["kind"] for f in revs[0]["files"]), ["native", "primary"])
    M.label_versions(revs)
    check("pickvers · newest row is current", revs[0]["version"], "current")
    check("pickvers · older row is original", revs[1]["version"], "original")

    # (b) THE REGRESSION. One revision, native only, path-form link — what a
    #     gallery item looks like. v1.1 discarded this link and downloaded
    #     nothing even in all-versions mode.
    html = _pickvers_html([
        ['<a href="/context/jesuit-gallery161/article/1000/type/native'
         '/viewcontent">GAS009628.CR2</a>']])
    revs = M.parse_pickvers(html)
    check("pickvers · native-only row yields one file",
          len(revs[0]["files"]), 1)
    check("pickvers · native-only row is kind native",
          revs[0]["files"][0]["kind"], "native")

    # (c) A cover letter still classifies as one, and an unknown type is
    #     labelled rather than dropped.
    html = _pickvers_html([
        ['<a href="/cgi/viewcoverletter.cgi?article=5&context=s">cover</a>',
         '<a href="/cgi/viewcontent.cgi?article=5&context=s'
         '&type=supplemental">extra</a>']])
    kinds = sorted(f["kind"] for f in M.parse_pickvers(html)[0]["files"])
    check("pickvers · coverletter and unknown type both kept",
          kinds, ["coverletter", "supplemental"])

    # (d) No table at all still returns None (login bounce upstream).
    check("pickvers · no table returns None",
          M.parse_pickvers("<html><body>nothing</body></html>"), None)


# ---------------------------------------------------------------------------
# 4 · parse_additional_files — index read from either shape, and a missing
#     index never discards the file.
# ---------------------------------------------------------------------------
def test_parse_additional_files():
    html = ("<html><table id='uploaded-files'>"
            "<tr><td><a href='/cgi/viewcontent.cgi?filename=2"
            "&type=additional&preview_mode=1'>scan.tif</a></td>"
            "<td><input type='checkbox' checked></td></tr>"
            "<tr><td><a href='/context/c/article/9/filename/3/type/"
            "additional/viewcontent'>data.csv</a></td>"
            "<td><input type='checkbox'></td></tr>"
            "<tr><td><a href='/cgi/viewcontent.cgi?article=9&context=c'>"
            "bare.jpg</a></td><td><input type='checkbox'></td></tr>"
            "</table></html>")
    supp = M.parse_additional_files(html)
    check("additional · all three rows kept", len(supp), 3)
    check("additional · query index", supp[0]["idx"], "2")
    check("additional · path index", supp[1]["idx"], "3")
    check("additional · missing index does not discard",
          (supp[2]["idx"], supp[2]["name"]), ("", "bare.jpg"))
    check("additional · shown flags",
          [f["shown"] for f in supp], [True, False, False])


# ---------------------------------------------------------------------------
# 5 · fetch_record_file — the fallback matrix, including call counts.
# ---------------------------------------------------------------------------
def test_fetch_record_file():
    calls = []

    def stub(results):
        def _fetch(url, cookie, timeout=0):
            calls.append(url)
            outcome = results.pop(0)
            if isinstance(outcome, Exception):
                raise outcome
            return outcome
        return _fetch

    real = M.fetch_file
    try:
        # (a) First call succeeds — the fallback must NOT be issued.
        calls.clear()
        M.fetch_file = stub([(b"%PDF-1.4", "a.pdf", "application/pdf")])
        got = M.fetch_record_file("URL", "c", "NATIVE")
        check("fetch · derivative succeeds", got[3], False)
        check("fetch · no speculative second request", len(calls), 1)

        # (b) Web page back, then the native answers.
        calls.clear()
        M.fetch_file = stub([
            ValueError("server returned a web page instead of a file"),
            (b"\xff\xd8\xff\xe0", "GAS1.CR2", "application/octet-stream")])
        data, name, ctype, used = M.fetch_record_file("URL", "c", "NATIVE")
        check("fetch · falls back to native", used, True)
        check("fetch · fallback hit the native URL", calls[1], "NATIVE")
        check("fetch · returns the native's filename", name, "GAS1.CR2")

        # (c) A 404 on the derivative also falls back.
        calls.clear()
        M.fetch_file = stub([
            urllib.error.HTTPError("URL", 404, "Not Found", {}, None),
            (b"\x89PNG\r\n\x1a\n", "", "image/png")])
        check("fetch · 404 falls back too",
              M.fetch_record_file("URL", "c", "NATIVE")[3], True)

        # (d) Both fail — one message that names BOTH attempts. v1.6 changed
        #     this wording deliberately: the old text reported only the
        #     derivative's error and dropped the native's, which is the
        #     fallback that was supposed to work and therefore the
        #     informative half.
        calls.clear()
        M.fetch_file = stub([ValueError("server returned a web page instead "
                                        "of a file"),
                             ValueError("empty response from the server")])
        try:
            M.fetch_record_file("URL", "c", "NATIVE")
            check("fetch · both fail raises NoFileAvailable", False, True)
        except M.NoFileAvailable as e:
            msg = str(e)
            check_true("fetch · both fail raises NoFileAvailable",
                       "no downloadable file" in msg, msg)
            check_true("fetch · the message names the derivative attempt",
                       "derivative: server returned a web page" in msg, msg)
            check_true("fetch · the message names the native attempt too",
                       "native: empty response from the server" in msg, msg)

        # (e) No native URL to try — the original error propagates unchanged.
        calls.clear()
        M.fetch_file = stub([ValueError("empty response from the server")])
        try:
            M.fetch_record_file("URL", "c", "")
            check("fetch · no native_url propagates original", False, True)
        except ValueError as e:
            check("fetch · no native_url propagates original",
                  str(e), "empty response from the server")

        # (f) LoginRequired must propagate WITHOUT a fallback attempt —
        #     the worker's session-expiry handling depends on seeing it.
        calls.clear()
        M.fetch_file = stub([M.LoginRequired()])
        try:
            M.fetch_record_file("URL", "c", "NATIVE")
            check("fetch · LoginRequired propagates", False, True)
        except M.LoginRequired:
            check("fetch · LoginRequired propagates", True, True)
            check("fetch · LoginRequired does not retry", len(calls), 1)
    finally:
        M.fetch_file = real


# ---------------------------------------------------------------------------
# 6 · sniff_extension — one fixture per format, plus the CR2/TIFF split.
# ---------------------------------------------------------------------------
def test_sniff_extension():
    tiff = b"II*\x00" + b"\x08\x00\x00\x00" + b"\x00\x00" + b"\x00" * 64
    cr2 = b"II*\x00" + b"\x10\x00\x00\x00" + b"CR\x02\x00" + b"\x00" * 64
    cases = [
        (".png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 32),
        (".jpg", b"\xff\xd8\xff\xe0" + b"\x00" * 32),
        (".gif", b"GIF89a" + b"\x00" * 32),
        (".tif", tiff),
        (".cr2", cr2),
        (".bmp", b"BM" + b"\x00" * 32),
        (".webp", b"RIFF\x24\x00\x00\x00WEBPVP8 "),
        (".wav", b"RIFF\x24\x00\x00\x00WAVEfmt "),
        (".pdf", b"%PDF-1.7\n%\xe2\xe3"),
        (".jp2", b"\x00\x00\x00\x0cjP  \r\n\x87\n" + b"\x00" * 16),
        (".ico", b"\x00\x00\x01\x00" + b"\x00" * 32),
        (".zip", b"PK\x03\x04" + b"\x00" * 40),
        (".docx", b"PK\x03\x04" + b"\x00" * 26 + b"word/document.xml"),
        (".xlsx", b"PK\x03\x04" + b"\x00" * 26 + b"xl/workbook.xml"),
        (".pptx", b"PK\x03\x04" + b"\x00" * 26 + b"ppt/presentation.xml"),
        (".doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32),
        (".mp3", b"ID3\x04\x00" + b"\x00" * 32),
        (".mp4", b"\x00\x00\x00\x20ftypisom" + b"\x00" * 16),
        (".txt", b"article,title,year\n1,Foo,2026\n"),
    ]
    for want, data in cases:
        check("sniff " + want, M.sniff_extension(data), want)

    check("sniff · CR2 is not mistaken for TIFF",
          (M.sniff_extension(cr2), M.sniff_extension(tiff)), (".cr2", ".tif"))
    check("sniff · HTML error page is not a file",
          M.sniff_extension(b"<!DOCTYPE html><html><body>error"), "")
    check("sniff · HTML without doctype is not a file",
          M.sniff_extension(b"  <html><head><title>Login"), "")
    check("sniff · empty is unknown", M.sniff_extension(b""), "")
    check("sniff · unknown binary is unknown",
          M.sniff_extension(b"\x07\x91\xfe\x13" + b"\x8a" * 64), "")


# ---------------------------------------------------------------------------
# 7 · build_filename — precedence, and the token shape downstream depends on.
# ---------------------------------------------------------------------------
def test_build_filename():
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
    cr2 = b"II*\x00" + b"\x10\x00\x00\x00" + b"CR\x02\x00" + b"\x00" * 64

    # Extensions are normalised to lower case (v1.1 behaviour, kept): the
    # point of this case is that .CR2 beat the octet-stream Content-Type,
    # not that the casing survived.
    check("filename · real extension wins over wrong Content-Type",
          M.build_filename("g", "1000", "native", "public", "current",
                           "GAS009628.CR2", "application/octet-stream", cr2),
          "g_1000_native_public_current_GAS009628.cr2")
    check("filename · no filename resolves from bytes",
          M.build_filename("g", "1000", "native", "public", "current",
                           "", "application/octet-stream", png),
          "g_1000_native_public_current_file.png")
    check("filename · .bin is re-resolved, not preserved",
          M.build_filename("g", "1", "native", "public", "current",
                           "old.bin", "", png),
          "g_1_native_public_current_old.png")
    check("filename · .native placeholder re-resolved",
          M.build_filename("g", "1", "native", "public", "current",
                           "x.native", "", cr2),
          "g_1_native_public_current_x.cr2")
    check("filename · Content-Type used when bytes are unknown",
          M.build_filename("g", "1", "primary", "public", "current",
                           "", "application/pdf; charset=binary",
                           b"\x07\x91\xfe\x13" + b"\x8a" * 64),
          "g_1_primary_public_current_file.pdf")
    check("filename · unknown everything falls back to .bin",
          M.build_filename("g", "1", "primary", "public", "current",
                           "", "application/octet-stream",
                           b"\x07\x91\xfe\x13" + b"\x8a" * 64),
          "g_1_primary_public_current_file.bin")

    # The Image Description Generator parses <ctx>_<article>_… back out of
    # this name. Extension may change; token order may not.
    name = M.build_filename("jesuit-gallery161", "1000", "native", "public",
                            "current", "", "", png)
    check_true("filename · leading tokens still <ctx>_<article>_<kind>",
               name.startswith("jesuit-gallery161_1000_native_"), name)


# ---------------------------------------------------------------------------
# 8 · dedupe_jobs — the double-download this release would otherwise create.
# ---------------------------------------------------------------------------
def test_native_file_url():
    """The shape is the whole point — see native_file_url's docstring.

    Only the path form serves the native; the query form answers "No PDF has
    been provided" whether or not preview_mode is set. Checked against a live
    native-only record on 2026-09-07.
    """
    check("native_file_url is the path form",
          M.native_file_url(BASE, "jesuit-gallery25", "1000"),
          BASE + "/context/jesuit-gallery25/article/1000/type/native"
                 "/viewcontent?preview_mode=1")
    check_true("native_file_url is not the query form",
               "viewcontent.cgi?" not in M.native_file_url(BASE, "c", "1"),
               M.native_file_url(BASE, "c", "1"))
    # preview_mode IS carried, deliberately: it is what reaches a record in
    # any state, and the path form accepts it and still returns the bytes.
    check_true("native_file_url carries preview_mode for unposted records",
               M.native_file_url(BASE, "c", "1").endswith("?preview_mode=1"),
               M.native_file_url(BASE, "c", "1"))
    check("native_file_url tolerates a trailing slash on base_url",
          M.native_file_url(BASE + "/", "c", "1"),
          BASE + "/context/c/article/1/type/native/viewcontent?preview_mode=1")
    # The primary URL is unchanged and still asks for the default stream.
    check("primary_file_url still asks for the derivative",
          M.primary_file_url(BASE, "c", "1"),
          BASE + "/cgi/viewcontent.cgi?article=1&context=c&preview_mode=1")


def test_dedupe_jobs():
    """v1.34.2: one FILE, not one URL.

    The old version of this test put a primary and a native on the SAME
    url and asserted one survived — which is how the defect lived: the
    real case is one file at three DIFFERENT addresses, which the test
    could not express.
    """
    short = BASE + "/context/g/article/1000/type/native/viewcontent" \
                   "?preview_mode=1"
    dated = (BASE + "/cgi/viewcontent.cgi?type=native&article=1000"
             "&unstamped=yes&date=1592928048&preview_mode=1&context=g"
             "&/1592928048-text.native")
    jobs = [
        {"kind": "native", "visibility": "public", "version": "current",
         "version_date": "", "url": short, "native_url": "",
         "orig_hint": ""},
        {"kind": "native", "visibility": "public", "version": "current",
         "version_date": "Tue Jun 23 09:00:00 2020", "url": dated,
         "native_url": "", "orig_hint": "GAS1.CR2"},
    ]
    out = M.dedupe_jobs(jobs)
    check("dedupe · one current native at two addresses is one job",
          len(out), 1)
    check("dedupe · the first address is kept", out[0]["url"], short)
    check("dedupe · hint is carried over", out[0]["orig_hint"], "GAS1.CR2")
    check("dedupe · the revision's date is carried over",
          out[0]["version_date"], "Tue Jun 23 09:00:00 2020")

    # Different versions are genuinely different files.
    jobs = [
        {"kind": "native", "visibility": "public", "version": "current",
         "version_date": "", "url": short, "native_url": "", "orig_hint": ""},
        {"kind": "native", "visibility": "hidden", "version": "original",
         "version_date": "Tue Jun 23 08:38:00 2020", "url": dated,
         "native_url": "", "orig_hint": ""},
    ]
    check("dedupe · different versions kept apart",
          len(M.dedupe_jobs(jobs)), 2)

    # The stamped default stream and the unstamped revision copy differ in
    # bytes (sspeach 1000: 254,280 against 235,659) and are both kept.
    jobs = [
        {"kind": "primary", "visibility": "public", "version": "current",
         "version_date": "",
         "url": BASE + "/cgi/viewcontent.cgi?article=1&context=s"
                       "&preview_mode=1",
         "native_url": "", "orig_hint": ""},
        {"kind": "primary", "visibility": "public",
         "version": M.VERSION_CURRENT_UNSTAMPED,
         "version_date": "Mon Oct  2 11:41:00 2023",
         "url": BASE + "/cgi/viewcontent.cgi?type=pdf&article=1"
                       "&unstamped=yes&date=1696272079&context=s",
         "native_url": "", "orig_hint": ""},
    ]
    check("dedupe · stamped and unstamped current PDFs kept apart",
          len(M.dedupe_jobs(jobs)), 2)
    check("dedupe · a primary is not a native",
          M.file_identity({"kind": "primary", "version": "current",
                           "url": short}) ==
          M.file_identity({"kind": "native", "version": "current",
                           "url": short}), False)


# ---------------------------------------------------------------------------
# 9 · plan_item_jobs — end to end over stubbed pages, incl. the native option.
# ---------------------------------------------------------------------------
def test_plan_item_jobs():
    opts_base = {"primary": True, "supp": False, "native": False,
                 "public": True, "hidden": True, "versions": "current",
                 "published": True, "unpublished": False}
    item = {"article": "1000", "title": "T", "af": 0}
    warns = []

    jobs = M.plan_item_jobs(BASE, "g", item, "published", True,
                            dict(opts_base), "cookie", warns.append)
    check("plan · default queues one primary", len(jobs), 1)
    check("plan · primary carries a path-form native fallback URL",
          jobs[0]["native_url"],
          BASE + "/context/g/article/1000/type/native/viewcontent?preview_mode=1")

    opts = dict(opts_base, native=True)
    jobs = M.plan_item_jobs(BASE, "g", item, "published", True, opts,
                            "cookie", warns.append)
    check("plan · native option adds a second job",
          [j["kind"] for j in jobs], ["primary", "native"])

    # All-versions over a native-only record: the pickvers native and the
    # primary's fallback must not both land.
    html = _pickvers_html([
        ['<a href="/context/g/article/1000/type/native/viewcontent">'
         'GAS1.CR2</a>']])
    real = M.fetch_html
    try:
        M.fetch_html = lambda url, cookie, timeout=0: html
        opts = dict(opts_base, versions="all")
        jobs = M.plan_item_jobs(BASE, "g", item, "published", True, opts,
                                "cookie", warns.append)
    finally:
        M.fetch_html = real
    kinds = [j["kind"] for j in jobs]
    check("plan · all-versions native-only yields primary + native",
          kinds, ["primary", "native"])
    check_true("plan · the two point at different URLs",
               jobs[0]["url"] != jobs[1]["url"],
               "{} vs {}".format(jobs[0]["url"], jobs[1]["url"]))
    check("plan · no warning raised for a native-only record", warns, [])


# ---------------------------------------------------------------------------
# 10 · Report shape — arity, and a round trip through a header-name map.
# ---------------------------------------------------------------------------
def test_report_shape():
    import inspect
    src = inspect.getsource(M.download_worker)
    # Both report_rows.append() tuples must carry len(REPORT_HEADERS) values.
    check("report · headers and widths agree",
          (len(M.REPORT_HEADERS), len(M.REPORT_WIDTHS)),
          (len(M.REPORT_HEADERS), len(M.REPORT_HEADERS)))
    # By NAME, not by position. This asserted REPORT_HEADERS[-1] until
    # 2026-09-09, which made it a test of "nothing has been appended"
    # wearing the name of a test that a column exists — and appending is
    # the one change the header comment explicitly calls safe.
    for name in ("Content Type", "Access", "Type / Release Option",
                 "Embargo", "Plan"):
        check_true("report · {} column present".format(name),
                   name in M.REPORT_HEADERS)
    check_true("report · appended columns come after Content Type",
               M.REPORT_HEADERS.index("Access")
               > M.REPORT_HEADERS.index("Content Type"))
    check_true("report · link columns still point at the URL columns",
               [M.REPORT_HEADERS[i - 1] for i in M.REPORT_LINK_COLS] ==
               ["File URL", "Public Record URL", "Admin Record URL"],
               str(M.REPORT_LINK_COLS))

    n = len(M.REPORT_HEADERS)
    rows = [tuple(["v"] * n), tuple(["w"] * n)]
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "r.xlsx")
        M.write_structure_report(path, rows, "1F4E79")
        from openpyxl import load_workbook
        wbk = load_workbook(path)
        ws = wbk.active
        # This mirrors dc_image_describer._header_map exactly: columns are
        # located by NAME, so appending a column is safe.
        hmap = {str(c.value).strip(): i + 1
                for i, c in enumerate(ws[1]) if c.value is not None}
        wbk.close()
        for needed in ("Saved Filename", "Article ID", "Title",
                       "Record State", "Admin Record URL"):
            check_true("report · image-describer column '{}' locatable"
                       .format(needed), needed in hmap, str(sorted(hmap)))
        check("report · round-trips every column", len(hmap), n)


# ---------------------------------------------------------------------------
# 11 · HTTP endpoints — the new file kind, and the standing guards.
# ---------------------------------------------------------------------------
def _post(port, path, body, origin=None, host=None):
    req = urllib.request.Request(
        "http://127.0.0.1:{}{}".format(port, path),
        data=json.dumps(body).encode(), method="POST",
        headers={"Content-Type": "application/json"})
    if origin:
        req.add_header("Origin", origin)
    if host:
        req.add_header("Host", host)
    try:
        with urllib.request.urlopen(req, timeout=10) as r:
            return r.status, json.loads(r.read() or b"{}")
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read() or b"{}")


def test_endpoints():
    session = M.load_session(Path("/nonexistent/session.json"))
    session["hub_url"] = "http://127.0.0.1:8750"
    srv = ThreadingHTTPServer(
        ("127.0.0.1", 0), M.make_handler(session, M.build_page(session)))
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        with tempfile.TemporaryDirectory() as d:
            base = {"out_dir": d, "mode": "all", "parents": [],
                    "primary": False, "supp": False, "native": False,
                    "public": True, "hidden": True, "versions": "current",
                    "published": True, "unpublished": False}

            code, j = _post(port, "/api/start", dict(base))
            check("api · no file kind checked is refused",
                  (code, j.get("error")),
                  (400, "Check at least one file kind."))

            # Native alone must clear the file-kind gate. With no hierarchy
            # loaded the next guard stops it, which is the proof it got past.
            code, j = _post(port, "/api/start", dict(base, native=True))
            check("api · native alone clears the file-kind gate",
                  (code, j.get("error")),
                  (400, "Load or map a hierarchy first."))

            code, j = _post(port, "/api/start", dict(base, primary=True,
                                                     public=False,
                                                     hidden=False))
            check("api · no visibility is still refused",
                  (code, j.get("error")),
                  (400, "Check at least one visibility."))

            code, j = _post(port, "/api/start",
                            dict(base, native=True, out_dir="/no/such/dir"))
            check("api · bad output folder is refused", code, 400)

            code, _ = _post(port, "/api/start", dict(base, native=True),
                            origin="https://evil.example")
            check("api · cross-origin request is refused", code, 403)

            code, _ = _post(port, "/api/start", dict(base, native=True),
                            host="evil.example")
            check("api · non-local Host is refused", code, 403)

            code, _ = _post(port, "/api/nope", {})
            check("api · unknown endpoint is 404", code, 404)

            code, _ = _post(port, "/api/stop", {})
            check("api · stop with no job running is 409", code, 409)

        page = urllib.request.urlopen(
            "http://127.0.0.1:{}/".format(port), timeout=10).read().decode()
        check_true("api · page renders the new file kind",
                   'id="knative"' in page, "knative checkbox missing")
        check_true("api · page explains automatic native download",
                   "downloaded automatically" in page, "hint text missing")
    finally:
        srv.shutdown()


# ---------------------------------------------------------------------------
# 12 · The listing stage.
#
# Added 2026-09-07, after Jeff reported that the revisions made to support
# book and image galleries appear to have broken the module for every
# structure type. This suite had 95 passing tests and could not have caught
# it: nothing here referenced parse_edikit_listing, parse_gallery_listing,
# _find_listing_table, _listing_rows, crawl_listing, admin_record_url or
# either listing-URL builder. Every test above covers the PER-ITEM stage —
# pickvers, filenames, extension sniffing, job planning. The stage that turns
# "a structure" into "a list of items" was untested, and gallery support is
# entirely in that stage.
#
# The fixtures below are synthetic, built from what the parsers require
# rather than from real pages, so they pin structure and not content. Real
# saved listing HTML replaces them as soon as it can be captured; until
# then, treat a green run here as "no NEW breakage", not as proof the
# regression is absent.
# ---------------------------------------------------------------------------
EDIKIT_LISTING = """<html><body>
<table>
  <tr><th>ID</th><th>Title</th><th>Last Event</th><th>Submitted</th>
      <th>Additional Files</th></tr>
  <tr><td>1001</td><td>First item</td><td>posted</td><td>2026-01-02</td>
      <td>2</td></tr>
  <tr><td>1002</td><td>Second item</td><td>posted</td><td>2026-01-03</td>
      <td>-</td></tr>
</table>
<a href="editor.cgi?context=series&amp;page=2">&gt;</a>
</body></html>"""

GALLERY_LISTING = """<html><body>
<table id="submission_list">
  <tr><th>ID</th><th>Title</th><th>Last Event</th><th>Submitted</th></tr>
  <tr><td>2001</td><td>Plate one</td><td>posted</td><td>2026-02-01</td></tr>
</table>
<a href="editor_gallery.cgi?context=gallery1&amp;page=2">Next</a>
</body></html>"""

# An EdiKit table whose first column is a select box rather than the ID.
# This is the shape that would explain "broken for every structure": see the
# hypothesis in test_listing_row_hazards.
EDIKIT_CHECKBOX_FIRST = """<html><body>
<table>
  <tr><th></th><th>ID</th><th>Title</th><th>Last Event</th>
      <th>Submitted</th></tr>
  <tr><td><input type="checkbox" name="pick"></td><td>1001</td>
      <td>First item</td><td>posted</td><td>2026-01-02</td></tr>
</table>
</body></html>"""

# The same two listings with no "next" link: a single-page structure.
EDIKIT_LISTING_LASTPAGE = EDIKIT_LISTING.replace(
    '<a href="editor.cgi?context=series&amp;page=2">&gt;</a>', "")
GALLERY_LISTING_LASTPAGE = GALLERY_LISTING.replace(
    '<a href="editor_gallery.cgi?context=gallery1&amp;page=2">Next</a>', "")
assert "<a href" not in EDIKIT_LISTING_LASTPAGE
assert "<a href" not in GALLERY_LISTING_LASTPAGE

LOGIN_PAGE = """<html><body><form action="/cgi/myaccount.cgi">
<input name="login" type="text"><input name="password" type="password">
</form></body></html>"""


def test_listing_urls():
    check("edikit_listing_url shape",
          M.edikit_listing_url(BASE, "series", "published"),
          BASE + "/cgi/editor.cgi?context=series&showstate=published"
                 "&status_filter=all&x_showall=1&x_field_1=abstract"
                 "&x_op_1=eq&x_value_1=")
    # Pinned separately with its reason, because the shape assertion above
    # would happily go along with dropping it. On the 3,383-record ETD
    # collection, paging 25 at a time is ~136 requests before a single file
    # is fetched - four server lockouts and about forty-five minutes spent
    # on nothing but listings.
    check_true("edikit listing asks for every record on one page",
               "x_showall=1" in M.edikit_listing_url(BASE, "c", "published"),
               M.edikit_listing_url(BASE, "c", "published"))
    check_true("gallery listing does not — it pages with page=N",
               "x_showall" not in M.gallery_listing_url(BASE, "c",
                                                        "published"),
               M.gallery_listing_url(BASE, "c", "published"))
    check("gallery_listing_url shape",
          M.gallery_listing_url(BASE, "gallery1", "published"),
          BASE + "/cgi/editor_gallery.cgi?context=gallery1"
                 "&show_state=published&search_field=abstract"
                 "&search_op=eq&search_value=")
    # The flavor flag reaches every per-item URL, so a spurious flip in
    # crawl_listing sends the whole run to the wrong editor.
    check("admin_record_url follows the gallery flag (gallery)",
          M.admin_record_url(BASE, "gallery1", "2001", True),
          BASE + "/cgi/editor_gallery.cgi?display=submission&article=2001"
                 "&context=gallery1")
    check("admin_record_url follows the gallery flag (edikit)",
          M.admin_record_url(BASE, "series", "1001", False),
          BASE + "/cgi/editor.cgi?window=abstract&article=1001"
                 "&context=series")


def test_parse_edikit_listing():
    page = M.parse_edikit_listing(EDIKIT_LISTING)
    check_true("edikit listing parses", page is not None)
    rows = page["rows"]
    check("edikit row count", len(rows), 2)
    check("edikit article ids", [r["article"] for r in rows],
          ["1001", "1002"])
    check("edikit titles", [r["title"] for r in rows],
          ["First item", "Second item"])
    check("edikit additional-files count parses", rows[0]["af"], 2)
    check("edikit additional-files dash means zero", rows[1]["af"], 0)
    check("edikit next link", page["next"],
          "editor.cgi?context=series&page=2")
    check_true("edikit parser rejects a page with no listing table",
               M.parse_edikit_listing("<html><body><p>nope</p></body></html>")
               is None)


def test_parse_gallery_listing():
    page = M.parse_gallery_listing(GALLERY_LISTING)
    check_true("gallery listing parses", page is not None)
    check("gallery row count", len(page["rows"]), 1)
    check("gallery article id", page["rows"][0]["article"], "2001")
    check("gallery title", page["rows"][0]["title"], "Plate one")
    # No Additional Files column at all — must be None, not 0. A 0 would
    # claim the record has no extra files; None says we do not know.
    check("gallery missing additional-files column is unknown",
          page["rows"][0]["af"], None)
    check("gallery next link", page["next"],
          "editor_gallery.cgi?context=gallery1&page=2")
    check_true("gallery parser rejects an EdiKit page",
               M.parse_gallery_listing(EDIKIT_LISTING) is None)
    check_true("gallery parser rejects a table with no id header",
               M.parse_gallery_listing(
                   "<table id='submission_list'><tr><th>Name</th></tr>"
                   "</table>") is None)


def test_flavor_detection_is_one_directional():
    """Characterization, not an endorsement.

    parse_gallery_listing is exclusive: it needs table#submission_list, so an
    EdiKit page returns None and the flip in crawl_listing fires. The EdiKit
    parser is NOT exclusive — _find_listing_table takes the first table whose
    header row carries "id" and "title", which a gallery table also has. So
    the wrong-flavor guess is only detectable in one direction: an EdiKit
    guess on a gallery page parses happily and never flips, leaving every
    per-item URL built for the wrong editor. Pinned so a change to either
    parser has to face this asymmetry deliberately.
    """
    check_true("edikit parser also parses a gallery table (asymmetry)",
               M.parse_edikit_listing(GALLERY_LISTING) is not None)


def test_listing_row_hazards():
    """Three things reading the code raised, none yet confirmed against real
    HTML. Recorded as expected-failures so they are neither forgotten nor
    treated as breakage. An XPASS here means the hypothesis was wrong."""
    # HAZARD 1 — the leading hypothesis for the reported regression.
    # _listing_rows is documented as serving "both flavors" but hard-codes
    # the article id to the FIRST cell — `"article": texts[0]` — and gates
    # every row on `texts[0].isdigit()`, while carrying an hmap that knows
    # which column "id" actually is. col("id") sits unused two lines below.
    # If one flavor puts a select box in column 0, every data row fails the
    # digit gate and the structure yields ZERO items: no error, no files,
    # which is what "broken" looks like from the outside.
    # FIXED in v1.5 (was an XFAIL): _listing_rows reads the id from the
    # column the header map names. Both real listings happen to put ID
    # first, so this is a latent-fragility fix rather than a live one — but
    # the failure mode it removes was "zero rows for the whole structure,
    # silently", which is not worth leaving to chance.
    page = M.parse_edikit_listing(EDIKIT_CHECKBOX_FIRST)
    rows = page["rows"] if page else []
    check("a leading select column still yields the row", len(rows), 1)
    check("the id comes from the ID column, not cell 0",
          rows[0]["article"] if rows else None, "1001")

    # HAZARD 2 — the row-width tolerance is a magic number with no comment
    # explaining the 2. Pinned as a question, not a claim: a flavor whose
    # data rows are narrower than its header row loses rows silently.
    narrow = ("<table><tr><th>ID</th><th>Title</th><th>Last Event</th>"
              "<th>Submitted</th><th>Additional Files</th><th>Notes</th>"
              "</tr><tr><td>1003</td><td>Third</td><td>posted</td></tr>"
              "</table>")
    page = M.parse_edikit_listing(narrow)
    check("a row narrower than len(hmap)-2 is dropped today",
          len(page["rows"]) if page else None, 0)


def _stub_fetch(pages, calls):
    """Replace fetch_html with a scripted dict of url -> html."""
    def fake(url, cookie, timeout=None):
        calls.append(url)
        if url not in pages:
            raise AssertionError("unscripted fetch: " + url)
        html = pages[url]
        lowered = html.lower()
        if 'name="password"' in lowered and 'name="login"' in lowered:
            raise M.LoginRequired()
        return html
    return fake


def test_crawl_listing():
    real_fetch, real_throttle = M.fetch_html, M._throttle
    M._throttle = lambda *a, **k: None
    try:
        # One page, supplied by the driver: no fetch at all.
        calls = []
        M.fetch_html = _stub_fetch({}, calls)
        rows, gallery = M.crawl_listing(
            BASE, "gallery1", True, "published", "c",
            first_html=GALLERY_LISTING_LASTPAGE)
        check("driver-supplied first page needs no fetch", calls, [])
        check("single-page crawl returns its rows", len(rows), 1)
        check_true("single-page crawl keeps the flavor", gallery is True)

        # Two pages, and the second repeats a row: dedupe by article, and
        # stop once a page adds nothing new.
        calls = []
        p2 = ("<table id='submission_list'><tr><th>ID</th><th>Title</th>"
              "</tr><tr><td>2001</td><td>Plate one</td></tr>"
              "<tr><td>2002</td><td>Plate two</td></tr></table>")
        M.fetch_html = _stub_fetch(
            {BASE + "/cgi/editor_gallery.cgi?context=gallery1&page=2": p2},
            calls)
        rows, _ = M.crawl_listing(BASE, "gallery1", True, "published", "c",
                                  first_html=GALLERY_LISTING)
        check("pagination follows the next link once", len(calls), 1)
        check("pagination dedupes by article id",
              [r["article"] for r in rows], ["2001", "2002"])

        # The flip: the caller guessed gallery, the page is EdiKit, and the
        # other parser understands it — flip without a refetch.
        calls = []
        M.fetch_html = _stub_fetch({}, calls)
        rows, gallery = M.crawl_listing(
            BASE, "series", True, "published", "c",
            first_html=EDIKIT_LISTING_LASTPAGE)
        check("a wrong flavor guess flips", gallery, False)
        check("the flip needs no second fetch", calls, [])
        check("the flip keeps the rows it already had", len(rows), 2)

        # HAZARD 3 — first_html comes from the driver and never passes
        # through fetch_html, so it is never checked for a login bounce.
        # looks_like_login_page() exists for exactly this and is not called
        # here. An expired session therefore presents as a flavor problem:
        # the flip fires, and if the refetch happens to succeed the run
        # continues with a spuriously flipped flag on every per-item URL.
        calls = []
        M.fetch_html = _stub_fetch(
            {M.edikit_listing_url(BASE, "gallery1", "published"):
             EDIKIT_LISTING_LASTPAGE}, calls)
        outcome = None
        try:
            rows, gallery = M.crawl_listing(BASE, "gallery1", True,
                                            "published", "c",
                                            first_html=LOGIN_PAGE)
            outcome = "returned %d row(s), gallery=%s" % (len(rows), gallery)
        except M.LoginRequired:
            outcome = "LoginRequired"
        except ValueError as e:
            outcome = "ValueError: %s" % e
        # Was an XFAIL — a suspected defect, carried for weeks. On
        # 2026-09-22 the same gap on the sibling path let an expired
        # session be reported as an empty record state, and a real
        # record as gone. A suspected defect that has since been
        # watched happening is not a hypothesis any more.
        check("login-bounce . driver HTML that is a login page raises "
              "LoginRequired rather than parsing as an empty state",
              outcome, "LoginRequired")
    finally:
        M.fetch_html, M._throttle = real_fetch, real_throttle


# ---------------------------------------------------------------------------
# 13 · The supplemental gate. Regression test for a confirmed data loss.
#
# ETD article 4395 holds two supplemental files, one shown and one hidden,
# while its listing cell reads "-" -> af=0. Until v1.5 that made
# plan_item_jobs skip the Supplemental content page entirely and both files
# left the run with no error. The listing count is a cross-check now, never
# a gate, and the page always wins.
# ---------------------------------------------------------------------------
SUPP_PAGE = """<html><body>
<table id="uploaded-files">
  <tr><th>Filename</th><th>Description</th><th>Upload new version</th>
      <th>Sort</th><th>Show</th></tr>
  <tr><td><a href="https://dc.example.edu/cgi/viewcontent.cgi?filename=0
&article=99&context=c&type=additional&preview_mode=1">notes-one.txt</a></td>
      <td></td><td></td><td></td>
      <td><input type="checkbox" name="show0" checked></td></tr>
  <tr><td><a href="https://dc.example.edu/cgi/viewcontent.cgi?filename=1
&article=99&context=c&type=additional&preview_mode=1">notes-two.txt</a></td>
      <td></td><td></td><td></td>
      <td><input type="checkbox" name="show1"></td></tr>
</table></body></html>""".replace("\n&article", "&article")

SUPP_OPTS = {"primary": False, "supp": True, "native": False,
             "public": True, "hidden": True, "versions": "current",
             "published": True, "unpublished": False}


def _plan_supp(af, page):
    real_fetch, real_throttle = M.fetch_html, M._throttle
    M._throttle = lambda *a, **k: None
    M.fetch_html = lambda url, cookie, timeout=None: page
    warns = []
    try:
        jobs = M.plan_item_jobs(BASE, "c", {"article": "99", "title": "T",
                                            "af": af},
                                "published", False, dict(SUPP_OPTS), "cookie",
                                warns.append)
    finally:
        M.fetch_html, M._throttle = real_fetch, real_throttle
    return jobs, warns


def test_supplemental_gate():
    # the real-world case: the listing says none, the page holds two
    jobs, warns = _plan_supp(0, SUPP_PAGE)
    check("supp · af=0 no longer skips the page", len(jobs), 2)
    check("supp · both files are queued",
          [j["kind"] for j in jobs], ["supp1", "supp2"])
    check("supp · shown file is public", jobs[0]["visibility"], "public")
    check("supp · unshown file is hidden", jobs[1]["visibility"], "hidden")
    check("supp · original names are carried",
          [j["orig_hint"] for j in jobs], ["notes-one.txt", "notes-two.txt"])
    check_true("supp · the discrepancy is reported",
               any("counts 0" in w and "lists 2" in w for w in warns), warns)

    # listing and page agree: no noise
    jobs, warns = _plan_supp(2, SUPP_PAGE)
    check("supp · agreement still queues both", len(jobs), 2)
    check("supp · agreement warns about nothing", warns, [])

    # the listing promised files and the page has no table: say so
    jobs, warns = _plan_supp(2, "<html><body><p>nothing</p></body></html>")
    check("supp · missing table queues nothing", len(jobs), 0)
    check_true("supp · missing table is reported when files were promised",
               any("no table" in w for w in warns), warns)

    # the ordinary case — no files claimed, no table, no complaint
    jobs, warns = _plan_supp(0, "<html><body><p>nothing</p></body></html>")
    check("supp · no files and no table queues nothing", len(jobs), 0)
    check("supp · no files and no table is silent", warns, [])


# ---------------------------------------------------------------------------
# 14 · describe_error — the report has to name the failure.
#
# Run A on 2026-09-08 reported "Failed: HTTPError" for three rows. The status
# code was never shown, so the only thing anyone could say was "4xx on
# everything", and that is not diagnosable. HTTPError already renders itself
# with the code; the old code excluded it from an isinstance whitelist and
# fell through to the class name.
# ---------------------------------------------------------------------------
def test_describe_error():
    import io
    import urllib.error

    def http(code, reason):
        return urllib.error.HTTPError("http://x/y", code, reason, {},
                                      io.BytesIO(b""))

    check("describe · 403 names the code and reason",
          M.describe_error(http(403, "Forbidden")), "HTTP 403 Forbidden")
    check("describe · 404 names the code and reason",
          M.describe_error(http(404, "Not Found")), "HTTP 404 Not Found")
    check("describe · a network failure names the reason",
          M.describe_error(urllib.error.URLError("connection reset")),
          "network error: connection reset")
    # our own exceptions already carry written messages; pass them through
    check("describe · our ValueError passes through unchanged",
          M.describe_error(ValueError("server returned a web page instead "
                                      "of a file")),
          "server returned a web page instead of a file")
    check("describe · NoFileAvailable passes through unchanged",
          M.describe_error(M.NoFileAvailable("no downloadable file — "
                                             "derivative: A | native: B")),
          "no downloadable file — derivative: A | native: B")
    # anything else gets its class AND its text, never the class alone
    check("describe · an unexpected error keeps its text",
          M.describe_error(RuntimeError("something odd")),
          "RuntimeError: something odd")
    check_true("describe · never returns a bare class name when text exists",
               M.describe_error(OSError("disk full")) != "OSError",
               M.describe_error(OSError("disk full")))


def test_no_file_available_names_both_attempts():
    """The native is the fallback that was supposed to work, so its error is
    the informative one. v1.5 reported the derivative's and dropped it.

    This test used to use a 403 on the native and expect NoFileAvailable.
    v1.8 changed that deliberately and the test was wrong, not the code: a
    403 means "you may not have this", never "this does not exist", and
    Digital Commons serves access-restricted content behind exactly that.
    Reporting a restricted ETD file as a record with no file is the same
    false statement this release exists to stop. So the absence case now
    uses a 404, and the 403 is asserted to propagate.
    """
    import io
    import urllib.error
    real_throttle = M._throttle
    M._throttle = lambda *a, **k: None
    real_fetch = M.fetch_file
    try:
        def fetch_with(native_exc):
            def fake_fetch(url, cookie, timeout=None):
                if "type/native" in url:
                    raise native_exc
                raise ValueError("server returned a web page instead "
                                 "of a file")
            M.fetch_file = fake_fetch
            try:
                M.fetch_record_file(BASE + "/cgi/viewcontent.cgi?article=1",
                                    "cookie",
                                    BASE + "/context/c/article/1/type/native"
                                           "/viewcontent?preview_mode=1")
                return "no exception"
            except Exception as e:
                return e

        got = fetch_with(urllib.error.HTTPError(
            "u", 404, "Not Found", {}, io.BytesIO(b"")))
        check_true("NoFileAvailable was raised",
                   isinstance(got, M.NoFileAvailable), "got {!r}".format(got))
        msg = str(got)
        check_true("both attempts are named",
                   "derivative:" in msg and "native:" in msg, msg)
        check_true("the native's status code survives",
                   "HTTP 404 Not Found" in msg, msg)

        got = fetch_with(urllib.error.HTTPError(
            "u", 403, "Forbidden", {}, io.BytesIO(b"")))
        check_true("a 403 on the native is not an absence verdict",
                   not isinstance(got, M.NoFileAvailable),
                   "got NoFileAvailable: {}".format(got))
        check_true("a 403 propagates with its status intact",
                   isinstance(got, urllib.error.HTTPError)
                   and got.code == 403, "got {!r}".format(got))
        check_true("a 403 row names the status, not an absence",
                   M.describe_error(got).endswith("— HTTP 403 Forbidden"),
                   M.describe_error(got))
        # v1.34.2: and names the native it came from, since the row's
        # File URL is the derivative's.
        check_true("a not-confirmed fallback names the native's URL",
                   "native fallback, " + BASE + "/context/c/article/1/type"
                   "/native/viewcontent?preview_mode=1," in
                   M.describe_error(got), M.describe_error(got))
    finally:
        M.fetch_file = real_fetch
        M._throttle = real_throttle


# ---------------------------------------------------------------------------
# Rate limiting. The defect these pin: on 2026-09-08 a run met HTTP 429 at
# its tail, the 429 fell through the native-fallback path, and twelve
# records that HAVE files were reported as "no downloadable file". A
# throttle says nothing about whether a file exists, so the two must never
# be able to collapse into one verdict again.
# ---------------------------------------------------------------------------
def _http(code, reason, headers=None):
    return urllib.error.HTTPError("http://x/y", code, reason,
                                  headers or {}, io.BytesIO(b""))


def test_retry_after_parsing():
    check("retry-after · absent header is 0",
          M._retry_after_seconds(_http(429, "Too Many Requests", {})), 0.0)
    check("retry-after · integer seconds are honored",
          M._retry_after_seconds(
              _http(429, "Too Many Requests", {"Retry-After": "12"})), 12.0)
    check("retry-after · garbage is 0, not an exception",
          M._retry_after_seconds(
              _http(429, "Too Many Requests", {"Retry-After": "soon"})), 0.0)
    # an HTTP-date is equally legal; a date in the past must not go negative
    past = M._retry_after_seconds(
        _http(429, "x", {"Retry-After": "Wed, 21 Oct 2015 07:28:00 GMT"}))
    check("retry-after · a past HTTP-date clamps to 0", past, 0.0)
    check_true("retry-after · an object with no headers is 0, not a crash",
               M._retry_after_seconds(object()) == 0.0)


def test_open_retries_only_rate_limits():
    """_open must retry 429/503 and pass every other status straight up."""
    real_urlopen, real_sleep = urllib.request.urlopen, M._sleep
    M._sleep = lambda secs: True   # _sleep is wall-clock; it has its own test
    calls = []
    try:
        # a 404 is answered once and re-raised as itself
        def once_404(req, timeout=None):
            calls.append("404")
            raise _http(404, "Not Found")
        urllib.request.urlopen = once_404
        M.reset_rate_state()
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("open · a 404 raises", False, "no exception")
        except urllib.error.HTTPError as e:
            check("open · a 404 is raised unchanged", e.code, 404)
        except Exception as e:
            check_true("open · a 404 is not converted",
                       False, "got {}".format(e.__class__.__name__))
        check("open · a 404 is attempted exactly once", len(calls), 1)
        check("open · a 404 does not count as a rate limit",
              M.rate_state()["hits"], 0)

        # a 429 that never clears becomes RateLimited, never HTTPError
        calls.clear()

        def always_429(req, timeout=None):
            calls.append("429")
            raise _http(429, "Too Many Requests")
        urllib.request.urlopen = always_429
        M.reset_rate_state()
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("open · a persistent 429 raises", False, "no exception")
        except M.RateLimited as e:
            check_true("open · a persistent 429 raises RateLimited", True)
            check_true("open · the message names the status",
                       "429" in str(e), str(e))
            check_true("open · the message says the file is not missing",
                       "not saying the file is missing" in str(e), str(e))
        except Exception as e:
            check_true("open · a 429 becomes RateLimited", False,
                       "got {}: {}".format(e.__class__.__name__, e))
        check("open · a 429 is retried, not attempted once",
              len(calls), M.RATE_RETRIES + 1)
        check("open · every refusal is counted",
              M.rate_state()["hits"], M.RATE_RETRIES + 1)

        # a 429 that clears returns the response
        calls.clear()
        seq = ["429", "429", "ok"]

        def clears(req, timeout=None):
            calls.append(seq[len(calls)])
            if calls[-1] == "429":
                raise _http(429, "Too Many Requests")
            return "RESPONSE"
        urllib.request.urlopen = clears
        M.reset_rate_state()
        check("open · a cleared 429 returns the response",
              M._open("http://x/y", "c=1", 5), "RESPONSE")
        check("open · it stopped retrying once it succeeded", len(calls), 3)

        # 503 is treated the same way
        calls.clear()
        urllib.request.urlopen = lambda req, timeout=None: (
            calls.append("503") or (_ for _ in ()).throw(_http(503, "Busy")))
        M.reset_rate_state()
        try:
            M._open("http://x/y", "c=1", 5)
        except M.RateLimited:
            check_true("open · 503 is retried like 429", len(calls) > 1,
                       "attempts: {}".format(len(calls)))
        except Exception as e:
            check_true("open · 503 becomes RateLimited", False,
                       "got {}".format(e.__class__.__name__))
    finally:
        urllib.request.urlopen = real_urlopen
        M._sleep = real_sleep
        M.reset_rate_state()


def test_rate_limit_never_becomes_no_file():
    """The regression itself: a throttled derivative is not an absence.

    fetch_record_file must re-raise RateLimited and must NOT spend a second
    request on the native, because the pair would be reported as
    "no downloadable file" for a record that has one.
    """
    real_fetch, real_throttle = M.fetch_file, M._throttle
    attempts = []
    try:
        def throttled(url, cookie, timeout=M.FILE_TIMEOUT):
            attempts.append(url)
            raise M.RateLimited("rate limited: HTTP 429 Too Many Requests "
                                "after 5 attempts")
        M.fetch_file = throttled
        M._throttle = lambda *a, **k: None
        try:
            M.fetch_record_file("http://x/derivative", "c=1",
                                "http://x/native")
            check_true("rate limit · fetch_record_file raises", False,
                       "no exception")
        except M.NoFileAvailable as e:
            check_true("rate limit · is NOT reported as an absence", False,
                       "got NoFileAvailable: {}".format(e))
        except M.RateLimited:
            check_true("rate limit · propagates as RateLimited", True)
        except Exception as e:
            check_true("rate limit · propagates as RateLimited", False,
                       "got {}: {}".format(e.__class__.__name__, e))
        check("rate limit · the native is not attempted after a throttle",
              attempts, ["http://x/derivative"])

        # The case that actually reached users, and the only one of the two
        # that a hollow test would miss: the derivative is legitimately
        # absent (a web page came back, which is normal for an image), and
        # then the NATIVE request is the one that gets throttled. The second
        # attempt's handler catches bare Exception, so without an explicit
        # re-raise this lands in NoFileAvailable and the record is reported
        # as empty. Planting that violation is how this test earned its
        # place: the first case above passes with the fix removed, because
        # RateLimited was never in the first handler's tuple anyway.
        attempts.clear()

        def absent_then_throttled(url, cookie, timeout=M.FILE_TIMEOUT):
            attempts.append(url)
            if url == "http://x/derivative":
                raise ValueError("server returned a web page instead "
                                 "of a file")
            raise M.RateLimited("rate limited: HTTP 429 Too Many Requests "
                                "after 5 attempts")
        M.fetch_file = absent_then_throttled
        try:
            M.fetch_record_file("http://x/derivative", "c=1",
                                "http://x/native")
            check_true("rate limit · a throttled native raises", False,
                       "no exception")
        except M.NoFileAvailable as e:
            check_true("rate limit · a throttled NATIVE is not an absence",
                       False, "got NoFileAvailable: {}".format(e))
        except M.RateLimited:
            check_true("rate limit · a throttled native propagates", True)
        except Exception as e:
            check_true("rate limit · a throttled native propagates", False,
                       "got {}: {}".format(e.__class__.__name__, e))
        check("rate limit · both URLs were tried in that order", attempts,
              ["http://x/derivative", "http://x/native"])

        # and the row text a user reads names the throttle, not an absence
        msg = M.describe_error(M.RateLimited(
            "rate limited: HTTP 429 Too Many Requests after 5 attempts"))
        check_true("rate limit · the report row says 'rate limited'",
                   "rate limited" in msg, msg)
        check_true("rate limit · the report row does not say 'no "
                   "downloadable file'",
                   "no downloadable file" not in msg, msg)
        check_true("rate limit · the row is never a bare class name",
                   msg != "RateLimited", msg)
    finally:
        M.fetch_file = real_fetch
        M._throttle = real_throttle


def test_adaptive_delay():
    """A throttled run slows down, within a ceiling, and does not leak the
    penalty into the next run."""
    M.reset_rate_state()
    check("delay · a fresh run starts with no added delay",
          M.rate_state()["extra"], 0.0)
    M._note_rate_limit()
    check("delay · one 429 adds one step",
          round(M.rate_state()["extra"], 6), round(M.RATE_DELAY_STEP, 6))
    for _ in range(200):
        M._note_rate_limit()
    check_true("delay · the slowdown is capped",
               M.rate_state()["extra"] <= M.RATE_DELAY_CAP,
               "extra={}".format(M.rate_state()["extra"]))
    check_true("delay · _throttle uses the added delay",
               True)
    M.reset_rate_state()
    st = M.rate_state()
    check("delay · reset clears the penalty for the next run",
          (st["extra"], st["hits"]), (0.0, 0))
    check("delay · reset also clears the cooldowns already waited out",
          st["cooldowns"], 0)
    check("delay · and restores the run's cooldown budget",
          st["budget"], M.COOLDOWN_MAX_WAITS)


# ---------------------------------------------------------------------------
# Only a real answer may become "no downloadable file". v1.7 fixed the 429
# it had seen and left the class open: on the next run Digital Commons
# pushed back with a 302 redirect loop instead, which is not in the
# rate-limit set, and records were once again reported as empty. These pin
# the distinction rather than the status code.
# ---------------------------------------------------------------------------
def test_pushback_classification():
    check_true("pushback . 429 is pushback", M.is_pushback(_http(429, "x")))
    check_true("pushback . 503 is pushback", M.is_pushback(_http(503, "x")))
    loop = _http(302, "The HTTP server returned a redirect error that would "
                      "lead to an infinite loop.")
    check_true("pushback . a 302 redirect loop is pushback",
               M.is_pushback(loop))
    check_true("pushback . a 307 loop is pushback too",
               M.is_pushback(_http(307, "infinite loop")))
    check_true("pushback . 404 is not pushback",
               not M.is_pushback(_http(404, "Not Found")))
    check_true("pushback . 403 is not pushback",
               not M.is_pushback(_http(403, "Forbidden")))
    check_true("pushback . a ValueError is not pushback",
               not M.is_pushback(ValueError("web page")))


def test_definite_absence_classification():
    check_true("absence . a web page where a file was expected is an answer",
               M.is_definite_absence(
                   ValueError("server returned a web page instead of a file")))
    check_true("absence . 404 is an answer",
               M.is_definite_absence(_http(404, "Not Found")))
    check_true("absence . 410 is an answer",
               M.is_definite_absence(_http(410, "Gone")))
    check_true("absence . 429 is NOT an answer",
               not M.is_definite_absence(_http(429, "Too Many Requests")))
    check_true("absence . a redirect loop is NOT an answer",
               not M.is_definite_absence(_http(302, "infinite loop")))
    check_true("absence . a timeout is NOT an answer",
               not M.is_definite_absence(urllib.error.URLError("timed out")))
    check_true("absence . a 500 is NOT an answer",
               not M.is_definite_absence(_http(500, "Server Error")))
    check_true("absence . a RateLimited is NOT an answer",
               not M.is_definite_absence(M.RateLimited("refused")))


def test_redirect_loop_is_retried_like_a_throttle():
    real_urlopen, real_sleep = urllib.request.urlopen, M._sleep
    M._sleep = lambda secs: True
    calls = []
    try:
        def always_loop(req, timeout=None):
            calls.append("302")
            raise _http(302, "The HTTP server returned a redirect error "
                             "that would lead to an infinite loop.")
        urllib.request.urlopen = always_loop
        M.reset_rate_state()
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("loop . a redirect loop raises", False, "no exception")
        except M.RateLimited as e:
            check_true("loop . a persistent redirect loop is RateLimited",
                       True)
            check_true("loop . the message says the file is not missing",
                       "not saying the file is missing" in str(e), str(e))
            check_true("loop . the message is one line, not urllib's three",
                       len(str(e).splitlines()) == 1, repr(str(e)))
        except Exception as e:
            check_true("loop . a redirect loop becomes RateLimited", False,
                       "got {}: {}".format(e.__class__.__name__, e))
        check_true("loop . it was retried, not attempted once",
                   len(calls) == M.RATE_RETRIES + 1,
                   "attempts: {}".format(len(calls)))
    finally:
        urllib.request.urlopen = real_urlopen
        M._sleep = real_sleep
        M.reset_rate_state()


def test_absence_requires_two_real_answers():
    """The defect this release fixes, in both directions.

    An absence verdict needs an answer from BOTH attempts. One answer plus
    one refusal is not an absence - that is how a throttled tail became
    records reported as empty. But two real answers must still produce the
    verdict, or a metadata-only record becomes a hard failure.
    """
    real_fetch, real_throttle = M.fetch_file, M._throttle
    M._throttle = lambda *a, **k: None
    try:
        def pair(derivative_exc, native_exc):
            def f(url, cookie, timeout=M.FILE_TIMEOUT):
                raise (derivative_exc if url.endswith("derivative")
                       else native_exc)
            M.fetch_file = f
            try:
                M.fetch_record_file("http://x/derivative", "c=1",
                                    "http://x/native")
                return "no exception"
            except Exception as e:
                return e

        page = ValueError("server returned a web page instead of a file")

        got = pair(page, _http(404, "Not Found"))
        check_true("absence . web page + 404 is still 'no downloadable file'",
                   isinstance(got, M.NoFileAvailable), "got {!r}".format(got))
        check_true("absence . and it still names both attempts",
                   "derivative:" in str(got) and "native:" in str(got),
                   str(got))

        got = pair(page, _http(302, "The HTTP server returned a redirect "
                                    "error that would lead to an infinite "
                                    "loop."))
        check_true("absence . web page + redirect loop is NOT an absence",
                   not isinstance(got, M.NoFileAvailable),
                   "got NoFileAvailable: {}".format(got))
        check_true("absence . the redirect loop is what gets reported",
                   isinstance(got, urllib.error.HTTPError)
                   and got.code == 302, "got {!r}".format(got))

        got = pair(_http(500, "Server Error"), page)
        check_true("absence . a 500 on the derivative is NOT an absence",
                   not isinstance(got, M.NoFileAvailable),
                   "got NoFileAvailable: {}".format(got))
        check_true("absence . the 500 is reported, not the page",
                   isinstance(got, urllib.error.HTTPError)
                   and got.code == 500, "got {!r}".format(got))

        msg = M.describe_error(pair(page, _http(302, "infinite loop")))
        check_true("absence . the row never says 'no downloadable file'",
                   "no downloadable file" not in msg, msg)
    finally:
        M.fetch_file = real_fetch
        M._throttle = real_throttle


def test_a_wait_is_announced():
    """A silent backoff is why a working run looked hung with no ETA."""
    real_urlopen, real_sleep = urllib.request.urlopen, M._sleep
    M._sleep = lambda secs: True
    said = []
    try:
        seq = ["429", "ok"]

        def clears(req, timeout=None):
            if seq.pop(0) == "429":
                raise _http(429, "Too Many Requests")
            return "RESPONSE"
        urllib.request.urlopen = clears
        M.reset_rate_state()
        M.set_note_hook(said.append)
        M._open("http://x/y", "c=1", 5)
        check_true("note . the wait is logged", len(said) == 1,
                   "said: {!r}".format(said))
        check_true("note . it names the status and the wait",
                   bool(said) and "429" in said[0] and "retrying" in said[0],
                   "said: {!r}".format(said))
        M.set_note_hook(lambda t: 1 / 0)
        seq[:] = ["429", "ok"]
        check("note . a raising hook does not break the fetch",
              M._open("http://x/y", "c=1", 5), "RESPONSE")
    finally:
        urllib.request.urlopen = real_urlopen
        M._sleep = real_sleep
        M.set_note_hook(None)
        M.reset_rate_state()


def test_stop_writes_what_it_has():
    """A stop mid-structure must keep the rows it already gathered.

    On 2026-09-08 a stop at item 33 of 38 wrote no report at all: 33 files
    sat on disk with nothing recording where they came from, because
    _PARTIAL only ever fired for a session expiry. download_worker needs
    Chrome and a hierarchy, so what is checkable here is the machinery the
    handler depends on - the column indices it counts with, and that they
    are derived rather than typed.
    """
    check("stop . the Status column is where the handler thinks",
          M.REPORT_HEADERS[M.COL_STATUS], "Status")
    check("stop . the Visibility column is where the handler thinks",
          M.REPORT_HEADERS[M.COL_VISIBILITY], "Visibility")

    # There are three `except StopRequested:` in the module - map_worker's,
    # an inner re-raise in the item loop, and the one that ends a download
    # run. Anchor on the run-ending one by the line only it logs, so this
    # test cannot quietly start inspecting a different handler.
    src = Path(M.__file__).read_text(encoding="utf-8")
    marker = 'log("Stop requested — {}.".format(\n                structures_done_text('
    check("stop . exactly one handler ends a download run",
          src.count(marker), 1)
    end = src.index(marker)
    start = src.rindex("    except StopRequested:", 0, end)
    handler = src[start:src.index("    except Exception as e:", end)]
    # The write lives in flush_pending() now, called from BOTH early-exit
    # paths. It was inline in the stop handler until 2026-09-08, which meant
    # an unexpected error - a full disk on an overnight run - lost the
    # record of everything already downloaded in that structure.
    body = src[src.index("def download_worker("):src.index("# Worker 3:")]
    flush = body[body.index("def flush_pending("):
                 body.index("def finish_master(")]
    check_true("stop . flush_pending writes a partial structure report",
               "write_structure_report" in flush, flush[:200])
    check_true("stop . the partial report is named _PARTIAL",
               "_PARTIAL.xlsx" in flush, flush[:200])
    check_true("stop . it counts with the derived columns, not literals",
               "COL_STATUS" in flush and "r[11]" not in flush, flush[:400])
    check_true("stop . a failed report write does not lose the master row",
               "could not write the partial report" in flush, flush[:400])
    check_true("stop . the stop path flushes",
               "flush_pending(" in handler, handler[:300])
    # the generic handler is the one that was missing it
    generic = body[body.rindex("except Exception as e:"):]
    check_true("stop . an unexpected error flushes too",
               "flush_pending(" in generic, generic[:300])
    check_true("stop . and it still writes the master log",
               "finish_master(" in generic, generic[:300])
    check("stop . both early exits flush, and only those two",
          body.count("flush_pending(\""), 2)

    # the row shape the handler indexes into must be the one the loop builds
    row = (1, "1037", "T", "Posted", "supp1", "hidden", "current", "",
           "a.jpg", "b.jpg", 10, "Downloaded", "u", "p", "a", "12:00:00",
           "image/jpg", "restricted", "Thesis (restricted to X)", "1-1-2030",
           "", "Chrome")
    check("stop . a real row's status reads back",
          row[M.COL_STATUS], "Downloaded")
    check("stop . a real row's visibility reads back",
          row[M.COL_VISIBILITY], "hidden")
    check("stop . a real row's access reads back",
          row[M.COL_ACCESS], "restricted")
    check("stop . the row is the width of the header",
          len(row), len(M.REPORT_HEADERS))


# ---------------------------------------------------------------------------
# Retry-After. On 2026-09-08 DC asked for longer than the 20s backoff cap,
# the cap silently overrode it, and three retries were burned against a wall
# whose height the server had already stated. The server's number wins up to
# a ceiling; past the ceiling, retrying is pointless and the figure is what
# the operator needs.
# ---------------------------------------------------------------------------
def test_retry_after_is_honored_not_clamped():
    real_urlopen, real_sleep = urllib.request.urlopen, M._sleep
    slept, calls = [], []
    M._sleep = lambda secs: (slept.append(secs), True)[1]
    try:
        def refuse(after):
            def f(req, timeout=None):
                calls.append(1)
                raise _http(429, "Too Many Requests",
                            {"Retry-After": str(after)})
            return f

        # a wait the run will sit through is used verbatim, not capped
        urllib.request.urlopen = refuse(45)
        M.reset_rate_state(); slept.clear(); calls.clear()
        try:
            M._open("http://x/y", "c=1", 5)
        except M.RateLimited:
            pass
        check("retry-after . 45s is honored verbatim, not cut to the cap",
              slept, [45.0, 45.0, 45.0])
        check_true("retry-after . 45s is above the backoff cap, so this "
                   "test would fail under the old clamp",
                   45 > M.RATE_BACKOFF_CAP)

        # a named lockout, with the run set not to wait: fail at once,
        # name the number, and do not retry - the wall does not move.
        urllib.request.urlopen = refuse(600)
        M.reset_rate_state(); M.set_cooldown_budget(0)
        slept.clear(); calls.clear()
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("retry-after . a lockout raises", False, "no exc")
        except M.RateLimited as e:
            check_true("retry-after . the message names the server's figure",
                       "600s" in str(e), str(e))
            check_true("retry-after . it says the file is not missing",
                       "not missing" in str(e), str(e))
            check_true("retry-after . it says why it did not wait",
                       "set not to wait" in str(e), str(e))
        except Exception as e:
            check_true("retry-after . a lockout is RateLimited", False,
                       "got {}: {}".format(e.__class__.__name__, e))
        check("retry-after . a lockout is not retried", len(calls), 1)
        check("retry-after . and nothing was slept", slept, [])

        # no header at all: the jittered backoff, still capped
        urllib.request.urlopen = refuse("")
        M.reset_rate_state(); slept.clear()
        try:
            M._open("http://x/y", "c=1", 5)
        except M.RateLimited:
            pass
        check_true("retry-after . with no header the backoff stays capped",
                   all(x <= M.RATE_BACKOFF_CAP + 1e-9 for x in slept),
                   "slept: {!r}".format(slept))
        check_true("retry-after . and it did retry", len(slept) >= 1,
                   "slept: {!r}".format(slept))
    finally:
        urllib.request.urlopen = real_urlopen
        M._sleep = real_sleep
        M.reset_rate_state()


def test_a_long_wait_stays_stoppable():
    """A 120s wait must not make Stop look broken for two minutes.

    Driven with a fake clock. A real one would make this test take as long
    as the wait it is checking, which is how the suite hung the first time
    it was written.
    """
    real_sleep, real_mono = M.time.sleep, M.time.monotonic
    naps, now = [], [1000.0]
    M.time.sleep = lambda s: (naps.append(s), now.__setitem__(0, now[0] + s))
    M.time.monotonic = lambda: now[0]
    try:
        M.set_stop_hook(lambda: False)
        check_true("sleep . a 5s wait completes", M._sleep(5) is True)
        check("sleep . it naps in one-second slices, not one long block",
              naps, [1.0] * 5)
        check_true("sleep . a zero wait naps not at all",
                   M._sleep(0) is True and len(naps) == 5)

        M.set_stop_hook(lambda: True)
        naps.clear()
        check_true("sleep . a stop cuts a two-minute wait short",
                   M._sleep(120) is False)
        check("sleep . and it does not nap at all once stopping", naps, [])

        # a stop that arrives partway through must still be noticed
        state = {"n": 0}

        def stop_after_three():
            state["n"] += 1
            return state["n"] > 3
        M.set_stop_hook(stop_after_three)
        naps.clear()
        check_true("sleep . a stop mid-wait ends it",
                   M._sleep(60) is False)
        check_true("sleep . after about three slices, not sixty",
                   len(naps) <= 4, "naps: {!r}".format(naps))

        M.set_stop_hook(lambda: 1 / 0)
        check_true("sleep . a raising stop hook is treated as 'not stopping'",
                   M._stopping() is False)
        M.set_stop_hook(None)
        check_true("sleep . no hook at all is 'not stopping'",
                   M._stopping() is False)
    finally:
        M.time.sleep, M.time.monotonic = real_sleep, real_mono
        M.set_stop_hook(None)


# ---------------------------------------------------------------------------
# Reports may live apart from the content. A run pulls hidden, unpublished
# and access-restricted files to disk; the reports carry only URLs, statuses
# and counts. The module had one folder for both, and the project's own
# instructions claimed otherwise - the operator had been hand-copying
# reports out of the download folder every run to work around it.
# ---------------------------------------------------------------------------
def test_report_destinations():
    D = M.report_destinations
    run, rpt, split = D("/dl", "", "S")
    check("dest . no report folder keeps the old single run folder",
          (run, rpt, split),
          (os.path.join("/dl", "DC_FileDownloads_S"),
           os.path.join("/dl", "DC_FileDownloads_S"), False))

    run, rpt, split = D("/dl", "/dl", "S")
    check_true("dest . naming the same folder is not a split", not split)
    check("dest . and the reports still land beside the files", run, rpt)

    run, rpt, split = D("/dl", "/dl/", "S")
    check_true("dest . a trailing slash is the same folder, not a split",
               not split, "run={} rpt={}".format(run, rpt))

    run, rpt, split = D("/dl", "/dl/../dl", "S")
    check_true("dest . a roundabout path to the same folder is not a split",
               not split, "run={} rpt={}".format(run, rpt))

    run, rpt, split = D("/dl", "/reports", "S")
    check_true("dest . a different folder splits", split)
    check("dest . files keep their run folder", run,
          os.path.join("/dl", "DC_FileDownloads_S"))
    check("dest . reports get a run folder of their own", rpt,
          os.path.join("/reports", "DC_FileDownloads_S"))
    check_true("dest . both run folders share the run's stamp",
               run.endswith("_S") and rpt.endswith("_S"))


def test_reports_are_written_where_the_split_says():
    """Source-level, because the write sites are inside download_worker.

    The hazard is specific: both writes used to take ctx_dir and run_dir,
    and a split that computes the right paths but is not used at the write
    sites would look correct in every unit test above.
    """
    src = Path(M.__file__).read_text(encoding="utf-8")
    body = src[src.index("def download_worker("):src.index("# Worker 3:")]

    check_true("split . download_worker takes a report folder",
               "report_dir=\"\"" in src[src.index("def download_worker("):
                                       src.index("def download_worker(") + 200],
               src[src.index("def download_worker("):][:200])
    check_true("split . the destinations come from the tested helper",
               "report_destinations(" in body)
    check_true("split . the master log is written to the report run folder",
               'os.path.join(report_run_dir, fname)' in body)
    check_true("split . the structure report is written to rpt_dir",
               "write_structure_report(os.path.join(rpt_dir, report_name)"
               in body)
    check_true("split . the master log no longer goes to the download folder",
               "os.path.join(run_dir, fname)" not in body)
    check_true("split . the stop handler writes to the report folder too",
               '"dir": rpt_dir' in body)
    check_true("split . map_worker follows the report folder as well",
               "def map_worker(session, out_dir, report_dir=\"\")" in src)
    # and the operator is told where each went, since it is now two places
    check_true("split . a split run says where each kind of output went",
               "Reports go to" in body)


# ---------------------------------------------------------------------------
# Waiting out a cooldown. The 2026-09-08 evidence: ten refusals reported
# 600s down to 563s, and time+Retry-After named the same instant in every
# one. The wall is fixed, so one wait clears the whole backlog - which is
# why the run waits rather than failing ten rows it could have fetched.
# ---------------------------------------------------------------------------
def test_a_cooldown_is_waited_out_within_budget():
    real_urlopen, real_sleep = urllib.request.urlopen, M._sleep
    slept, calls = [], []
    M._sleep = lambda secs: (slept.append(secs), True)[1]
    said = []
    try:
        # refuses once with a 600s lockout, then answers
        def once_then_ok(req, timeout=None):
            calls.append(1)
            if len(calls) == 1:
                raise _http(429, "Too Many Requests",
                            {"Retry-After": "600"})
            return "RESPONSE"
        urllib.request.urlopen = once_then_ok
        M.reset_rate_state(); M.set_cooldown_budget(3)
        M.set_note_hook(said.append)
        check("cooldown . the request succeeds after the wait",
              M._open("http://x/y", "c=1", 5), "RESPONSE")
        check("cooldown . one wait was spent",
              M.rate_state()["cooldowns"], 1)
        check_true("cooldown . it waited about the stated interval",
                   abs(sum(slept) - (600 + M.COOLDOWN_GRACE)) < 1e-6,
                   "slept {} total: {!r}".format(sum(slept), slept))
        check_true("cooldown . in slices, so it can be interrupted and "
                   "can report progress",
                   all(x <= M.COOLDOWN_TICK + 1e-9 for x in slept)
                   and len(slept) > 1, "slept: {!r}".format(slept))
        check_true("cooldown . the wait is announced on the way in",
                   any("locked this run out" in m for m in said), "{!r}".format(said[:3]))
        check_true("cooldown . and counts down while it waits",
                   sum(1 for m in said if "to go" in m) > 1,
                   [m for m in said if "to go" in m][:3])
        check_true("cooldown . and says when it resumes",
                   any("resuming" in m for m in said), "{!r}".format(said[-3:]))
        check_true("cooldown . the wait did not consume a retry",
                   len(calls) == 2, "attempts: {}".format(len(calls)))

        # budget exhausted: the next lockout is reported, not waited
        def always_locked(req, timeout=None):
            calls.append(1)
            raise _http(429, "x", {"Retry-After": "600"})
        urllib.request.urlopen = always_locked
        M.reset_rate_state(); M.set_cooldown_budget(2)
        calls.clear(); slept.clear()
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("cooldown . an endless lockout raises", False, "no exc")
        except M.RateLimited as e:
            check_true("cooldown . it stops after the budget",
                       "already waited out 2 cooldown(s)" in str(e), str(e))
        except Exception as e:
            check_true("cooldown . an endless lockout is RateLimited", False,
                       "got {}: {}".format(e.__class__.__name__, e))
        check("cooldown . exactly the budgeted number of waits were spent",
              M.rate_state()["cooldowns"], 2)

        # a lockout longer than the module will ever sit out
        M.reset_rate_state(); M.set_cooldown_budget(5)
        urllib.request.urlopen = lambda req, timeout=None: (_ for _ in ()) \
            .throw(_http(429, "x", {"Retry-After": "99999"}))
        slept.clear()
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("cooldown . an absurd lockout raises", False, "no exc")
        except M.RateLimited as e:
            check_true("cooldown . an absurd lockout is refused outright",
                       "longer than this run will ever wait" in str(e),
                       str(e))
        except Exception as e:
            check_true("cooldown . an absurd lockout is RateLimited", False,
                       "got {}".format(e.__class__.__name__))
        check("cooldown . and no budget was spent on it",
              M.rate_state()["cooldowns"], 0)
        check("cooldown . and nothing was slept", slept, [])
    finally:
        urllib.request.urlopen = real_urlopen
        M._sleep = real_sleep
        M.set_note_hook(None)
        M.reset_rate_state()


def test_a_cooldown_wait_is_stoppable():
    """Ten minutes is far too long for Stop to be unresponsive."""
    real_urlopen, real_sleep = urllib.request.urlopen, M._sleep
    said = []
    M._sleep = lambda secs: False          # as if Stop were pressed
    try:
        urllib.request.urlopen = lambda req, timeout=None: (_ for _ in ()) \
            .throw(_http(429, "x", {"Retry-After": "600"}))
        M.reset_rate_state(); M.set_cooldown_budget(3)
        M.set_note_hook(said.append)
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("cooldown-stop . it raises", False, "no exception")
        except M.RateLimited:
            check_true("cooldown-stop . a stop ends the wait and the "
                       "request, without retrying", True)
        except Exception as e:
            check_true("cooldown-stop . a stop yields RateLimited", False,
                       "got {}".format(e.__class__.__name__))
        check_true("cooldown-stop . the interruption is said out loud",
                   any("interrupted" in m for m in said), "{!r}".format(said))
    finally:
        urllib.request.urlopen = real_urlopen
        M._sleep = real_sleep
        M.set_note_hook(None)
        M.reset_rate_state()


def test_the_retry_loop_terminates_whatever_the_budget():
    """The loop must end on constants alone, not on shared mutable state.

    Written because planting a violation in the budget check did not fail
    this suite - it HUNG it. _open span forever against a server that never
    yielded, with no timeout and nothing to show. A bound that lives only
    in _RATE_STATE is one edit away from that, so there is now a second
    bound that depends on nothing but the constants, and this test drives
    the loop with the budget deliberately maxed out.
    """
    real_urlopen, real_sleep = urllib.request.urlopen, M._sleep
    M._sleep = lambda secs: True
    calls = []
    try:
        check("budget . a caller cannot ask for unbounded waiting",
              (M.set_cooldown_budget(10 ** 9), M.rate_state()["budget"])[1],
              M.COOLDOWN_HARD_MAX)
        check("budget . nor for a negative one",
              (M.set_cooldown_budget(-5), M.rate_state()["budget"])[1], 0)

        def never_yields(req, timeout=None):
            calls.append(1)
            raise _http(429, "Too Many Requests", {"Retry-After": "600"})
        urllib.request.urlopen = never_yields
        M.reset_rate_state()
        M.set_cooldown_budget(10 ** 9)          # as permissive as possible
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("budget . it raises rather than looping", False,
                       "no exception")
        except M.RateLimited:
            check_true("budget . the loop terminates", True)
        except Exception as e:
            check_true("budget . the loop terminates with RateLimited",
                       False, "got {}".format(e.__class__.__name__))
        bound = M.RATE_RETRIES + 2 * M.COOLDOWN_HARD_MAX + 4
        check_true("budget . and in a bounded number of attempts",
                   len(calls) <= bound,
                   "attempts {} vs bound {}".format(len(calls), bound))
        check_true("budget . the bound is small enough to be useful",
                   bound < 100, "bound is {}".format(bound))

        # The clamp above means the budget path ends the loop first, so the
        # hard bound is unreachable through the public API - which makes it
        # exactly the kind of guard that quietly stops working. Reach past
        # the clamp, as a wrong budget or a bad edit would, and check that
        # the loop still ends. Removing the hard bound must fail this.
        M.reset_rate_state()
        with M._RATE_LOCK:
            M._RATE_STATE["budget"] = 10 ** 6      # past what any caller can set
        calls.clear()
        try:
            M._open("http://x/y", "c=1", 5)
            check_true("bound . it raises rather than looping", False,
                       "no exception")
        except M.RateLimited as e:
            check_true("bound . a budget past the clamp still terminates",
                       "kept refusing" in str(e), str(e))
        except Exception as e:
            check_true("bound . a budget past the clamp still terminates",
                       False, "got {}".format(e.__class__.__name__))
        check_true("bound . and it is the constant bound that stopped it",
                   len(calls) <= bound,
                   "attempts {} vs bound {}".format(len(calls), bound))
        check_true("bound . the budget was not what ended it",
                   M.rate_state()["cooldowns"] > M.COOLDOWN_HARD_MAX,
                   "cooldowns spent: {}".format(
                       M.rate_state()["cooldowns"]))
    finally:
        urllib.request.urlopen = real_urlopen
        M._sleep = real_sleep
        M.reset_rate_state()


# ---------------------------------------------------------------------------
# Re-running only what failed. Built because DC locks out a sustained run,
# so a large collection has to be worked in chunks, and re-fetching what is
# already on disk is what provokes the next lockout.
# ---------------------------------------------------------------------------
# The module's own neutral fallback, never a brand literal. Hard-coding
# the institution's primary hex here failed the --generic build's
# identity check on 2026-09-08 - a test helper is still source, and
# source may not carry the institution's colors. The marker list is
# why that was caught, so the hex is deliberately not repeated even in
# this comment. Asking the module for the value beats pinning one.
def _brand():
    return M.FALLBACK_BRAND["primary"].lstrip("#")


def _fake_report(folder, name, rows):
    """Write a real report workbook, then read it back through the module."""
    path = os.path.join(folder, name)
    M.write_structure_report(path, rows, _brand())
    return path


def _row(order, article, kind, vis, version, status, url, adm,
         orig="", saved="", size="", ctype="", access="", release="",
         embargo="", plan="", via=""):
    # Built from REPORT_HEADERS rather than as a hand-counted literal. The
    # literal version silently became the wrong width the moment a column
    # was appended, and appending is documented as the safe change.
    row = {"Order": order, "Article ID": article, "Title": "A title",
           "Record State": "Posted", "File Kind": kind, "Visibility": vis,
           "Version": version, "Version Date": "", "Original Filename": orig,
           "Saved Filename": saved, "Size": size, "Status": status,
           "File URL": url, "Public Record URL": "http://pub/" + article,
           "Admin Record URL": adm, "Time": "12:00:00",
           "Content Type": ctype, "Access": access,
           "Type / Release Option": release, "Embargo": embargo, "Plan": plan,
           "Fetched via": via}
    missing = [h for h in M.REPORT_HEADERS if h not in row]
    if missing:
        raise AssertionError("_row does not build: " + ", ".join(missing))
    return tuple(row[h] for h in M.REPORT_HEADERS)


def test_every_workbook_this_module_opens_can_be_closed():
    """Windows CI, 2026-09-21 — the first Windows run this project had.

    Nine tests failed at `shutil.rmtree`, cleaning up a temp directory
    Windows would not let them delete, because a report workbook was
    still open. `load_workbook(...).active` leaves the Workbook itself
    unreferenced, and with read_only=True openpyxl holds the file until
    something closes it — which nothing could, because nothing had it.

    POSIX unlinks open files without complaint, so six months of Linux
    CI and every run on a Mac said nothing. In production on Windows it
    is worse than a test failure: every re-run reads a report through
    that function, so afterwards the operator cannot move, rename or
    overwrite their own report until the module process exits.

    Structural, with `ast`, because the defect is invisible to a grep
    for `close`: the file HAD a close, on the other workbook, and the
    count matched the number of readers. What it did not have was a
    NAME on this one.
    """
    import ast
    # BOTH FILES. Checking only the module is the same mistake one
    # layer up: the first fix closed the module's reader and three
    # Windows failures survived it, because THIS suite leaked the
    # identical way in a helper the failing tests all went through.
    # When the group is larger than a pair, count the group.
    here = os.path.dirname(os.path.abspath(__file__))
    opens, unclosable = 0, []
    for fname in ("dc_file_downloader.py", os.path.basename(__file__)):
        tree = ast.parse(open(os.path.join(here, fname),
                              encoding="utf-8").read())
        for n in ast.walk(tree):
            if isinstance(n, ast.Call):
                g = n.func
                if (isinstance(g, ast.Name) and g.id == "load_workbook") or \
                        (isinstance(g, ast.Attribute)
                         and g.attr == "load_workbook"):
                    opens += 1
            # load_workbook(...).<attr> — the Workbook is discarded at
            # once, so nothing holds it and nothing can close it.
            if isinstance(n, ast.Attribute) and isinstance(n.value, ast.Call):
                g = n.value.func
                if (isinstance(g, ast.Name) and g.id == "load_workbook") or \
                        (isinstance(g, ast.Attribute)
                         and g.attr == "load_workbook"):
                    unclosable.append("{}:{}".format(fname,
                                                     getattr(n, "lineno", 0)))
    check_true("handles . workbooks are opened somewhere, so this checks "
               "something", opens >= 2, "found {}".format(opens))
    check("handles . none is opened in a form that cannot be closed",
          unclosable, [])


def test_read_failed_rows():
    import tempfile
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article=7&context=etd"
    with tempfile.TemporaryDirectory() as d:
        rows = [
            _row(1, "7", "primary", "public", "current", "Downloaded",
                 base + "/cgi/viewcontent.cgi?article=7&context=etd", adm,
                 orig="a.pdf", saved="etd_7.pdf", size=10),
            _row(2, "7", "supp1", "hidden", "current",
                 "Failed: HTTP 429 Too Many Requests",
                 base + "/cgi/viewcontent.cgi?article=7&context=etd"
                        "&type=additional&filename=1", adm),
            _row(3, "8", "primary", "public", "current",
                 "Failed: server locked this run out for 600s",
                 base + "/cgi/viewcontent.cgi?article=8&context=etd",
                 base + "/cgi/editor.cgi?article=8&context=etd"),
            # Article 9 is FINISHED — every row downloaded, so it has
            # nothing outstanding at all. It is still part of the plan, and
            # it is the only row here that can tell the plan's record set
            # apart from the outstanding rows: article 7 has a supplemental
            # still failing, so it appears in both either way.
            _row(4, "9", "primary", "public", "current", "Downloaded",
                 base + "/cgi/viewcontent.cgi?article=9&context=etd",
                 base + "/cgi/editor.cgi?article=9&context=etd",
                 orig="c.pdf", saved="etd_9.pdf", size=11),
        ]
        p = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)

        got, skipped, known, = None, None, None
        got, skipped, _absent, known = M.read_failed_rows(p)
        check("retry-read . only the failed rows come back", len(got), 2)

        # `known` is every record the PLAN covers, per structure — which
        # is NOT the records in `got`. Article 7's primary downloaded, so
        # it is not outstanding, but it is still part of the plan. Built
        # from the outstanding rows instead, this would omit every record
        # already finished, and reconciliation would then report each of
        # them as "added since the plan was taken". That is what produced
        # "256 record(s) have been added" on 2026-09-11.
        # A lone session report does NOT establish the job's scope. It
        # records what that session had outstanding, and "in the plan but
        # not in this folder" is indistinguishable from "added since the
        # plan" — which is what produced two wrong drift lines on
        # 2026-09-11. Only a plan workbook enumerates the job.
        check("retry-read . a session report does not claim to be the plan",
              known, {})

        # With a plan workbook present, the scope IS known, and it spans
        # records with nothing outstanding: article 9 downloaded cleanly
        # and is still part of the job.
        p3 = _fake_report(d, "DC_FileDownload_Plan_etd_20260101_000000.xlsx",
                          rows)
        _got3, _s3, _a3, known3 = M.read_failed_rows(d)
        check("retry-read . a plan workbook does establish it",
              known3, {"etd": {"7", "8", "9"}})
        check_true("retry-read . including a record with nothing outstanding",
                   "9" in known3.get("etd", set())
                   and not any(r["article"] == "9" for r in _got3),
                   "%s / %s" % (known3, [(r["article"], r["kind"])
                                         for r in _got3]))
        os.remove(p3)
        check("retry-read . nothing was skipped", skipped, [])
        check("retry-read . the downloaded row is left alone",
              [g["article"] for g in got], ["7", "8"])
        check("retry-read . context is recovered from the admin URL",
              sorted({g["context"] for g in got}), ["etd"])
        check("retry-read . the kind survives", got[0]["kind"], "supp1")
        check("retry-read . the visibility survives",
              got[0]["visibility"], "hidden")
        check_true("retry-read . the exact file URL survives",
                   got[0]["url"].endswith("type=additional&filename=1"),
                   got[0]["url"])
        check_true("retry-read . the old status is kept for reference",
                   "429" in got[0]["was"], got[0]["was"])

        # a whole run folder, not just one report
        _fake_report(d, "DC_FileDownload_Report_series_S.xlsx", [
            _row(1, "9", "primary", "public", "current",
                 "Failed: network error: reset",
                 base + "/cgi/viewcontent.cgi?article=9&context=series",
                 base + "/cgi/editor.cgi?article=9&context=series")])
        got, skipped, _absent, _known = M.read_failed_rows(d)
        check("retry-read . a folder picks up every report in it", len(got), 3)
        check("retry-read . across both structures",
              sorted({g["context"] for g in got}), ["etd", "series"])

        # a row that cannot be addressed is skipped, never guessed at
        p2 = _fake_report(d, "DC_FileDownload_Report_bad_S.xlsx", [
            _row(1, "10", "primary", "public", "current",
                 "Failed: something", "", "")])
        got, skipped, _absent, _known = M.read_failed_rows(p2)
        check("retry-read . an unaddressable row yields no request", got, [])
        check("retry-read . and is reported as skipped", len(skipped), 1)
        check_true("retry-read . naming what was missing",
                   "file URL" in skipped[0], "{!r}".format(skipped[0]))

        # not a report at all
        junk = os.path.join(d, "notareport.xlsx")
        M.write_master_workbook(junk, [(1, "x", "y", 0, 0, 0, 0, "s",
                                        "12:00:00", "")], _brand())
        try:
            M.read_failed_rows(junk)
            check_true("retry-read . a non-report is refused", False,
                       "no exception")
        except ValueError as e:
            check_true("retry-read . a non-report names the missing columns",
                       "missing column" in str(e), str(e))
        try:
            M.read_failed_rows(os.path.join(d, "nope.xlsx"))
            check_true("retry-read . a missing path is refused", False,
                       "no exception")
        except ValueError as e:
            check_true("retry-read . a missing path says so",
                       "No such report" in str(e), str(e))


def test_retry_native_fallback_is_offered_only_where_it_is_valid():
    """A wrong fallback here would file the wrong bytes under this row.

    The supplemental and revision URLs each name one exact file. Falling
    back from them to the record's current native would fetch a different
    file and save it under this row's name and label - a silent mix-up far
    worse than a reported failure.
    """
    B = "https://dc.example.edu"

    def nf(kind, version):
        return M.retry_native_url(B, {"kind": kind, "version": version,
                                      "context": "etd", "article": "7"})

    got = nf("primary", "current")
    check_true("retry-native . a current primary gets the native fallback",
               got.endswith("/type/native/viewcontent?preview_mode=1"), got)
    check_true("retry-native . and it is the path form DC actually serves",
               "/context/etd/article/7/" in got, got)
    check("retry-native . a row already native gets none",
          nf("native", "current"), "")
    check("retry-native . a supplemental gets none",
          nf("supp1", "current"), "")
    check("retry-native . a second supplemental gets none",
          nf("supp2", "current"), "")
    check("retry-native . a dated revision gets none",
          nf("primary", "original"), "")
    check("retry-native . nor an all-versions revision",
          nf("primary", "2026-01-01"), "")


def test_the_retry_endpoint_is_wired():
    """Source-level: the worker is unreachable without Chrome and a report."""
    src = Path(M.__file__).read_text(encoding="utf-8")
    check_true("retry-wire . there is an endpoint", '"/api/retry"' in src)
    # By what the endpoint DOES, not by how the spawn is spelled. This
    # asserted "target=retry_worker" and broke when starting a run became
    # an atomic claim rather than a bare threading.Thread — a test of the
    # spelling wearing the name of a test of the wiring.
    _blk = src[src.index('if self.path == "/api/retry":'):]
    _blk = _blk[:_blk.index('if self.path == "/api/start":')]
    check_true("retry-wire . it starts the retry worker",
               "retry_worker" in _blk, _blk[-300:])
    check_true("retry-wire . and claims the run slot before doing so",
               "_claim_and_start(" in _blk, _blk[-300:])
    check_true("retry-wire . it refuses a path that does not exist",
               "No such report or folder: " in src)
    check_true("retry-wire . the page has the control",
               'id="retryrpt"' in src and 'id="retrygo"' in src)
    check_true("retry-wire . and the button posts to it",
               "/api/retry" in src.split("<script>")[1])
    # just retry_worker's own body, so a match elsewhere in the file
    # cannot stand in for one here
    body = src[src.index("def retry_worker("):]
    body = body[:body.index("# UI — branded page")]
    check_true("retry-wire . the retry run honors the report folder",
               "report_destinations(" in body)
    check_true("retry-wire . and the cooldown budget",
               "set_cooldown_budget(" in body)
    check_true("retry-wire . a mid-run expiry still writes a partial",
               "_PARTIAL.xlsx" in body)
    check_true("retry-wire . it does NOT need a hierarchy",
               "MODEL" not in body and "resolve_targets" not in body)


def test_the_run_sets_its_own_pace():
    """Pacing is a per-run setting, not a constant to be hand-edited.

    The protocol for the deciding experiment used to say "edit REQUEST_DELAY
    from 0.3 to 5.0". That would have produced a report traceable to no
    commit, which is the one thing a test report must not be - and it asked
    the operator to edit a tree he has asked not to touch.
    """
    M.reset_rate_state()
    check("pace . a fresh run uses the default",
          M.request_delay(), M.REQUEST_DELAY)

    # Derived from the default so it can never coincide with it. A literal
    # 5.0 here silently stopped testing anything the moment the default
    # became 5.0 - which happened the same day this test was written.
    other = M.REQUEST_DELAY + 3.0
    M.set_request_delay(other)
    check("pace . the run's setting takes effect", M.request_delay(), other)
    check_true("pace . and it differs from the default, so this is a real "
               "change", other != M.REQUEST_DELAY)

    # The default is measured, not chosen: 0.3s met a lockout after 26-36
    # records on four attempts against one structure; 5.0s fetched all 40
    # rows with no pushback, 78 downloads in 687s. Pinned with its reason
    # so a revert to a burst pace has to argue with this.
    check("pace . the shipped default is the figure Run C measured",
          M.REQUEST_DELAY, 5.0)
    check_true("pace . which is slower than the rate that provoked lockouts",
               M.REQUEST_DELAY > 1.0, "{}".format(M.REQUEST_DELAY))

    # the adaptive slowdown stacks on the run's pace, not on the constant
    M._note_rate_limit()
    check("pace . pushback adds to the run's pace, not to the default",
          round(M.request_delay(), 6), round(other + M.RATE_DELAY_STEP, 6))
    check_true("pace . and the run's pace is not the default here, so that "
               "assertion means something",
               abs(other - M.REQUEST_DELAY) > 1e-9)

    # Clamping is a property of the run's BASE pace. Asserting it through
    # request_delay() was wrong: the adaptive penalty added just above is
    # part of that sum, so every clamp check read 0.4s too high. Assert the
    # base, and keep a separate check that the sum still includes both.
    def base():
        return M.rate_state()["base"]

    check("pace . zero is allowed, for the impatient",
          (M.set_request_delay(0), base())[1], 0.0)
    check("pace . a negative pace is clamped to zero",
          (M.set_request_delay(-3), base())[1], 0.0)
    check("pace . an absurd pace is clamped to the ceiling",
          (M.set_request_delay(10 ** 6), base())[1], M.REQUEST_DELAY_MAX)
    check_true("pace . the ceiling is a real limit, not a formality",
               M.REQUEST_DELAY_MAX > M.REQUEST_DELAY,
               "{} vs {}".format(M.REQUEST_DELAY_MAX, M.REQUEST_DELAY))

    M.reset_rate_state()
    check("pace . reset restores the default for the next run",
          M.request_delay(), M.REQUEST_DELAY)

    # _throttle must consult the run, not the constant
    real_sleep = M.time.sleep
    naps = []
    M.time.sleep = lambda x: naps.append(x)
    try:
        M.set_request_delay(2.5)
        M._throttle()
        check("pace . _throttle sleeps the run's gap", naps, [2.5])
        naps.clear()
        M.set_request_delay(0)
        M._throttle()
        check("pace . and does not sleep at all at zero", naps, [])
    finally:
        M.time.sleep = real_sleep
        M.reset_rate_state()

    src = Path(M.__file__).read_text(encoding="utf-8")
    check_true("pace . the page offers the control",
               'id="reqdelay"' in src)
    check_true("pace . both workers accept it",
               src.count('set_request_delay(opts.get("delay"') == 2)
    check_true("pace . the page offers the admin-page control",
               'id="pagedelay"' in src)
    check_true("pace . both workers accept the page pace",
               src.count('set_page_delay(opts.get("page_delay"') == 2)

    # Range-checking is asserted by CALLING the endpoint, not by grepping
    # the source for the message. The grep version searched for a
    # contiguous string that the source wraps across two lines, so it
    # failed on a message that was present and correct — a test of the
    # formatting wearing the name of a test of the behavior.
    import threading as _t
    session = M.load_session(Path("/nonexistent/session.json"))
    session["hub_url"] = "http://127.0.0.1:8750"
    srv = ThreadingHTTPServer(
        ("127.0.0.1", 0), M.make_handler(session, M.build_page(session)))
    port = srv.server_address[1]
    _t.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        base = {"out_dir": "/nonexistent", "mode": "all", "parents": [],
                "primary": True, "versions": "current", "cooldowns": 1}
        for field, label in (("delay", "file download"),
                             ("page_delay", "admin page")):
            for bad in (-1, M.REQUEST_DELAY_MAX + 1):
                body = dict(base); body[field] = bad
                code, out = _post(port, "/api/start", body)
                check("pace . {}={} is refused".format(field, bad), code, 400)
                check_true("pace . and the {} error names that control"
                           .format(label),
                           label in str(out.get("error", "")),
                           str(out.get("error", "")))
        # A value inside the range must NOT be refused for being out of range.
        body = dict(base); body["delay"] = 5.0; body["page_delay"] = 1.0
        code, out = _post(port, "/api/start", body)
        check_true("pace . a legal pair passes the range check",
                   "must be between" not in str(out.get("error", "")),
                   str(out.get("error", "")))
    finally:
        srv.shutdown()


def test_the_paged_fallback_still_works_if_showall_is_ignored():
    """x_showall=1 is unverified against the live server, so the fallback
    is the whole safety argument for shipping it — and an argument nobody
    has checked is not a safety argument.

    If DC ignores the parameter, the page comes back with its 25 rows and
    its ">" link, and crawl_listing must page on exactly as before. This
    drives that path, and checks the page count is reported either way so
    the real answer lands in the next run's log instead of being assumed.
    """
    real_fetch, real_throttle = M.fetch_html, M._throttle
    real_hook = M._NOTE_HOOK[0]
    said = []
    M._throttle = lambda *a, **k: None
    M.set_note_hook(said.append)
    try:
        def page(n, last=False, count=25):
            rows = "".join(
                '<tr><td>{0}</td><td><a href="editor.cgi?article={0}">'
                'Title {0}</a></td><td>Posted</td><td>-</td></tr>'
                .format(1000 + n * 25 + i) for i in range(count))
            nxt = ("" if last else
                   '<a href="editor.cgi?x_start={}">&gt;</a>'
                   .format((n + 1) * 25))
            return ('<html><body><h1>EdiKit</h1><table>'
                    '<tr><th>ID</th><th>Title</th><th>State</th>'
                    '<th>Additional Files</th></tr>'
                    + rows + "</table>" + nxt + "</body></html>")

        # the server ignores x_showall: three paged responses
        served = [page(0), page(1), page(2, last=True)]
        calls = []

        def fake_fetch(url, cookie, timeout=None):
            calls.append(url)
            return served[len(calls) - 1]
        M.fetch_html = fake_fetch

        said.clear()
        rows, gallery = M.crawl_listing(BASE, "etd", False, "published",
                                        "cookie")
        check("fallback . every page was walked", len(calls), 3)
        check("fallback . and every row collected", len(rows), 75)
        check_true("fallback . rows are not duplicated",
                   len({r["article"] for r in rows}) == 75)
        check_true("fallback . the page count is reported",
                   any("over 3 page(s)" in m for m in said), "{!r}".format(said))
        # It must NOT claim a request count. The first listing page of
        # the first state comes from the driver navigation, not from
        # urllib, so "3 requests spent" - which this line used to say -
        # was wrong by one on every run.
        check_true("fallback . it does not claim a request count",
                   not any("request" in m for m in said),
                   "{!r}".format(said))

        # The server honors it: one page holding more than a page's worth.
        served[:] = [page(0, last=True, count=60)]
        calls.clear(); said.clear()
        rows, gallery = M.crawl_listing(BASE, "etd", False, "published",
                                        "cookie")
        check("showall . one request only", len(calls), 1)
        check("showall . all rows in it", len(rows), 60)
        check_true("showall . and the log says so plainly",
                   any("x_showall worked" in m for m in said),
                   "{!r}".format(said))

        # Everything below is the defect Jeff saw within minutes of v1.13
        # shipping: the line fired where it meant nothing, and claimed a
        # success that had not happened.
        #
        # A gallery listing carries no x_showall in its URL at all, so it
        # can never be credited with one. Driven with the suite's real
        # gallery fixture: a hand-built one shaped like an EdiKit table
        # gets flavour-flipped to EdiKit by crawl_listing, quite correctly,
        # and then the assertion tests nothing. The flip means `gallery`
        # inside the loop reflects what the page actually is, which is the
        # flag the note must consult.
        served[:] = [GALLERY_LISTING_LASTPAGE]
        calls.clear(); said.clear()
        rows, gallery = M.crawl_listing(BASE, "g", True, "published",
                                        "cookie")
        check_true("note . the fixture really is read as a gallery", gallery,
                   "flipped to EdiKit; the fixture is wrong, not the code")
        check_true("note . a gallery is never credited with x_showall",
                   not any("x_showall" in m for m in said),
                   "{!r}".format(said))

        # An empty state comes back in one page with no rows. Five states
        # per structure meant four of these per run, each announcing a
        # success that never happened.
        served[:] = ['<html><body><h1>EdiKit</h1><table>'
                     '<tr><th>ID</th><th>Title</th><th>State</th>'
                     '<th>Additional Files</th></tr></table></body></html>']
        calls.clear(); said.clear()
        rows, gallery = M.crawl_listing(BASE, "etd", False, "pending",
                                        "cookie")
        check("note . an empty state yields no rows", len(rows), 0)
        check("note . and says nothing at all", said, [])

        # A structure that fits in one page anyway proves nothing about
        # x_showall either way, so it is silent too.
        served[:] = [page(0, last=True, count=12)]
        calls.clear(); said.clear()
        rows, gallery = M.crawl_listing(BASE, "etd", False, "published",
                                        "cookie")
        check("note . a structure smaller than one page is silent", said, [])
        check("note . though its rows still come back", len(rows), 12)
    finally:
        M.fetch_html = real_fetch
        M._throttle = real_throttle
        M.set_note_hook(real_hook)


def test_progress_and_eta_measure_records_not_structures():
    """Both read 'nothing is happening' on a single-structure run.

    Jeff, mid-run on 2026-09-08: "the progress bar is full blue." It was:
    the item loop passed the STRUCTURE index as the bar's position and the
    structure COUNT as its maximum, so one structure read 1 of 1 from the
    first record onward. The same mistake gave `ETA --:--` on every run of
    this module ever made - eta_text(s_idx - 1, ...) is eta_text(0, ...)
    for a single structure, and eta_text returns --:-- when done is zero.
    I had blamed the cooldown for that; the cooldown only made an
    already-broken figure look broken for longer.
    """
    P = M.run_position
    check("progress . a single structure starts at zero, not full",
          P(0, 0, 38), 0.0)
    check_true("progress . and climbs as records finish",
               0.4 < P(0, 16, 38) < 0.5, "{}".format(P(0, 16, 38)))
    check("progress . reaching one only at the last record",
          P(0, 38, 38), 1.0)
    check_true("progress . the old behavior would have read full at record 1",
               P(0, 1, 38) < 0.1, "{}".format(P(0, 1, 38)))
    check("progress . a second structure starts where the first ended",
          P(1, 0, 20), 1.0)
    check("progress . and finishes at two", P(1, 20, 20), 2.0)
    check("progress . an empty structure does not divide by zero",
          P(3, 0, 0), 3.0)
    check_true("progress . a miscount cannot push the bar past its own max",
               P(0, 99, 38) == 1.0, "{}".format(P(0, 99, 38)))

    # the root cause of ETA --:--, pinned so it cannot come back
    check("eta . nothing done yields no estimate",
          M.eta_text(0, 38, M.time.time() - 60), "ETA --:--")
    check_true("eta . something done yields a real figure",
               M.eta_text(19, 38, M.time.time() - 190) != "ETA --:--",
               M.eta_text(19, 38, M.time.time() - 190))
    check("eta . and it is roughly right — half done in 190s, ~190s to go",
          M.eta_text(19, 38, M.time.time() - 190), "ETA 03:10")

    # and the call site must pass records, not structures
    src = Path(M.__file__).read_text(encoding="utf-8")
    body = src[src.index("def download_worker("):src.index("# Worker 3:")]
    loop = body[body.index("for i_idx, (item, state) in enumerate"):]
    check_true("progress . the item loop positions the bar by record",
               "run_position(s_idx - 1, i_idx, len(items))" in loop,
               loop[:300])
    check_true("eta . the item loop estimates from the structure's own pace",
               "eta_text(i_idx, len(items), t_struct)" in loop, loop[:300])
    check_true("eta . and no longer from the structure count",
               "eta_text(s_idx - 1, total, t0)" not in loop, loop[:300])


# ---------------------------------------------------------------------------
# Run B, first attempt: 'etd' listed 3,106 records under `published` with
# x_showall, then answered HTTP 400 for another record state - and the whole
# structure was skipped, discarding all 3,106. Two defects, both here.
# ---------------------------------------------------------------------------
def test_a_refused_showall_falls_back_to_paging():
    real_fetch, real_throttle = M.fetch_html, M._throttle
    real_hook = M._NOTE_HOOK[0]
    said, seen = [], []
    M._throttle = lambda *a, **k: None
    M.set_note_hook(said.append)
    try:
        def refuse_showall(url, cookie, timeout=None):
            seen.append(url)
            if "x_showall=1" in url:
                raise _http(400, "Bad Request")
            return EDIKIT_LISTING_LASTPAGE
        M.fetch_html = refuse_showall

        rows, gallery = M.crawl_listing(BASE, "etd", False, "published",
                                        "cookie")
        check_true("showall-400 . the listing still comes back",
                   len(rows) > 0, "{} row(s)".format(len(rows)))
        check("showall-400 . it asked twice: with, then without",
              len(seen), 2)
        check_true("showall-400 . the first attempt carried x_showall",
                   "x_showall=1" in seen[0], seen[0])
        check_true("showall-400 . the second did not",
                   "x_showall=1" not in seen[1], seen[1])
        check_true("showall-400 . and every other parameter survived",
                   "status_filter=all" in seen[1]
                   and "showstate=published" in seen[1], seen[1])
        check_true("showall-400 . the fallback is said out loud",
                   any("refused x_showall" in m for m in said),
                   "{!r}".format(said))

        # a 400 on a URL that never had x_showall is a real error, not
        # something to paper over
        seen.clear()

        def always_400(url, cookie, timeout=None):
            seen.append(url)
            raise _http(400, "Bad Request")
        M.fetch_html = always_400
        try:
            M.crawl_listing(BASE, "g", True, "published", "cookie")
            check_true("showall-400 . a gallery 400 still raises", False,
                       "no exception")
        except urllib.error.HTTPError as e:
            check("showall-400 . a gallery 400 is raised, not retried",
                  (e.code, len(seen)), (400, 1))
        except Exception as e:
            check_true("showall-400 . a gallery 400 is raised as itself",
                       False, "got {}".format(e.__class__.__name__))

        # and a 403 with x_showall present is not a showall problem
        seen.clear()

        def always_403(url, cookie, timeout=None):
            seen.append(url)
            raise _http(403, "Forbidden")
        M.fetch_html = always_403
        try:
            M.crawl_listing(BASE, "etd", False, "published", "cookie")
            check_true("showall-400 . a 403 raises", False, "no exception")
        except urllib.error.HTTPError as e:
            check("showall-400 . only a 400 triggers the fallback",
                  (e.code, len(seen)), (403, 1))
        except Exception as e:
            check_true("showall-400 . a 403 is raised as itself", False,
                       "got {}".format(e.__class__.__name__))
    finally:
        M.fetch_html = real_fetch
        M._throttle = real_throttle
        M.set_note_hook(real_hook)


def test_one_bad_state_does_not_discard_the_others():
    """Source-level: the state loop lives inside download_worker.

    The defect was structural - a try/except around the whole loop rather
    than around each state - so what has to be checked is the shape. It
    cost 3,106 already-listed records the first time Run B was attempted.
    """
    src = Path(M.__file__).read_text(encoding="utf-8")
    body = src[src.index("def download_worker("):src.index("# Worker 3:")]
    loop = body[body.index("# -- crawl each requested state"):
                body.index("items.sort(")]

    check_true("states . each state's crawl has its own handler",
               loop.count("except Exception as e:") >= 2, loop[:200])
    check_true("states . a failing state is recorded and the loop continues",
               "state_errs.append(" in loop and "continue" in loop,
               loop[:200])
    check_true("states . the operator is told which state failed",
               "continuing with the other record states" in loop,
               loop[:200])
    check_true("states . a structure is only skipped when it yielded nothing",
               "state_errs and not items" in loop, loop[:200])
    check_true("states . login expiry still aborts rather than being "
               "swallowed as a state error",
               "except (LoginRequired, StopRequested):" in loop, loop[:200])
    # Indexing split()[1] blindly raised IndexError when the guard above
    # was absent, so a missing fix crashed the suite instead of reporting.
    # Same lesson as the list-as-detail bug earlier today: a check must
    # still be able to describe its own failure.
    marker = "except (LoginRequired, StopRequested):"
    after = loop.split(marker)[1][:40] if marker in loop else ""
    check_true("states . and a stop is never turned into a listing error",
               "raise" in after, "nothing re-raised after the guard")


# ---------------------------------------------------------------------------
# Links are resolved against the page they came from.
#
# Run B, 2026-09-08: every revision row failed, and the report showed why -
#   current:  https://host/cgi/viewcontent.cgi?article=1000&context=etd&...
#   revision: https://host/viewcontent.cgi?type=pdf&article=1000&...
# The /cgi/ is missing. The revisions table serves DOCUMENT-relative hrefs
# ("viewcontent.cgi?article=..." with no leading slash) and they were joined
# onto the site root. So all-versions mode has never produced a working URL.
#
# It survived 380 tests because every fixture here used a root-relative or
# absolute href, which joins correctly against either base. That is open
# item 7 - synthetic fixtures pin assumptions, real HTML pins behavior - and
# this is the second defect to reach production through that same gap, after
# the &type=native URL shape that 95 green tests certified.
#
# The three href forms below are the ones the captured admin pages actually
# contain.
# ---------------------------------------------------------------------------
def test_revision_links_resolve_against_their_page():
    opts = {"primary": True, "supp": False, "native": False,
            "public": True, "hidden": True, "versions": "all",
            "published": True, "unpublished": False}
    item = {"article": "1000", "title": "T", "af": 0}
    warns = []
    real_fetch, real_throttle = M.fetch_html, M._throttle
    M._throttle = lambda *a, **k: None
    try:
        def plan_with(href):
            """The URL of the job that came from the revisions table.

            jobs[0] is the record's CURRENT primary, built from
            primary_file_url and nothing to do with the revisions page -
            asserting on it tested the wrong thing entirely. The revision
            jobs are the ones carrying a version_date.
            """
            # v1.34.2: the link under test sits in an EARLIER revision.
            # As the only revision it would be the current one, and a bare
            # current link IS the default stream the primary job already
            # fetches — so identity dedupe rightly folds it away, and this
            # test of link resolution would be reading the wrong job.
            html = _pickvers_html([
                ['<a href="viewcontent.cgi?type=pdf&article=1000'
                 '&unstamped=yes&date=1999999999&context=etd">PDF</a>'],
                ['<a href="{}">PDF</a>'.format(href)]])
            M.fetch_html = lambda url, cookie, timeout=0: html
            jobs = M.plan_item_jobs(BASE, "etd", item, "published", False,
                                    dict(opts), "cookie", warns.append)
            dated = [j for j in jobs if j["version"] == "original"]
            check_true("links . a revision job was planned at all",
                       len(dated) == 1,
                       "{} of {} job(s) carried a date".format(
                           len(dated), len(jobs)))
            return dated

        # (a) document-relative — what the real revisions table serves
        jobs = plan_with("viewcontent.cgi?type=pdf&article=1000"
                         "&unstamped=yes&date=1374154673&context=etd")
        check("links . a document-relative revision href keeps /cgi/",
              jobs[0]["url"],
              BASE + "/cgi/viewcontent.cgi?type=pdf&article=1000"
                     "&unstamped=yes&date=1374154673&context=etd")
        check_true("links . and does not land at the site root",
                   "/cgi/viewcontent.cgi" in jobs[0]["url"]
                   and not jobs[0]["url"].startswith(BASE + "/viewcontent"),
                   jobs[0]["url"])

        # (b) root-relative — also on the real pages, must be unchanged
        jobs = plan_with("/cgi/viewcontent.cgi?article=1000&context=etd")
        check("links . a root-relative href is unchanged",
              jobs[0]["url"], BASE + "/cgi/viewcontent.cgi?article=1000"
                                     "&context=etd")

        # (c) absolute — what the supplemental page serves
        jobs = plan_with(BASE + "/cgi/viewcontent.cgi?article=1000"
                                "&context=etd&type=pdf")
        check("links . an absolute href is unchanged",
              jobs[0]["url"], BASE + "/cgi/viewcontent.cgi?article=1000"
                                     "&context=etd&type=pdf")

        # (d) the path form, from a gallery revision row
        jobs = plan_with("/context/etd/article/1000/type/native/viewcontent")
        check("links . the path form survives too",
              jobs[0]["url"],
              BASE + "/context/etd/article/1000/type/native/viewcontent")

        # (e) the SUPPLEMENTAL page, same join, same hazard. Its hrefs are
        #     absolute on the pages captured so far, which is exactly why
        #     joining them onto the site root worked and why fixing it was
        #     untested until this case existed: an absolute href resolves
        #     identically against either base, so the fixture could not
        #     tell a correct join from a broken one.
        sopts = {"primary": False, "supp": True, "native": False,
                 "public": True, "hidden": True, "versions": "current",
                 "published": True, "unpublished": False}

        def supp_with(href):
            html = ("<html><table id='uploaded-files'>"
                    "<tr><td><a href='{}'>extra.pdf</a></td>"
                    "<td><input type='checkbox' checked></td></tr>"
                    "</table></html>".format(href))
            M.fetch_html = lambda url, cookie, timeout=0: html
            return M.plan_item_jobs(BASE, "etd", item, "published", False,
                                    dict(sopts), "cookie", warns.append)

        jobs = supp_with("viewcontent.cgi?filename=0&article=1000"
                         "&context=etd&type=additional")
        check("links . a document-relative supplemental href keeps /cgi/",
              jobs[0]["url"],
              BASE + "/cgi/viewcontent.cgi?filename=0&article=1000"
                     "&context=etd&type=additional")
        jobs = supp_with(BASE + "/cgi/viewcontent.cgi?filename=0"
                                "&article=1000&context=etd&type=additional")
        check("links . an absolute supplemental href is unchanged",
              jobs[0]["url"],
              BASE + "/cgi/viewcontent.cgi?filename=0&article=1000"
                     "&context=etd&type=additional")
    finally:
        M.fetch_html = real_fetch
        M._throttle = real_throttle


def test_a_regeneration_link_is_never_harvested():
    """`force=yes` regenerates the record's PDF on the server.

    It is a WRITE wearing a viewcontent URL, so the "viewcontent in href"
    test both parsers use cannot tell it apart from a download. It sits on
    the submission details page, which nothing parses yet - so this is a
    guard, not a fix, and it is worth having before open item 6 parses that
    page rather than after.
    """
    html = _pickvers_html([[
        '<a href="viewcontent.cgi?article=4395&context=etd'
        '&preview_mode=1&force=yes">Regenerate</a>',
        '<a href="viewcontent.cgi?article=4395&context=etd'
        '&preview_mode=1&z=1777415248">PDF</a>']])
    revs = M.parse_pickvers(html)
    hrefs = [f["href"] for f in revs[0]["files"]]
    check("force . the regeneration link is not harvested", len(hrefs), 1)
    check_true("force . and the real content link still is",
               "z=1777415248" in hrefs[0], hrefs[0])
    check_true("force . nothing harvested mentions force=yes",
               not any("force=yes" in h for h in hrefs), "{!r}".format(hrefs))

    check_true("force . the predicate refuses it in any case",
               not M.harvestable("VIEWCONTENT.CGI?FORCE=YES"))
    check_true("force . and allows an ordinary content link",
               M.harvestable("/cgi/viewcontent.cgi?article=1&context=etd"))

    # a row holding nothing but the regeneration link contributes no files
    html = _pickvers_html([[
        '<a href="viewcontent.cgi?article=1&force=yes">Regenerate</a>']])
    revs = M.parse_pickvers(html)
    check("force . a row with only that link yields no files",
          revs[0]["files"] if revs else "no revisions", [])


def test_the_family_is_not_re_guessed_once_known():
    """An unreadable record state is not evidence of the other family.

    'etd' listed 3,106 records as an EdiKit structure, then a later state
    returned a page neither parser recognised - and the flip fired, asking
    editor_gallery.cgi for an ETD collection. DC answered HTTP 400, and
    that nonsensical request's failure then masked whatever the state had
    actually returned.
    """
    real_fetch, real_throttle = M.fetch_html, M._throttle
    M._throttle = lambda *a, **k: None
    seen = []
    try:
        junk = "<html><body><p>Nothing here.</p></body></html>"
        M.fetch_html = lambda url, cookie, timeout=0: (seen.append(url), junk)[1]

        # allow_flip=False: no second URL is tried, and the error names the
        # state rather than the structure
        seen.clear()
        try:
            M.crawl_listing(BASE, "etd", False, "pending", "cookie",
                            allow_flip=False)
            check_true("flip . an unreadable state raises", False, "no exc")
        except ValueError as e:
            check_true("flip . the error is about this record state",
                       "record state" in str(e), str(e))
            check_true("flip . and says the family is already known",
                       "already known" in str(e), str(e))
        except Exception as e:
            check_true("flip . an unreadable state raises ValueError", False,
                       "got {}".format(e.__class__.__name__))
        check("flip . and no second request was made", len(seen), 1)
        check_true("flip . the one request was the EdiKit listing",
                   "editor.cgi" in seen[0] and "editor_gallery" not in seen[0],
                   seen[0])

        # allow_flip=True (the first state of a structure) still probes
        seen.clear()
        try:
            M.crawl_listing(BASE, "etd", False, "published", "cookie",
                            allow_flip=True)
        except ValueError:
            pass
        check("flip . the first state may still try the other family",
              len(seen), 2)
        check_true("flip . which is where editor_gallery.cgi comes from",
                   "editor_gallery" in seen[1], seen[1])

        # and the worker only allows it until a state has parsed
        src = Path(M.__file__).read_text(encoding="utf-8")
        body = src[src.index("def download_worker("):
                   src.index("# UI — branded")]
        check_true("flip . the worker stops allowing it once a state parsed",
                   "allow_flip=not flavour_known" in body
                   and "flavour_known = True" in body, "not wired")
    finally:
        M.fetch_html = real_fetch
        M._throttle = real_throttle


def test_a_long_eta_is_readable():
    """"ETA 1940:39" was correct and unreadable.

    Driven with a frozen clock. eta_text calls time.time() itself, so
    sampling `now` separately raced it by a minute - a test that fails
    depending on which second it runs in is worse than no test.
    """
    real = M.time.time
    M.time.time = lambda: 1_000_000.0
    now = 1_000_000.0
    try:
        check("eta . under an hour stays mm:ss",
              M.eta_text(1, 4, now - 600), "ETA 30:00")
        check("eta . past an hour switches to hours",
              M.eta_text(1, 3, now - 3600), "ETA 2h 00m")
        # Run B's own numbers: 2 of 3,383 records in 66 seconds.
        # 66/2 x 3,381 = 111,573s = 30h 59m. The run displayed this as
        # "ETA 1940:39" - the same figure, unreadable.
        check("eta . a very long run says hours, not 1940 minutes",
              M.eta_text(2, 3383, now - 66), "ETA 30h 59m")
        check_true("eta . and never shows a minute count above 59",
                   "h" in M.eta_text(2, 3383, now - 66),
                   M.eta_text(2, 3383, now - 66))
        check("eta . exactly one hour reads as hours",
              M.eta_text(1, 2, now - 3600), "ETA 1h 00m")
        check("eta . just under reads as minutes",
              M.eta_text(1, 2, now - 3599), "ETA 59:59")
        check("eta . nothing done is still unknown",
              M.eta_text(0, 10, now - 10), "ETA --:--")
    finally:
        M.time.time = real


def test_an_empty_record_state_is_not_a_failure():
    """Four warnings that read like breakage, for the ordinary case.

    A series with records in one state produced, on 2026-09-08:
      listing for 'Queued for update' records failed - no recognizable
      item listing for this record state
    ...four times. True of the markup and wrong about the cause: those
    states simply held no records. window.pageData is on every admin page
    - verified across every captured sample - so it distinguishes a page
    with nothing in it from a page we could not read.
    """
    real_fetch, real_throttle = M.fetch_html, M._throttle
    real_hook = M._NOTE_HOOK[0]
    said = []
    M._throttle = lambda *a, **k: None
    M.set_note_hook(said.append)
    try:
        empty_admin = ("<html><head><script>window.pageData = "
                       "{\"page\":{\"name\":\"editor:index\"}};</script>"
                       "</head><body><p>No submissions.</p></body></html>")
        M.fetch_html = lambda url, cookie, timeout=0: empty_admin
        rows, gallery = M.crawl_listing(BASE, "s", False, "pending",
                                        "cookie", allow_flip=False)
        check("empty . an empty admin listing yields no rows", rows, [])
        check_true("empty . and does not raise", True)
        check_true("empty . it says the state is empty, not unreadable",
                   any("no records in the 'Not yet posted' state" in m
                       for m in said), "{!r}".format(said))
        check_true("empty . and never calls it unrecognizable",
                   not any("recognizable" in m for m in said),
                   "{!r}".format(said))

        # a page that is NOT an admin page is still an error
        said.clear()
        M.fetch_html = lambda url, cookie, timeout=0: (
            "<html><body><h1>Gateway Timeout</h1></body></html>")
        try:
            M.crawl_listing(BASE, "s", False, "pending", "cookie",
                            allow_flip=False)
            check_true("empty . a non-admin page still raises", False,
                       "no exception")
        except ValueError as e:
            check_true("empty . a non-admin page is still unreadable",
                       "recognizable" in str(e), str(e))
        except Exception as e:
            check_true("empty . a non-admin page raises ValueError", False,
                       "got {}".format(e.__class__.__name__))

        check_true("empty . the predicate finds pageData",
                   M.looks_like_admin_page(empty_admin))
        check_true("empty . and rejects a page without it",
                   not M.looks_like_admin_page("<html><body>x</body></html>"))
        check_true("empty . an empty string is not an admin page",
                   not M.looks_like_admin_page(""))
    finally:
        M.fetch_html = real_fetch
        M._throttle = real_throttle
        M.set_note_hook(real_hook)


# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# 26 · Access state (v1.20).
#
# The listing's Type column says whether a record may leave the institution.
# Until v1.20 nothing read it, so an access-restricted ETD downloaded into a
# batch the report called `public`.
# ---------------------------------------------------------------------------
def test_access_state_is_read_and_never_guessed():
    check("access . open access", M.access_state("Thesis (open access)"),
          M.ACCESS_OPEN)
    check("access . restricted",
          M.access_state("Thesis (restricted to Some University)"),
          M.ACCESS_RESTRICTED)
    check("access . embargo wording counts as restricted",
          M.access_state("Dissertation (embargoed)"), M.ACCESS_RESTRICTED)
    # The two that matter: neither may come back as `open`.
    check("access . an empty cell is unknown, not open",
          M.access_state(""), M.ACCESS_UNKNOWN)
    check("access . an unrecognized cell is unknown, not open",
          M.access_state("Dissertation/Thesis"), M.ACCESS_UNKNOWN)
    check("access . None is unknown, not open",
          M.access_state(None), M.ACCESS_UNKNOWN)
    # Restricted wins a cell that somehow says both.
    check("access . restricted beats open when a cell says both",
          M.access_state("Thesis (open access, restricted until 2030)"),
          M.ACCESS_RESTRICTED)

    # It must classify on the RELEASE WORDING, never on an institution name.
    # A house-standard name in shared source is what --generic refuses.
    #
    # THE MARKERS COME FROM THE BUILD, not from a copy written here. Until
    # 2026-09-21 this line spelled them out — which put the institution's
    # name in shared source and made the --generic build fail on the very
    # test that exists to protect it. Two weeks of a red release gate, and
    # the gate was right. Importing the build's own reader also means the
    # two cannot drift: a marker the build enforces is a marker this
    # checks, by construction rather than by memory.
    import inspect
    src = inspect.getsource(M)
    markers = _identity_markers()
    if markers:
        for marker in markers:
            check_true("access . no identity marker {!r} in the module"
                       .format(marker), marker not in src)
        check_true("access . and it had a real list to check against "
                   "({} marker(s) from the profile)".format(len(markers)),
                   len(markers) >= 3, repr(markers))
    else:
        # A generic copy has no named profile, so there is nothing to
        # search FOR. Say that. Printing nothing here would read exactly
        # like a clean scan, and the build refuses for the same reason.
        check_true("access . NOTHING WAS CHECKED — no institution profile, "
                   "so no marker to search for", True)
    check("access . the same wording classifies for any institution",
          {M.access_state("Thesis (restricted to " + who + ")")
           for who in ("A University", "B College", "C Institute")},
          {M.ACCESS_RESTRICTED})


def test_the_identity_gate_and_this_suite_read_one_list():
    """Open item 8, closed 2026-09-21 — and the shape of the fix checked.

    The gate refused every --generic build for two weeks because this
    suite had to contain the institution's name in order to search the
    module for it. The names moved to `profiles/`; the build derives the
    rest from the settings already there; both read one function.

    Asserted STRUCTURALLY, against the profile on disk, because writing
    the expected markers out here would put them back in shared source
    and recreate the defect exactly.
    """
    import json
    here = os.path.dirname(os.path.abspath(__file__))
    folder = os.path.join(os.path.dirname(here), "profiles")
    named = []
    if os.path.isdir(folder):
        named = [os.path.join(folder, f) for f in sorted(os.listdir(folder))
                 if f.endswith(".json")]
    markers = _identity_markers()

    # A profile on disk and no markers means the reader is broken or was
    # never called — which would make the identity test above pass while
    # searching for nothing at all.
    check("identity . markers exist exactly when a named profile does",
          bool(markers), bool(named))
    if not named:
        return

    for path in named:
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
        if not isinstance(data, dict):
            continue
        tag = os.path.basename(path)
        for declared in data.get("identity_markers") or ():
            check_true("identity . {} declares {!r} and the gate has it"
                       .format(tag, declared), declared in markers,
                       "declared in the profile and not enforced")
        for key in ("crossref_prefix", "prefix"):
            val = str(data.get(key) or "").strip()
            if val:
                check_true("identity . {} the DOI prefix is derived, not "
                           "copied".format(tag), val in markers, val)
        base = str(data.get("base_url") or "").strip()
        if base:
            from urllib.parse import urlparse
            host = urlparse(base).netloc
            check_true("identity . {} the instance host is derived"
                       .format(tag), host in markers, host)
        colors = ((data.get("branding") or {}).get("colors") or {})
        for val in colors.values():
            val = str(val or "").strip()
            if val:
                check_true("identity . {} the brand color {} is derived, "
                           "with and without the hash".format(tag, val),
                           val in markers and val.lstrip("#") in markers)

    # And the module itself is clean against the whole list, which is the
    # claim the release gate makes.
    src = open(os.path.join(here, "dc_file_downloader.py"),
               encoding="utf-8").read()
    dirty = sorted(m for m in markers if m in src)
    check("identity . the module names none of them", dirty, [])


def test_the_listing_carries_the_release_option():
    # Header set copied from a real EdiKit listing (11 columns, Type 7th).
    html = """<table><tr>
      <th>ID</th><th>Author</th><th>Title</th><th>Last Event</th>
      <th>Date of Last Event</th><th>Waiting for Administrator</th>
      <th>Type</th><th>Locked by Administrator</th><th>Administrator</th>
      <th>Submitted</th><th>Additional Files</th></tr>
      <tr><td>4484</td><td>A</td><td>T1</td><td>Published to web</td>
        <td>d</td><td></td><td>Thesis (open access)</td><td></td>
        <td>Grad</td><td>2026-07-15</td><td>-</td></tr>
      <tr><td>4483</td><td>B</td><td>T2</td><td>Withdrawn</td>
        <td>d</td><td></td><td>Thesis (restricted to Some University)</td>
        <td></td><td>Grad</td><td>2026-07-13</td><td>2</td></tr>
      <tr><td>4482</td><td>C</td><td>T3</td><td>Published to web</td>
        <td>d</td><td></td><td></td><td></td><td>Grad</td>
        <td>2026-07-08</td><td>1</td></tr></table>"""
    res = M.parse_edikit_listing(html)
    check_true("listing-access . the listing still parses", res is not None)
    rows = res["rows"]
    check("listing-access . every row is read", len(rows), 3)
    check("listing-access . access per row",
          [r["access"] for r in rows],
          [M.ACCESS_OPEN, M.ACCESS_RESTRICTED, M.ACCESS_UNKNOWN])
    check("listing-access . the verbatim wording is kept for the report",
          rows[1]["release"], "Thesis (restricted to Some University)")
    # The fields that were there before must be untouched.
    check("listing-access . af still read", [r["af"] for r in rows],
          [0, 2, 1])
    check("listing-access . titles still read",
          [r["title"] for r in rows], ["T1", "T2", "T3"])

    # A gallery listing has no Type column at all — that must be `unknown`,
    # not an exception and not `open`.
    gal = """<table id="submission_list"><tr><th>ID</th><th>Title</th>
      <th>Last Event</th></tr>
      <tr><td>1037</td><td>G</td><td>Published to web</td></tr></table>"""
    grows = M.parse_gallery_listing(gal)["rows"]
    check("listing-access . a gallery has no Type column, so: unknown",
          grows[0]["access"], M.ACCESS_UNKNOWN)


def test_a_type_column_is_not_always_a_release_column():
    """The defect this file's own first version shipped, same day.

    "Type" means "Document Type and Release Option" on an ETD collection and
    a plain document type on a faculty series. Reading a cell on its own
    cannot tell those apart, so the first cut called 170 newsletters
    `unknown` and advised treating 176 public files as restricted.
    """
    def item(article, release):
        return {"article": article, "release": release,
                "access": M.access_state(release)}

    # A faculty series: the column is a document type. Nothing here says
    # anything about access, so nothing may be flagged.
    series = [item("1", "Newsletter"), item("2", "Newsletter"),
              item("3", "Article")]
    check("access-na . a document-type column is recognized as such",
          M.resolve_listing_access(series), "type-only")
    check("access-na . and every row becomes n/a, not unknown",
          {i["access"] for i in series}, {M.ACCESS_NA})
    check_true("access-na . n/a is not unknown",
               M.ACCESS_NA != M.ACCESS_UNKNOWN)

    # An ETD collection: the column IS a release column, so a row that does
    # not parse is genuinely unknown and must stay flagged.
    etd = [item("1", "Thesis (open access)"),
           item("2", "Thesis (restricted to Some University)"),
           item("3", "Dissertation/Thesis"),
           item("4", "")]
    check("access-na . a release column is recognized as such",
          M.resolve_listing_access(etd), "release")
    check("access-na . open and restricted survive",
          [etd[0]["access"], etd[1]["access"]],
          [M.ACCESS_OPEN, M.ACCESS_RESTRICTED])
    check("access-na . the unreadable rows stay unknown, and stay flagged",
          [etd[2]["access"], etd[3]["access"]],
          [M.ACCESS_UNKNOWN, M.ACCESS_UNKNOWN])

    # A gallery has no Type column at all — also n/a, never unknown.
    gallery = [item("1", ""), item("2", "")]
    check("access-na . no Type column at all is n/a too",
          M.resolve_listing_access(gallery), "type-only")
    check("access-na . an empty structure does not crash",
          M.resolve_listing_access([]), "type-only")

    # One restricted row anywhere makes the whole column a release column.
    mixed = [item("1", "Newsletter"), item("2", "Article"),
             item("3", "Thesis (restricted to Some University)")]
    check("access-na . one release row settles it for the structure",
          M.resolve_listing_access(mixed), "release")
    check("access-na . and the document-type rows are then truly unknown",
          mixed[0]["access"], M.ACCESS_UNKNOWN)

    # The worker must settle this once per structure, with all records in
    # hand — not per row and not per page.
    import inspect
    body = inspect.getsource(M.download_worker)
    check("access-na . the worker resolves it exactly once",
          body.count("resolve_listing_access("), 1)
    check_true("access-na . after the records are gathered and sorted",
               body.index("items.sort(")
               < body.index("resolve_listing_access("), "wrong order")
    # Only a genuine unknown may raise the summary warning.
    check_true("access-na . only `unknown` is counted for the warning",
               "ACCESS_UNKNOWN" in body and "== ACCESS_NA" not in body)


def test_the_embargo_parser_harvests_no_links():
    html = """<html><body><table>
      <tr><th scope="row">Document Type and Release Option</th>
          <td>Thesis (restricted to Some University)</td></tr>
      <tr><th scope="row">OTHER_LICENSE</th>
          <td><span class="empty-value">-&nbsp;empty&nbsp;-</span></td></tr>
      <tr><th scope="row">Embargo Period</th><td>11-27-2025</td></tr>
      <tr><th scope="row">Regenerate</th>
          <td><a href="/cgi/viewcontent.cgi?article=1&force=yes">go</a></td></tr>
      <tr><th scope="row">Reviewers</th>
          <td><a href="/cgi/editor.cgi?window=reviewers">list</a></td></tr>
      </table></body></html>"""
    meta = M.parse_record_metadata(html)
    check("embargo . the date is read", M.embargo_from_metadata(meta),
          "11-27-2025")
    check("embargo . '- empty -' becomes empty, not the literal text",
          meta["OTHER_LICENSE"], "")
    check("embargo . a page with no table gives {} not None",
          M.parse_record_metadata("<p>nothing</p>"), {})
    check("embargo . no metadata means no embargo, not an exception",
          M.embargo_from_metadata({}), "")
    check("embargo . the label is matched case-insensitively",
          M.embargo_from_metadata({"embargo period": "1-1-2030"}), "1-1-2030")
    # This parser must never become a link harvester. force=yes is a WRITE.
    values = " ".join(meta.values()).lower()
    check_true("embargo . the regeneration link cannot escape through it",
               "force=yes" not in values, values[:200])
    check_true("embargo . no viewcontent link escapes through it",
               "viewcontent" not in values, values[:200])


# ---------------------------------------------------------------------------
# 27 · An absence is not a failure (v1.20).
#
# The first on-campus run asked for `native` on 86 records of 2017 PDF
# uploads. Every one answered 404 — the server saying, definitely, that
# there is no separate native — and the log filled with WARNING lines on a
# run where nothing had broken.
# ---------------------------------------------------------------------------
def test_an_absence_is_reported_as_an_absence():
    import urllib.error
    absent = urllib.error.HTTPError("u", 404, "Not Found", {}, None)
    gone = urllib.error.HTTPError("u", 410, "Gone", {}, None)
    refused = urllib.error.HTTPError("u", 403, "Forbidden", {}, None)
    throttled = urllib.error.HTTPError("u", 429, "Too Many", {}, None)
    check_true("absence . 404 is an absence", M.is_definite_absence(absent))
    check_true("absence . 410 is an absence", M.is_definite_absence(gone))
    # 403 means "you may not have this", never "this does not exist".
    check_true("absence . 403 is NOT an absence",
               not M.is_definite_absence(refused))
    check_true("absence . 429 is NOT an absence",
               not M.is_definite_absence(throttled))

    check_true("absence . a native absence says so in its own words",
               M.ABSENT_LABELS["native"].startswith("No native file"))
    check_true("absence . 'No native file: ...' is recognized as an absence",
               ("{}: HTTP 404 Not Found".format(M.ABSENT_LABELS["native"]))
               .startswith(M.ABSENT_PREFIXES))
    check_true("absence . 'Failed: ...' is NOT recognized as an absence",
               not "Failed: HTTP 403 Forbidden".startswith(M.ABSENT_PREFIXES))
    check_true("absence . 'Downloaded' is not an absence either",
               not "Downloaded".startswith(M.ABSENT_PREFIXES))

    # The worker must branch on it rather than calling everything Failed.
    import inspect
    body = inspect.getsource(M.download_worker)
    check_true("absence . the download worker distinguishes the two",
               "is_definite_absence(e)" in body, body[:200])
    check_true("absence . and does not warn for an absence",
               body.count("n_absent += 1") == 1)
    retry = inspect.getsource(M.retry_worker)
    check_true("absence . the retry worker draws the same line",
               "is_definite_absence(e)" in retry)


def test_a_definite_absence_is_not_retried():
    import tempfile
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article=7&context=etd"
    with tempfile.TemporaryDirectory() as d:
        p = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", [
            _row(1, "7", "native", "public", "current",
                 "No native file: HTTP 404 Not Found",
                 base + "/cgi/viewcontent.cgi?article=7&context=etd", adm),
            _row(2, "8", "primary", "public", "current",
                 "Failed: HTTP 429 Too Many Requests",
                 base + "/cgi/viewcontent.cgi?article=8&context=etd", adm),
        ])
        got, skipped, absent, _known = M.read_failed_rows(p)
        check("retry-absence . only the real failure is retried",
              [g["article"] for g in got], ["8"])
        check("retry-absence . the absence is counted, not dropped silently",
              absent, 1)
        check("retry-absence . and it is not reported as a problem",
              skipped, [])


# ---------------------------------------------------------------------------
# 28 · Records the run never reached (v1.20).
#
# A run that ended early recorded only the rows it ATTEMPTED, so the re-run
# control could not resume it: it recovers failures, and a record never
# tried is an absence, not a failure.
# ---------------------------------------------------------------------------
def test_the_plan_round_trips():
    opts = {"primary": True, "supp": True, "native": False, "public": True,
            "hidden": True, "published": True, "unpublished": False,
            "embargo": False, "versions": "all"}
    spec = M.plan_spec(opts)
    back = M.parse_plan_spec(spec)
    for k, v in opts.items():
        check("plan . {} survives the round trip".format(k), back[k], v)
    check("plan . an empty cell means no plan", M.parse_plan_spec(""), {})
    check("plan . whitespace means no plan", M.parse_plan_spec("   "), {})
    # It must never invent a run. An unreadable plan is no plan.
    check("plan . an unrecognized cell means no plan, not a default run",
          M.parse_plan_spec("nonsense,rubbish"), {})
    check("plan . a bad versions value does not become a plan on its own",
          M.parse_plan_spec("versions=sideways"), {})
    check("plan . versions defaults to current when the plan omits it",
          M.parse_plan_spec("primary")["versions"], "current")


def test_unreached_records_are_written_and_resumable():
    items = [({"article": "1", "title": "A",
               "release": "Thesis (open access)"}, "published"),
             ({"article": "2", "title": "B",
               "release": "Thesis (restricted to Some University)"},
              "published"),
             ({"article": "3", "title": "C", "release": ""}, "pending")]
    plan = M.plan_spec({"primary": True, "supp": True, "public": True,
                        "hidden": True, "published": True,
                        "versions": "all"})
    rows = M.not_attempted_rows(items, 1, 10, "https://dc.example.edu",
                                "etd", False, plan, when="00:00:00")
    i = {h: M.REPORT_HEADERS.index(h) for h in M.REPORT_HEADERS}
    check("unreached . one row per record never reached", len(rows), 2)
    check("unreached . every row is the width of the header",
          {len(r) for r in rows}, {len(M.REPORT_HEADERS)})
    check("unreached . it starts at the first un-reached record",
          [r[i["Article ID"]] for r in rows], ["2", "3"])
    check("unreached . the order continues from the rows already written",
          [r[i["Order"]] for r in rows], [11, 12])
    check("unreached . the status says what it is",
          rows[0][i["Status"]], M.NOT_ATTEMPTED)
    check("unreached . the kind marks it as a record, not a file",
          rows[0][i["File Kind"]], M.RECORD_KIND)
    check("unreached . access is carried even for a record never fetched",
          [r[i["Access"]] for r in rows],
          [M.ACCESS_RESTRICTED, M.ACCESS_UNKNOWN])
    check("unreached . the record state is carried",
          [r[i["Record State"]] for r in rows],
          [M.STATE_LABELS["published"], M.STATE_LABELS["pending"]])
    check("unreached . the plan is carried", rows[0][i["Plan"]], plan)
    check_true("unreached . the URL names the record and its context",
               "article=2" in rows[0][i["File URL"]]
               and "context=etd" in rows[0][i["File URL"]],
               rows[0][i["File URL"]])
    # Nothing un-reached must produce nothing at all.
    check("unreached . a completed structure writes no such rows",
          M.not_attempted_rows(items, 3, 0, "https://d", "etd", False, plan),
          [])
    check("unreached . and neither does an empty structure",
          M.not_attempted_rows([], 0, 0, "https://d", "etd", False, plan), [])

    # A gallery record must get the gallery admin URL, not the EdiKit one.
    grow = M.not_attempted_rows(items, 2, 0, "https://d", "gal", True, plan)
    check_true("unreached . a gallery record gets the gallery admin URL",
               "editor_gallery.cgi" in grow[0][i["Admin Record URL"]],
               grow[0][i["Admin Record URL"]])


def test_an_unreached_record_is_replanned_not_refetched():
    import tempfile
    base = "https://dc.example.edu"
    plan = M.plan_spec({"primary": True, "supp": True, "public": True,
                        "hidden": True, "published": True,
                        "versions": "current"})
    adm = base + "/cgi/editor.cgi?window=abstract&article=9&context=etd"
    with tempfile.TemporaryDirectory() as d:
        p = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", [
            _row(1, "9", M.RECORD_KIND, "", "", M.NOT_ATTEMPTED, adm, adm,
                 plan=plan),
        ])
        got, skipped, _absent, _known = M.read_failed_rows(p)
        check("replan . the un-reached record comes back", len(got), 1)
        check_true("replan . and is marked for re-planning",
                   got[0].get("replan") is True)
        check("replan . with the plan the interrupted run recorded",
              got[0]["plan"]["versions"], "current")
        check_true("replan . and its primary/supp intent",
                   got[0]["plan"]["primary"] and got[0]["plan"]["supp"])
        check("replan . the record state comes back for re-planning",
              got[0]["record_state"], "Posted")

        # PLANT THE VIOLATION: a record row with NO plan must not be
        # fetched. Its File URL is an admin HTML page, and fetching it
        # would save a web page under a name claiming to be the file.
        p2 = _fake_report(d, "DC_FileDownload_Report_x_S.xlsx", [
            _row(1, "9", M.RECORD_KIND, "", "", M.NOT_ATTEMPTED, adm, adm),
        ])
        got2, skipped2, _a2, _known = M.read_failed_rows(p2)
        check("replan . a record row with no plan is NOT queued for fetch",
              got2, [])
        check("replan . it is reported rather than silently dropped",
              len(skipped2), 1)
        check_true("replan . and the reason says what to do instead",
                   "Re-run the structure" in skipped2[0], skipped2[0])

    # The worker must reproduce the report's plan, not the current form.
    import inspect
    retry = inspect.getsource(M.retry_worker)
    check_true("replan . the retry worker plans from the row's own plan",
               'row["plan"]' in retry, retry[:200])
    check_true("replan . and calls the same planner the download run uses",
               "plan_item_jobs(" in retry)
    check_true("replan . the editor family comes from the admin URL, "
               "not a guess", "editor_gallery.cgi" in retry)


def test_the_early_exits_record_what_they_never_reached():
    import inspect
    body = inspect.getsource(M.download_worker)
    flush = body[body.index("def flush_pending"):body.index("def finish_master")]
    check_true("unreached . the stop/error flush writes them",
               "not_attempted_rows(" in flush, flush[:300])
    check_true("unreached . it reads the live position, not a stale copy",
               'pending["pos"][0]' in flush, flush[:400])
    # A session expiry ends a run just as a stop does.
    check_true("unreached . a session expiry writes them too",
               body.count("not_attempted_rows(") == 2, body[:200])
    check_true("unreached . the position only moves once a record is done",
               'pending["pos"][0] = i_idx' in body)


# ---------------------------------------------------------------------------
# 29 · The report writers refuse a row of the wrong width (v1.20).
#
# Appending a column means touching every append site. A missed one does not
# raise: openpyxl writes the short row and every value after the gap lands
# under the wrong heading — a silent, plausible, wrong report.
# ---------------------------------------------------------------------------
def test_a_failed_one_page_listing_falls_back_to_paging():
    """2026-09-10, on campus: the ETD `published` state answered

        RemoteDisconnected: Remote end closed connection without response

    after about sixty seconds. 3,106 records, ~4.8 MB — and the very same
    request had succeeded off campus the evening before, so this is a
    marginal cost of asking for everything at once, not a property of the
    state.

    x_showall is an OPTIMIZATION worth ~125 requests. Losing 3,106 records
    rather than spending two minutes paging is the wrong trade, and it is
    what happened: the structure carried on with 270 records from another
    state and the run silently became a much smaller test than intended.
    """
    import http.client
    real_fetch, real_throttle, real_note = M.fetch_html, M._throttle, M._note
    try:
        M._throttle = lambda *a, **k: None
        notes = []
        M._note = lambda m: notes.append(m)
        seen = []

        def one_page_dies(url, cookie, timeout=None):
            seen.append(url)
            if "x_showall=1" in url:
                raise http.client.RemoteDisconnected(
                    "Remote end closed connection without response")
            return "<html>paged</html>"

        M.fetch_html = one_page_dies
        url = M.edikit_listing_url("https://h", "etd", "published")
        html, used = M._fetch_listing(url, "c")
        check_true("listing-fallback . it pages instead of giving up",
                   "x_showall" not in used, used)
        check("listing-fallback . and returns the paged body",
              html, "<html>paged</html>")
        check_true("listing-fallback . the one-page request was tried first",
                   "x_showall=1" in seen[0], str(seen[:2]))
        check_true("listing-fallback . the log names the actual cause",
                   notes and "RemoteDisconnected" in notes[0], str(notes))
        check_true("listing-fallback . and says paging loses nothing",
                   notes and "loses nothing" in notes[0], str(notes))

        # A timeout is the same class of problem.
        seen.clear(); notes.clear()

        def one_page_times_out(url, cookie, timeout=None):
            seen.append(url)
            if "x_showall=1" in url:
                raise OSError("timed out")
            return "<html>paged</html>"
        M.fetch_html = one_page_times_out
        html, used = M._fetch_listing(url, "c")
        check_true("listing-fallback . a timeout falls back too",
                   "x_showall" not in used, used)

        # HTTP 400 — the case that already worked — must still work.
        seen.clear()

        def four_hundred(url, cookie, timeout=None):
            seen.append(url)
            if "x_showall=1" in url:
                raise urllib.error.HTTPError(url, 400, "Bad", {}, None)
            return "<html>paged</html>"
        M.fetch_html = four_hundred
        html, used = M._fetch_listing(url, "c")
        check_true("listing-fallback . HTTP 400 still falls back",
                   "x_showall" not in used, used)

        # But a page that has NO x_showall to drop must raise, not loop.
        def always_dies(url, cookie, timeout=None):
            raise http.client.RemoteDisconnected("nope")
        M.fetch_html = always_dies
        try:
            M._fetch_listing("https://h/cgi/editor.cgi?context=etd", "c")
            check_true("listing-fallback . a paged URL that fails raises",
                       False, "it did not raise")
        except http.client.RemoteDisconnected:
            check_true("listing-fallback . a paged URL that fails raises",
                       True)

        # A session expiry and a throttle are NOT the page being too big.
        for exc, name in ((M.LoginRequired(), "LoginRequired"),
                          (M.RateLimited("locked out"), "RateLimited")):
            def raiser(url, cookie, timeout=None, _e=exc):
                raise _e
            M.fetch_html = raiser
            try:
                M._fetch_listing(url, "c")
                check_true("listing-fallback . {} propagates".format(name),
                           False, "it was swallowed")
            except type(exc):
                check_true("listing-fallback . {} propagates".format(name),
                           True)
    finally:
        M.fetch_html, M._throttle, M._note = real_fetch, real_throttle, real_note


def test_a_run_of_403s_is_a_block_not_eighty_problems():
    """Measured off campus 2026-09-09, and again on campus 2026-09-11.

    09-09: 99 files downloaded, then 80 refusals in an unbroken row, every
    one on a record whose release option reads "open access", with four
    ACCESS-RESTRICTED files already downloaded — so the refusal has nothing
    to do with what the records permit.

    09-11: a RE-RUN met sixty refusals, printed sixty identical WARNING
    lines and never stopped. The detection existed only in download_worker;
    retry_worker had no copy of it, and a re-run does not go through
    download_worker. Eighth instance of "fixed in one worker, not the
    other".

    **This test used to be nine greps of download_worker's source.** Every
    one of them passed on the day the re-run failed, because they asked
    whether a function CONTAINED some text, not whether the behavior
    happened — and they only ever looked at one of the two functions. The
    logic is now an object, so the rules are exercised directly, and the
    worker that actually failed is driven end to end below.
    """
    W = M.BlockWatch

    # --- the run means "refusals since the last file served" --------------
    w = W(say_at=3, stop_at=5)
    for _ in range(2):
        w.refused("7", M.ACCESS_OPEN)
    check("403-run . refusals accumulate", w.run, 2)
    check_true("403-run . two is not yet a block", not w.is_block())

    w.served()
    check("403-run . a served file ends the run", w.run, 0)
    check_true("403-run . and it is no longer a block", not w.is_block())

    # A SUCCESS resetting the run is the half that never existed: the old
    # counter lived in an `except` clause, which a success never reaches.
    # So refusals scattered among good downloads accumulated forever.
    w2 = W(say_at=3, stop_at=5)
    for _ in range(10):
        w2.refused("7", M.ACCESS_OPEN)
        w2.served()
    check("403-run . scattered refusals never reach a block", w2.run, 0)
    check_true("403-run . and none of them is reported as one",
               not w2.is_block())
    check("403-run . though all of them are still counted", w2.total, 10)

    # An ABSENCE is neutral. The old counter reset on it, which meant a
    # block interleaved with missing natives could stay invisible; a 404
    # hands over no content, so it is no evidence the server has anything
    # left to give.
    w3 = W(say_at=3, stop_at=99)
    w3.refused("7"); w3.absent(); w3.refused("8"); w3.absent()
    check("403-run . an absence neither counts nor resets", w3.run, 2)
    w3.refused("9")
    check_true("403-run . so a block through absences is still seen",
               w3.is_block(), "run=%s" % w3.run)

    # --- said once, not per row -------------------------------------------
    w4 = W(say_at=2, stop_at=99)
    w4.refused("7", M.ACCESS_OPEN)
    check("403-run . nothing is said before the threshold",
          w4.take_message(), "")
    check_true("403-run . and saying nothing does not arm it",
               not w4.said)
    w4.refused("8", M.ACCESS_OPEN)
    first = w4.take_message()
    check_true("403-run . the block is announced once it is one", bool(first))
    check("403-run . and never a second time", w4.take_message(), "")
    for _ in range(50):
        w4.refused("9", M.ACCESS_OPEN)
        check_true("403-run . fifty more refusals stay silent",
                   w4.take_message() == "")
    check_true("403-run . but every one is still a block", w4.is_block())

    # --- what the message says --------------------------------------------
    for phrase in ("since the last file it served", "record(s)",
                   "block on this session or address",
                   "not a permission problem"):
        check_true("403-run . the message says {!r}".format(phrase[:28]),
                   phrase in first, first)
    check_true("403-run . it does not claim to know the cause",
               "because" not in first.lower(), first)
    check_true("403-run . open-access records are named as the tell",
               "open access" in first, first)
    check_true("403-run . and counted, not just mentioned",
               "2 marked" in first, first)

    w5 = W(say_at=1, stop_at=99)
    w5.refused("7", "restricted")
    check_true("403-run . with no open-access rows it says nothing of them",
               "open access" not in w5.take_message())

    # --- stopping ---------------------------------------------------------
    w6 = W(say_at=2, stop_at=4)
    for n in range(1, 4):
        w6.refused(str(n))
        check_true("403-run . {} refusals is not yet a stop".format(n),
                   not w6.should_stop())
    w6.refused("4")
    check_true("403-run . the stop threshold stops", w6.should_stop())
    check_true("403-run . and the control can switch it off",
               not w6.should_stop(False))
    check_true("403-run . the stop reason counts the records",
               "across 4 record(s)" in w6.stop_reason(), w6.stop_reason())

    check_true("403-run . the defaults come from the module constants",
               W().say_at == M.BLOCK_RUN_403
               and W().stop_at == M.BLOCK_STOP_403)
    check_true("403-run . and the stop threshold is the higher of the two",
               M.BLOCK_STOP_403 > M.BLOCK_RUN_403,
               "%s vs %s" % (M.BLOCK_STOP_403, M.BLOCK_RUN_403))

    # --- a 403 is never an absence ----------------------------------------
    import urllib.error
    f = urllib.error.HTTPError("u", 403, "Forbidden", {}, None)
    check_true("403-run . a 403 is still not an absence",
               not M.is_definite_absence(f))


def test_both_workers_watch_for_the_block():
    """The defect of 2026-09-11: retry_worker had no block detection.

    Six local variables in download_worker, and a re-run — which does not
    go through download_worker at all — logged sixty identical warnings and
    ground on. The standing rule is "when you fix a failure mode in one
    worker, check the other one", and this is the check that makes the rule
    mechanical rather than remembered.

    A source check, and recorded as such: it proves each worker names the
    calls, not that a block is detected. The behavior is proved for the
    watcher above and for retry_worker end to end below. download_worker is
    the one path still asserted only structurally.
    """
    import inspect
    for name in ("download_worker", "retry_worker"):
        body = inspect.getsource(getattr(M, name))
        code = "\n".join(l for l in body.splitlines()
                          if not l.strip().startswith("#"))
        check_true("both-watch . {} builds a watcher".format(name),
                   "BlockWatch()" in code, code[:200])
        for call in (".refused(", ".served(", ".absent(",
                     ".is_block()", ".should_stop(", ".take_message()"):
            check_true("both-watch . {} calls {}".format(name, call),
                       "watch" + call in code, name)
        check_true("both-watch . {} raises ServerRefusing".format(name),
                   "raise ServerRefusing(" in code, name)
        check_true("both-watch . {} honors the control".format(name),
                   'opts.get("stop_when_blocked", True)' in code, name)
        check_true("both-watch . {} reports the refusal count".format(name),
                   "refused with HTTP 403 across" in code, name)
        # The row that PROVED the block must be in the report. Raising
        # inside the except clause would stop one row short of its own
        # reason, so the decision is taken there and acted on after the
        # append.
        # Guarded rather than indexed straight: a missing anchor must
        # FAIL this check, not crash the test with a ValueError that
        # happens to look like a refusal.
        has_both = ("stop_after_row =" in code
                    and "raise ServerRefusing(" in code)
        check_true("both-watch . {} records the row before stopping"
                   .format(name),
                   has_both and (code.index("stop_after_row =")
                                 < code.index("raise ServerRefusing(")),
                   name)

    # There must be exactly one implementation, or this whole class of
    # defect comes back the next time one of them is edited.
    src = inspect.getsource(M)
    check("both-watch . the logic exists once",
          src.count("class BlockWatch"), 1)
    check("both-watch . and is constructed once per worker",
          src.count("BlockWatch()"), 2)


def test_a_blocked_rerun_stops_and_says_so_once():
    """The 2026-09-11 failure, driven end to end through retry_worker.

    Patches the fetch to refuse everything with 403 and runs the real
    worker over a real report. Asserts what the log and the report say, not
    what the source contains — the source version of this test passed on
    the morning the run failed.
    """
    import tempfile, urllib.error

    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article={}&context=etd"
    rows = [_row(i, str(i), "primary", "public", "current",
                 "Failed: HTTP 429 Too Many Requests",
                 base + "/cgi/viewcontent.cgi?article={}&context=etd".format(i),
                 adm.format(i), access=M.ACCESS_OPEN)
            for i in range(1, 41)]

    class _FakeDriver:
        page_source = "<html><body>admin</body></html>"
        current_url = base + "/"
        # A real driver always has a title; v1.34.1 reads it to prove the
        # tab answers before the run begins (ready_the_tab).
        title = "Digital Commons"

        def get(self, url):
            pass

    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "fetch_record_file",
              "LISTING_SETTLE", "STATE")}
    logs = []
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: _FakeDriver()
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0

            def _refuse(url, cookie, native_url=None, **kw):
                raise urllib.error.HTTPError(url, 403, "Forbidden", {}, None)
            M.fetch_record_file = _refuse

            real_log = M.log
            M.log = lambda msg: (logs.append(str(msg)), real_log(str(msg)))[0]
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}},
                           out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            # These exercise worker logic against a
                            # stubbed fetch_record_file, so they name
                            # the network path explicitly rather than
                            # inheriting whatever the default is. The
                            # browser path has worker-level tests of
                            # its own further down.
                            "fetch_path": M.FETCH_NETWORK,
                            "stop_when_blocked": True,
                            # Reconciliation is exercised by its own
                            # tests. Left on here it would reach for a real
                            # listing at dc.example.edu — a suite that
                            # touches the network is a suite that fails for
                            # reasons that have nothing to do with the code.
                            "reconcile": False,
                            },
                           report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            M.log = real_log

    text = "\n".join(logs)
    block_lines = [l for l in logs if "block on this session or address" in l]
    check("blocked-rerun . the block is announced exactly once",
          len(block_lines), 1)
    per_row = [l for l in logs if "HTTP 403 Forbidden" in l
               and "block on this session" not in l]
    check_true("blocked-rerun . and not once per refused row",
               len(per_row) <= M.BLOCK_RUN_403,
               "%s line(s): %s" % (len(per_row), per_row[:3]))

    check_true("blocked-rerun . the run stops instead of grinding on",
               any("refusing file requests" in l for l in logs), text[-900:])
    state = M.STATE if isinstance(M.STATE, dict) else {}
    summary = str(state.get("summary", ""))
    check_true("blocked-rerun . the summary says the SERVER stopped it",
               "refusing" in summary, summary)
    check_true("blocked-rerun . not that the user did",
               "Stopped early" not in summary, summary)

    # It must stop near the threshold, not after all forty rows.
    check_true("blocked-rerun . it stops on the evidence, not at the end",
               M.BLOCK_STOP_403 <= len(rows),
               "the fixture is too short to prove a stop")
    # Everything fetched before the block has to survive it — the whole
    # point of stopping is that the run is resumable.
    check_true("blocked-rerun . a partial report is still written",
               "PARTIAL" in summary, summary)
    check_true("blocked-rerun . and the summary counts what was attempted",
               "of {} row(s)".format(len(rows)) in summary, summary)
    check_true("blocked-rerun . it stopped well short of the whole report",
               "({} of {} row(s)".format(M.BLOCK_STOP_403, len(rows))
               in summary, summary)


def test_a_bounce_is_recognized_by_where_the_bytes_came_from():
    """Measured 2026-09-22 — and it had been checked the weaker way.

    An expired session bounced an admin request to /cgi/login.cgi.
    urllib followed it, so the body read was the login page's — but the
    only test applied to it was a search for two field names, and on
    the real page that search did not match. HTTP 200, a readable page,
    and a crawl that reported "no records in this state".

    The address the bytes came from cannot be mistaken for an admin
    page, and `looks_like_login_page` already knew what to do with it.
    Nothing was showing it.
    """
    import contextlib

    class _Resp:
        def __init__(self, body, landed):
            self._body, self._landed = body, landed
            self.headers = {}
        def read(self):
            return self._body
        def geturl(self):
            return self._landed
        def get(self, k, d=None):
            return d

    @contextlib.contextmanager
    def _fake(body, landed):
        yield _Resp(body, landed)

    real = M._open
    try:
        # A login page that this module's OLD body test cannot see: no
        # field named "login", no field named "password". Reskinned
        # markup is not an exotic case; it is a vendor release.
        body = ("<html><body><h1>Sign in</h1>"
                "<form><input name='user'><input name='pw' "
                "type='password'></form></body></html>")
        M._open = lambda *a, **k: _fake(
            body.encode("utf-8"),
            "https://dc.example.edu/cgi/login.cgi?return_to=x")
        try:
            M.fetch_html("https://dc.example.edu/cgi/editor.cgi?context=etd",
                         "c")
            check_true("bounce . a request that landed on the login page "
                       "raises LoginRequired", False,
                       "it returned the login page as if it were an admin "
                       "page, which is how a real record was reported gone")
        except M.LoginRequired:
            check_true("bounce . a request that landed on the login page "
                       "raises LoginRequired", True)

        # And an ordinary admin page still comes back as a page.
        M._open = lambda *a, **k: _fake(
            b"<html><body><table>a listing</table></body></html>",
            "https://dc.example.edu/cgi/editor.cgi?context=etd")
        try:
            got = M.fetch_html(
                "https://dc.example.edu/cgi/editor.cgi?context=etd", "c")
            check_true("bounce . an ordinary admin page is still returned",
                       "a listing" in got, got[:80])
        except M.LoginRequired:
            check_true("bounce . an ordinary admin page is still returned",
                       False, "a working page was called a login bounce")

        # A file request that lands there is the same answer.
        M._open = lambda *a, **k: _fake(
            b"<html><body>Sign in</body></html>",
            "https://dc.example.edu/cgi/login.cgi?return_to=y")
        try:
            M.fetch_file("https://dc.example.edu/cgi/viewcontent.cgi?a=1", "c")
            check_true("bounce . a FILE request that lands there does too",
                       False, "reported as a web page instead of a session "
                              "expiry, and the row would read as an absence")
        except M.LoginRequired:
            check_true("bounce . a FILE request that lands there does too",
                       True)
        except ValueError as e:
            check_true("bounce . a FILE request that lands there does too",
                       False, "ValueError: {}".format(e))
    finally:
        M._open = real


def test_an_empty_listing_is_not_every_record_going_away():
    """Measured 2026-09-22, on a one-record re-run that should have been
    boring.

    The session had expired. The listing bounced to the login page, which
    was read as "no records in the 'Posted' state", which arrived at
    reconcile_rows as an empty live set — and every outstanding row was
    marked departed:

        listing drift - 1 row(s) name a record the listing no longer
        carries (1845) - those are skipped, not requested

    The record was there. One row that time; it would have been all 51
    on the next run, and every row of an ETD re-run after that, each one
    reported as a fact about the repository into a report the NEXT
    re-run reads.

    The asymmetry is what decides it, and it is not close. Requesting a
    record that really has gone costs one request and gets an honest
    per-record answer. Skipping one that is still there loses it
    silently and permanently.
    """
    rows = [{"article": str(1000 + i), "access": "open", "release": "",
             "record_state": "Posted"} for i in range(51)]
    rep = M.reconcile_rows([dict(r) for r in rows], {})
    check("empty-listing . nothing is called departed", rep["gone"], [])
    check_true("empty-listing . and it says it could not check",
               rep.get("unreadable") is True, str(rep))

    # The rows must still be REQUESTED. A row marked drift='gone' is
    # skipped by the caller, so this is the assertion that matters.
    marked = [dict(r) for r in rows]
    M.reconcile_rows(marked, {})
    check("empty-listing . every row is still requested",
          len([r for r in marked if r.get("drift") == "gone"]), 0)

    line = M.drift_summary(rep)
    check_true("empty-listing . the log says nothing could be checked",
               "no records at all" in line and "not a record going away"
               in line, line)
    check_true("empty-listing . and points at the likely cause",
               "log in" in line.lower(), line)

    # A record genuinely missing from a listing that DID come back is
    # still departed — the guard must not swallow the real case.
    live = {"1001": {"access": "open", "release": "",
                     "record_state": "Posted"}}
    rows2 = [{"article": "1000", "access": "open", "release": "",
              "record_state": "Posted"},
             {"article": "1001", "access": "open", "release": "",
              "record_state": "Posted"}]
    rep2 = M.reconcile_rows(rows2, live)
    check("empty-listing . a real departure is still reported",
          rep2["gone"], ["1000"])
    check("empty-listing . and the row that is there is not",
          rows2[1].get("drift"), None)


def test_the_browser_never_loads_the_big_listing():
    """Run E died two minutes in, on its first navigation.

        ReadTimeoutError: HTTPConnectionPool(host='localhost', port=56920)
        Read timed out. (read timeout=120)

    That is chromedriver, not Digital Commons. The first driver.get() was
    the ETD listing WITH x_showall=1 — 3,383 rows, about 4.8 MB — pulled
    through the DevTools channel, past Selenium's 120s client default that
    nothing was raising.

    The navigation exists only to validate the session and refresh
    cookies; every listing after it is fetched over urllib, where the
    timeout is ours. So it asks for the paged page and lets crawl_listing
    request x_showall over the transport that can afford it.
    """
    check_true("big-listing . the listing URL can be asked for paged",
               "x_showall=1" not in
               M.edikit_listing_url("https://h", "etd", "published",
                                    showall=False))
    check("big-listing . and the two differ ONLY by that parameter",
          M.edikit_listing_url("https://h", "etd", "published")
          .replace("&x_showall=1", ""),
          M.edikit_listing_url("https://h", "etd", "published",
                               showall=False))
    check_true("big-listing . showall is still the default",
               "x_showall=1" in
               M.edikit_listing_url("https://h", "etd", "published"))

    import inspect, ast, tempfile
    body = inspect.getsource(M.download_worker)
    # v1.28 sent the browser to the PAGED listing and two campus runs still
    # died at exactly 120.0s, so the invariant is stronger than "a smaller
    # admin page": the browser does not navigate to a cgi admin page at
    # all. Its only job is to look at the site page; everything else goes
    # over urllib, where the timeout is ours.
    #
    # Checked BEHAVIORALLY since 2026-09-22. It used to grep the worker's
    # source for `driver.get(` lines — and when that one navigation moved
    # into session_still_open(), where both workers could share it, the
    # grep found nothing and would have gone on passing while guarding an
    # empty string. Drive the thing and read where it went.
    with tempfile.TemporaryDirectory() as d:
        f, drv = _browser_fetcher(os.path.join(d, "_incoming"), [])
        html, bounced = M.session_still_open(drv, "https://dc")
        check("big-listing . the browser navigates to the site root, and "
              "nowhere else", drv.navigations, ["https://dc/"])
        check("big-listing . and nothing is reported as a bounce",
              bounced, False)
    # Neither worker may navigate on its own account any more: the one
    # navigation lives in the helper both of them call.
    for worker in (M.download_worker, M.retry_worker):
        tree = ast.parse(inspect.getsource(worker).lstrip())
        direct = [n.lineno for n in ast.walk(tree)
                  if isinstance(n, ast.Call)
                  and isinstance(n.func, ast.Attribute)
                  and n.func.attr == "get"
                  and getattr(n.func.value, "id", "") == "driver"]
        check("big-listing . {} drives the browser only through the one "
              "helper".format(worker.__name__), direct, [])

    # And the driver's page must NOT be handed to crawl_listing, or the
    # first state of every structure silently gives up x_showall.
    crawl = body[body.index("rows, gallery = crawl_listing("):]
    crawl = crawl[:crawl.index("flavour_known = True")]
    check_true("big-listing . the driver's page is not reused as the listing",
               "first_html" not in crawl, crawl[:300])

    # Bounds, so a slow page is a slow page and not an unhandled error.
    check_true("big-listing . a listing gets its own, longer timeout",
               M.LISTING_TIMEOUT > M.PAGE_TIMEOUT)
    src = Path(M.__file__).read_text(encoding="utf-8")
    fetcher = src[src.index("def _fetch_listing("):
                  src.index("def looks_like_admin_page(")]
    check("big-listing . and actually uses it, on both paths",
          fetcher.count("LISTING_TIMEOUT"), 2)
    check_true("big-listing . the driver command wait is raised above 120s",
               M.DRIVER_COMMAND_TIMEOUT > 120)
    attach = src[src.index("def attach_chrome("):]
    attach = attach[:attach.index("# ---")]
    check_true("big-listing . attach_chrome sets the command timeout",
               "RemoteConnection.set_timeout" in attach, attach[:400])
    check_true("big-listing . and a page-load timeout",
               "set_page_load_timeout" in attach, attach[:400])
    check_true("big-listing . every timeout lever is guarded",
               attach.count("except Exception") >= 3, attach[:800])
    # v1.28 wrapped the raise in a bare `except: pass`, so when it did not
    # take effect nothing said so — and two runs then failed at exactly
    # 120.0s, which IS the tell. A guard that says nothing turns a bug into
    # a mystery.
    check_true("big-listing . and reports what it actually achieved",
               "_dc_timeout_report" in attach, attach[-600:])
    # The PROPERTY, not the old comment text: when the raise fails, the
    # failure is recorded so the report can say so. Checking for a literal
    # `except: pass` string meant the check stopped matching the moment the
    # code was reformatted, which is not the same as the fix being present.
    # v1.34 reworded it ("unavailable", not "FAILED", in front of what is
    # usually a success) — the property is unchanged: the exception's
    # class is recorded in the report rather than swallowed.
    check_true("big-listing . a failed raise is recorded, not swallowed",
               'set_via = "class-level lever unavailable ({})".format('
               in attach, attach[:900])
    # v1.29 READ the instance timeout without ever SETTING it, so the
    # connection went on reporting 120s — honestly, which is how the log
    # named the cause. Reading a value is not setting it.
    check_true("big-listing . the instance timeout is SET, not just read",
               "conn._timeout = DRIVER_COMMAND_TIMEOUT" in attach,
               attach[:1200])
    check_true("big-listing . every lever that exists is tried",
               attach.count("DRIVER_COMMAND_TIMEOUT") >= 4, attach[:1200])
    check_true("big-listing . and each one reports whether it worked",
               attach.count("levers.append") >= 3, attach[:1200])
    # The one that actually bit: a guard on `is not None` that skipped
    # silently when the attribute was None.
    check_true("big-listing . a skipped lever is named, not silent",
               "FAILED (" in attach and "levers" in attach, attach[:1200])
    src2 = Path(M.__file__).read_text(encoding="utf-8")
    body2 = src2[src2.index("def download_worker("):]
    body2 = body2[:body2.index("def read_failed_rows(")]
    check_true("big-listing . and the worker logs that report",
               "_dc_timeout_report" in body2, body2[:600])


def test_the_rerun_summary_counts_what_happened():
    """v1.25 reported "recovered 153 of 75 failed row(s) . -78 still failing".

    Re-planning grows the row count — 75 un-reached records became 228 file
    rows — and the summary was still dividing by the ORIGINAL 75, so the
    denominator was stale and the remainder went negative. Nothing in the
    run was wrong; the arithmetic describing it was. A number derived by
    subtraction from a figure that moves is not a count.
    """
    import inspect
    body = inspect.getsource(M.retry_worker)
    done = body[body.index("write_master_workbook(mpath"):]
    done = done[:done.index("except StopRequested")]
    # Comments are prose ABOUT the code, not the code. The first version of
    # the check below tripped on its own explanatory comment, which quotes
    # the wording it forbids.
    code = "\n".join(l for l in done.splitlines()
                     if not l.strip().startswith("#"))

    check_true("rerun-sum . the denominator is the rows actually attempted",
               "tot_rows" in done, done[:400])
    check_true("rerun-sum . not the original row count",
               "of {} row(s)" in done and ", total," not in done, done[:400])
    check_true("rerun-sum . absences are named rather than folded into it",
               "tot_absent" in done, done[:500])
    check_true("rerun-sum . and the remainder cannot go negative",
               "still > 0" in done, done[:600])
    check_true("rerun-sum . the accumulators are counted, not derived",
               "tot_rows += n_rows" in body and "tot_absent += n_absent"
               in body, body[:200])
    # Records nobody attempted are not failures, so they are not called that.
    check_true("rerun-sum . un-attempted rows are not called failures",
               "failed row(s)" not in code, code[:400])

    # The arithmetic itself, on the shape of Jeff's actual run.
    rows, recovered, absent = 228, 153, 75
    still = rows - recovered - absent
    check("rerun-sum . the real run's remainder is zero, not negative",
          still, 0)
    check_true("rerun-sum . and the old formula is what produced -78",
               75 - recovered == -78)


def test_pages_and_downloads_are_paced_apart():
    """Digital Commons counts downloads, not admin pages.

    Evidence: the Hierarchy Mapping Tool sleeps 1.0s between admin page
    loads and made ~1,152 of them in 28 minutes with no lockout, after
    which a download run fetched 37 of 38 files unpenalised. One natural
    experiment, so the paces are controls with the proven value as their
    default — and pushback is attributed, so a lockout says which pace
    provoked it rather than leaving the question open.
    """
    slept = []
    real_sleep = M.time.sleep
    try:
        M.reset_rate_state()
        M.set_request_delay(5.0)
        M.set_page_delay(1.0)
        check("pace-split . the download pace is its own", M.request_delay(), 5.0)
        check("pace-split . and the page pace is its own", M.page_delay(), 1.0)

        # The one that matters: _throttle must actually WAIT the right
        # amount. Everything else here — two settings, labelled call sites,
        # attribution — is decoration if the throttle ignores the kind and
        # sleeps the download gap for everything. Planting exactly that
        # passed a suite of twelve other assertions about this feature.
        M.time.sleep = lambda n: slept.append(n)
        try:
            M._throttle("page")
            M._throttle("file")
            M._throttle()          # the default is the cautious one
        finally:
            M.time.sleep = real_sleep
        check("pace-split . a page request waits the page gap",
              slept[0], 1.0)
        check("pace-split . a download waits the download gap",
              slept[1], 5.0)
        check("pace-split . an unlabelled call gets the download gap",
              slept[2], 5.0)
        check_true("pace-split . and the two are not the same wait",
                   slept[0] != slept[1], str(slept))

        # The accumulated penalty is global: a lockout is a statement about
        # the session, not about one kind of request.
        M._note_rate_limit("file")
        check_true("pace-split . pushback slows downloads",
                   M.request_delay() > 5.0)
        check_true("pace-split . and slows admin pages too",
                   M.page_delay() > 1.0)
        st = M.rate_state()
        check("pace-split . the hit is attributed to downloads",
              (st["hits_file"], st["hits_page"]), (1, 0))
        M._note_rate_limit("page")
        st = M.rate_state()
        check("pace-split . and a page hit is attributed to pages",
              (st["hits_file"], st["hits_page"]), (1, 1))
        check("pace-split . the total still counts both", st["hits"], 2)

        M.reset_rate_state()
        st = M.rate_state()
        check("pace-split . a reset clears the attribution too",
              (st["hits"], st["hits_file"], st["hits_page"]), (0, 0, 0))
    finally:
        M.time.sleep = real_sleep

    # Every throttle site must say which kind it is. An unlabelled one
    # silently gets the download pace, which is the safe direction but
    # quietly forfeits the saving this whole change is for.
    src = Path(M.__file__).read_text(encoding="utf-8")
    import re
    sites = re.findall(r"^\s*_throttle\((.*?)\)\s*$", src, re.M)
    check_true("pace-split . every throttle site names its kind",
               all(t.strip() in ('"file"', '"page"') for t in sites),
               "unlabelled: %r" % [t for t in sites
                                   if t.strip() not in ('"file"', '"page"')])
    check_true("pace-split . there are sites of both kinds",
               '"file"' in sites and '"page"' in sites, str(sites))

    # And the fetch wrappers must carry the kind down to where a 429 is
    # recorded, or the attribution is decoration.
    html_fn = src[src.index("def fetch_html("):]
    html_fn = html_fn[:html_fn.index("def fetch_file(")]
    check_true("pace-split . fetching a page is recorded as a page",
               '_open(url, cookie, timeout, "page")' in html_fn, html_fn[:300])
    file_fn = src[src.index("def fetch_file("):]
    file_fn = file_fn[:file_fn.index("class NoFileAvailable")]
    check_true("pace-split . fetching a file keeps the download default",
               '"page"' not in file_fn, file_fn[:300])
    check_true("pace-split . and _open attributes what it caught",
               "_note_rate_limit(kind)" in src)

    # The result has to reach the person reading the log, in both
    # directions: pushback attributed, and its ABSENCE stated, because at
    # two different paces "no lockout" is the finding.
    body = src[src.index("def download_worker("):]
    body = body[:body.index("def read_failed_rows(")]
    check_true("pace-split . the summary attributes pushback",
               "on admin pages at" in body, body[-1200:])
    check_true("pace-split . and states its absence as a result",
               "no server pushback at" in body, body[-1200:])


def test_a_session_is_capped_in_both_units():
    """Jeff's design, 2026-09-11.

    Which unit Digital Commons rations is not known. Measured on campus:
    379 files / 322 MB across four runs unblocked; 167 files / 442 MB
    blocked; 158 files / 766 MB blocked. A file count does not predict it —
    176 files went through in one run on 09-09 and 158 blocked the next
    day. Megabytes fit those three better, and off campus blocked lower in
    both units, so the address matters as well.

    Rather than pick a theory and ship it, the module carries both limits,
    stops on whichever binds first, and lets the operator change either in
    the plan. This test holds that contract, including the part that makes
    it an instrument: a limit the operator edits is an experiment, and the
    plan workbook is its record.
    """
    # --- the table -------------------------------------------------------
    lim = M.session_limits()
    keys = [d["key"] for d in lim]
    check_true("limits . both units are present",
               "max_files" in keys and "max_mb" in keys, str(keys))
    for d in lim:
        # v1.34: the defaults are 0 — no limit. They were sized for a
        # download ceiling that turned out not to exist.
        check_true("limits . {} defaults to no limit".format(d["key"]),
                   isinstance(d["value"], int) and d["value"] == 0, str(d))
        check_true("limits . {} says where the number came from"
                   .format(d["key"]), len(d["basis"]) > 30, str(d))
        check_true("limits . {} does not claim the unit is known"
                   .format(d["key"]),
                   "not known" in d["basis"].lower()
                   or "not proven" in d["basis"].lower(),
                   d["basis"])
        check_true("limits . {} names its unit for the sheet".format(d["key"]),
                   d["unit"] in ("files", "MB"), str(d))

    # The old numbers are kept, as history, where the basis is shown —
    # replaced rather than deleted, because they were measured.
    for d in lim:
        check_true("limits . {} keeps the retired number as history"
                   .format(d["key"]), "Formerly" in d["basis"], d["basis"])
    vals = M.limit_values(lim)

    # --- what the operator sets wins, including "no limit" ---------------
    chosen = M.limit_values(M.session_limits(chosen={"max_files": 50}))
    check("limits . the plan's value wins", chosen["max_files"], 50)
    check("limits . and the other limit is left alone",
          chosen["max_mb"], vals["max_mb"])
    off = M.session_limits(chosen={"max_mb": 0})
    check("limits . zero means no limit, not fall back to the default",
          M.limit_values(off)["max_mb"], 0)
    check_true("limits . and the sheet says that is deliberate",
               "no limit" in [d for d in off if d["key"] == "max_mb"][0]["basis"],
               str(off))

    # --- self-calibration: what this install has seen the server do ------
    seen = M.session_limits(observed={"max_files": 200})
    check("limits . an observed block resizes the limit under it",
          M.limit_values(seen)["max_files"], 170)
    check_true("limits . and says it is this install's measurement",
               "this install" in
               [d for d in seen if d["key"] == "max_files"][0]["basis"],
               str(seen))
    both = M.session_limits(observed={"max_files": 200},
                            chosen={"max_files": 12})
    check("limits . but the operator still outranks the measurement",
          M.limit_values(both)["max_files"], 12)

    # --- the running budget, which is the half that is not an estimate ---
    b = M.SessionBudget(max_files=3, max_mb=0)
    for _ in range(2):
        b.spend(1024)
    check_true("budget . two of three files is not full", not b.exhausted())
    b.spend(1024)
    check("budget . the third fills it", b.exhausted(), "max_files")
    check_true("budget . and it says which limit and what it fetched",
               "3 file(s)" in b.reason() and "MB" in b.reason(), b.reason())

    b2 = M.SessionBudget(max_files=0, max_mb=2)
    b2.spend(1048576)
    check_true("budget . one MB of a two MB limit is not full",
               not b2.exhausted())
    b2.spend(1048576)
    check("budget . two fills it", b2.exhausted(), "max_mb")
    check_true("budget . and the reason names the volume",
               "2 MB" in b2.reason(), b2.reason())

    # Whichever binds FIRST, which is the whole point of carrying two.
    small = M.SessionBudget(max_files=100, max_mb=1)
    small.spend(2 * 1048576)
    check("budget . a few huge files end a session on volume",
          small.exhausted(), "max_mb")
    many = M.SessionBudget(max_files=2, max_mb=500)
    many.spend(1); many.spend(1)
    check("budget . many tiny ones end it on count",
          many.exhausted(), "max_files")

    off_b = M.SessionBudget(0, 0)
    for _ in range(500):
        off_b.spend(10 * 1048576)
    check("budget . with both off, nothing ends the session",
          off_b.exhausted(), "")

    # --- the plan-time split honours both too ----------------------------
    rows = [{"article": str(i)} for i in range(1, 101)]
    take, left = M.session_slice(rows, 40, max_mb=10, avg_mb=1.0)
    check("cap . the smaller of the two caps decides the slice", len(take), 10)
    check("cap . and the remainder is what is left", left, 90)
    take, left = M.session_slice(rows, 40, max_mb=1000, avg_mb=1.0)
    check("cap . a loose volume cap leaves the file cap in charge",
          len(take), 40)
    take, left = M.session_slice(rows, 0, max_mb=25, avg_mb=0.5)
    check("cap . volume alone can cap it", len(take), 50)
    take, left = M.session_slice(rows, 0, max_mb=25, avg_mb=0.0)
    check_true("cap . an unmeasured average cannot cap anything at plan time",
               len(take) == 100 and left == 0, "%s/%s" % (len(take), left))

    # --- the average is measured, never assumed --------------------------
    check("avg . no sizes means no average", M.average_file_mb([]), 0.0)
    check_true("avg . blanks and text are skipped, not counted as zero",
               M.average_file_mb([{"size": ""}, {"size": "n/a"},
                                  {"size": 2097152}]) == 2.0,
               str(M.average_file_mb([{"size": ""}, {"size": 2097152}])))


def test_a_full_session_is_not_reported_as_an_interruption():
    """A session that did exactly what the plan said must not read like one
    that went wrong. Driven end to end through retry_worker.

    The distinction is the operator's next move: 'session complete' means
    start the next one when the allowance refreshes; 'stopped early' means
    find out what happened first.
    """
    import tempfile

    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article={}&context=etd"
    rows = [_row(i, str(i), "primary", "public", "current",
                 "Failed: HTTP 429 Too Many Requests",
                 base + "/cgi/viewcontent.cgi?article={}&context=etd".format(i),
                 adm.format(i), access=M.ACCESS_OPEN)
            for i in range(1, 21)]

    class _FakeDriver:
        page_source = "<html><body>admin</body></html>"
        current_url = base + "/"
        # A real driver always has a title; v1.34.1 reads it to prove the
        # tab answers before the run begins (ready_the_tab).
        title = "Digital Commons"

        def get(self, url):
            pass

    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "fetch_record_file",
              "LISTING_SETTLE")}
    logs = []
    real_log = M.log
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: _FakeDriver()
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            # One megabyte a file, so the two-megabyte limit binds at
            # the third row. A kilobyte fixture against an MB limit is a
            # test that never reaches the thing it is testing.
            M.fetch_record_file = (
                lambda url, cookie, native_url=None, **kw:
                (b"x" * 1048576, "f.pdf", "application/pdf", False))
            M.log = lambda msg: (logs.append(str(msg)), real_log(str(msg)))[0]
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}},
                           out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            # These exercise worker logic against a
                            # stubbed fetch_record_file, so they name
                            # the network path explicitly rather than
                            # inheriting whatever the default is. The
                            # browser path has worker-level tests of
                            # its own further down.
                            "fetch_path": M.FETCH_NETWORK,
                            # The VOLUME cap, with no file cap, so the plan
                            # hands the session all 20 rows and the budget
                            # ends it at 2 — which is the case this whole
                            # mechanism exists for: a plan-time estimate
                            # that did not see how big the files were.
                            "max_files": 0, "max_mb": 2,
                            # Reconciliation is exercised by its own
                            # tests. Left on here it would reach for a real
                            # listing at dc.example.edu — a suite that
                            # touches the network is a suite that fails for
                            # reasons that have nothing to do with the code.
                            "reconcile": False,
                            },
                           report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            M.log = real_log

    summary = str(M.STATE.get("summary", ""))
    check_true("session-full . the session stops at its limit",
               "(2 of 20 row(s)" in summary, summary)
    check_true("session-full . and counts what it actually recovered",
               "2 recovered" in summary,
               "the count comes from an accumulator that had not caught up")
    check_true("session-full . and is reported as complete, not interrupted",
               "Session complete" in summary, summary)
    check_true("session-full . never as 'stopped early'",
               "Stopped early" not in summary, summary)
    check_true("session-full . the log says what to do next",
               any("allowance refreshes" in l for l in logs),
               "\n".join(logs[-4:]))
    check_true("session-full . and it does not read as a server refusal",
               not any("refusing" in l for l in logs), "\n".join(logs[-4:]))
    check_true("session-full . the rows fetched are still written out",
               "PARTIAL" in summary, summary)


def _check_job_mode(page):
    """Execute the page's applyJobMode() in node and report what it did.

    A stub `document` records `hidden` and `textContent` per element id, so
    this asserts the OUTCOME — which sections are shown and how the visible
    steps are numbered — rather than the presence of some source text.
    """
    import json as _json
    import shutil as _shutil
    import subprocess as _sp
    import tempfile as _tf

    node = _shutil.which("node")
    if not node:
        check_true("rerun-ui . job-mode behavior (node not installed)",
                   True, "skipped")
        return

    start = page.find("const jobNew=")
    end = page.find("jobNew.addEventListener")
    if start < 0 or end < 0:
        check_true("rerun-ui . the job-mode switch is in the page", False,
                   "applyJobMode() was not found")
        return
    js = page[start:end]

    harness = """
    const EL = {};
    function el(id){ return EL[id] || (EL[id] = {id:id, hidden:false,
                                                 textContent:'',
                                                 checked:false}); }
    ['sec1','sec2','sec3','secresume','leg1','leg2','leg3','legfetch',
     'leglimits','leg4','legresume','jobnew','jobresume'].forEach(el);
    const document = { getElementById: (id)=> EL[id] || null };
    const go = el('go'), retrygo = el('retrygo');
    __JS__
    const out = {};
    for (const mode of ['new','resume']) {
      el('jobnew').checked   = (mode === 'new');
      el('jobresume').checked = (mode === 'resume');
      applyJobMode();
      out[mode] = {
        hidden: Object.keys(EL).filter(k => EL[k].hidden),
        legends: Object.fromEntries(
          ['leg1','leg2','leg3','legfetch','leglimits','leg4','legresume']
            .map(k => [k, EL[k].textContent])),
      };
    }
    console.log(JSON.stringify(out));
    """.replace("__JS__", js)

    with _tf.NamedTemporaryFile("w", suffix=".mjs", delete=False) as fh:
        fh.write(harness)
        path = fh.name
    try:
        proc = _sp.run([node, path], capture_output=True, text=True,
                       encoding="utf-8", errors="replace",
                       timeout=30)
    finally:
        os.unlink(path)
    if proc.returncode != 0:
        check_true("rerun-ui . the job-mode switch runs at all", False,
                   (proc.stderr or "")[-400:])
        return
    got = _json.loads(proc.stdout.strip().splitlines()[-1])

    new, res = got["new"], got["resume"]
    check_true("rerun-ui . a new job shows all the numbered steps",
               not ({"sec1", "sec2", "sec3"} & set(new["hidden"])),
               str(new["hidden"]))
    check_true("rerun-ui . and hides the previous-report section",
               "secresume" in new["hidden"], str(new["hidden"]))
    check_true("rerun-ui . re-run mode hides the steps it does not use",
               {"sec1", "sec2", "sec3"} <= set(res["hidden"]),
               "hidden in resume mode: " + str(res["hidden"]))
    check_true("rerun-ui . and shows the one it does",
               "secresume" not in res["hidden"], str(res["hidden"]))
    check_true("rerun-ui . each mode offers only the button it can use",
               "retrygo" in new["hidden"] and "go" in res["hidden"],
               "new=%s resume=%s" % (new["hidden"], res["hidden"]))

    # "1, 4, 5" reads as three steps that went missing.
    def numbers(legends):
        return [v.split(" ")[0] for v in legends.values() if v]
    # The hidden section keeps its name and loses its number, so filter to
    # the numbered ones — a legend reading "Previous report" in new-job mode
    # is the correct result, not a step called "Previous".
    # Six since v1.33: "How files are fetched" is a step in both modes,
    # because the re-run fetches through the same path and meets the same
    # verification prompts as a new job.
    check("rerun-ui . a new job numbers its steps 1 to 6",
          sorted(n for n in numbers(new["legends"]) if n.isdigit()),
          ["1", "2", "3", "4", "5", "6"])
    check("rerun-ui . a re-run numbers the steps it actually shows",
          sorted(n for n in numbers(res["legends"]) if n.isdigit()),
          ["1", "2", "3", "4"])
    check_true("rerun-ui . the hidden steps carry no number at all",
               all(not res["legends"][k][:1].isdigit()
                   for k in ("leg1", "leg2", "leg3")),
               str(res["legends"]))


def test_the_page_does_not_shout_its_evidence_at_staff():
    """Jeff, 2026-09-11: "the dev UI has gotten long with development notes,
    and the organization may confuse staff in a few places if it ships."

    Every control's hint had grown into the full measurement narrative that
    justified its default — genuinely worth keeping, and not what a person
    about to press Start needs in front of them. The instruction stays
    visible; the evidence moves behind a disclosure. Nothing is deleted:
    losing the derivations would be worse than the wall of text, because a
    number whose basis is gone is a number nobody can check.
    """
    page = M.build_page(
        M.load_session(Path("/nonexistent/session.json"))).decode("utf-8")

    check_true("ui-len . the evidence is kept, behind a disclosure",
               page.count("<details>") >= 4,
               "%s disclosure(s)" % page.count("<details>"))
    check_true("ui-len . each one says what it holds",
               page.count("<summary>") == page.count("<details>"))

    # The visible hints are the instruction, not the argument for it.
    import re
    visible = re.sub(r"<details>.*?</details>", "", page, flags=re.S)
    body = visible[visible.index("<fieldset"):visible.index("</main>")]
    hints = re.findall(r'<p class="hint">(.*?)</p>', body, re.S)
    longest = max((len(re.sub(r"\s+", " ", h)) for h in hints), default=0)
    check_true("ui-len . no visible hint runs to an essay",
               longest <= 700, "longest visible hint is %s chars" % longest)

    # The message callout scrolls itself into view. At the FOOT of the page
    # that meant every scan result dragged the reader to the bottom, away
    # from the control they had just used — Jeff's third point.
    check_true("ui-scroll . the message callout sits above the form",
               page.index('id="cerr"') < page.index("<fieldset"),
               "it is still below the form, so it scrolls the page away")
    check_true("ui-scroll . and it still scrolls itself into view",
               "cerrBox.scrollIntoView" in page)

    # Numbers in the page come from the constants, never typed twice.
    check_true("ui-len . no unreplaced token reached the page",
               not re.search(r"__[A-Z]+__", page),
               str(re.findall(r"__[A-Z]+__", page))[:200])


def test_a_session_re_reads_the_listing_before_it_fetches():
    """The plan is a snapshot, and Jeff was explicit that it should stay one.

    "Snapshot sessions are fine so long as they don't choke on records that
    are removed or change status since the plan was created (I'm not
    concerned about missing new records)." — 2026-09-10.

    Not choking was mostly already handled: a removed record's URL answers
    404 and a refused one answers 403, and neither stops a session. **The
    hazard was never the crash. It was the label** — a record restricted
    since the plan was taken would be downloaded and then recorded with the
    plan's old Access value, and that report is the only thing marking which
    files on disk are access-restricted.
    """
    def row(article, access="open", release="Thesis (open access)",
            state="Posted"):
        return {"article": article, "access": access, "release": release,
                "record_state": state, "kind": "primary"}

    # --- nothing moved ---------------------------------------------------
    rows = [row("1"), row("2")]
    live = {"1": {"access": "open", "release": "Thesis (open access)",
                  "record_state": "Posted"},
            "2": {"access": "open", "release": "Thesis (open access)",
                  "record_state": "Posted"}}
    rep = M.reconcile_rows(rows, live, planned={"1", "2"})
    check("drift . a quiet collection reports nothing gone", rep["gone"], [])
    check("drift . nothing changed", rep["changed"], [])
    check("drift . and nothing added", rep["added"], 0)

    # `rows` is what is still OUTSTANDING — a session's slice of it — and
    # the plan is the whole job. Measuring additions against the slice made
    # every record the session was not touching look new: on 2026-09-11 a
    # session holding 14 records out of a 270-record listing announced
    # "256 record(s) have been added since the plan was taken", and none
    # had been. Without the plan the count is not computable, and that is
    # reported as unknown rather than as zero.
    blind = M.reconcile_rows([dict(r) for r in rows], live)
    check("drift . with no plan to compare against, no number is invented",
          blind["added"], None)
    # Silence would be wrong too: "no records were added" and "I cannot
    # tell whether any were" are different claims, and a log that omits the
    # second reads as the first.
    check_true("drift . and the log says the question is unanswerable",
               "cannot be counted" in M.drift_summary(blind),
               M.drift_summary(blind))
    check_true("drift . naming what is missing, so it can be fixed",
               "plan workbook" in M.drift_summary(blind),
               M.drift_summary(blind))
    check_true("drift . while a known plan says nothing of the sort",
               "cannot be counted" not in M.drift_summary(rep),
               M.drift_summary(rep))
    slice_of = M.reconcile_rows([dict(rows[0])], live,
                                planned={"1", "2"})
    check("drift . a session's slice is not mistaken for the whole plan",
          slice_of["added"], 0)
    check_true("drift . no row is flagged",
               not any("drift" in r for r in rows), str(rows))
    # An absence is a result: "checked and found nothing" and "never
    # checked" look identical in a log otherwise, and they are not the
    # same claim.
    line = M.drift_summary(rep)
    check_true("drift . and the session SAYS it checked",
               "checked 2 planned row(s)" in line and "nothing has been"
               in line, line)

    # --- a record restricted since the plan ------------------------------
    rows = [row("1")]
    live = {"1": {"access": "restricted",
                  "release": "Thesis (restricted to A University)",
                  "record_state": "Posted"}}
    rep = M.reconcile_rows(rows, live)
    check("drift . the change is reported", len(rep["changed"]), 2)
    check("drift . and the row is re-labelled from the listing",
          rows[0]["access"], "restricted")
    check_true("drift . the release option travels with it",
               "restricted to" in rows[0]["release"], str(rows[0]))
    check("drift . the row is flagged as changed", rows[0]["drift"], "changed")
    check_true("drift . and the log names the fields, not just a count",
               "access" in M.drift_summary(rep)
               and "release" in M.drift_summary(rep), M.drift_summary(rep))
    # One record whose access AND release moved is two changes and ONE
    # record. Counting the change tuples reported it as "2 row(s) changed".
    check_true("drift . one record changed twice is still one record",
               "1 record(s) changed" in M.drift_summary(rep),
               M.drift_summary(rep))

    # --- a record removed since the plan ---------------------------------
    rows = [row("1"), row("2")]
    rep = M.reconcile_rows(rows, {"1": live["1"]})
    check("drift . a departed record is named", rep["gone"], ["2"])
    check("drift . and flagged so it is not requested",
          rows[1]["drift"], "gone")
    check_true("drift . the log says those are skipped, not requested",
               "skipped, not requested" in M.drift_summary(rep),
               M.drift_summary(rep))

    # --- records added since the plan: counted, and that is all ----------
    rows = [row("1")]
    # A FRESH reading. Reusing the dict an earlier case re-labelled through
    # would have this assert that two records changed, which is the fixture
    # talking, not the code.
    same = {"access": "open", "release": "Thesis (open access)",
            "record_state": "Posted"}
    rep = M.reconcile_rows(rows, {"1": dict(same), "9": dict(same),
                                  "10": dict(same)}, planned={"1"})
    check("drift . new records are counted", rep["added"], 2)
    check("drift . and change nothing else",
          (rep["gone"], rep["changed"]), ([], []))
    check_true("drift . the log says they are not part of the plan",
               "not part of it" in M.drift_summary(rep), M.drift_summary(rep))

    # --- an empty reading is not a reading of "none" ---------------------
    # A faculty series' Type column carries a document type and no access
    # information at all. Blanking the plan's Access from a listing that
    # does not express access would destroy the only record of it.
    rows = [row("1", access="restricted", release="Thesis (restricted)")]
    rep = M.reconcile_rows(rows, {"1": {"access": "", "release": "",
                                        "record_state": ""}})
    check("drift . a silent listing does not blank the plan's access",
          rows[0]["access"], "restricted")
    check("drift . nor its release option",
          rows[0]["release"], "Thesis (restricted)")
    check("drift . and that is not reported as a change", rep["changed"], [])

    # --- the check is not a precondition for downloading -----------------
    import inspect
    body = inspect.getsource(M.live_listing)
    check_true("drift . a listing that will not load is not fatal",
               "failed.append(label)" in body and "continue" in body,
               body[:400])
    check_true("drift . and the session says which states went unread",
               "the plan's labels." in body, body[:900])
    check_true("drift . a login bounce still propagates",
               "except (LoginRequired, StopRequested):" in body, body[:600])

    worker = inspect.getsource(M.retry_worker)
    # Not a grep for "unread_labels": the first version of this check
    # matched the variable's ASSIGNMENT and stayed green when its only USE
    # was deleted — "a check matching a variable's declaration rather than
    # its use", verbatim from the standing list. The behavior is driven in
    # test_a_session_acts_on_what_the_listing_says_now below.
    # Anchored on fetch_one(), which is where BOTH workers now fetch. It
    # said fetch_record_file( until v1.33 moved the retry worker onto the
    # fetcher object; the name it greps for has to be the call the worker
    # actually makes, or the ordering it asserts is an ordering of
    # something else.
    check_true("drift . the session reconciles BEFORE it fetches",
               worker.index("reconcile_rows(")
               < worker.index("got = fetch_one("),
               "re-labelling after fetching is re-labelling too late")
    check_true("drift . only the states the plan used are re-read",
               "STATE_BY_LABEL.get(r.get(\"record_state\"" in worker,
               "crawling states outside the job would call a re-stated "
               "record removed")
    check_true("drift . a departed record is not re-planned either",
               'row.get("drift") == "gone"' in worker
               and worker.count('drift") == "gone"') >= 2,
               "re-planning costs two admin pages for a record that is gone")
    check_true("drift . it is paced as an admin page, not a download",
               '_throttle("page")' in body, body[-400:])
    check_true("drift . and the operator can turn it off",
               'opts.get("reconcile", True)' in worker, worker[:400])


def _rows_by_article(run_dir):
    """{article: (Status, Access)} from every structure report under a run.

    Raises if it finds nothing, rather than returning {} — an empty result
    makes every assertion built on it fail for the same uninformative
    reason, which is how a broken fixture wears the costume of a broken
    feature.
    """
    import glob as _g
    import openpyxl
    got = {}
    found = 0
    for f in _g.glob(os.path.join(run_dir, "**", "*.xlsx"), recursive=True):
        if "MasterLog" in os.path.basename(f):
            continue
        found += 1
        # Bind and close. The module's own reader was fixed for this on
        # 2026-09-21; the harness had the identical defect in a helper,
        # which is why three Windows failures survived that fix — and
        # they were the three whose reports this helper reads.
        wbk = openpyxl.load_workbook(f, read_only=True)
        try:
            data = list(wbk.active.values)
        finally:
            wbk.close()
        hdr = [str(h) for h in data[0]]
        I = {h: i for i, h in enumerate(hdr)}
        for r in data[1:]:
            got[str(r[I["Article ID"]])] = (str(r[I["Status"]] or ""),
                                            str(r[I["Access"]] or ""))
    if not found:
        raise AssertionError("no structure report under " + run_dir)
    return got


def test_a_session_acts_on_what_the_listing_says_now():
    """Drives a re-run against a listing that has moved since the plan.

    Three things have to happen and only one of them is visible in the pure
    function: the departed record must not be REQUESTED (a capped allowance
    must not be spent being told 404 by a server that already said so for
    free), the restricted record must be written with its NEW access, and a
    record state whose listing will not load must not make its records look
    removed.
    """
    import tempfile
    import urllib.error

    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article={}&context=etd"
    rows = [_row(i, str(i), "primary", "public", "current",
                 "Failed: HTTP 429 Too Many Requests",
                 base + "/cgi/viewcontent.cgi?article={}&context=etd".format(i),
                 adm.format(i), access=M.ACCESS_OPEN,
                 release="Thesis (open access)")
            for i in (1, 2, 3)]

    class _FakeDriver:
        page_source = "<html><body>admin</body></html>"
        current_url = base + "/"
        # A real driver always has a title; v1.34.1 reads it to prove the
        # tab answers before the run begins (ready_the_tab).
        title = "Digital Commons"

        def get(self, url):
            pass

    asked = []

    def _listing(base_url, ctx, gallery, state, cookie, first_html=None,
                 allow_flip=True):
        # Article 3 has gone; article 2 has been restricted since the plan.
        return ([{"article": "1", "release": "Thesis (open access)",
                  "access": M.ACCESS_OPEN},
                 {"article": "2",
                  "release": "Thesis (restricted to A University)",
                  "access": M.ACCESS_RESTRICTED}], gallery)

    def _fetch(url, cookie, native_url=None, **kw):
        asked.append(url)
        return (b"x" * 16, "f.pdf", "application/pdf", False)

    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "fetch_record_file",
              "crawl_listing", "LISTING_SETTLE")}
    logs = []
    real_log = M.log
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: _FakeDriver()
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            M.crawl_listing = _listing
            M.fetch_record_file = _fetch
            M.log = lambda msg: (logs.append(str(msg)), real_log(str(msg)))[0]
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}},
                           out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            # These exercise worker logic against a
                            # stubbed fetch_record_file, so they name
                            # the network path explicitly rather than
                            # inheriting whatever the default is. The
                            # browser path has worker-level tests of
                            # its own further down.
                            "fetch_path": M.FETCH_NETWORK,
                            "reconcile": True},
                           report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            M.log = real_log
        # Read the reports INSIDE the temporary directory's lifetime. The
        # first version of this read them after the `with` block, by which
        # point the folder was gone and every row silently read as absent —
        # a test that proved nothing and looked like a failing feature.
        got = _rows_by_article(out)

    check("drift-live . the departed record is never requested",
          len([u for u in asked if "article=3" in u]), 0)
    check("drift-live . the two that remain are", len(asked), 2)
    check_true("drift-live . and the session says what moved",
               any("listing drift" in l for l in logs),
               "\n".join(logs[-5:])[:400])

    # The report is the only thing marking which files on disk are
    # access-restricted, so the NEW value has to reach it.
    check("drift-live . the restricted record is labelled restricted now",
          got.get("2", ("", ""))[1], M.ACCESS_RESTRICTED)
    check("drift-live . the one that did not move keeps its label",
          got.get("1", ("", ""))[1], M.ACCESS_OPEN)
    check_true("drift-live . the departed record is recorded, not dropped",
               "3" in got, str(sorted(got)))
    check_true("drift-live . and recorded as departed, not as a failure",
               got.get("3", ("", ""))[0] == M.DRIFT_GONE,
               str(got.get("3")))
    check_true("drift-live . which is not counted among the failures",
               not got.get("3", ("", ""))[0].startswith("Failed"),
               str(got.get("3")))

    # --- a listing that will not load --------------------------------------
    # Reconciliation is a labelling improvement, not a precondition for
    # downloading. A record state whose listing refuses must leave its
    # records with the plan's labels — reconciling them against a listing
    # that was never read would report every one of them as removed and
    # skip the lot.
    rows2 = [_row(1, "1", "primary", "public", "current",
                  "Failed: HTTP 429 Too Many Requests",
                  base + "/cgi/viewcontent.cgi?article=1&context=etd",
                  adm.format(1), access=M.ACCESS_OPEN),
             _row(2, "7", "primary", "hidden", "current",
                  "Failed: HTTP 429 Too Many Requests",
                  base + "/cgi/viewcontent.cgi?article=7&context=etd",
                  adm.format(7), access=M.ACCESS_OPEN)]
    # The second row is in a different record state, and that state's
    # listing is the one that will not load.
    rows2[1] = tuple("Not yet posted" if i == M.REPORT_HEADERS.index(
        "Record State") else v for i, v in enumerate(rows2[1]))

    asked2 = []

    def _listing2(base_url, ctx, gallery, state, cookie, first_html=None,
                  allow_flip=True):
        if state == "pending":
            raise urllib.error.HTTPError(base, 400, "Bad Request", {}, None)
        return ([{"article": "1", "release": "Thesis (open access)",
                  "access": M.ACCESS_OPEN}], gallery)

    saved2 = {k: getattr(M, k) for k in
              ("_attach_or_fail", "cookie_header", "fetch_record_file",
               "crawl_listing", "LISTING_SETTLE")}
    logs2 = []
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows2)
        out2 = os.path.join(d, "out")
        os.makedirs(out2, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: _FakeDriver()
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            M.crawl_listing = _listing2
            M.fetch_record_file = (
                lambda url, cookie, native_url=None, **kw:
                (asked2.append(url),
                 (b"x" * 16, "f.pdf", "application/pdf", False))[1])
            M.log = lambda msg: (logs2.append(str(msg)), real_log(str(msg)))[0]
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}},
                           out2, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            # These exercise worker logic against a
                            # stubbed fetch_record_file, so they name
                            # the network path explicitly rather than
                            # inheriting whatever the default is. The
                            # browser path has worker-level tests of
                            # its own further down.
                            "fetch_path": M.FETCH_NETWORK,
                            "reconcile": True},
                           report_dir=out2)
        finally:
            for k, v in saved2.items():
                setattr(M, k, v)
            M.log = real_log
        got2 = _rows_by_article(out2)

    check("drift-live . an unread state's record is still requested",
          len([u for u in asked2 if "article=7" in u]), 1)
    check_true("drift-live . and is not recorded as departed",
               got2.get("7", ("", ""))[0] != M.DRIFT_GONE, str(got2.get("7")))
    check("drift-live . the state that DID load is reconciled as normal",
          len([u for u in asked2 if "article=1" in u]), 1)
    check_true("drift-live . and the session says which listing it could "
               "not read",
               any("could not re-read" in l for l in logs2),
               "\n".join(logs2[-6:])[:500])
    check_true("drift-live . naming what that means for those rows",
               any("keep the plan's labels" in l.replace("\n", " ")
                   for l in logs2),
               "\n".join(logs2[-6:])[:500])


def test_a_stopped_rerun_records_what_it_never_reached():
    """2026-09-11: a session of 448 rows stopped at 249, and the other 199
    were in no file anywhere.

    v1.20 gave download_worker the rule that a run ending early records the
    records it never reached, with the plan needed to resume them.
    retry_worker's _PARTIAL wrote only the rows it had attempted. **Ninth
    instance** of a failure mode fixed in one worker and not the other — and
    this one is load-bearing for the whole sessions architecture, in which
    session 2 reads session 1's folder to learn what is left. Rows missing
    from that folder are rows silently dropped from the job.
    """
    hdr = M.REPORT_HEADERS
    src = [{"article": str(i), "title": "T%s" % i, "record_state": "Posted",
            "kind": "primary", "visibility": "public", "version": "current",
            "version_date": "", "orig_hint": "a.pdf",
            "url": "http://dc/x?article=%s" % i, "pub": "http://pub/%s" % i,
            "adm": "http://adm/%s" % i, "access": "open",
            "release": "Thesis (open access)", "embargo": "",
            "plan": {"primary": True, "supp": False, "native": False,
                     "embargo": False, "versions": "current"}}
           for i in range(1, 6)]

    got = M.retry_not_attempted_rows(src, 2, 10)
    check("rerun-never . every row past the marker is written", len(got), 3)
    check("rerun-never . and none before it", [r[hdr.index("Article ID")]
                                               for r in got],
          ["3", "4", "5"])
    check("rerun-never . numbering continues from the rows already written",
          [r[0] for r in got], [11, 12, 13])
    for r in got:
        check("rerun-never . the status says nobody tried",
              r[hdr.index("Status")], M.NOT_ATTEMPTED)
        check_true("rerun-never . the file URL survives, so it is resumable",
                   r[hdr.index("File URL")].startswith("http://dc/"), str(r))
        check_true("rerun-never . so does the access it was labelled with",
                   r[hdr.index("Access")] == "open", str(r))
    # The Plan cell is read back by the next session. read_failed_rows hands
    # the worker a PARSED plan, and openpyxl refuses a dict outright — which
    # is the only reason a column of "{}" was caught rather than shipped.
    cell = got[0][hdr.index("Plan")]
    check_true("rerun-never . the plan travels as its token, not a dict",
               isinstance(cell, str) and "primary" in cell, repr(cell))
    check_true("rerun-never . and parses back to the job it came from",
               M.parse_plan_spec(cell).get("primary") is True, repr(cell))

    check("rerun-never . a row width that a report will accept",
          {len(r) for r in got}, {len(hdr)})
    check("rerun-never . nothing to write is not an error",
          M.retry_not_attempted_rows(src, None, 0), [])
    check("rerun-never . a session that reached the end writes none",
          M.retry_not_attempted_rows(src, 5, 0), [])

    # The row in flight when a session ends is inside the un-reached slice
    # even though its outcome was recorded a moment earlier. Writing it
    # again would put "Failed: HTTP 403" and "Not attempted" in the same
    # report for the same file.
    already = [tuple(
        {"Order": 1, "Article ID": "3", "File Kind": "primary",
         "File URL": "http://dc/x?article=3", "Status": "Failed: HTTP 403",
         "Title": "", "Record State": "", "Visibility": "", "Version": "",
         "Version Date": "", "Original Filename": "", "Saved Filename": "",
         "Size": "", "Public Record URL": "", "Admin Record URL": "",
         "Time": "", "Content Type": "", "Access": "",
         "Type / Release Option": "", "Embargo": "", "Plan": "",
         "Fetched via": ""}[h]
        for h in M.REPORT_HEADERS)]
    dedup = M.retry_not_attempted_rows(src, 2, 1, written=already)
    check("rerun-never . a row whose outcome is recorded is not re-written",
          [r[hdr.index("Article ID")] for r in dedup], ["4", "5"])
    check("rerun-never . and the numbering has no gap where it was",
          [r[0] for r in dedup], [2, 3])

    # A RECORD row is superseded by the FILE rows that planning it
    # produced, whose URLs differ by construction — so identity alone
    # cannot drop it, and a record fully expanded before the session ended
    # would come back as "Not attempted" beside its own results.
    rec = [dict(src[2], kind=M.RECORD_KIND, replan=True,
                url="http://adm/3")]
    files = [tuple(
        {"Order": 1, "Article ID": "3", "File Kind": "primary",
         "File URL": "http://dc/x?article=3&type=native",
         "Status": "Downloaded", "Title": "", "Record State": "",
         "Visibility": "", "Version": "", "Version Date": "",
         "Original Filename": "", "Saved Filename": "", "Size": "",
         "Public Record URL": "", "Admin Record URL": "", "Time": "",
         "Content Type": "", "Access": "", "Type / Release Option": "",
         "Embargo": "", "Plan": "", "Fetched via": ""}[h]
        for h in M.REPORT_HEADERS)]
    check("rerun-never . a record row is superseded by its own file rows",
          M.retry_not_attempted_rows(rec, 0, 0, written=files), [])
    check_true("rerun-never . but a record with no file rows is still named",
               len(M.retry_not_attempted_rows(rec, 0, 0, written=[])) == 1)

    # And the worker actually uses it. Driven end to end: the run is blocked
    # after 20 of 40 rows, so 20 must come back as Not attempted.
    import tempfile, urllib.error
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article={}&context=etd"
    rows = [_row(i, str(i), "primary", "public", "current",
                 "Failed: HTTP 429 Too Many Requests",
                 base + "/cgi/viewcontent.cgi?article={}&context=etd".format(i),
                 adm.format(i), access=M.ACCESS_OPEN)
            for i in range(1, 41)]

    class _FakeDriver:
        page_source = "<html><body>admin</body></html>"
        current_url = base + "/"
        # A real driver always has a title; v1.34.1 reads it to prove the
        # tab answers before the run begins (ready_the_tab).
        title = "Digital Commons"

        def get(self, url):
            pass

    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "fetch_record_file",
              "LISTING_SETTLE")}
    real_log = M.log
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: _FakeDriver()
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0

            def _refuse(url, cookie, native_url=None, **kw):
                raise urllib.error.HTTPError(url, 403, "Forbidden", {}, None)
            M.fetch_record_file = _refuse
            M.log = lambda msg: msg
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}}, out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            # These exercise worker logic against a
                            # stubbed fetch_record_file, so they name
                            # the network path explicitly rather than
                            # inheriting whatever the default is. The
                            # browser path has worker-level tests of
                            # its own further down.
                            "fetch_path": M.FETCH_NETWORK,
                            "stop_when_blocked": True, "reconcile": False},
                           report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            M.log = real_log
        written = _rows_by_article(out)

    check("rerun-never . every planned row is in the partial, not just the "
          "attempted ones", len(written), 40)
    never = [a for a, (st, _acc) in written.items()
             if st == M.NOT_ATTEMPTED]
    check("rerun-never . the rows the session never reached are named",
          len(never), 20)
    check_true("rerun-never . and the next session can therefore see them",
               all(str(i) in written for i in range(1, 41)),
               "missing: %s" % sorted(set(str(i) for i in range(1, 41))
                                      - set(written)))


def test_a_blocked_run_stops_instead_of_grinding():
    """Measured 2026-09-10: after the ceiling, 24 more minutes and 285 more
    requests were spent being told no. On an overnight run that is hours."""
    import inspect
    check_true("blocked . ServerRefusing IS a StopRequested",
               issubclass(M.ServerRefusing, M.StopRequested))
    check_true("blocked . so every existing stop path handles it",
               isinstance(M.ServerRefusing("x"), M.StopRequested))
    check_true("blocked . the stop threshold is higher than the report one",
               M.BLOCK_STOP_403 > M.BLOCK_RUN_403,
               "%s vs %s" % (M.BLOCK_STOP_403, M.BLOCK_RUN_403))

    body = inspect.getsource(M.download_worker)
    code = "\n".join(l for l in body.splitlines()
                     if not l.strip().startswith("#"))
    check_true("blocked . the worker raises it on a long run of refusals",
               "raise ServerRefusing(" in code, code[:200])
    check_true("blocked . gated on the run length and the control",
               "watch.should_stop(" in code
               and 'opts.get("stop_when_blocked"' in code, code[:200])
    check_true("blocked . the control defaults to on",
               'opts.get("stop_when_blocked", True)' in code, code[:200])
    check_true("blocked . the summary says the server refused, not that "
               "the user stopped",
               "the server is refusing" in body, body[-2500:])
    check_true("blocked . and says the rest are resumable",
               "resumable" in body, body[-2500:])

    # The page must quote the code's threshold, not a hand-typed one.
    page = M.build_page(
        M.load_session(Path("/nonexistent/session.json"))).decode("utf-8")
    check_true("blocked . the page offers the control", 'id="stopblocked"' in page)
    check_true("blocked . checked by default",
               'id="stopblocked" checked' in page, page[:0])
    check_true("blocked . and quotes the real threshold",
               str(M.BLOCK_STOP_403) in page)
    check_true("blocked . no token left unreplaced", "__BLOCKSTOP__" not in page)


def test_inventory_mode_records_without_fetching():
    """The ceiling makes a full download impossible; it does not make a full
    INVENTORY impossible, because admin pages are not what is rationed."""
    import inspect
    body = inspect.getsource(M.download_worker)
    code = "\n".join(l for l in body.splitlines()
                     if not l.strip().startswith("#"))

    check_true("inventory . the worker has a no-fetch path",
               'opts.get("inventory")' in code, code[:200])
    # The row must be written BEFORE any fetch, and skip the download gap:
    # waiting 5s to not make a request is the whole cost of the mode.
    inv = code[code.index('if opts.get("inventory"):'):]
    inv = inv[:inv.index("_throttle(\"file\")")]
    check_true("inventory . it records a row", "report_rows.append" in inv, inv[:300])
    check_true("inventory . with the INVENTORIED status",
               "INVENTORIED" in inv, inv[:300])
    check_true("inventory . and continues without fetching",
               "continue" in inv and "fetch_record_file" not in inv, inv[:300])
    check_true("inventory . the download gap is not spent",
               '_throttle("file")' not in inv, inv[:300])

    # It must carry the metadata that makes the report worth having.
    for field in ("job[\"access\"]", "job[\"release\"]", "job[\"embargo\"]",
                  "job[\"version_date\"]", "job[\"url\"]"):
        check_true("inventory . the row carries %s" % field.split('"')[1],
                   field in inv, inv[:400])

    # And the resulting report must re-run — that is the entire point.
    check_true("inventory . INVENTORIED is not an absence",
               not M.INVENTORIED.startswith(M.ABSENT_PREFIXES))
    check_true("inventory . nor mistaken for a download",
               not M.INVENTORIED.startswith("Downloaded"))
    src = Path(M.__file__).read_text(encoding="utf-8")
    rfr = src[src.index("def read_failed_rows("):src.index("def retry_native_url(")]
    check_true("inventory . so the re-run picks it up",
               "INVENTORIED" not in rfr, rfr[:400])

    # Messages must not call it a download.
    check_true("inventory . the structure line says inventoried",
               "file(s) inventoried" in body, body[-3000:])
    check_true("inventory . the summary says nothing was downloaded",
               "none downloaded" in body, body[-3000:])

    page = M.build_page(
        M.load_session(Path("/nonexistent/session.json"))).decode("utf-8")
    check_true("inventory . the page offers it", 'id="kinventory"' in page)
    check_true("inventory . off by default",
               'id="kinventory" checked' not in page)


def test_the_plan_workbook_carries_the_job_that_made_it():
    """The spec must travel with the evidence it describes.

    Otherwise session 5 is a fresh trip through the controls and can
    quietly be a different job from session 1 — native unticked,
    all-versions off — and nothing would notice. The report would simply
    show fewer files and look like progress.
    """
    import tempfile
    from openpyxl import load_workbook
    spec = {"primary": True, "supp": True, "native": True, "public": True,
            "hidden": True, "published": True, "unpublished": True,
            "embargo": True, "versions": "all"}
    counts = {"records": 3383, "files": 9800, "restricted": 773,
              "unknown": 3, "session_size": 132, "sessions": 75,
              "basis": "85% of the 156 file(s) ..."}
    rows = [tuple(str(i) for i in range(len(M.REPORT_HEADERS)))]
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "DC_FileDownload_Plan_etd_S.xlsx")
        M.write_plan_workbook(path, rows, spec, counts, _brand())
        wb = load_workbook(path)
        check("plan . two sheets, inventory and job",
              wb.sheetnames, ["Inventory", M.PLAN_SHEET])
        inv = wb["Inventory"]
        check("plan . the inventory keeps the report's headers",
              [c.value for c in inv[1]], M.REPORT_HEADERS)
        job = {r[0].value: r[1].value for r in wb[M.PLAN_SHEET].iter_rows()
               if r[0].value}
        # The one thing read back by code is the spec token.
        check("plan . the job spec round-trips",
              M.parse_plan_spec(job["Job spec"])["versions"], "all")
        for flag in ("primary", "supp", "native", "embargo"):
            check_true("plan . the spec keeps %s" % flag,
                       M.parse_plan_spec(job["Job spec"])[flag], job["Job spec"])
        # And the rest is for a person deciding how many nights this is.
        for label in ("Taken", "File kinds", "Versions", "Records",
                      "Files found", "Access-restricted",
                      "SESSION LIMITS — edit either value",
                      "Files per session (files)",
                      "Megabytes per session (MB)",
                      "Sessions needed",
                      "How the split was sized"):
            check_true("plan . the job sheet states %r" % label, label in job,
                       str(sorted(job))[:300])
        # Numbers stay numbers in the sheet — a librarian sorting or
        # summing this column should not meet text.
        check("plan . the counts are the inventory's, not invented",
              (job["Records"], job["Files found"],
               job["Access-restricted"]), (3383, 9800, 773))

    # Sizing: both limits apply and the smaller wins, because which unit
    # the server counts is not known.
    # Explicit limits: the defaults are 0 since v1.34, and sizing is what
    # happens when an operator sets one.
    lim = M.session_limits(chosen={"max_files": 134, "max_mb": 375})
    size, sessions, basis = M.suggest_session_size(9800, lim, avg_mb=2.64)
    by_files = M.limit_values(lim)["max_files"]
    by_mb = int(M.limit_values(lim)["max_mb"] / 2.64)
    check("plan . the smaller of the two limits wins",
          size, min(by_files, by_mb))
    check_true("plan . and the basis names both, not just the winner",
               "file(s) per session" in basis or "MB per session" in basis,
               basis)
    check_true("plan . it says why there are two",
               "not known" in basis, basis)
    check("plan . the session count covers the whole job",
          sessions, -(-9800 // size))

    # Big files, so volume binds first; small files, so the count does.
    big, _, b_basis = M.suggest_session_size(9800, lim, avg_mb=20.0)
    small, _, s_basis = M.suggest_session_size(9800, lim, avg_mb=0.85)
    check_true("plan . large files make the MB limit bind",
               big < by_files and "MB per session" in b_basis, b_basis)
    check("plan . small files make the file limit bind", small, by_files)

    # With no measured average the MB limit cannot be turned into rows, and
    # the plan must say that rather than implying it was honoured.
    n_size, n_sessions, n_basis = M.suggest_session_size(
        50, [{"key": "max_mb", "value": 375, "label": "", "unit": "MB",
              "basis": ""}], avg_mb=0.0)
    check("plan . an unmeasurable limit does not silently split", n_size, 0)
    check_true("plan . and the plan says the run still enforces it",
               "during the run" in n_basis, n_basis)

    # The plan workbook must re-run like any other report, or the whole
    # workflow stops at step 2.
    src = Path(M.__file__).read_text(encoding="utf-8")
    rfr = src[src.index("def read_failed_rows("):src.index("def retry_native_url(")]
    check_true("plan . a folder picks the plan workbook up",
               "DC_FileDownload_Plan_*.xlsx" in rfr, rfr[:700])

    body = inspect.getsource(M.download_worker)
    check_true("plan . an inventory run writes a plan, not a bare report",
               "write_plan_workbook(" in body, body[:200])


def test_a_session_takes_its_share_and_says_what_is_left():
    rows = [{"article": str(i)} for i in range(1, 11)]
    take, left = M.session_slice(rows, 4)
    check("cap . a session takes its share", len(take), 4)
    check("cap . and says what remains", left, 6)
    check("cap . taken from the front, in order",
          [r["article"] for r in take], ["1", "2", "3", "4"])

    take, left = M.session_slice(rows, 0)
    check("cap . 0 means everything", (len(take), left), (10, 0))
    take, left = M.session_slice(rows, -5)
    check("cap . a negative cap is not a trap either", (len(take), left),
          (10, 0))
    take, left = M.session_slice(rows, 99)
    check("cap . a cap larger than the work takes it all",
          (len(take), left), (10, 0))
    check("cap . and an empty plan is not an error",
          M.session_slice([], 10), ([], 0))
    # It must not hand back the caller's own list to mutate.
    take, _ = M.session_slice(rows, 0)
    take.append({"article": "x"})
    check("cap . the slice is a copy", len(rows), 10)

    import inspect
    body = inspect.getsource(M.retry_worker)
    code = "\n".join(l for l in body.splitlines()
                     if not l.strip().startswith("#"))
    check_true("cap . the retry worker slices before it counts",
               code.index("session_slice(") < code.index("nonlocal_total"),
               "the slice must come first or the totals describe the plan")
    check_true("cap . it reports what is left in the PLAN, not the session",
               "still outstanding in this plan" in body, body[-2000:])
    check_true("cap . and says so before fetching, not only after",
               "remain after this one" in body, body[:4000])

    page = M.build_page(
        M.load_session(Path("/nonexistent/session.json"))).decode("utf-8")
    check_true("cap . the page offers the file control",
               'id="sessioncap"' in page)
    check_true("cap . and the volume control beside it",
               'id="sessionmb"' in page)
    # The default is no longer 0. A default of "no limit" put every operator
    # one forgotten box away from a run that spends the day's allowance and
    # ends on a wall of refusals — which is what happened on 2026-09-10 and
    # 09-11. The page now opens with the measured figures in both boxes.
    check_true("cap . the file box opens at the measured default",
               'id="sessioncap" value="{}"'.format(M.SESSION_LIMIT_FILES)
               in page.replace("\n", " ").replace("  ", " ")
               or 'value="{}"'.format(M.SESSION_LIMIT_FILES) in page,
               "file default is not %s" % M.SESSION_LIMIT_FILES)
    check_true("cap . the volume box opens at the measured default",
               'value="{}"'.format(M.SESSION_LIMIT_MB) in page,
               "MB default is not %s" % M.SESSION_LIMIT_MB)
    # A quoted number must come from the constant, never be typed twice.
    check_true("cap . the page never quotes a limit the code does not use",
               "__LIMFILES__" not in page and "__LIMMB__" not in page,
               "an unreplaced token reached the served page")


def test_sessions_do_not_redo_each_others_work():
    """A long job is worked in sessions because DC caps files per day.

    After four nights the folder holds four reports, and "what is still
    outstanding?" is answered by all four together. Without the merge,
    session 4 re-requests what session 3 downloaded — on a capped
    allowance that is the whole day spent on files already on disk, and it
    is invisible, because every one of those requests succeeds.
    """
    def frow(article, kind, status, url=None, version="current"):
        return {"context": "etd", "article": article, "kind": kind,
                "version": version, "version_date": "", "was": status,
                "url": url or "https://h/f/%s/%s" % (article, kind),
                "adm": "https://h/cgi/editor.cgi?article=%s&context=etd" % article,
                "plan": {}, "replan": False}

    def rrow(article, plan="primary,versions=current"):
        return {"context": "etd", "article": article, "kind": M.RECORD_KIND,
                "version": "", "version_date": "", "was": M.NOT_ATTEMPTED,
                "url": "https://h/cgi/editor.cgi?article=%s&context=etd" % article,
                "adm": "https://h/cgi/editor.cgi?article=%s&context=etd" % article,
                "plan": {"primary": True}, "replan": True}

    # A file failed in session 1 and downloaded in session 2.
    merged = M.merge_rows([
        ("s1", [frow("1", "primary", "Failed: HTTP 403 Forbidden")]),
        ("s2", [frow("1", "primary", "Downloaded")]),
    ])
    check("merge . the later status wins", [r["was"] for r in merged],
          ["Downloaded"])

    # And the other way round: a success then a later failure is a failure.
    merged = M.merge_rows([
        ("s1", [frow("1", "primary", "Downloaded")]),
        ("s2", [frow("1", "primary", "Failed: HTTP 403 Forbidden")]),
    ])
    check("merge . order is chronological, not preferential",
          [r["was"] for r in merged], ["Failed: HTTP 403 Forbidden"])

    # Different files of the same record are different rows.
    merged = M.merge_rows([
        ("s1", [frow("1", "primary", "Downloaded"),
                frow("1", "native", "Failed: HTTP 403 Forbidden")]),
    ])
    check("merge . kind is part of a file's identity", len(merged), 2)
    merged = M.merge_rows([
        ("s1", [frow("1", "primary", "Downloaded"),
                frow("1", "primary", "Downloaded", version="original")]),
    ])
    check("merge . so is the version", len(merged), 2)

    # THE ONE URL-KEYING ALONE WOULD MISS. A record row is superseded by
    # the FILE rows that replanning it produced — whose URLs are different
    # by construction, because the record row's URL is an admin page.
    merged = M.merge_rows([
        ("s1", [rrow("7")]),
        ("s2", [frow("7", "primary", "Downloaded"),
                frow("7", "native", "Downloaded")]),
    ])
    kinds = sorted(r["kind"] for r in merged)
    check("merge . a replanned record row is superseded by its files",
          kinds, ["native", "primary"])
    check_true("merge . and does not survive to be replanned again",
               M.RECORD_KIND not in kinds, str(kinds))

    # But a record nobody has reached yet must survive.
    merged = M.merge_rows([
        ("s1", [rrow("7"), rrow("8")]),
        ("s2", [frow("7", "primary", "Downloaded")]),
    ])
    arts = sorted((r["article"], r["kind"]) for r in merged)
    check("merge . an un-reached record still survives",
          arts, [("7", "primary"), ("8", M.RECORD_KIND)])

    # Three sessions, the realistic shape.
    merged = M.merge_rows([
        ("s1", [frow("1", "primary", "Downloaded"),
                frow("2", "primary", "Failed: HTTP 403 Forbidden"),
                rrow("3")]),
        ("s2", [frow("2", "primary", "Downloaded"),
                frow("3", "primary", "Failed: HTTP 403 Forbidden")]),
        ("s3", [frow("3", "primary", "Downloaded")]),
    ])
    check("merge . three sessions leave nothing outstanding",
          sorted((r["article"], r["was"]) for r in merged),
          [("1", "Downloaded"), ("2", "Downloaded"), ("3", "Downloaded")])

    # Reports are read oldest first, by the stamp in the filename.
    names = ["DC_FileDownload_Retry_etd_20260911_090000.xlsx",
             "DC_FileDownload_Report_etd_20260910_153905.xlsx",
             "DC_FileDownload_Retry_etd_20260910_170440.xlsx"]
    check("merge . reports sort chronologically by their stamp",
          [n.split("_")[-2] for n in sorted(names, key=M.report_sort_key)],
          ["20260910", "20260910", "20260911"])

    # A real folder, because the source greps for this were both wrong:
    # "DC_FileDownload_Retry_*.xlsx" also appears in the error message, and
    # "key=report_sort_key" still matches when reverse=True is appended to
    # it. Two more checks reading NEAR the thing instead of at it.
    import tempfile
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article=7&context=etd"
    url = base + "/cgi/viewcontent.cgi?article=7&context=etd"
    with tempfile.TemporaryDirectory() as d:
        _fake_report(d, "DC_FileDownload_Report_etd_20260910_120000.xlsx",
                     [_row(1, "7", "primary", "public", "current",
                           "Failed: HTTP 403 Forbidden", url, adm)])
        _fake_report(d, "DC_FileDownload_Retry_etd_20260911_120000.xlsx",
                     [_row(1, "7", "primary", "public", "current",
                           "Downloaded", url, adm,
                           orig="a.pdf", saved="etd_7.pdf", size=10)])
        got, skipped, _absent, _known = M.read_failed_rows(d)
        check("merge . a file recovered in a later session is not re-asked",
              [g["article"] for g in got], [])
        check("merge . and nothing is reported as unaddressable", skipped, [])

        # The reverse case proves the ordering is real and not an accident
        # of which file happens to be read second.
        _fake_report(d, "DC_FileDownload_Retry_etd_20260912_120000.xlsx",
                     [_row(1, "7", "primary", "public", "current",
                           "Failed: HTTP 403 Forbidden", url, adm)])
        got, _s, _a, _known = M.read_failed_rows(d)
        check("merge . a file that failed again IS outstanding",
              [g["article"] for g in got], ["7"])

    src = Path(M.__file__).read_text(encoding="utf-8")
    rfr = src[src.index("def read_failed_rows("):src.index("def retry_native_url(")]
    check_true("merge . the rows go through merge_rows",
               "merge_rows(sources)" in rfr, rfr[:1500])


def _stop_handler_at(body):
    """Index of the StopRequested handler, however it is spelled.

    It became `except StopRequested as _stop:` in v1.31 and two tests that
    hard-coded the bare form broke — a check coupled to punctuation rather
    than to the thing it tests.
    """
    for form in ("except StopRequested as _stop:", "except StopRequested:"):
        if form in body:
            return body.index(form)
    raise AssertionError("no StopRequested handler found")


def test_no_name_is_used_without_being_bound():
    """The defect source-slicing produces, and it produced it again.

    A replacement anchored on a comment moved `blocked = ...` into
    retry_worker while download_worker went on USING `blocked` — a
    NameError on the next Stop, in the handler whose whole job is not
    losing work. Compiling does not catch it and neither did 667 tests.

    So: walk the AST of the long workers and check every plain name that
    is read is also bound somewhere in that function, or is a global.
    """
    import ast
    import builtins
    src = Path(M.__file__).read_text(encoding="utf-8")
    tree = ast.parse(src)
    module_names = {n.id for node in tree.body
                    if isinstance(node, ast.Assign)
                    for t in [node.targets[0]] if isinstance(t, ast.Name)
                    for n in [t]}
    module_names |= {n.name for n in tree.body
                     if isinstance(n, (ast.FunctionDef, ast.ClassDef))}
    module_names |= {a.asname or a.name.split(".")[0]
                     for n in ast.walk(tree) if isinstance(n, ast.Import)
                     for a in n.names}
    module_names |= {a.asname or a.name
                     for n in ast.walk(tree)
                     if isinstance(n, ast.ImportFrom) for a in n.names}
    module_names |= set(dir(builtins))

    for fn in ("download_worker", "retry_worker", "map_worker"):
        node = next(n for n in tree.body
                    if isinstance(n, ast.FunctionDef) and n.name == fn)
        bound = {a.arg for a in node.args.args}
        for sub in ast.walk(node):
            if isinstance(sub, ast.Name) and isinstance(sub.ctx, ast.Store):
                bound.add(sub.id)
            elif isinstance(sub, ast.ExceptHandler) and sub.name:
                bound.add(sub.name)
            elif isinstance(sub, (ast.FunctionDef, ast.AsyncFunctionDef)):
                bound.add(sub.name)
                bound |= {a.arg for a in sub.args.args}
            elif isinstance(sub, ast.Lambda):
                bound |= {a.arg for a in sub.args.args}
            elif isinstance(sub, ast.Import):
                bound |= {a.asname or a.name.split(".")[0] for a in sub.names}
            elif isinstance(sub, ast.ImportFrom):
                bound |= {a.asname or a.name for a in sub.names}
            elif isinstance(sub, ast.comprehension):
                for t in ast.walk(sub.target):
                    if isinstance(t, ast.Name):
                        bound.add(t.id)
        read = {n.id for n in ast.walk(node)
                if isinstance(n, ast.Name) and isinstance(n.ctx, ast.Load)}
        unbound = sorted(read - bound - module_names)
        check("bound . every name %s reads is bound" % fn, unbound, [])


def test_a_stopped_rerun_writes_what_it_had():
    """retry_worker's stop handler wrote no report at all.

    download_worker was given flush_pending in v1.20 for exactly this — a
    stop at row 33 of 38 left 33 files on disk and nothing recording where
    they came from. retry_worker never got the same treatment, and making
    the re-run long enough that someone would want to stop it is what
    exposed the gap.
    """
    import tempfile
    brand = _brand()
    row = tuple(str(i) for i in range(len(M.REPORT_HEADERS)))
    with tempfile.TemporaryDirectory() as d:
        name = M._flush_retry({"ctx": "etd", "dir": d, "rows": [row]},
                              "20260909_120000", brand)
        check_true("retry-stop . a partial retry report is written",
                   bool(name), "nothing was written")
        check_true("retry-stop . named _PARTIAL so it reads as incomplete",
                   name.endswith("_PARTIAL.xlsx"), name)
        check_true("retry-stop . and it is on disk",
                   os.path.isfile(os.path.join(d, name)), name)
        check_true("retry-stop . the context is in the filename",
                   "etd" in name, name)

        # Nothing gathered yet is not an error, and must not fabricate a file.
        check("retry-stop . no rows means no report",
              M._flush_retry({"ctx": "etd", "dir": d, "rows": []},
                             "s", brand), "")
        check("retry-stop . and no pending at all is handled",
              M._flush_retry(None, "s", brand), "")

        # A failed write must not swallow the exception that brought us here.
        bad = M._flush_retry(
            {"ctx": "etd", "dir": os.path.join(d, "nope"), "rows": [row]},
            "s", brand)
        check("retry-stop . a failed write reports itself rather than raising",
              bad, "")

    src = Path(M.__file__).read_text(encoding="utf-8")
    body = src[src.index("def retry_worker("):]
    # Slice the TWO handlers apart. Slicing from "except StopRequested:" to
    # the end of the function covers the except-Exception block as well, so
    # gutting the stop handler still found a _flush_retry below it and the
    # check passed. Caught by planting; the third check of this shape today.
    tail = body[_stop_handler_at(body):]
    split = tail.index("except Exception as e:")
    stop, err = tail[:split], tail[split:]
    check_true("retry-stop . the stop handler flushes",
               "_flush_retry(" in stop, stop[:400])
    check_true("retry-stop . the unexpected-error path flushes too",
               "_flush_retry(" in err, err[:400])
    check_true("retry-stop . and the summary names the partial report",
               "Partial report" in stop, stop[:600])


def test_the_rerun_plans_as_it_goes():
    """Source-level, and labelled as such.

    The first cut built the whole work list before fetching anything: for
    75 un-reached records that is 150 throttled requests, about twelve
    minutes, with the progress bar frozen on "Attaching to Chrome…", no
    file written, and a Stop discarding all of it. download_worker plans
    inside its loop; hoisting it out of the retry loop reintroduced a
    problem this module had already solved.

    This cannot be exercised without Chrome and a live session, so it
    checks shape, not behavior. Recorded as half a check.
    """
    src = Path(M.__file__).read_text(encoding="utf-8")
    body = src[src.index("def retry_worker("):]
    body = body[:body.index("PAGE = r")]
    check_true("retry-lazy . planning is a generator",
               "def expand(rows):" in body and "yield" in body, body[:200])
    check_true("retry-lazy . the loop consumes it directly",
               "enumerate(expand(ctx_rows), 1)" in body, body[:200])
    check_true("retry-lazy . no work list is materialized first",
               "work = []" not in body, "the list is being built up front")
    check_true("retry-lazy . the planner is called inside the generator",
               body.index("def expand(rows):")
               < body.index("plan_item_jobs(")
               < body.index("enumerate(expand("), "planning moved out again")
    # The row total still grows as records become files — it is what the
    # summary counts — but since v1.34 PROGRESS reads records, whose
    # number the report fixes: an ETA from a growing denominator meant
    # nothing (887 → 1,485 rows on 2026-09-28).
    check_true("retry-lazy . the row total grows as records expand",
               "nonlocal_total[0] +=" in body)
    check_true("retry-lazy . and progress reads records, not the growing "
               "row total", "show_progress(ctx)" in body
               and "set_progress(done, nonlocal_total[0]" not in body,
               body[:200])


def test_two_simultaneous_starts_produce_one_run():
    """2026-09-09: a re-run was started twice and two workers ran at once.

    Two copies of every log line, both writing into the same run folder,
    and between them hitting a server measured to allow one download every
    8.8 seconds at roughly 2.5s intervals.

    This races the real endpoint rather than reading the source, because
    the defect WAS invisible to reading: the check was there, it was
    honest, and it was stale by the time it was used. Every worker sets
    phase="starting" as its first statement, inside the new thread, so
    between the endpoint answering and the thread being scheduled the
    phase still reads "idle".

    On the retry path that window is not microseconds: read_failed_rows()
    parses an openpyxl workbook 27 lines before phase is set, so it is as
    wide as loading the report takes. Two deliberate clicks a second apart
    were enough, which is exactly how it was found.
    """
    import inspect
    import threading as _t
    import time
    session = M.load_session(Path("/nonexistent/session.json"))
    session["hub_url"] = "http://127.0.0.1:8750"
    srv = ThreadingHTTPServer(
        ("127.0.0.1", 0), M.make_handler(session, M.build_page(session)))
    port = srv.server_address[1]
    _t.Thread(target=srv.serve_forever, daemon=True).start()

    started = []
    gate = _t.Event()

    def slow_worker(*_a, **_k):
        started.append(1)
        gate.wait(5.0)          # hold the slot open across both requests

    real = M.map_worker
    M.map_worker = slow_worker
    try:
        with tempfile.TemporaryDirectory() as d:
            results = []
            lock = _t.Lock()

            def fire():
                code, body = _post(port, "/api/map", {"out_dir": d})
                with lock:
                    results.append(code)

            # Fire together, the way a double-click does.
            threads = [_t.Thread(target=fire) for _ in range(6)]
            for t in threads:
                t.start()
            for t in threads:
                t.join(10)
            gate.set()
            time.sleep(0.4)

            check("one-run . exactly one request is accepted",
                  results.count(200), 1)
            check("one-run . the others are refused with 409",
                  results.count(409), len(results) - 1)
            check("one-run . and exactly one worker ever started",
                  len(started), 1)
    finally:
        M.map_worker = real
        gate.set()
        srv.shutdown()

    # The slot must be released even by a worker that dies immediately,
    # or the module wedges until it is restarted.
    check_true("one-run . the claim is released in a finally",
               "finally:" in inspect.getsource(M.make_handler)
               and 'STATE["claimed"] = False' in
               inspect.getsource(M.make_handler))
    check_true("one-run . and the guard reads the claim, not just the phase",
               'STATE["claimed"]' in inspect.getsource(M.make_handler))

    # Client side: one slot, so any start button locks all of them.
    page = M.build_page(session).decode("utf-8")
    check_true("one-run . the page defines a single lock for all starts",
               "function lockStarts()" in page)
    # Two traps here, both of which this check fell into first time.
    #  * Searching near the endpoint string matched `function lockStarts(){`
    #    — the definition satisfying a test meant for the call. It passed
    #    with the fix removed.
    #  * "go.addEventListener" is a substring of "retrygo.addEventListener",
    #    so the `go` case silently inspected the retrygo handler. Anchored
    #    on a newline so each button finds its own.
    for btn, endpoint in (("go", "/api/start"),
                          ("retrygo", "/api/retry"),
                          ("mapBtn", "/api/map")):
        start = page.find("\n" + btn + ".addEventListener")
        if start < 0:
            check_true("one-run . {} has its own click handler".format(btn),
                       False, "handler not found")
            continue
        body = page[start:page.index("});", start)]
        check_true("one-run . {} handler is the right one".format(btn),
                   endpoint in body, body[:200])
        call = body.find("lockStarts();")
        fetch = body.find("await api(")
        check_true("one-run . {} calls lockStarts() in its own handler"
                   .format(btn), call >= 0, body[:220])
        check_true("one-run . {} locks BEFORE it awaits".format(btn),
                   0 <= call < fetch, body[:220])
    check_true("one-run . nothing calls a repaint that does not exist",
               "paintSteps" not in page)


def test_the_rerun_control_is_findable_and_honestly_named():
    """Jeff asked how to fill in "1 . Hierarchy source" for a re-run.

    He asked because the fieldsets are numbered 1 to 4 and the re-run field
    was buried at the bottom of 4, so the page implied three mandatory steps
    in front of a path that needs none of them. /api/retry takes only a
    download folder, the report, and the pacing controls. A control someone
    cannot find, or believes is unavailable, is not shipped.
    """
    session = M.load_session(Path("/nonexistent/session.json"))
    page = M.build_page(session).decode("utf-8")

    def at(needle, label):
        """Index of `needle`, or None with a stated failure — never a raise.

        An assertion that dies with ValueError says nothing about what is
        wrong; this file has shipped two of those before.
        """
        i = page.find(needle)
        if i < 0:
            check_true("rerun-ui . the page contains {}".format(label), False,
                       "not found in the served page: " + repr(needle))
            return None
        return i

    # Until 2026-09-11 the re-run controls sat AFTER four numbered
    # sections, none of which a re-run uses, and the page said which ones
    # applied only in a hint at the very bottom. Jeff asked twice whether
    # to fill in a hierarchy source before re-running, which makes it the
    # page's fault. The mode is now the first thing asked.
    mode = at('id="sec0"', "the mode question")
    hier = at('id="sec1"', "the hierarchy fieldset")
    alt = at('id="secresume"', "the re-run fieldset")
    btn = at('id="retrygo"', "the re-run button")
    if None in (mode, hier, alt, btn):
        return
    check_true("rerun-ui . the mode is asked before anything else",
               mode < hier, "the numbered steps come first")
    check_true("rerun-ui . both modes are offered as a choice",
               'id="jobnew"' in page and 'id="jobresume"' in page)
    check_true("rerun-ui . and it sits before the buttons", alt < btn)
    check_true("rerun-ui . the field itself moved with it",
               'id="retryrpt"' in page[alt:btn], "field left behind")

    # Hiding is what answers "does this section apply to me?" — a hint at
    # the foot of the page did not, twice.
    #
    # RUN the function rather than grepping for it. The grep version of
    # this check looked for ".hidden=resume" and was satisfied by
    # "go.hidden=resume" on a different line — the same adjacent match as
    # "go.addEventListener" inside "retrygo.addEventListener". Gutting
    # applyJobMode's body left it green. So the page's own JavaScript is
    # executed against a stub document and the result is read off.
    _check_job_mode(page)

    block = page[alt:btn]

    # The name must match what the control does. It resumes records that
    # were never attempted, and so never failed; "Re-run failures" was a
    # button whose label asserted something untrue about half its job.
    import re
    m = re.search(r'id="retrygo"[^>]*>(.*?)</button>', page, re.S)
    label = re.sub(r"\s+", " ", m.group(1)).strip()
    check("rerun-ui . the button is named for what it does",
          label, "Re-run from report")
    check_true("rerun-ui . no control is still called 'Re-run failures'",
               "Re-run failures" not in page, "stale control name in page")
    check_true("rerun-ui . the instructions name the button that exists",
               label in block, "the hint names a button that is not there")
    check_true("rerun-ui . and it explains that absences are not re-asked",
               "not</em> re-requested" in block or
               "not re-requested" in block, block[-500:])


def test_a_stopped_run_accounts_for_every_row():
    """v1.21 wrote 110 rows and logged "24 file(s)".

    The other 86 — 11 files the server said do not exist, and 75 records
    never reached — were correct in the report and absent from the log.
    Someone reads the log, not the workbook.
    """
    import inspect
    body = inspect.getsource(M.download_worker)
    flush = body[body.index("def flush_pending"):
                 body.index("def finish_master")]
    for name, token in (("downloaded", '== "Downloaded"'),
                        ("not present", "ABSENT_PREFIXES"),
                        ("never reached", "NOT_ATTEMPTED")):
        check_true("stop-count . the flush counts {}".format(name),
                   token in flush, flush[:400])
    check_true("stop-count . and every row is accounted for, not just files",
               "len(rows) - done - gone - never" in flush, flush[:600])
    check_true("stop-count . the breakdown reaches the log",
               "breakdown" in flush and "log(" in flush)
    check_true("stop-count . and the master row, not only the log",
               flush.count("breakdown") >= 3, flush[-800:])
    # The summary must tell the reader the run can be resumed at all.
    handler = body[_stop_handler_at(body):]
    check_true("stop-count . the stop summary says how many were unreached",
               "never" in handler and "re-run this report" in handler,
               handler[:600])
    check_true("stop-count . the flush hands the count back for it",
               "never = flush_pending(" in handler, handler[:300])


def test_a_row_of_the_wrong_width_is_refused():
    import tempfile
    brand = _brand()
    good = [tuple(str(i) for i in range(len(M.REPORT_HEADERS)))]
    with tempfile.TemporaryDirectory() as d:
        # control: the well-formed row still writes
        try:
            M.write_structure_report(os.path.join(d, "ok.xlsx"), good, brand)
            check_true("arity . a well-formed report still writes", True)
        except Exception as e:
            check_true("arity . a well-formed report still writes", False,
                       repr(e))
        # PLANT THE VIOLATION, in both directions
        for label, n in (("short", -1), ("long", +1)):
            rows = [tuple(str(i) for i in
                          range(len(M.REPORT_HEADERS) + n))]
            try:
                M.write_structure_report(
                    os.path.join(d, label + ".xlsx"), rows, brand)
                check_true("arity . a {} row is refused".format(label),
                           False, "it was written")
            except ValueError as e:
                check_true("arity . a {} row is refused".format(label), True)
                check_true("arity . the {} failure names the row and the "
                           "counts".format(label),
                           "row 1" in str(e) and "column" in str(e), str(e))
            except Exception as e:
                check_true("arity . a {} row is refused with a clear "
                           "error".format(label), False, repr(e))
        # the master log is guarded the same way
        try:
            M.write_master_workbook(
                os.path.join(d, "m.xlsx"),
                [tuple(range(len(M.MASTER_HEADERS) - 1))], brand)
            check_true("arity . a short master row is refused", False,
                       "it was written")
        except ValueError:
            check_true("arity . a short master row is refused", True)

    check("arity . headers and widths agree, report",
          len(M.REPORT_HEADERS), len(M.REPORT_WIDTHS))
    check("arity . headers and widths agree, master",
          len(M.MASTER_HEADERS), len(M.MASTER_WIDTHS))
    check("arity . COL_ACCESS is derived, never typed",
          M.COL_ACCESS, M.REPORT_HEADERS.index("Access"))


# ---------------------------------------------------------------------------
# v1.33 — the browser fetch path and the verification prompt
#
# Behavioral wherever it can be: a fake driver that actually writes files
# into an incoming folder exercises the real DirectoryDownloadWatch, the
# real BrowserFetcher and the real fetch_one, so "a download completed" is
# a file transitioning on disk rather than a string found in the source.
# ---------------------------------------------------------------------------
class _Tab:
    """One scripted navigation outcome for the fake driver."""

    def __init__(self, title="", source="", drop=None, partial=None,
                 url="", stall=False, drop_on_clear=None):
        self.title = title
        self.source = source
        self.drop = drop          # (name, bytes) written as a completed file
        self.partial = partial    # name left as a stranded .crdownload
        self.url = url
        # A navigation that times out WITHOUT delivering a file — which
        # is what happens when the interstitial is slow: Chrome has not
        # finished loading a page and no download has begun. This is
        # the path on which a prompt can only be noticed by watching,
        # and it is the path the first live run actually took.
        self.stall = stall
        # The file that lands at the moment somebody clears the prompt,
        # because clearing returns the browser to the URL it asked for.
        self.drop_on_clear = drop_on_clear
        # (seconds, name, bytes): a person clears the prompt after that
        # long and the file follows — ON A TIMER, not on the module
        # looking. That distinction is the whole point: a real person
        # clears a prompt whether or not the run glanced at the tab,
        # and a fixture in which clearing is TRIGGERED by looking makes
        # "we stopped looking" fail for the wrong reason.
        self.timed_clear = None
        # (starts_after, finishes_after, name, bytes): a transfer that
        # BEGINS at one moment and COMPLETES at a later one, with the
        # tab frozen on the challenge page throughout. Without this the
        # fake's downloads are instantaneous, `settle` finds the file on
        # its first look and never glances at the tab at all — which is
        # why the off-campus run of 2026-09-18 recorded a prompt the
        # suite could not reproduce.
        self.timed_transfer = None
        # (seconds, title, source): a page that goes on screen AFTER the
        # navigation has already timed out — a new document, served for
        # this request, just late. This is the only legitimate way for
        # the post-wait branch of `_await` to meet an absence page, and
        # before v1.34 the suite reached that branch only through a tab
        # PARKED on an absence page from an earlier request, which is
        # the stale-tab defect itself dressed as a fixture.
        self.late_render = None
        # (seconds, title, source): the SERVER takes this long to answer
        # with a page (v1.34.2). Chromedriver stops the load when the
        # page-load timeout set for this navigation expires, so the page
        # renders only when that allowance is at least `seconds`;
        # otherwise the navigation stalls with nothing on screen. That is
        # the missing native measured on 2026-10-01: 404 at 15.2 s
        # against a 15.0 s allowance, and a fixture that rendered
        # regardless of the allowance could not tell the two apart.
        self.answers_after = None


class _FakeBrowser:
    """A driver that behaves the way Chrome does on this path.

    It writes into the download folder the module hands it, completes a
    download by renaming a .crdownload, raises the page-load timeout that
    a real download always raises, and remembers every CDP call so the
    restore can be asserted.
    """

    class TimeoutException(Exception):
        pass

    def __init__(self, script, clock=None, clear_after=None,
                 clear_after_reads=None, at=None):
        self.script = list(script)
        self.visited = []        # file requests only
        self.navigations = []    # EVERY get(), including the site root.
                                 # `visited` filters to viewcontent, so a
                                 # check that "nothing was navigated"
                                 # could not see a root navigation at all.
        self.cdp = []
        self.timeout_calls = []
        self.folder = ""
        # WHERE THE TAB ALREADY IS. A real tab is never blank: it is
        # sitting on whatever the last navigation left on screen, and a
        # download leaves it there untouched. A fake that starts blank
        # makes every navigation look like a move, which is precisely
        # the comparison the module uses to decide what the tab is
        # evidence of — so a blank start hides every defect that only
        # appears when the tab lands where it already was. Measured
        # 2026-09-21: 51 such rows in one run, none reproducible here.
        at_title, at_url, at_source = at or ("", "", "")
        self._title = at_title
        self._source = at_source
        self.current_url = at_url
        self._raise_on_restore = False
        # A person clearing the prompt. Without this the tab still reads
        # as a challenge for as long as the hold watches it, which is a
        # test that waits out the whole window rather than one that
        # exercises the clearance.
        self.clock = clock
        self.clear_after = clear_after
        self._asked_at = None
        # Clearing after N reads of the tab, for the worker-level tests.
        # They run the REAL hold, which sleeps in real seconds and cannot
        # be handed a fake clock, so "a person cleared it" has to be
        # expressed in something the hold actually does: looking.
        self.clear_after_reads = clear_after_reads
        self.reads = 0
        # WHICH DOCUMENT IS ON SCREEN. Bumped only when a page is
        # actually rendered, never for a download or a stall — which is
        # the distinction the module now reads, and which a fake that
        # only tracks title and URL cannot express at all. Two
        # successive renders of the SAME page are two documents, and
        # that is the whole case this exists for.
        self._doc = 1
        # A browser that will not answer the probe, so the fallback to
        # the appearance test has a check of its own.
        self._no_script = False
        self._pending_drop = None
        self._timed_cleared = False
        self._timers = []
        # v1.34: "navigation accepted, nothing downloads, nothing renders"
        # — what a Chrome with downloads blocked by policy looks like.
        # Every scripted tab before this dropped a file, stalled or
        # rendered, so the preflight's failure case was inexpressible.
        self.downloads_blocked = False
        self.preflights = []

    def _cleared(self):
        if self._timed_cleared:
            return True
        if self._asked_at is None and self.clear_after_reads is None:
            return False
        if self.clear_after_reads is not None:
            done = self.reads >= self.clear_after_reads
            if done and self._pending_drop:
                # Clearing the prompt is what starts the download.
                name, data = self._pending_drop
                self._pending_drop = None
                tmp = os.path.join(self.folder, name + ".crdownload")
                with open(tmp, "wb") as fh:
                    fh.write(data)
                os.replace(tmp, os.path.join(self.folder, name))
            return done
        return (self.clear_after is not None and self.clock is not None
                and self._asked_at is not None
                and self.clock.t - self._asked_at >= self.clear_after)

    @property
    def title(self):
        cleared = self._cleared()
        self.reads += 1
        return "Digital Commons" if cleared else self._title

    @property
    def page_source(self):
        return "<html>ok</html>" if self._cleared() else self._source

    def execute_script(self, js, *args):
        """Only the document-identity probe is scripted here. Anything
        else raises, so a future caller cannot quietly get a stub answer
        and build a claim on it."""
        if "performance.timeOrigin" in js:
            if self._no_script:
                raise RuntimeError("javascript is unavailable")
            return "{}|{}".format(self._doc, self.current_url)
        raise RuntimeError("the fake does not run: {!r}".format(js))

    def execute_cdp_cmd(self, cmd, params):
        self.cdp.append((cmd, dict(params)))
        if params.get("behavior") == "allow":
            self.folder = params["downloadPath"]
        if params.get("behavior") == "default" and self._raise_on_restore:
            raise RuntimeError("devtools went away")
        return {}

    class _Timeouts:
        # Selenium exposes the live timeouts here. The fake reports a
        # distinctive 42 so a test can tell "put back what was there"
        # apart from "put back the module's constant" — otherwise both
        # are 300 and the difference goes unnoticed.
        page_load = 42.0

    timeouts = _Timeouts()

    def set_page_load_timeout(self, seconds):
        self.timeout_calls.append(seconds)

    def get(self, url):
        # v1.34: the download-path preflight. A download never moves the
        # tab, so it changes nothing on screen — and when downloads are
        # BLOCKED (the failure the preflight exists for) the navigation
        # is accepted and nothing arrives at all. Kept out of
        # `navigations`, which older tests count.
        if M.PREFLIGHT_PATH in url:
            self.preflights.append(url)
            if not self.downloads_blocked and self.folder:
                name = url.rsplit("/", 1)[-1]
                tmp = os.path.join(self.folder, name + ".crdownload")
                with open(tmp, "wb") as fh:
                    fh.write(M.PREFLIGHT_BYTES)
                os.replace(tmp, os.path.join(self.folder, name))
            raise _FakeBrowser.TimeoutException("page load timeout")
        self.navigations.append(url)
        # The workers navigate the site root once to prove the session is
        # live. That is not a file request and must not consume a scripted
        # outcome — the first version of this fake let it eat the first
        # tab, and every assertion after it was about the wrong
        # navigation.
        # "Looks at this prompt", not "looks at anything ever": the
        # counter is per-navigation. Without the reset, one glance at
        # the session-validating page left the fake reporting every
        # later prompt as already cleared — and the worker test then
        # passed while recording no prompt at all.
        self.reads = 0
        if "viewcontent" not in url:
            self._title = "Digital Commons"
            self._source = "<html><body>admin</body></html>"
            self.current_url = url
            self._asked_at = None
            return
        self.visited.append(url)
        tab = self.script.pop(0) if self.script else _Tab(title="nothing")
        if tab.answers_after:
            after, a_title, a_source = tab.answers_after
            allowed = self.timeout_calls[-1] if self.timeout_calls else 300
            if allowed >= after:
                tab.title, tab.source = a_title, a_source
            else:
                tab.stall = True
        # **A DOWNLOAD DOES NOT NAVIGATE THE TAB.** Chrome hands the
        # bytes to its own network stack and the page stays exactly
        # where it was. The first version of this fake moved the tab on
        # every get(), which is why 1,189 tests passed while the live
        # module read a frozen Cloudflare page as a fresh prompt 145
        # times. A fixture that behaves better than the world hides
        # exactly the defects the world produces.
        if not tab.title and not tab.source:
            # Nothing rendered, so nothing changed. True of a download
            # (Chrome takes the bytes and leaves the page alone) and of
            # a navigation that stalls. Only a tab that names a page
            # replaces what is on screen.
            pass
        else:
            self._title = tab.title
            self._source = tab.source
            self.current_url = tab.url or url
            # A rendered page is a NEW document even when it is the same
            # page — that is exactly the live case: redirect to the home
            # page, from the home page.
            self._doc += 1
        body = (tab.title + " " + tab.source).lower()
        self._asked_at = (self.clock.t if self.clock is not None
                          and any(m in body for m in M.CHALLENGE_MARKERS)
                          else None)
        if tab.drop_on_clear:
            self._pending_drop = tab.drop_on_clear
        if tab.timed_transfer:
            import threading as _th
            begins, ends, name, data = tab.timed_transfer
            part = os.path.join(self.folder, name + ".crdownload")

            def _begin():
                try:
                    with open(part, "wb") as fh:
                        fh.write(data[:1])
                except OSError:
                    pass

            def _end():
                try:
                    with open(part, "wb") as fh:
                        fh.write(data)
                    os.replace(part, os.path.join(self.folder, name))
                except OSError:
                    pass

            for delay, fn in ((begins, _begin), (ends, _end)):
                t = _th.Timer(delay, fn)
                t.daemon = True
                t.start()
                self._timers.append(t)
        if tab.late_render:
            import threading as _th
            delay, l_title, l_source = tab.late_render
            l_url = tab.url or url

            def _late():
                self._title = l_title
                self._source = l_source
                self.current_url = l_url
                self._doc += 1

            t = _th.Timer(delay, _late)
            t.daemon = True
            t.start()
            self._timers.append(t)
        if tab.timed_clear:
            import threading as _th
            delay, name, data = tab.timed_clear

            def _clear():
                # The person clears the prompt and the download starts —
                # **but the tab stays on the Cloudflare page**, because
                # a download never navigates it. Setting _timed_cleared
                # here made the fake tidy the tab up by itself, which is
                # the one thing the real browser does not do, and it is
                # what hid the defect that produced 145 fictional
                # prompts. Guarded: the thread can outlive the temporary
                # directory the test built, and a traceback from a fake
                # reads like a broken module.
                try:
                    tmp = os.path.join(self.folder, name + ".crdownload")
                    with open(tmp, "wb") as fh:
                        fh.write(data)
                    os.replace(tmp, os.path.join(self.folder, name))
                except OSError:
                    pass

            t = _th.Timer(delay, _clear)
            t.daemon = True
            t.start()
            self._timers.append(t)
        if tab.stall:
            # No page, no file — just a navigation that ran out of
            # patience, which is exactly what a slow interstitial looks
            # like from here.
            raise _FakeBrowser.TimeoutException("page load timeout")
        if tab.partial:
            with open(os.path.join(self.folder, tab.partial + ".crdownload"),
                      "wb") as fh:
                fh.write(b"half")
        if tab.drop:
            name, data = tab.drop
            tmp = os.path.join(self.folder, name + ".crdownload")
            with open(tmp, "wb") as fh:
                fh.write(data)
            os.replace(tmp, os.path.join(self.folder, name))
            # A real download never finishes its navigation.
            raise _FakeBrowser.TimeoutException("page load timeout")


def _browser_fetcher(folder, script, clock=None, clear_after=None,
                     clear_after_reads=None, at=None):
    drv = _FakeBrowser(script, clock, clear_after, clear_after_reads, at)
    f = M.BrowserFetcher(folder)
    # Short waits, so a suite of "nothing arrived" cases does not spend
    # the module's real eight-second patience on each of them. The
    # production values are asserted separately, below.
    f.start_within = 0.2
    f.settle_timeout = 1.0
    # OFF here, ON in production (asserted separately). Every test built
    # on this helper before v1.34 is about what the BROWSER path does —
    # the tab, the hold, the second opinion — and most of them ask for a
    # native URL. With the network asked first, a 404 fixture would be
    # answered before the browser was ever sent, and those tests would
    # pass while testing nothing they are named for. The tests of the
    # status-first path switch it on explicitly.
    f.status_first = False
    f.open(drv)
    return f, drv


def test_a_download_is_a_transition_not_an_appearance():
    """Completion is .crdownload → final name.

    Probe 2 counted a .DS_Store that Finder wrote as a 0.01 MB download
    taking 39.9 seconds, because it asked whether a new file had appeared.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        w = M.DirectoryDownloadWatch(d)
        w.begin()
        # A dotfile appears. It is not a download.
        with open(os.path.join(d, ".DS_Store"), "wb") as fh:
            fh.write(b"\x00" * 64)
        name, note = w.settle(start_within=0.05, timeout=0.4)
        check("browser-watch . a dotfile is not a download", name, "")
        check_true("browser-watch . and the watcher says what it saw",
                   "nothing began arriving" in note, note)

        # A real one: partial first, then the rename.
        w.begin()
        tmp = os.path.join(d, "thesis.pdf.crdownload")
        with open(tmp, "wb") as fh:
            fh.write(b"%PDF-1.4 body")
        check("browser-watch . a transfer in flight is a partial",
              w.partials(), ["thesis.pdf.crdownload"])
        os.replace(tmp, os.path.join(d, "thesis.pdf"))
        name, note = w.settle(start_within=1.0, timeout=2.0)
        check("browser-watch . the rename is the completion", name,
              "thesis.pdf")

        # Stranded: a partial that never becomes anything.
        w.begin()
        with open(os.path.join(d, "stuck.pdf.crdownload"), "wb") as fh:
            fh.write(b"half")
        name, note = w.settle(start_within=0.1, timeout=0.5)
        check("browser-watch . a partial alone is not a download", name, "")
        check_true("browser-watch . and it is named, not diagnosed",
                   "stuck.pdf.crdownload" in note, note)
        # A finished file NEXT TO one still arriving is not a completed
        # fetch. Without this case, "completion is a transition" and
        # "completion is a new file appearing" behave identically on
        # every fixture above, and the difference between them is the
        # whole point of the watcher.
        w.begin()
        with open(os.path.join(d, "done.pdf"), "wb") as fh:
            fh.write(b"%PDF-done")
        name, note = w.settle(start_within=0.1, timeout=0.5)
        check("browser-watch . a finished file beside one in flight is "
              "not a completion", name, "")
        gone, growing = w.clear_partials(settle=0)
        check("browser-watch . clearing removes the stranded partial", gone,
              ["stuck.pdf.crdownload"])
        check("browser-watch . and nothing was growing", growing, [])
        check_true("browser-watch . and leaves the finished file alone",
                   os.path.exists(os.path.join(d, "thesis.pdf")))


def test_the_browser_path_tells_the_three_outcomes_apart():
    """Challenge, absence and success, from what the tab shows.

    There are no status codes on this path. Getting this vocabulary wrong
    is how an absence becomes a failure, or — far worse — how a refusal
    becomes a claim that a record has no file.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        # 1. success
        f, drv = _browser_fetcher(inc, [_Tab(drop=("a.pdf", b"%PDF-x"))])
        got = f.fetch("https://dc.example.edu/cgi/viewcontent.cgi?article=1")
        check("browser-vocab . a file that lands is a success", got.name,
              "a.pdf")
        check("browser-vocab . and its size is the file's", got.nbytes, 6)
        check("browser-vocab . head() reads the leading bytes",
              got.head(4), b"%PDF")

        # 2. challenge
        f, drv = _browser_fetcher(
            inc, [_Tab(title="Just a moment...",
                       source="<html>Enable JavaScript and cookies</html>")])
        try:
            f.fetch("https://dc.example.edu/cgi/viewcontent.cgi?article=2")
            check_true("browser-vocab . a challenge raises", False)
        except M.VerificationRequired as e:
            check("browser-vocab . a challenge raises VerificationRequired",
                  e.observed, "Just a moment...")
            check_true("browser-vocab . and never reads as an absence",
                       not M.is_definite_absence(e))

        # 3. definite absence
        f, drv = _browser_fetcher(
            inc, [_Tab(title="Digital Commons",
                       source="<p>No PDF has been provided.</p>")])
        try:
            f.fetch("https://dc.example.edu/cgi/viewcontent.cgi?article=3")
            check_true("browser-vocab . an absence raises", False)
        except M.NoFileAvailable as e:
            check_true("browser-vocab . an absence page is NoFileAvailable",
                       "no file to download" in str(e), str(e))

        # 4. neither: observed, not diagnosed. The page has to render AT
        # the requested URL — a page at a DIFFERENT URL is the server
        # redirecting us away, which since v1.33.3 is a definite absence.
        # This fixture used to land on /other and so tested that instead.
        unexplained = "https://dc.example.edu/cgi/viewcontent.cgi?article=4"
        f, drv = _browser_fetcher(
            inc, [_Tab(title="Something else", url=unexplained)])
        try:
            f.fetch(unexplained)
            check_true("browser-vocab . an unexplained page raises", False)
        except M.BrowserFetchFailed as e:
            msg = str(e)
            check_true("browser-vocab . it reports the tab and the URL",
                       "Something else" in msg and unexplained in msg,
                       msg)
            check_true("browser-vocab . it offers no cause",
                       "probably" not in msg.lower()
                       and "because" not in msg.lower(), msg)
            check_true("browser-vocab . and it is NOT a definite absence",
                       not M.is_definite_absence(e),
                       "'no file arrived' must never become 'the record "
                       "has no file'")


def test_a_challenge_does_not_spend_a_second_request_on_the_native():
    """A challenge is not an answer about this URL, so the native fallback
    must not fire on it — that would spend another request to be asked the
    same question, and could be read as an absence."""
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [_Tab(title="Verify you are human"),
                  _Tab(drop=("native.docx", b"PK\x03\x04"))])
        try:
            f.fetch("https://dc/viewcontent.cgi?article=9",
                    "https://dc/viewcontent.cgi?article=9&type=native")
        except M.VerificationRequired:
            pass
        check("browser-native . a challenge stops at one navigation",
              len(drv.visited), 1)

        # And the ordinary case still falls back.
        f, drv = _browser_fetcher(
            inc, [_Tab(title="DC", source="No PDF has been provided"),
                  _Tab(drop=("n.docx", b"PK\x03\x04zz"))])
        M.set_request_delay(0)
        got = f.fetch("https://dc/viewcontent.cgi?article=9",
                      "https://dc/viewcontent.cgi?article=9&type=native")
        check("browser-native . an absence falls back to the native",
              got.used_native, True)
        check("browser-native . and reports the native's file", got.name,
              "n.docx")

        # Both absent is an absence naming both; one unexplained is not.
        f, drv = _browser_fetcher(
            inc, [_Tab(title="DC", source="No PDF has been provided"),
                  _Tab(title="DC", source="No PDF has been provided")])
        try:
            f.fetch("https://dc/viewcontent.cgi?article=10",
                    "https://dc/viewcontent.cgi?article=10&type=native")
            check_true("browser-native . both absent raises", False)
        except M.NoFileAvailable as e:
            check_true("browser-native . both absent names both attempts",
                       "derivative:" in str(e) and "native:" in str(e),
                       str(e))
        f, drv = _browser_fetcher(
            inc, [_Tab(title="DC", source="No PDF has been provided"),
                  _Tab(title="who knows")])
        try:
            f.fetch("https://dc/viewcontent.cgi?article=11",
                    "https://dc/viewcontent.cgi?article=11&type=native")
            check_true("browser-native . one unanswered raises", False)
        except Exception as e:
            check_true("browser-native . an unanswered second attempt is "
                       "not folded into an absence",
                       isinstance(e, M.BrowserFetchFailed), repr(e))
            # v1.34.2: and the row says it was the NATIVE, at its URL —
            # the row's File URL is the derivative's, which did answer.
            check_true("browser-native . and names the native it was",
                       "native fallback, https://dc/viewcontent.cgi?"
                       "article=11&type=native," in M.describe_error(e),
                       M.describe_error(e))


def test_a_missing_native_is_given_time_to_say_so():
    """v1.34.2: the native page-load allowance, against the measurement.

    A HAR of 2026-10-01: a missing native's URL redirects in 93 ms to the
    CGI form, which answers 404 — "Sorry, that file doesn't exist" — after
    15,210 ms. The module allowed a native navigation 15.0 s, so the load
    was stopped just before the answer and the row ended "not confirmed".
    Driven with a fake that renders only when the allowance covers the
    server's answer time.
    """
    import tempfile
    sorry = (15.2, "Sorry, that file doesn't exist",
             "<html><body>Sorry, that file doesn't exist.</body></html>")
    native = ("https://dc/context/c/article/1/type/native/viewcontent"
              "?preview_mode=1")
    with tempfile.TemporaryDirectory() as d:
        tab = _Tab()
        tab.answers_after = sorry
        f, drv = _browser_fetcher(os.path.join(d, "_incoming"), [tab],
                                  at=("Digital Commons", "https://dc/", ""))
        try:
            f.fetch(native)
            check_true("native-allowance . a missing native is an absence",
                       False, "a file was returned")
        except Exception as e:
            check_true("native-allowance . a missing native is an absence",
                       isinstance(e, M.NoFileAvailable), repr(e))
        check_true("native-allowance . the native navigation was allowed "
                   "30 s", M.BROWSER_NATIVE_PAGE_TIMEOUT in drv.timeout_calls
                   and M.BROWSER_NATIVE_PAGE_TIMEOUT >= 30,
                   str(drv.timeout_calls))
        check_true("native-allowance . which covers the measured 15.2 s "
                   "with room", M.BROWSER_NATIVE_PAGE_TIMEOUT >= 2 * 15.2 - 1,
                   str(M.BROWSER_NATIVE_PAGE_TIMEOUT))

        # Every sharer of the 15 s allowance, checked: a derivative keeps
        # it, the dated query-form native gets the longer one, and the
        # site-root navigation normalize_tab() makes keeps the short one.
        check("native-allowance . a derivative keeps the short allowance",
              f.page_timeout_for("https://dc/cgi/viewcontent.cgi?article=1"
                                 "&context=c&preview_mode=1"),
              M.BROWSER_PAGE_TIMEOUT)
        check("native-allowance . a dated revision native gets the long one",
              f.page_timeout_for("https://dc/cgi/viewcontent.cgi?type=native"
                                 "&article=1&unstamped=yes&date=1&context=c"),
              M.BROWSER_NATIVE_PAGE_TIMEOUT)
        check("native-allowance . a supplemental keeps the short one",
              f.page_timeout_for("https://dc/cgi/viewcontent.cgi?filename=0"
                                 "&article=1&context=c&type=additional"),
              M.BROWSER_PAGE_TIMEOUT)
        drv.timeout_calls[:] = []
        drv._title = "Just a moment..."
        f.normalize_tab("https://dc/")
        check_true("native-allowance . the site-root navigation keeps the "
                   "short one", M.BROWSER_NATIVE_PAGE_TIMEOUT
                   not in drv.timeout_calls, str(drv.timeout_calls))

        # And a file that arrives is not held to the allowance at all: the
        # fake, like Chrome, completes no navigation for a download.
        f, drv = _browser_fetcher(os.path.join(d, "_in2"),
                                  [_Tab(drop=("n.docx", b"PK\x03\x04x"))])
        t0 = time.time()
        got = f.fetch(native)
        check_true("native-allowance . an existing native still downloads "
                   "at once", got.nbytes > 0 and time.time() - t0 < 5,
                   "{:.1f}s".format(time.time() - t0))


def test_the_browser_gets_its_download_folder_back():
    """Module contract 5. The failure mode is the person's next manual
    download landing silently in a run folder."""
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(inc, [_Tab(drop=("a.pdf", b"x"))])
        check("browser-restore . opening points Chrome at the run folder",
              drv.folder, inc)
        check_true("browser-restore . and asks for download events",
                   drv.cdp[0][1].get("eventsEnabled") is True,
                   str(drv.cdp[0]))
        f.fetch("https://dc/viewcontent.cgi?article=1")
        check_true("browser-restore . the page timeout is put back after "
                   "each navigation",
                   drv.timeout_calls[-1] == f.restore_page_timeout,
                   str(drv.timeout_calls))
        check_true("browser-restore . and it is put back to what was "
                   "there, not to a constant",
                   f.restore_page_timeout == 42.0,
                   "the Chrome bridge sets its own page-load timeout at "
                   "attach time; overruling it with a constant after "
                   "every navigation is a silent change to someone "
                   "else's setting (read {!r})".format(
                       f.restore_page_timeout))
        check_true("browser-restore . and it was short during it",
                   M.BROWSER_PAGE_TIMEOUT in drv.timeout_calls,
                   str(drv.timeout_calls))
        notes = f.close()
        check_true("browser-restore . close sets the folder back",
                   any(c[1].get("behavior") == "default" for c in drv.cdp),
                   str(drv.cdp))
        check_true("browser-restore . and says so", any("set back" in n
                                                        for n in notes),
                   str(notes))

        # It must report rather than raise, and loudly.
        f, drv = _browser_fetcher(inc, [])
        drv._raise_on_restore = True
        notes = f.close()
        check_true("browser-restore . a failed restore does not raise",
                   True)
        check_true("browser-restore . a failed restore is said loudly",
                   any("COULD NOT restore" in n for n in notes), str(notes))

        # A stranded partial is cleared on the way out.
        f, drv = _browser_fetcher(inc, [])
        with open(os.path.join(inc, "left.pdf.crdownload"), "wb") as fh:
            fh.write(b"half")
        notes = f.close()
        check_true("browser-restore . a stranded partial is cleared on exit",
                   not os.path.exists(os.path.join(inc,
                                                   "left.pdf.crdownload")))
        check_true("browser-restore . and the clearing is reported",
                   any("unfinished transfer" in n for n in notes),
                   str(notes))


def test_a_fetched_file_is_moved_not_copied_and_never_empty():
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        os.makedirs(inc)
        src = os.path.join(inc, "a.pdf")
        with open(src, "wb") as fh:
            fh.write(b"%PDF-1.4 hello")
        got = M.Fetched(path=src, name="a.pdf")
        dest = os.path.join(d, "etd_1_primary_public_current_a.pdf")
        got.save_as(dest)
        check_true("fetched . the file is moved to its final name",
                   os.path.exists(dest))
        check_true("fetched . and the incoming folder is left empty",
                   os.listdir(inc) == [], str(os.listdir(inc)))
        check("fetched . nbytes is the file's size", got.nbytes, 14)
        try:
            M.Fetched()
            check_true("fetched . an empty Fetched is refused", False,
                       "a helper that returns something empty is how a "
                       "zero-byte file gets a real record's name")
        except ValueError:
            check_true("fetched . an empty Fetched is refused", True)


class _Clock:
    def __init__(self):
        self.t = 1000.0

    def now(self):
        return self.t

    def sleep(self, seconds):
        self.t += seconds
        return True


def test_the_hold_waits_passively_and_says_how_long_is_left():
    """The hold must not navigate: retrying the URL every few seconds yanks
    the page out from under whoever is clicking the checkbox."""
    clock = _Clock()
    said = []
    asking = [True]
    v = M.VerificationWatch(window=300, poll=2, tick=30,
                            sleep=clock.sleep, stopping=lambda: False,
                            now=clock.now, announce=said.append)

    def still_asking():
        # Cleared after 10 seconds of waiting.
        return clock.t < 1010.0

    took, why = v.hold(still_asking)
    check_true("hold . it returns when the page stops asking",
               9.0 <= took <= 12.0, str(took))
    check("hold . and says the page was what cleared", why, "cleared")
    check_true("hold . the banner states the window as a number",
               any("5 minutes" in m for m in said), str(said[:1]))
    check_true("hold . and never prints the format spec itself",
               not any("{:" in m for m in said), str(said))

    # A file arriving is the other passive signal — clearing the prompt
    # usually returns the browser to the URL it was asked for.
    clock = _Clock()
    landed = [False]
    v = M.VerificationWatch(window=300, poll=2, tick=30, sleep=clock.sleep,
                            stopping=lambda: False, now=clock.now)

    def arrived():
        if clock.t > 1006.0:
            landed[0] = True
        return landed[0]

    took, why = v.hold(lambda: True, arrived)
    check_true("hold . a file arriving also ends the hold", took > 0,
               str(took))
    check("hold . and says so, because bytes mean do not ask again", why,
          "arrived")

    # Nobody comes.
    clock = _Clock()
    said = []
    v = M.VerificationWatch(window=120, poll=2, tick=30, sleep=clock.sleep,
                            stopping=lambda: False, now=clock.now,
                            announce=said.append)
    took, why = v.hold(lambda: True)
    check("hold . an unanswered prompt returns -1", took, -1.0)
    check("hold . with no reason to give", why, "")
    check_true("hold . and it counted down along the way",
               any("s left" in m for m in said), str(said))
    check_true("hold . the countdown carries a real number",
               any(re.search(r"\d+s left", m) for m in said), str(said))

    # Window 0 means do not wait at all.
    clock = _Clock()
    v = M.VerificationWatch(window=0, sleep=clock.sleep,
                            stopping=lambda: False, now=clock.now)
    check("hold . a zero window does not wait", v.hold(lambda: True)[0],
          -1.0)
    check("hold . and no time passed", clock.t, 1000.0)


def test_stop_works_while_the_run_is_holding():
    """A person who would rather end the run than clear the prompt must not
    have to clear it first."""
    clock = _Clock()
    shown = []
    v = M.VerificationWatch(window=300, poll=2, sleep=clock.sleep,
                            stopping=lambda: clock.t > 1004.0,
                            now=clock.now,
                            on_hold=lambda h, l, n: shown.append(h))
    try:
        v.hold(lambda: True)
        check_true("hold . Stop ends the hold", False, "it kept waiting")
    except M.StopRequested:
        check_true("hold . Stop ends the hold", True)
    check_true("hold . and the page stops saying a person is awaited",
               shown and shown[0] is True and shown[-1] is False,
               str(shown))


def test_the_prompt_record_is_the_audit_trail():
    clock = _Clock()
    v = M.VerificationWatch(window=60, poll=2, sleep=clock.sleep,
                            stopping=lambda: False, now=clock.now)
    # "Seen", not "requested": the run can only claim what it observed,
    # and on 2026-09-14 it claimed the stronger thing about a run in
    # which a prompt had been answered before it looked.
    check_true("audit . with nothing asked it says so, not nothing",
               "No verification prompt was seen" in v.summary(),
               v.summary())
    v.context = "honors-theses"
    v.requests = 257
    v.record(257, 256, 2.0, "Just a moment...")
    v.requests = 500
    v.record(500, 243, -1)
    rows = v.rows()
    check("audit . one row per prompt", len(rows), 2)
    check("audit . a row is the width of the sheet", len(rows[0]),
          len(M.VERIFY_HEADERS))
    check("audit . it records the structure", rows[0][2], "honors-theses")
    check("audit . where in the run it arrived", rows[0][3], 257)
    check("audit . and how much ran between prompts", rows[0][4], 256)
    check("audit . a cleared prompt says how long it took", rows[0][5], 2.0)
    check("audit . an uncleared prompt is named as such", rows[1][6],
          "Not cleared")
    # Driven through requested(), not by setting the counter: the reset
    # only shows when something has actually accumulated since the last
    # prompt. The first version of this check set v.requests directly,
    # so `since` was zero either way and deleting the reset changed
    # nothing it looked at.
    v.requested()
    v.requested()
    v.requested()
    check("audit . a prompt's count of preceding work is live",
          v.since, 3)
    v.record(503, v.since, 1.0)
    check("audit . recording a prompt resets the count of what follows",
          v.since, 0)
    v.requested()
    check("audit . which then counts only the work since that prompt",
          v.since, 1)
    s = v.summary()
    check_true("audit . the summary counts every prompt and its outcome",
               "3 verification prompt(s)" in s and "2 cleared" in s
               and "1 not cleared" in s, s)
    check_true("audit . and does not read as a failure",
               "fail" not in s.lower() and "error" not in s.lower(), s)


def test_a_prompt_is_retried_properly_and_bounded():
    """Probe 3 lost article 1495: after a two-second clearance the single
    immediate retry still saw the challenge and the file was written off.
    So the retry re-enters the hold — but not forever."""
    import tempfile
    clock = _Clock()

    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        # Challenge, then still challenging on the retry, then the file.
        f, drv = _browser_fetcher(
            inc, [_Tab(title="Just a moment..."),
                  _Tab(title="Just a moment..."),
                  _Tab(drop=("late.pdf", b"%PDF-late"))],
            clock=clock, clear_after=4.0)
        v = M.VerificationWatch(window=300, poll=2, sleep=clock.sleep,
                                stopping=lambda: False, now=clock.now)
        M.set_request_delay(0)
        got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1495")
        check("retry . the interrupted file is fetched in the end",
              got.name, "late.pdf")
        check("retry . and it took two prompts to get there",
              len(v.prompts), 2)
        check_true("retry . both prompts were recorded as cleared",
                   len(v.cleared) == 2, str(v.prompts))
        check("retry . the file was requested three times",
              len(drv.visited), 3)

    # Bounded: a site that asks again every time is not answered forever.
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        # ONE clock. Two of them is a hold whose `now` never advances,
        # which is an infinite wait rather than a failing test.
        c2 = _Clock()
        f, drv = _browser_fetcher(
            inc, [_Tab(title="Just a moment...") for _ in range(20)],
            clock=c2, clear_after=4.0)
        v = M.VerificationWatch(window=300, poll=2, sleep=c2.sleep,
                                stopping=lambda: False, now=c2.now)
        try:
            M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1")
            check_true("retry . an endless prompt stops the run", False,
                       "it kept answering")
        except M.NotVerified as e:
            check_true("retry . an endless prompt stops the run", True)
            check_true("retry . and says that is why",
                       "returned" in str(e), str(e))
        except Exception as e:
            # Deliberately broad. With the bound removed the loop does
            # not hang here — it exhausts the scripted tabs and fails on
            # something else entirely — and the count below is then the
            # assertion that actually speaks to the bound. Letting that
            # exception out would have the test fail under its own name
            # rather than under the check written for this.
            check_true("retry . an endless prompt stops the run", False,
                       "ended as {}: {}".format(e.__class__.__name__, e))
        check_true("retry . it did not loop unboundedly",
                   len(drv.visited) <= M.VERIFY_MAX_HOLDS + 1,
                   "{} navigations".format(len(drv.visited)))


def test_nobody_confirming_is_a_stop_reason_not_a_failure():
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(inc, [_Tab(title="Just a moment...")])
        clock = _Clock()
        v = M.VerificationWatch(window=10, poll=2, sleep=clock.sleep,
                                stopping=lambda: False, now=clock.now)
        try:
            M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1")
            check_true("notverified . it raises", False)
        except M.NotVerified as e:
            check_true("notverified . NotVerified is a stop reason",
                       isinstance(e, M.StopRequested),
                       "every existing handler flushes rows on "
                       "StopRequested; a new sibling would silently skip "
                       "all of it")
            check_true("notverified . and it is not a failure word",
                       "fail" not in str(e).lower(), str(e))
        check("notverified . the prompt is still recorded",
              len(v.prompts), 1)
        check("notverified . as not cleared", v.prompts[0]["cleared"], None)


def test_a_partial_is_cleared_before_the_retry():
    """Whether a challenge can interrupt a transfer mid-flight is NOT
    measured — probe 3's one prompt arrived between files. Handled
    defensively, and a partial must never be counted as a finished file."""
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            # The stranded partial is named differently from the file
            # the retry fetches. Same name and the retry's own rename
            # would tidy it away, and the test would pass whether or not
            # anything cleared it.
            inc, [_Tab(title="Just a moment...", partial="other-half.pdf"),
                  _Tab(drop=("thesis.pdf", b"%PDF-whole"))])
        clock = _Clock()
        v = M.VerificationWatch(window=300, poll=2, sleep=clock.sleep,
                                stopping=lambda: False, now=clock.now)
        got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1")
        check("partial . the retry returns the whole file", got.nbytes, 10)
        check("partial . and no partial is left behind",
              [n for n in os.listdir(inc) if n.endswith(".crdownload")], [])


def test_the_in_flight_records_remaining_files_are_not_lost():
    """Tenth instance of "a partial failure must not escalate into total
    loss", and the first found by reading rather than by losing data.

    A run stopping part-way through a record wrote file rows for what it
    had attempted and a record row for the record — and merge_rows drops
    that record row, because a record with file rows has by definition
    been planned. Everything the run had not reached was therefore in no
    file anywhere.
    """
    ctx = "honors-theses"
    item = {"article": "1495", "title": "A thesis"}
    jobs = [{"kind": "primary", "visibility": "public", "version": "current",
             "version_date": "", "orig_hint": "a.pdf", "access": "open",
             "release": "Thesis (open access)", "embargo": "",
             "url": "https://dc/viewcontent.cgi?article=1495"},
            {"kind": "supp1", "visibility": "public", "version": "current",
             "version_date": "", "orig_hint": "b.xlsx", "access": "open",
             "release": "Thesis (open access)", "embargo": "",
             "url": "https://dc/viewcontent.cgi?article=1495&type=additional"}]
    left = M.unreached_job_rows(jobs, 1, 1, item, "Posted",
                                "https://dc/pub/1495", "https://dc/adm/1495")
    check("inflight . one row per job not reached", len(left), 1)
    check("inflight . the row is the width of the report", len(left[0]),
          len(M.REPORT_HEADERS))
    check("inflight . it names the file, not the record",
          left[0][M.REPORT_HEADERS.index("File Kind")], "supp1")
    check("inflight . and carries the file's own URL",
          left[0][M.REPORT_HEADERS.index("File URL")], jobs[1]["url"])
    check("inflight . recorded as not attempted",
          left[0][M.COL_STATUS], M.NOT_ATTEMPTED)

    # The defect itself: these rows survive the merge, and a record row
    # in their place does not.
    def as_dict(row):
        d = {h: row[n] for n, h in enumerate(M.REPORT_HEADERS)}
        return {"context": ctx, "article": d["Article ID"],
                "kind": d["File Kind"], "version": d["Version"],
                "version_date": d["Version Date"], "url": d["File URL"],
                "was": d["Status"]}

    downloaded = {"context": ctx, "article": "1495", "kind": "primary",
                  "version": "current", "version_date": "",
                  "url": jobs[0]["url"], "was": "Downloaded"}
    record_row = {"context": ctx, "article": "1495", "kind": M.RECORD_KIND,
                  "version": "", "version_date": "",
                  "url": "https://dc/adm/1495", "was": M.NOT_ATTEMPTED}
    merged = M.merge_rows([("r.xlsx", [downloaded, record_row])])
    check_true("inflight . a record row cannot carry the remainder",
               not any(r["kind"] == M.RECORD_KIND for r in merged),
               "this is the defect: merge_rows correctly drops it")
    merged = M.merge_rows([("r.xlsx", [downloaded, as_dict(left[0])])])
    kinds = sorted(r["kind"] for r in merged)
    check("inflight . file rows for the remainder do survive it", kinds,
          ["primary", "supp1"])


def _calls_in(func, name):
    """Every call to `name` inside `func`, found with ast rather than grep.

    A grep for a call name matches it in a comment, in a docstring, and
    inside a longer name — all three have produced green checks over
    hollow code in this module, most recently when a comment of MINE
    mentioning fetch_one() satisfied an ordering check meant for the call.
    """
    import ast as _ast
    tree = _ast.parse(textwrap.dedent(inspect.getsource(func)))
    found = []
    for node in _ast.walk(tree):
        if not isinstance(node, _ast.Call):
            continue
        fn = node.func
        label = (fn.attr if isinstance(fn, _ast.Attribute)
                 else getattr(fn, "id", ""))
        if label == name:
            found.append(node)
    return found


def test_both_workers_fetch_the_same_way():
    """The parity rule, checked as parity rather than as two greps.

    Nine defects in this module are "fixed in one worker and not the
    other", three of them in this pair of functions, and the nine checks
    guarding the last one were every one a grep of download_worker — all
    nine green the morning a re-run failed in the way they existed to
    prevent. So this asserts the thing that makes parity structural: both
    workers fetch through fetch_one(), and neither reaches past it.
    """
    for worker in (M.download_worker, M.retry_worker):
        n = worker.__name__
        check("parity . {} fetches through fetch_one".format(n),
              len(_calls_in(worker, "fetch_one")), 1)
        check("parity . {} does not fetch around it".format(n),
              len(_calls_in(worker, "fetch_record_file")), 0)
        check("parity . {} builds a fetcher".format(n),
              len(_calls_in(worker, "make_fetcher")), 1)
        check("parity . {} holds a VerificationWatch".format(n),
              len(_calls_in(worker, "VerificationWatch")), 1)
        src = textwrap.dedent(inspect.getsource(worker))
        check_true("parity . {} gives the browser back in a finally"
                   .format(n),
                   "    finally:" in src and "fetcher.close()" in src,
                   "module contract 5, on every exit path")
        check_true("parity . {} reports NotVerified as its own stop reason"
                   .format(n), "NotVerified)" in src,
                   "a run nobody verified must not read as a failure")
        check_true("parity . {} lets a stop out of the fetch, not into a "
                   "row".format(n), _stop_escapes_the_fetch(worker),
                   "NotVerified inherits Exception, so a bare `except "
                   "Exception` around the fetch swallows it and records "
                   "'Failed:' on a file that was never refused")
    check("parity . the in-flight record's remaining files are written",
          len(_calls_in(M.download_worker, "unreached_job_rows")), 1)


def _stop_escapes_the_fetch(worker):
    """True when the try that fetches re-raises StopRequested first.

    Structural, and precisely so: it finds the `try` whose body actually
    contains the fetch_one() call, then asks whether a StopRequested
    handler precedes any handler that would also catch it. A grep for
    "except StopRequested" would pass on one anywhere in the function,
    including one sitting after the `except Exception` that had already
    swallowed it.
    """
    import ast as _ast
    tree = _ast.parse(textwrap.dedent(inspect.getsource(worker)))
    # The INNERMOST try containing the fetch. A worker's whole body sits
    # inside an outer try whose handlers flush reports and write
    # summaries; asking that one whether it re-raises answers a different
    # question, and answers it wrongly. Innermost = latest to start.
    holders = []
    for node in _ast.walk(tree):
        if not isinstance(node, _ast.Try):
            continue
        inner = _ast.Module(body=node.body, type_ignores=[])
        if any(isinstance(c, _ast.Call)
               and getattr(c.func, "id", "") == "fetch_one"
               for c in _ast.walk(inner)):
            holders.append(node)
    for node in sorted(holders, key=lambda t: -t.lineno)[:1]:
        for handler in node.handlers:
            names = []
            if isinstance(handler.type, _ast.Name):
                names = [handler.type.id]
            elif isinstance(handler.type, _ast.Tuple):
                names = [getattr(e, "id", "") for e in handler.type.elts]
            if "StopRequested" in names:
                return any(isinstance(st, _ast.Raise)
                           for st in handler.body)
            if "Exception" in names or handler.type is None:
                return False       # swallowed before it could get out
    return False



def test_a_rerun_through_the_browser_lands_files_and_an_audit_trail():
    """Worker level, on the default path, driving the real fetcher."""
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article=7&context=etd"
    rows = [_row(1, str(7 + i), "primary", "public", "current",
                 "Failed: HTTP 500 Server Error",
                 base + "/cgi/viewcontent.cgi?article={}&context=etd"
                 .format(7 + i), adm)
            for i in range(3)]

    script = [_Tab(drop=("a.pdf", b"%PDF-one")),
              _Tab(title="Just a moment..."),          # prompt mid-run
              _Tab(drop=("b.pdf", b"%PDF-two")),       # the retry
              _Tab(drop=("c.pdf", b"%PDF-three"))]
    # The worker runs the REAL hold, in real seconds, so the person
    # clears the prompt after one look rather than after a fake interval.
    drv = _FakeBrowser(script, clear_after_reads=1)

    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "LISTING_SETTLE",
              "BROWSER_START_WITHIN")}
    logs = []
    real_log = M.log
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: drv
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            # The worker builds its own fetcher, which reads this at
            # construction. Left at the real eight seconds the suite
            # spends most of its time asleep.
            M.BROWSER_START_WITHIN = 0.3
            M.log = lambda msg: (logs.append(str(msg)), real_log(str(msg)))[0]
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}}, out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            "reconcile": False,
                            "fetch_path": M.FETCH_BROWSER,
                            "verify_window": 300},
                           report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            M.log = real_log

        landed = []
        for root, _dirs, files in os.walk(out):
            landed += [f for f in files if f.endswith(".pdf")]
        check("rerun-browser . every file was fetched through Chrome",
              len(landed), 3)
        # v1.34.1: and the report says so, row by row.
        import glob as _g
        import openpyxl
        vias = []
        for rp in _g.glob(os.path.join(out, "**", "*.xlsx"), recursive=True):
            if "MasterLog" in os.path.basename(rp):
                continue
            wbk = openpyxl.load_workbook(rp, read_only=True)
            try:
                rows_ = list(wbk.worksheets[0].iter_rows(values_only=True))
            finally:
                wbk.close()
            hi = rows_[0].index("Fetched via")
            si = rows_[0].index("Status")
            vias += [r[hi] for r in rows_[1:] if r[si] == "Downloaded"]
        check("rerun-browser . each downloaded row says Chrome fetched it",
              vias, ["Chrome"] * 3)
        check_true("rerun-browser . and nothing is left in _incoming",
                   not os.path.isdir(os.path.join(out, "_incoming"))
                   or not [f for f in os.listdir(
                       os.path.join(out, "_incoming"))
                       if not f.startswith(".")],
                   "a landed file must be moved, not copied")

        # The run writes into a run folder under `out`, so look for it
        # rather than assuming it is at the top — a test that reads the
        # wrong folder fails for a reason that has nothing to do with
        # the feature.
        master = []
        for root, _dirs, files in os.walk(out):
            master += [os.path.join(root, f) for f in files
                       if "MasterLog" in f]
        check("rerun-browser . a master log was written", len(master), 1)
        from openpyxl import load_workbook
        wb = load_workbook(master[0])
        check_true("rerun-browser . it carries the verification sheet",
                   M.VERIFY_SHEET in wb.sheetnames, str(wb.sheetnames))
        vs = wb[M.VERIFY_SHEET]
        vals = [[c for c in r] for r in vs.values]
        check("rerun-browser . the sheet's header is the audit schema",
              list(vals[0]), M.VERIFY_HEADERS)
        check("rerun-browser . one row for the one prompt", len(vals) > 1
              and str(vals[1][6]), "Cleared while the run waited")
        # v1.34.1: and what was checked before any request.
        check_true("rerun-browser . it carries the Checks sheet",
                   M.CHECKS_SHEET in wb.sheetnames, str(wb.sheetnames))
        if M.CHECKS_SHEET in wb.sheetnames:
            cv = [list(r) for r in wb[M.CHECKS_SHEET].values]
            check("rerun-browser . the Checks header", cv[0],
                  M.CHECKS_HEADERS)
            # Free space is this machine's real disk, so "note" (under
            # 10 GB free) is as right as "passed"; the other two are ours.
            got_ = [(r[0], r[1]) for r in cv[1:]]
            # v1.34.2: the tab check is recorded first — it runs first.
            check("rerun-browser . the tab, depth, space and the download "
                  "test, in that order", [g[0] for g in got_],
                  ["Tab ready", "Destination depth", "Free space",
                   "Download test"])
            check_true("rerun-browser . each with its result",
                       len(got_) == 4 and got_[0][1] == "passed"
                       and got_[1][1] == "passed"
                       and got_[2][1] in ("passed", "note")
                       and got_[3][1] == "passed", str(got_))
            check_true("rerun-browser . the tab check says how long it took",
                       " answered in " in str(cv[1][2])
                       and " s (title: " in str(cv[1][2]), str(cv[1]))
            check_true("rerun-browser . the download test says what arrived",
                       "2 of 2 test file(s)" in str(cv[-1][2]), str(cv[-1]))

    check_true("rerun-browser . the prompt was announced in words a "
               "person can act on",
               any("VERIFICATION NEEDED" in l for l in logs), "\n".join(
                   l for l in logs if "VERIF" in l.upper())[:400])
    check_true("rerun-browser . the summary counts the prompt",
               any("verification prompt(s)" in l for l in logs),
               "\n".join(logs[-3:]))
    check_true("rerun-browser . Chrome's download folder was set back",
               any(c[1].get("behavior") == "default" for c in drv.cdp),
               str(drv.cdp))


def test_a_rerun_nobody_verifies_stops_without_calling_it_a_failure():
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article=7&context=etd"
    rows = [_row(1, str(7 + i), "primary", "public", "current",
                 "Failed: HTTP 500 Server Error",
                 base + "/cgi/viewcontent.cgi?article={}&context=etd"
                 .format(7 + i), adm)
            for i in range(4)]
    # One file, then a prompt nobody ever clears.
    script = [_Tab(drop=("a.pdf", b"%PDF-one"))] + \
             [_Tab(title="Just a moment...") for _ in range(8)]
    drv = _FakeBrowser(script)          # no clock: nobody clears it

    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "LISTING_SETTLE",
              "BROWSER_START_WITHIN")}
    logs = []
    real_log = M.log
    summary = ""
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: drv
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            # The worker builds its own fetcher, which reads this at
            # construction. Left at the real eight seconds the suite
            # spends most of its time asleep.
            M.BROWSER_START_WITHIN = 0.3
            M.log = lambda msg: (logs.append(str(msg)), real_log(str(msg)))[0]
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}}, out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            "reconcile": False,
                            "fetch_path": M.FETCH_BROWSER,
                            # Zero: do not wait at all, so the suite does
                            # not sit out a real five minutes to learn
                            # what an unanswered prompt does.
                            "verify_window": 0},
                           report_dir=out)
            summary = M.STATE.get("summary", "")
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            M.log = real_log
        partial = []
        for root, _dirs, files in os.walk(out):
            partial += [f for f in files if "PARTIAL" in f]

    check_true("notverified-rerun . it says nobody confirmed",
               "nobody confirmed" in summary.lower(), summary)
    check_true("notverified-rerun . it does not read as a failure",
               "fail" not in summary.lower(), summary)
    check_true("notverified-rerun . nor as a server refusal",
               "refusing" not in summary.lower(), summary)
    check_true("notverified-rerun . the rows fetched are still written out",
               partial, str(partial))
    check_true("notverified-rerun . and the log says the rest were never "
               "asked for",
               any("not requested" in l for l in logs),
               "\n".join(logs[-4:]))
    check_true("notverified-rerun . the download folder is still set back",
               any(c[1].get("behavior") == "default" for c in drv.cdp),
               "a run that ended badly must still give Chrome back")


def test_the_page_offers_the_fetch_path_and_the_window():
    page = M.build_page(
        M.load_session(Path("/nonexistent/session.json"))).decode("utf-8")
    check_true("fetch-ui . the page offers a fetch path",
               'id="fetchbrowser"' in page and 'id="fetchnetwork"' in page)
    check_true("fetch-ui . Chrome is the default",
               'id="fetchbrowser"' in page
               and page.split('id="fetchbrowser"')[1].split(">")[0]
               .find("checked") >= 0,
               "the only path that can answer a verification prompt must "
               "be the one a run takes without being told")
    check_true("fetch-ui . the window is editable",
               'id="verifywin"' in page)
    check_true("fetch-ui . and it is rendered from the constant, not typed",
               'value="{:.0f}"'.format(M.VERIFY_WINDOW / 60.0) in page,
               "a page that quotes a number the code does not use is a "
               "number nobody can check")
    check_true("fetch-ui . no token survived into the served bytes",
               "__VERIFYMIN__" not in page and "__VERIFYMAXMIN__" not in page)
    check_true("fetch-ui . the long explanation is behind a disclosure",
               "<details><summary>What the two paths do differently"
               in page)
    check_true("fetch-ui . the tab title changes for a prompt",
               _title_rule_behaves(page),
               "on a second screen the tab title is what gets noticed")
    # The CALL SITE, separately, and source-checked because poll() cannot
    # be executed without a server to answer /api/state. Asserted on the
    # live value rather than on the function's name: a call that passes a
    # literal would satisfy any check that only asked whether the title
    # code was still there.
    check_true("fetch-ui . and the title follows the run's actual state",
               "applyVerifyTitle(!!v.holding);" in page,
               "a title driven by a literal announces nothing")


def _run_node(node, js, timeout=30):
    """Run a script in node; returns a CompletedProcess with UTF-8 text.

    The script goes in a UTF-8 file, not on the command line, and node's
    output is decoded as UTF-8, not with the locale's codepage. CI's first
    run of 775b4a2 failed three checks on both Windows jobs because
    text=True decoded node's "—" as cp1252; and under a non-UTF-8 locale
    the script could not even be passed as an argument, because it
    carries the page's own non-ASCII text.
    """
    import subprocess as _sp
    fd, path = tempfile.mkstemp(suffix=".js")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(js)
        return _sp.run([node, path], capture_output=True, text=True,
                       encoding="utf-8", errors="replace", timeout=timeout)
    finally:
        os.unlink(path)


def _title_rule_behaves(page):
    """Run applyVerifyTitle() in node and watch document.title.

    Behavioral: it asserts the title changes while a prompt is waiting
    and goes back afterwards, rather than asserting the function exists.
    """
    import shutil as _shutil
    import subprocess as _sp

    node = _shutil.which("node")
    if not node:
        return True                        # nothing to say without node
    start = page.find("const PAGETITLE=")
    end = page.find("function idleStatus()")
    if start < 0 or end < 0:
        return False
    js = page[start:end]
    harness = ("const document = {title: 'Batch File Downloader'};\n"
               + js
               + "\nconst before = document.title;\n"
                 "applyVerifyTitle(true); const during = document.title;\n"
                 "applyVerifyTitle(false); const after = document.title;\n"
                 "console.log(JSON.stringify([before, during, after]));\n")
    r = _run_node(node, harness)
    if r.returncode != 0:
        return False
    import json as _json
    before, during, after = _json.loads(r.stdout.strip())
    return during != before and after == before and before in during


def test_the_modules_wording_names_no_vendor_and_no_agreement():
    """Mechanism only, everywhere.

    The page was once about to explain the prompts by citing the
    arrangement that produces them — text that would have shipped in the
    public Distribution copy. What the module says is true of any
    repository and useful to any institution: you may be asked to verify
    that you are present, the run pauses until you do, each prompt is
    recorded.
    """
    src = Path(M.__file__).read_text(encoding="utf-8")
    lowered = src.lower()
    # Word boundaries, not substrings. The first version of this check
    # failed on "standalone" (it contains the three letters of the
    # licence term), "module contract 5" and "no
    # agreement" — a check verifying the adjacent thing, written into the
    # very test meant to catch adjacent matches.
    #
    # "contract" and "agreement" are deliberately absent from this list:
    # both are used in the module for its own internal contract, and a
    # check that cries wolf on them is a check somebody switches off.
    # What may never appear is a vendor, or the vocabulary of a licence.
    #
    # Since 2026-10-03 the words come from the institution profile, read
    # by the public build's own reader (make_distribution.private_terms),
    # which scans every file of the public copy, this one included.
    # "bepress" is not a private term (it names the product), so it is
    # checked here, for this module only.
    build = _private_terms()
    terms = build.private_terms()
    profiles = os.path.join(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))), "profiles")
    named = os.path.isdir(profiles) and any(
        f.endswith(".json") for f in os.listdir(profiles))
    check("wording . private terms are declared exactly when a named "
          "profile exists", bool(terms), bool(named))
    if terms:
        for pattern, why in terms:
            check_true("wording . the source never says {!r} ({})".format(
                           pattern, why),
                       not build.private_term_hits(src, [(pattern, why)]),
                       "a vendor or an institution's terms in public source "
                       "is what the public build refuses")
    else:
        # A public copy declares nothing private, so there is nothing to
        # search FOR. Say that, as the identity check does.
        check_true("wording . NOTHING WAS CHECKED for private terms — no "
                   "profile declares any", True)
    check_true("wording . the source never says 'bepress'",
               re.search(r"\bbepress\b", lowered) is None, "")
    check_true("wording . and it does say what actually happens",
               "verify that you are present" in lowered, "")


def test_both_endpoints_refuse_a_fetch_path_they_cannot_take():
    """One validator, called by both. /api/start and /api/retry have
    drifted apart before."""
    check("endpoint-fetch . a good pair passes",
          M._fetch_option_error({"fetch_path": M.FETCH_BROWSER,
                                 "verify_window": 300}), "")
    bad = M._fetch_option_error({"fetch_path": "curl",
                                 "verify_window": 300})
    check_true("endpoint-fetch . an unknown path is refused", bool(bad), bad)
    check_true("endpoint-fetch . and the refusal names the choice",
               "Chrome" in bad, bad)
    long_wait = M._fetch_option_error(
        {"fetch_path": M.FETCH_BROWSER,
         "verify_window": M.VERIFY_WINDOW_MAX + 1})
    check_true("endpoint-fetch . an unbounded window is refused",
               bool(long_wait), long_wait)
    check("endpoint-fetch . zero is allowed — it means do not wait",
          M._fetch_option_error({"fetch_path": M.FETCH_BROWSER,
                                 "verify_window": 0}), "")
    check_true("endpoint-fetch . a negative window is refused",
               bool(M._fetch_option_error({"fetch_path": M.FETCH_BROWSER,
                                           "verify_window": -1})))
    # Both endpoints, by ast: the validator is called in each branch.
    src = textwrap.dedent(inspect.getsource(M.make_handler))
    check("endpoint-fetch . both endpoints call the one validator",
          src.count("_fetch_option_error("), 2)
    check_true("endpoint-fetch . and both forward the choice to the worker",
               src.count('"fetch_path"') >= 2
               and src.count('"verify_window"') >= 2, "")


def test_a_prompt_is_seen_while_the_run_waits_for_the_file():
    """The defect the first live run exposed, and the reason for
    BROWSER_OBSERVE.

    2026-09-14: a real run met a verification prompt, the operator
    cleared it in about two seconds, the file arrived, and the run
    reported "No verification was requested during this run (70 file
    requests)". Nothing had looked at the tab, because the tab was only
    read after the wait for a file GAVE UP — and it never gave up, the
    file came. Every one of the 1,114 tests passed that morning; none
    of them had a prompt that resolved by itself.

    Two things make this test the real one, both learned by planting
    the violation and watching nothing go red:

      * **the navigation STALLS rather than rendering.** A rendered
        challenge page is read immediately by a different branch, so
        the first version stayed green with the watching deleted;
      * **the person clears the prompt on a TIMER**, not in response to
        being looked at. With clearing triggered by looking, deleting
        the observation made the fetch fail — which is a red test for
        the wrong reason, and hides that the live failure was a fetch
        that SUCCEEDED while losing the prompt.

    With the tab watched, the prompt is recorded. Without it, the file
    still arrives and the prompt is gone — which is exactly what
    happened on the live run.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        first = _Tab(title="Just a moment...", stall=True)
        first.timed_clear = (0.6, "late.pdf", b"%PDF-late")
        f, drv = _browser_fetcher(inc, [first,
                                        _Tab(drop=("late2.pdf", b"%PDF-2"))])
        f.settle_timeout = M.BROWSER_SETTLE_TIMEOUT  # as production (v1.34.1)
        # Long enough that the file arrives inside the wait — which is
        # the situation in which the old code never looked at all.
        f.start_within = 5.0
        v = M.VerificationWatch(
            window=30, poll=0.2,
            sleep=lambda secs: (time.sleep(min(secs, 0.2)) or True),
            stopping=lambda: False, now=time.time)
        got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=7")
        check_true("seen . the file still arrives",
                   got.name in ("late.pdf", "late2.pdf"), got.name)
        check("seen . the prompt was RECORDED, not lost", len(v.prompts), 1)
        # Guarded, so a missing prompt fails the check above and does
        # not ALSO raise an IndexError under the test's own name — two
        # reports of one defect, the second of them noise.
        check_true("seen . with how long it took to clear",
                   bool(v.prompts) and v.prompts[0]["cleared"] is not None,
                   str(v.prompts))
        check_true("seen . the summary no longer denies it",
                   "No verification prompt was seen" not in v.summary(),
                   v.summary())
        check_true("seen . and it counts the prompt",
                   "1 verification prompt(s)" in v.summary(), v.summary())


def test_a_prompt_cleared_in_the_gap_is_still_recorded():
    """The narrow race: showing when glanced at, gone when read properly.

    The run must not treat that as an unexplained failure, and must not
    forget the prompt either — it is the clearest evidence a person was
    present.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(inc, [_Tab(title="Just a moment...")],
                                  clear_after_reads=0)
        # clear_after_reads=0 means the very first read already shows it
        # cleared, so settle's glance sees a challenge only if it looks
        # before that. Drive the observation directly instead, which is
        # what makes this a test of the RECORDING rather than of timing.
        f.watch.begin()
        f.watch.challenge_at = 0.4
        f.watch.challenge_gone_at = 1.9
        f._last_title = "Just a moment..."
        got = M.Fetched(data=b"%PDF-x", name="a.pdf",
                        prompt={"observed": f._last_title,
                                "cleared_in": 1.9})
        clock = _Clock()
        v = M.VerificationWatch(window=300, poll=2, sleep=clock.sleep,
                                stopping=lambda: False, now=clock.now)

        class _Once:
            path = M.FETCH_BROWSER

            def fetch(self, url, native_url=""):
                return got

        M.fetch_one(_Once(), v, "https://dc/viewcontent.cgi?article=7")
        check("gap . the prompt is recorded", len(v.prompts), 1)
        check("gap . with how long it took to clear",
              v.prompts[0]["cleared"] if v.prompts else None, 1.9)
        check("gap . and it says the run never stopped",
              M._prompt_outcome(v.prompts[0]) if v.prompts else None,
              "Cleared without stopping the run")
        check_true("gap . which the summary distinguishes",
                   "without stopping the run" in v.summary(), v.summary())


def test_the_outcome_says_what_was_observed_not_who_acted():
    cleared_held = {"cleared": 2.0, "paused": True}
    cleared_free = {"cleared": 1.2, "paused": False}
    never = {"cleared": None, "paused": True}
    check("outcome . a prompt that stopped the run",
          M._prompt_outcome(cleared_held), "Cleared while the run waited")
    check("outcome . a prompt that did not",
          M._prompt_outcome(cleared_free),
          "Cleared without stopping the run")
    check("outcome . and one nobody answered",
          M._prompt_outcome(never), "Not cleared")
    for p in (cleared_held, cleared_free, never):
        check_true("outcome . none of them claims to know who acted",
                   "person" not in M._prompt_outcome(p),
                   "the run knows the page stopped asking, not who made "
                   "it stop")


def test_the_page_between_structures_is_checked_too():
    """The second blind spot.

    Each worker loads the repository home page once per structure to
    prove the session is live, and nothing looked at what came back. A
    prompt served there was invisible — one of the two ways the first
    live run could meet a prompt and report none.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(inc, [], clear_after_reads=2)
        drv._title = "Just a moment..."
        drv._source = "<html>Enable JavaScript and cookies</html>"
        drv._asked_at = 0
        clock = _Clock()
        said = []
        v = M.VerificationWatch(window=300, poll=2, sleep=clock.sleep,
                                stopping=lambda: False, now=clock.now)
        held = M.hold_if_challenged(f, v, "the site page", said.append)
        check("between . a prompt on the site page is caught", held, True)
        check("between . and recorded", len(v.prompts), 1)
        check_true("between . naming where it happened",
                   "site page" in v.prompts[0]["observed"],
                   str(v.prompts[0]))
        check_true("between . and announced to the operator",
                   any("waiting" in m for m in said), str(said))

        # A quiet page costs one title read and says nothing.
        f2, drv2 = _browser_fetcher(inc, [])
        drv2._title = "A Repository"
        drv2._source = "<html>ok</html>"
        v2 = M.VerificationWatch(window=300, sleep=_Clock().sleep,
                                 stopping=lambda: False, now=_Clock().now)
        check("between . an ordinary page is not a prompt",
              M.hold_if_challenged(f2, v2, "the site page"), False)
        check("between . and nothing is recorded for it", len(v2.prompts), 0)

        # Nobody answers: the run stops, and it stops as NotVerified.
        f3, drv3 = _browser_fetcher(inc, [])
        drv3._title = "Just a moment..."
        drv3._asked_at = 0
        c3 = _Clock()
        v3 = M.VerificationWatch(window=0, sleep=c3.sleep,
                                 stopping=lambda: False, now=c3.now)
        try:
            M.hold_if_challenged(f3, v3, "the site page")
            check_true("between . an unanswered prompt stops the run",
                       False)
        except M.NotVerified:
            check_true("between . an unanswered prompt stops the run", True)


def test_a_rendered_page_is_read_at_once_not_after_the_file_wait():
    """Measured cost, from the first live run: 44 seconds a record on a
    structure whose records are link-only, almost all of it spent
    waiting for a file the page had already said was not there."""
    import tempfile
    import time as _time
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [_Tab(title="sorry",
                       source="<p>No PDF has been provided.</p>")])
        # The real patience, so the test measures the shortcut rather
        # than the harness.
        f.start_within = 3.0
        t0 = _time.time()
        try:
            f.fetch("https://dc/viewcontent.cgi?article=1000")
            check_true("fast-absence . it raises", False)
        except M.NoFileAvailable:
            pass
        took = _time.time() - t0
        check_true("fast-absence . a rendered absence is read immediately",
                   took < 1.0,
                   "{:.1f}s — the page had already answered".format(took))

        # ... but bytes on disk overrule the page. A download and a
        # rendered page are not mutually exclusive, and the absence
        # branch must never discard a file that is arriving.
        f2, drv2 = _browser_fetcher(inc, [])
        f2.watch.begin()
        with open(os.path.join(inc, "arriving.pdf.crdownload"), "wb") as fh:
            fh.write(b"half")
        check("fast-absence . bytes on disk overrule the page",
              f2.watch.anything_started(), True)


def test_a_prompt_on_the_derivative_survives_the_native_fallback():
    """A prompt met while asking for the derivative is still a prompt,
    even when the derivative then turns out not to exist and the native
    is what answers.

    The first version of this test used a rendered challenge, which is
    handled by the hold long before the fallback runs — so it recorded
    the prompt by another route entirely and stayed green when the
    carry-over was deleted. This one drives the fallback directly,
    which is the only thing that exercises that line.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(inc, [])
        calls = []

        def _attempt(url):
            calls.append(url)
            if len(calls) == 1:
                # A prompt appeared during this attempt and cleared
                # itself; the derivative then turned out not to exist.
                f.watch.challenge_at = 0.5
                f.watch.challenge_gone_at = 2.1
                f._last_title = "Just a moment..."
                raise M.NoFileAvailable("no derivative")
            return M.Fetched(data=b"PK\x03\x04zz", name="n.docx")

        f._attempt = _attempt
        M.set_request_delay(0)
        got = f.fetch("https://dc/viewcontent.cgi?article=9",
                      "https://dc/viewcontent.cgi?article=9&type=native")
        check("fallback . the native answered", got.name, "n.docx")
        check("fallback . both URLs were tried", len(calls), 2)
        check_true("fallback . the prompt is not lost to the fallback",
                   got.prompt is not None,
                   "a prompt met on the derivative is still a prompt")
        check("fallback . and it carries what was observed",
              got.prompt["observed"], "Just a moment...")


def test_both_workers_check_the_page_between_structures():
    for worker in (M.download_worker, M.retry_worker):
        check("parity . {} checks a non-file navigation too"
              .format(worker.__name__),
              len(_calls_in(worker, "hold_if_challenged")), 1)


def test_a_record_whose_planning_fails_still_gets_a_row():
    """Eleventh instance of "a partial failure must not escalate into
    total loss", caught by the live run of 2026-09-14.

    Article 1222's admin page answered HTTP 502 while the run was
    planning it. The log said "skipped" and the record then appeared in
    NO row of the report: not downloaded, not absent, not un-reached.
    1,093 of 1,094 records accounted for — and a re-run reads the
    report, so the missing one could never be recovered by any means
    the module offers.

    Both workers dropped it, in the same way, which is the shape the
    standing rule names.
    """
    src = inspect.getsource(M.download_worker)
    # The except clause itself, not the region around it. The first
    # version of this sliced from the plan call to the job loop and then
    # asserted the SLICE ended with "plan" — which is a claim about
    # where the slice stopped, not about what the row carries.
    body = src[src.index("except Exception as e:\n                    # **Eleventh"):]
    body = body[:body.index("continue")]
    check_true("planning-fail . the download worker records the record",
               "report_rows.append(" in body,
               "a record that was reached and could not be planned must "
               "still be accounted for")
    check_true("planning-fail . as a record row, so a re-run re-plans it",
               "RECORD_KIND" in body, body[-400:])
    check_true("planning-fail . carrying the plan it would have used",
               # Plan is the second-last cell since v1.34.1 appended
               # "Fetched via", which a record row leaves blank.
               '"", plan, ""))' in body,
               "without the plan cell the re-run reports it unretryable")
    check_true("planning-fail . and it is not ALSO called un-reached",
               'pending["pos"][0] = i_idx' in body,
               "two rows making two different claims about one record")

    # The re-run worker, behaviorally: a source row whose re-planning
    # raises must come out of expand() rather than vanishing inside it.
    retry = inspect.getsource(M.retry_worker)
    # Anchored on the CODE. "while re-planning" also occurs inside a
    # v1.25 comment about a stale denominator, and slicing from there
    # swept in half the worker — a source check tripping on prose,
    # which is item six on the standing list.
    seg = retry[retry.index('"article {}: {} while re-planning — recorded "'):]
    seg = seg[:seg.index("nonlocal_total[0] +=")]
    check_true("planning-fail . the re-run worker hands the row on",
               "yield skipped" in seg, seg[:400])
    check_true("planning-fail . rather than swallowing it",
               seg.count("continue") == 1, seg[:400])
    # "and writes it out without fetching it" is driven for real in
    # test_a_record_that_will_not_replan_is_still_reported below. The
    # grep that used to stand here looked for `out_rows.append(` and
    # `skip_status` inside the loop — both of which survive turning the
    # branch off, so it passed with the record dropped again.


def test_a_stopped_run_still_reports_what_the_monitoring_saw():
    """The live run held for two verification prompts and its on-screen
    summary mentioned neither.

    The line was appended where a run FINISHES. A long run almost always
    ends on Stop instead — which is how a run with a complete audit
    trail in its log and its workbook still told the operator nothing.
    """
    # Sliced to the NEXT handler, not to the end of the function: both
    # workers call verify.summary() again in their `finally`, and a
    # grep of the whole tail was satisfied by that — so deleting the
    # line from the summary itself changed nothing it looked at.
    for worker in (M.download_worker, M.retry_worker):
        src = inspect.getsource(worker)
        stop = src[src.index("except StopRequested as _stop:"):]
        stop = stop[:stop.index("    except Exception as e:")]
        check_true("stop-summary . {} says it in the stop summary"
                   .format(worker.__name__),
                   "verify.summary()" in stop,
                   "the monitoring record has to be where the operator "
                   "actually looks, and a long run ends here")
        check_true("stop-summary . {} adds it to the summary, not the log"
                   .format(worker.__name__),
                   'summary += " " + verify.summary()' in stop
                   or 'summary += " · " + verify.summary()' in stop,
                   "logging it is what the live run already did")


def test_a_record_that_will_not_replan_is_still_reported():
    """Driven, not grepped. A source row whose re-planning raises must
    come out of the re-run WITH a row — and must not be fetched, since
    its URL is an admin page.

    The grep this replaces asked whether the loop still contained
    `out_rows.append(` and `skip_status`; both survive switching the
    branch off, so it stayed green with the record dropped.
    """
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article=7&context=etd"
    rows = [_row(1, "7", M.RECORD_KIND, "", "", M.NOT_ATTEMPTED, adm, adm,
                 plan="primary,supp")]
    drv = _FakeBrowser([])
    fetched = []

    def _boom(*a, **kw):
        raise urllib.error.HTTPError(adm, 502, "Bad Gateway", {}, None)

    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "LISTING_SETTLE",
              "plan_item_jobs", "fetch_one")}
    real_log = M.log
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: drv
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            M.plan_item_jobs = _boom
            M.fetch_one = lambda *a, **kw: fetched.append(a[2])
            M.log = lambda msg: msg
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}}, out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            "reconcile": False,
                            "fetch_path": M.FETCH_BROWSER,
                            "verify_window": 0},
                           report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            M.log = real_log
        written = _rows_by_article(out)

    check_true("replan-fail . the record still has a row",
               "7" in written,
               "a record the re-run reached and could not plan must not "
               "vanish from the report a later re-run reads")
    check_true("replan-fail . and the row says what happened",
               "7" in written and "502" in written["7"][0],
               str(written.get("7")))
    check("replan-fail . and nothing was fetched for it", fetched, [])


def test_a_redirect_to_the_site_root_is_noticed_not_concluded():
    """CORRECTED in v1.34 — read this first.

    Until v1.34 this test asserted that a redirect away IS an absence.
    That rested on the reading below, which v1.33.9 showed was never
    measured, and on 2026-09-29 a redirect to a "500 Internal Server
    Error" page was written down as "No native file" because of it. A
    redirect is now noticed at once and judged by the page it lands on
    (test_a_redirect_is_judged_by_the_page_it_lands_on). The original
    text is kept below as the history of the claim.

    Measured on the live run of 2026-09-15.

    A native file was requested for every record in honors-theses, where
    most uploads were already PDFs and have no separate native. The
    content URL for a file that does not exist **redirects to the
    repository's home page** — no status code, no marker phrase, just a
    tab that ends up somewhere other than the URL that was asked for:

        BrowserFetchFailed: nothing began arriving within 8s —
        tab: '<the repository>',
        landed on: <the repository home page>

    Recorded as `Failed:` that meant one warning per record, a summary
    counting a thousand failures, and a report telling every future
    re-run that a thousand non-existent files "exist, re-run to collect
    them". A wrong claim at scale, and unrepairable afterwards without
    hand-editing the workbook.
    """
    base = "https://dc.example.edu"
    req = base + "/cgi/viewcontent.cgi?article=1003&context=etd&type=native"
    check("redirect . landing on the site root is a redirect away",
          M.redirected_away(req, base + "/"), True)
    check("redirect . landing on the requested page is not",
          M.redirected_away(req, req), False)
    check("redirect . nor is the same path with a different query",
          M.redirected_away(req, base + "/cgi/viewcontent.cgi?article=9"),
          False)
    check("redirect . a trailing slash is not a redirect",
          M.redirected_away(base + "/cgi/x/", base + "/cgi/x"), False)
    check("redirect . an unreadable tab supports no claim at all",
          M.redirected_away(req, ""), False)
    check("redirect . and neither does an unknown request",
          M.redirected_away("", base + "/"), False)

    # Behavioral, through the real fetcher: the tab renders the home page
    # and no file arrives.
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [_Tab(title="A Repository", source="<html>home</html>",
                       url=base + "/")])
        f.start_within = 3.0
        t0 = time.time()
        try:
            f.fetch(req)
            check_true("redirect . it raises", False)
        except M.NoFileAvailable as e:
            check_true("redirect . a redirect alone is not an absence",
                       False, "concluded from the redirect: {}".format(e))
        except M.BrowserFetchFailed as e:
            check_true("redirect . a redirect alone is not an absence",
                       "does not say the file is absent" in str(e), str(e))
            check_true("redirect . and it names where it landed",
                       base + "/" in str(e), str(e))
        took = time.time() - t0
        check_true("redirect . and it is answered without the file wait",
                   took < 1.0,
                   "{:.1f}s — eight seconds times a thousand records is "
                   "most of a morning".format(took))

    # A transfer already under way must NOT be called an absence just
    # because the tab has moved on. A partial is not an absence.
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [_Tab(title="A Repository", url=base + "/",
                       partial="thesis.pdf")])
        f.start_within = 0.3
        f.settle_timeout = 0.8
        try:
            f.fetch(req)
            check_true("redirect . bytes in flight are not an absence",
                       False, "it returned something")
        except M.NoFileAvailable as e:
            check_true("redirect . bytes in flight are not an absence",
                       False,
                       "a file was arriving and it was called missing: "
                       "{}".format(e))
        except M.BrowserFetchFailed as e:
            check_true("redirect . bytes in flight are not an absence",
                       True)



def test_a_redirect_from_the_page_it_redirects_to_is_still_a_redirect():
    """Measured on the live off-campus run of 2026-09-21: 51 rows.

        BrowserFetchFailed: nothing began arriving within 8s
        — tab: '<the repository>',
          landed on: <the repository home page>

    v1.33.3 named this answer and v1.33.5 then made it unreadable: the
    redirect check was put behind "did the tab move?", and a redirect
    to the home page LEAVES THE TAB ON THE HOME PAGE. So the second
    such answer looks identical to the first, a download never moves
    the tab back, and the tab rests there from the start — which is
    why not one of the 51 was named, including the first.

    Two cases, differing in exactly one thing: where the tab was
    sitting before the navigation. Both are the same request, the same
    server answer and the same absent file, so both must reach the same
    verdict. Before this release the first one did not.
    """
    import tempfile
    base = "https://dc.example.edu"
    req = base + "/cgi/viewcontent.cgi?article=1003&context=etd&type=native"
    home = _Tab(title="A Repository", source="<html>home</html>",
                url=base + "/")

    for where, at in (("tab already on the home page",
                       ("A Repository", base + "/", "<html>home</html>")),
                      ("tab on some other page",
                       ("A Thesis", base + "/etd/7/", "<html>rec</html>"))):
        with tempfile.TemporaryDirectory() as d:
            inc = os.path.join(d, "_incoming")
            f, drv = _browser_fetcher(inc, [home], at=at)
            # THE FIXTURE ITSELF IS THE EXPERIMENT, so assert it took.
            # A fake that quietly starts blank makes every navigation a
            # move and this whole test passes while testing nothing.
            check("same-page . {} . the tab starts there".format(where),
                  drv.current_url, at[1])
            # Long enough that the two branches are TELLABLE APART. With
            # a fifth of a second of patience, answering after the wait
            # looks the same as answering at once, and planting showed
            # the post-wait branch silently covering for the fast one.
            f.start_within = 3.0
            f.settle_timeout = 3.5
            t0 = time.time()
            try:
                f.fetch(req)
                check_true("same-page . {} . it raises".format(where), False)
            except M.NoFileAvailable as e:
                check_true("same-page . {} . a redirect is noticed, and is "
                           "not by itself an absence".format(where), False,
                           "concluded from the redirect: {}".format(e))
            except M.BrowserFetchFailed as e:
                check_true("same-page . {} . a redirect is noticed, and is "
                           "not by itself an absence".format(where),
                           "does not say the file is absent" in str(e),
                           str(e))
            took = time.time() - t0
            check_true("same-page . {} . and answered without waiting for "
                       "a file the redirect ruled out".format(where),
                       took < 1.0,
                       "{:.1f}s — the rendered page was read only after "
                       "the wait".format(took))

    # THE SAME, AFTER THE WAIT. A navigation that stalls and then shows
    # the home page takes the other branch, and the two branches asked
    # the question separately before this release.
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [_Tab(title="A Repository", source="<html>home</html>",
                       url=base + "/", stall=True)],
            at=("A Repository", base + "/", "<html>home</html>"))
        f.start_within = 0.2
        f.settle_timeout = 0.6
        try:
            f.fetch(req)
            check_true("same-page . after the wait . it raises", False)
        except M.NoFileAvailable as e:
            check_true("same-page . after the wait . a redirect is noticed, "
                       "and is not by itself an absence", False,
                       "concluded from the redirect: {}".format(e))
        except M.BrowserFetchFailed as e:
            check_true("same-page . after the wait . a redirect is noticed, "
                       "and is not by itself an absence",
                       "does not say the file is absent" in str(e), str(e))

    # A TRANSFER ALREADY UNDER WAY IS NOT AN ABSENCE, and it must not
    # become one now that a fresh document no longer needs the tab to
    # look different. This is the guard the wider rule leans on.
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [_Tab(title="A Repository", source="<html>home</html>",
                       url=base + "/", partial="thesis.pdf")],
            at=("A Repository", base + "/", "<html>home</html>"))
        f.start_within = 0.3
        f.settle_timeout = 0.8
        try:
            f.fetch(req)
            check_true("same-page . bytes in flight are not an absence",
                       False, "it returned something")
        except M.NoFileAvailable as e:
            check_true("same-page . bytes in flight are not an absence",
                       False,
                       "a file was arriving and it was called missing: "
                       "{}".format(e))
        except M.BrowserFetchFailed:
            check_true("same-page . bytes in flight are not an absence", True)


def test_a_browser_that_will_not_say_falls_back_rather_than_claims():
    """Document identity is a question the browser can refuse.

    When it does, the answer must be the older, weaker appearance test
    — not a guess. A tab that did not move then supports no claim, the
    way it did before this release, and the run reports what it saw
    instead of concluding an absence it cannot see.
    """
    import tempfile
    base = "https://dc.example.edu"
    req = base + "/cgi/viewcontent.cgi?article=1003&context=etd&type=native"
    home = _Tab(title="A Repository", source="<html>home</html>",
                url=base + "/")

    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [home], at=("A Repository", base + "/", "<html>home</html>"))
        drv._no_script = True
        f.start_within = 0.2
        f.settle_timeout = 0.5
        check("fallback . the probe returns nothing at all", f.doc_id(), "")
        try:
            f.fetch(req)
            check_true("fallback . it raises", False)
        except M.NoFileAvailable as e:
            check_true("fallback . no identity and no movement supports no "
                       "claim", False,
                       "claimed an absence with nothing to go on: "
                       "{}".format(e))
        except M.BrowserFetchFailed:
            check_true("fallback . no identity and no movement supports no "
                       "claim", True)

    # And the appearance test still works when it can see something.
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [home], at=("A Thesis", base + "/etd/7/", "<html>r</html>"))
        drv._no_script = True
        f.start_within = 0.2
        f.settle_timeout = 0.5
        try:
            f.fetch(req)
            check_true("fallback . moved . it raises", False)
        except M.NoFileAvailable as e:
            check_true("fallback . a tab that moved is still readable "
                       "without the probe", False,
                       "concluded from the redirect: {}".format(e))
        except M.BrowserFetchFailed as e:
            # v1.34: a redirect is noticed and named, then judged by the
            # page — which here says nothing, so nothing is concluded.
            check_true("fallback . a tab that moved is still readable "
                       "without the probe",
                       "sent the request to" in str(e), str(e))


def test_a_download_leaves_the_document_alone():
    """The premise the whole rule rests on, asserted rather than assumed.

    If a download bumped the document the way a page does, every
    successful fetch would look like a fresh answer from the server and
    the redirect rule would fire on all of them.
    """
    import tempfile
    base = "https://dc.example.edu"
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [_Tab(drop=("thesis.pdf", b"%PDF-1.4 x"))],
            at=("A Repository", base + "/", "<html>home</html>"))
        f.start_within = 1.0
        f.settle_timeout = 1.0
        was = f.doc_id()
        got = f.fetch(base + "/cgi/viewcontent.cgi?article=1003")
        check("download . the file is returned", got.name, "thesis.pdf")
        check("download . and the document never changed", f.doc_id(), was)


def test_a_prompt_that_interrupts_a_transfer_does_not_cost_the_transfer():
    """Open item 4, measured at last — and the module failed it.

    2026-09-15, article 1014, a video. The prompt appeared, the operator
    cleared it in 97 seconds, and clearing it restarted the download.
    The hold ended because bytes had begun arriving — and the next thing
    the code did was DELETE them and navigate again:

        Verification cleared in 97s — resuming.
        Cleared 1 unfinished transfer(s) before retrying:
            MalloryTaylor_Capstone.mp4.crdownload

    The fresh request drew a fresh interstitial, twice, two seconds
    apart, and the bound stopped the run — 1,084 records — on a
    situation the module had created. The audit trail recorded three
    prompts where a person had answered one.

    Two rules come out of it, and both are asserted here: a hold that
    ends in bytes is followed by WAITING, and a partial that is still
    growing is a live transfer that must not be deleted.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        os.makedirs(inc)
        w = M.DirectoryDownloadWatch(inc)
        w.begin()
        part = os.path.join(inc, "video.mp4.crdownload")
        with open(part, "wb") as fh:
            fh.write(b"x" * 1000)

        # A growing partial is a transfer, not litter.
        import threading
        def _grow():
            time.sleep(0.4)
            with open(part, "ab") as fh:
                fh.write(b"y" * 5000)
        t = threading.Thread(target=_grow, daemon=True)
        t.start()
        gone, growing = w.clear_partials(settle=1.2)
        t.join()
        check("midflight . a growing partial is not deleted", gone, [])
        check("midflight . it is named as still arriving", growing,
              ["video.mp4.crdownload"])
        check_true("midflight . and it is still on disk",
                   os.path.exists(part),
                   "deleting it is what cost the run")

        # One that has stopped growing is stranded and goes.
        gone, growing = w.clear_partials(settle=0.3)
        check("midflight . a partial that stopped growing is stranded",
              gone, ["video.mp4.crdownload"])
        check("midflight . and nothing is left growing", growing, [])


def test_a_hold_that_ends_in_bytes_waits_instead_of_asking_again():
    """The whole failure, driven end to end through fetch_one.

    The navigation stalls and a prompt appears; a person clears it after
    a moment, which starts the download; the file then completes. The
    run must come back with the file and ONE prompt recorded — not
    re-navigate, not delete the partial, and not count a prompt per
    retry.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        first = _Tab(title="Just a moment...", stall=True)
        # v1.34.1: a transfer that TAKES TIME, and no total ceiling on the
        # wait — which is what production has had since v1.34
        # (BROWSER_SETTLE_TIMEOUT = None). This test failed once on a
        # Windows runner (2026-09-30, "nothing completed within 1s", on a
        # SECOND navigation): the helper's 1 s ceiling, older than v1.34,
        # let a slow transfer outlast the resume, and the module asked for
        # the file again. Planting the old ceiling back with this
        # three-second transfer makes that happen every time, which is
        # the "requested only once" check below going red.
        first.timed_transfer = (0.6, 3.0, "video.mp4",
                                b"%PDF-video-stands-in")
        f, drv = _browser_fetcher(inc, [first])
        f.start_within = 5.0
        f.settle_timeout = M.BROWSER_SETTLE_TIMEOUT
        v = M.VerificationWatch(
            window=30, poll=0.2,
            sleep=lambda secs: (time.sleep(min(secs, 0.2)) or True),
            stopping=lambda: False, now=time.time)
        got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1014")
        check("bytes-hold . the interrupted file is returned", got.name,
              "video.mp4")
        check("bytes-hold . exactly one prompt is recorded",
              len(v.prompts), 1)
        check("bytes-hold . and the URL was requested only once",
              len(drv.visited), 1)
        check_true("bytes-hold . the file is on disk, not deleted",
                   got.nbytes > 0, str(got.nbytes))


def test_a_redirect_is_judged_by_the_page_it_lands_on():
    """v1.34: what a redirect says is on the page, not in the redirect.

    Measured 2026-09-29. A native with no file lands on the CGI form of
    its own URL, which says — title and body worded differently —
    "Sorry, that file does not exist" / "Sorry, that file doesn't
    exist." That IS an absence, in the page's own words. And article
    1905's native landed on a "500 Internal Server Error" page and was
    written "No native file" because it was a redirect. A server error is
    not an absence.
    """
    import tempfile
    base = "https://dc.example.edu"
    req = base + "/context/etd/article/1905/type/native/viewcontent"
    cgi = (base + "/cgi/viewcontent.cgi?params=/context/etd/article/1905/"
           "type/native/&path_info=&preview_mode=1")
    cases = (
        ("title only", "Sorry, that file does not exist",
         "<html><body>Search</body></html>", True),
        ("body only", "Sorry",
         "<html><body>Sorry, that file doesn't exist.</body></html>", True),
        ("a server error", "500 Internal Server Error",
         "<html><body>Internal Server Error</body></html>", False),
    )
    with tempfile.TemporaryDirectory() as d:
        for n, (what, title, source, absent) in enumerate(cases):
            f, drv = _browser_fetcher(
                os.path.join(d, "_{}".format(n)),
                [_Tab(title=title, source=source, url=cgi)])
            try:
                f.fetch(req)
                check_true("redirect-page . {} . it raises".format(what),
                           False)
            except M.NoFileAvailable as e:
                check_true("redirect-page . {} . {}".format(
                    what, "is an absence in the page's words" if absent
                    else "is NOT an absence"), absent, str(e))
            except M.BrowserFetchFailed as e:
                check_true("redirect-page . {} . {}".format(
                    what, "is an absence in the page's words" if absent
                    else "is NOT an absence"), not absent, str(e))


def test_a_redirect_found_after_the_wait_is_noticed_too():
    """The same rule, on the path where no page rendered.

    Planting showed the post-wait redirect check was unguarded: the only
    fixture for it rendered a page immediately, which is read by the
    fast branch, so deleting the other check changed nothing. A
    navigation that STALLS and then shows the home page has to reach the
    same verdict.
    """
    import tempfile
    base = "https://dc.example.edu"
    req = base + "/cgi/viewcontent.cgi?article=1003&context=etd&type=native"
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc, [_Tab(title="A Repository", source="<html>home</html>",
                       url=base + "/", stall=True)])
        f.start_within = 0.2
        f.settle_timeout = 0.6
        try:
            f.fetch(req)
            check_true("redirect-late . it raises", False)
        except M.NoFileAvailable as e:
            check_true("redirect-late . a stalled navigation that landed "
                       "on another page is noticed, and is not an absence",
                       False, "concluded from the redirect: {}".format(e))
        except M.BrowserFetchFailed as e:
            check_true("redirect-late . a stalled navigation that landed "
                       "on another page is noticed, and is not an absence",
                       "does not say the file is absent" in str(e), str(e))


def test_arriving_means_since_this_fetch_began():
    """Not "is the incoming folder non-empty".

    A leftover from anything else would otherwise read as "the file is
    coming", which ends a hold that nobody has answered — the run would
    resume on the strength of a file that has nothing to do with it.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        os.makedirs(inc)
        # Something is already sitting there when this fetch starts.
        with open(os.path.join(inc, "leftover.pdf"), "wb") as fh:
            fh.write(b"%PDF-old")
        f, drv = _browser_fetcher(inc, [])
        f.watch.begin()
        check("arriving . a leftover file is not this fetch's file",
              f.arriving(), False)

        clock = _Clock()
        v = M.VerificationWatch(window=10, poll=2, sleep=clock.sleep,
                                stopping=lambda: False, now=clock.now)
        took, why = v.hold(lambda: True, f.arriving)
        check("arriving . so it does not end a hold nobody answered",
              took, -1.0)
        check("arriving . and no reason is claimed", why, "")

        # A file that arrives AFTER begin() does end it.
        with open(os.path.join(inc, "new.pdf.crdownload"), "wb") as fh:
            fh.write(b"half")
        check("arriving . a new transfer does count", f.arriving(), True)


def test_whatever_ended_the_hold_a_transfer_in_progress_is_waited_for():
    """One rule, exercised on both ways a hold can end.

    The first version had two branches — "the hold ended in bytes" and
    "a partial is still growing" — and planting showed the second was
    all but unreachable. Both cases are now the same question: is
    anything coming? The assertion that matters either way is that the
    URL is requested ONCE.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        first = _Tab(title="Just a moment...", stall=True)
        first.timed_clear = (0.5, "big.mp4", b"%PDF-stands-in-for-a-video")
        f, drv = _browser_fetcher(inc, [first])
        f.settle_timeout = M.BROWSER_SETTLE_TIMEOUT  # as production (v1.34.1)
        f.start_within = 5.0
        v = M.VerificationWatch(
            window=30, poll=0.2,
            sleep=lambda secs: (time.sleep(min(secs, 0.2)) or True),
            stopping=lambda: False, now=time.time)
        got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1014")
        check("one-rule . the transfer is waited for, not re-requested",
              len(drv.visited), 1)
        check("one-rule . and it is what comes back", got.name, "big.mp4")
        check("one-rule . with one prompt recorded", len(v.prompts), 1)


def test_a_frozen_challenge_page_is_not_a_new_prompt_every_time():
    """The most expensive defect of the week, and the cheapest to state.

    2026-09-15, honors-theses: 148 verification prompts recorded in 858
    file requests. Three had work between them — 239, 241 and 230 files
    — and were real. **The other 145 showed "0 files since previous"
    and "cleared in 2s", one per download**, and none of them happened.

    A download does not navigate the tab. After the third real
    challenge the tab was left on Cloudflare's own page and stayed
    there, and watching the tab while waiting read that frozen page as
    a fresh prompt every single time. The audit trail — the artifact
    this feature exists to produce — was 98% fiction.

    The rule: a challenge already on screen BEFORE the navigation is
    not news about this request. A real one still cannot be missed,
    because it stops the file arriving and the check after the wait is
    decisive.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        # The tab is parked on a challenge page, as it is for the whole
        # rest of a run once one has been served. Ten ordinary
        # downloads follow; the page never moves, because downloads
        # never move it.
        # The file must arrive after a beat. With an instant drop,
        # settle finds it on its first look and never glances at the
        # tab at all — so the test could not tell a watched tab from an
        # unwatched one, and the plant that removed the watching passed.
        tabs = []
        for n in range(10):
            t = _Tab(stall=True)
            t.timed_clear = (0.25, "f{}.pdf".format(n),
                             b"%PDF-" + bytes([65 + n]))
            tabs.append(t)
        f, drv = _browser_fetcher(inc, tabs)
        f.settle_timeout = M.BROWSER_SETTLE_TIMEOUT  # as production (v1.34.1)
        f.start_within = 5.0
        drv._title = "Just a moment..."
        drv._source = "<html>Enable JavaScript and cookies</html>"
        drv.current_url = "https://challenges.example/cdn-cgi/x"
        v = M.VerificationWatch(
            window=30, poll=0.2,
            sleep=lambda secs: (time.sleep(min(secs, 0.2)) or True),
            stopping=lambda: False, now=time.time)
        M.set_request_delay(0)
        for n in range(10):
            got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article={}"
                              .format(1000 + n))
            check_true("stale . download {} still arrives".format(n),
                       got.name == "f{}.pdf".format(n), got.name)
        check("stale . and NOT ONE of them is recorded as a prompt",
              len(v.prompts), 0)
        check_true("stale . so the summary does not invent them",
                   "No verification prompt was seen" in v.summary(),
                   v.summary())
        check("stale . and the tab was navigated once per file, no more",
              len(drv.navigations), 10)


def test_a_real_prompt_is_still_caught_while_the_tab_is_frozen():
    """The other half, and the one that makes the rule safe.

    With the tab already parked on a challenge page, a genuinely
    challenged file cannot be told apart by LOOKING — but it can be
    told apart by what happens: no file arrives. The check after the
    wait is decisive precisely because a challenge blocks the transfer.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        # No clearing: the point is that the tab STAYS a challenge and
        # no file arrives. Setting clear_after_reads here made the fake
        # clear the page under the test, which is the fixture answering
        # the question instead of the code.
        f, drv = _browser_fetcher(inc, [_Tab(stall=True)])
        f.start_within = 0.2
        f.settle_timeout = 0.5
        drv._title = "Just a moment..."
        drv._source = "<html>Enable JavaScript and cookies</html>"
        drv.current_url = "https://challenges.example/cdn-cgi/x"
        try:
            f.fetch("https://dc/viewcontent.cgi?article=1")
            check_true("stale-real . a blocked file still raises", False)
        except M.VerificationRequired:
            check_true("stale-real . a real prompt is caught even with a "
                       "frozen tab", True)
        except Exception as e:
            check_true("stale-real . a real prompt is caught even with a "
                       "frozen tab", False,
                       "raised {} instead".format(e.__class__.__name__))


def test_a_tab_that_never_moved_is_not_a_redirect():
    """The latent half of Friday's redirect rule.

    "The tab is on a different URL than I asked for" is true of EVERY
    download attempt, because the tab never moves for a download. Read
    as a redirect it would call any file that failed to arrive an
    absence — the server saying "there is nothing here" when it had
    said nothing at all. Only the MOVE carries the meaning.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(inc, [_Tab(stall=True)])
        f.start_within = 0.2
        f.settle_timeout = 0.5
        # Parked on some earlier page, and this navigation moves nothing.
        drv._title = "Some earlier page"
        drv._source = "<html>whatever</html>"
        drv.current_url = "https://dc/cgi/viewcontent.cgi?article=999"
        try:
            f.fetch("https://dc/cgi/viewcontent.cgi?article=1000")
            check_true("no-move . it raises", False)
        except M.NoFileAvailable as e:
            check_true("no-move . a tab that never moved is not an absence",
                       False,
                       "called it missing on the strength of a page that "
                       "belongs to the previous request: {}".format(e))
        except M.BrowserFetchFailed:
            check_true("no-move . a tab that never moved is not an absence",
                       True)


def test_the_tab_is_put_back_after_a_prompt_so_it_can_speak_again():
    """Otherwise the detector goes deaf for the rest of the run.

    Once the tab is parked on a challenge page, the only way a later
    prompt is noticed is the file failing to arrive. Moving the tab off
    it — after the file is safely in hand, never during a transfer —
    restores the cheaper signal.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(inc, [_Tab(title="A Repository",
                                             source="<html>home</html>",
                                             url="https://dc/")])
        drv._title = "Just a moment..."
        drv._source = "<html>Enable JavaScript and cookies</html>"
        check("normalize . the tab starts stuck",
              f.tab_is_stuck_on_a_challenge(), True)
        note = f.normalize_tab("https://dc/")
        check("normalize . moving it off reports no problem", note, "")
        check("normalize . and it is no longer stuck",
              f.tab_is_stuck_on_a_challenge(), False)

        # It must not act when there is nothing to fix.
        f2, drv2 = _browser_fetcher(inc, [])
        drv2._title = "An ordinary page"
        check("normalize . an ordinary tab is left alone",
              f2.normalize_tab("https://dc/"), "")
        # EVERY navigation, not only the file ones: `visited` filters to
        # viewcontent URLs, so it could not see a root navigation at all
        # and this check passed with normalize_tab moving every tab.
        check("normalize . with no navigation at all",
              len(drv2.navigations), 0)


def test_a_prompt_cleared_while_waiting_gives_the_file_another_chance():
    """The third way a prompt ends, and it had no check at all.

    Found 2026-09-22 by planting: disabling the branch left every one
    of 1,263 tests green.

    The other two endings are covered — bytes arrive, or the hold ends
    and the file is re-requested. This is the one INSIDE a single
    attempt: the prompt was never raised, because it appeared while
    `settle` was already watching and was answered before the wait ran
    out. The page is no longer a challenge and no file came, which
    looks exactly like an absence — and the clearance is precisely the
    thing that gives the file a fresh chance to arrive.

    Getting it wrong costs the file AND the prompt: the run would call
    a present file missing, and the audit trail would not record that
    anybody was asked.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        # A challenge that renders nothing, cleared after a few looks —
        # and the file lands AT THE MOMENT of clearing, which is what
        # answering the question actually does.
        tab = _Tab(title="Just a moment...", stall=True,
                   drop_on_clear=("c.pdf", b"%PDF-c"))
        # ONE look. `settle` returns the moment its observer sees a
        # challenge, so the clearance has to land between that look and
        # the decisive read — which is exactly the window this branch
        # exists for, and why a larger number puts the test on the
        # other path entirely.
        f, drv = _browser_fetcher(inc, [tab], clear_after_reads=1)
        f.start_within = 1.5
        f.settle_timeout = 2.0
        got = f.fetch("https://dc/viewcontent.cgi?article=5")
        check("cleared-waiting . the file is returned", got.name, "c.pdf")
        check("cleared-waiting . and it took ONE navigation, because the "
              "clearance came during the wait", len(drv.visited), 1)
        check_true("cleared-waiting . the prompt is carried out, not lost",
                   got.prompt is not None,
                   "a prompt answered during the wait still happened, and "
                   "an audit trail that omits it is the v1.33.1 defect")
        if got.prompt:
            check_true("cleared-waiting . and it says how long it took",
                       got.prompt.get("cleared_in") is not None,
                       str(got.prompt))


def test_both_workers_put_the_tab_back_after_a_prompt():
    check("unstick . fetch_one takes the page to return to",
          "home_url" in inspect.signature(M.fetch_one).parameters, True)
    for worker in (M.download_worker, M.retry_worker):
        src = inspect.getsource(worker)
        # A window after the CALL, anchored on the assignment so a
        # comment mentioning fetch_one() cannot satisfy it — that exact
        # mistake is item seventeen on the standing list. The first
        # version sliced to the next ")", which is the call's own first
        # argument.
        seg = src[src.index("= fetch_one("):][:700]
        check_true("unstick . {} passes it".format(worker.__name__),
                   "home_url=base_url" in seg,
                   "without it the tab stays on the verification page "
                   "and the cheap signal goes deaf for the rest of the "
                   "run")

    # Behavioral. A HOLD ENDS TWO WAYS and both are legitimate: the bytes
    # start arriving, or the page stops being a challenge. Until
    # 2026-09-22 only the first was ever exercised — the second was a
    # RACE this test happened to lose on fast machines and win on a
    # Windows runner, where it surfaced as
    #
    #   NoFileAvailable: the server sent the request to https://dc/
    #                    instead of serving a file
    #
    # which read like v1.33.7 calling a parked tab a redirect. It was
    # not. `normalize_tab()` goes to a NON-viewcontent URL and the fake
    # does not spend a scripted tab on those, so the second tab —
    # written for the put-back — sat unconsumed until a retry reached
    # it and was told the site root. A leftover scripted answer is a
    # landmine that only the other timing steps on.
    #
    # Both endings are now driven deliberately, so neither depends on
    # which of two timers wins.
    import tempfile

    # (1) the hold ends because BYTES ARRIVED. No retry, so one
    #     navigation, and the tab is still on the challenge until the
    #     put-back moves it — which is what this test is named for.
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        first = _Tab(title="Just a moment...", stall=True)
        # v1.34.1: a transfer that takes time, and production's wait. This
        # case failed once on Windows/3.13 (f91b54f, the same "nothing
        # completed within 1s" on a second navigation) under the helper's
        # old 1 s ceiling — the same family as the bytes-hold test.
        first.timed_transfer = (0.4, 3.0, "a.pdf", b"%PDF-a")
        f, drv = _browser_fetcher(inc, [first])
        f.settle_timeout = M.BROWSER_SETTLE_TIMEOUT  # as production (v1.34.1)
        f.start_within = 5.0
        v = M.VerificationWatch(
            window=30, poll=0.2,
            sleep=lambda secs: (time.sleep(min(secs, 0.2)) or True),
            stopping=lambda: False, now=time.time)
        got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1",
                          home_url="https://dc/")
        check("unstick . bytes ended the hold . the file is returned",
              got.name, "a.pdf")
        check("unstick . bytes ended the hold . the file was requested "
              "once", len(drv.visited), 1)
        check_true("unstick . bytes ended the hold . the tab no longer "
                   "reads as a challenge",
                   not f.tab_is_stuck_on_a_challenge(),
                   "left stuck, the next 145 downloads each look like a "
                   "fresh prompt")

    # (2) the hold ends because THE PAGE CLEARED and nothing has
    #     arrived. The module re-requests the file, which is correct,
    #     and the server serves it — which is what a real one does
    #     after a person answers. The fixture used to answer that
    #     retry with the home page and the run was told the file did
    #     not exist.
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(
            inc,
            [_Tab(title="Just a moment...", stall=True),
             _Tab(drop=("a.pdf", b"%PDF-a"))],
            clear_after_reads=3)
        f.start_within = 0.5
        f.settle_timeout = 1.0
        v = M.VerificationWatch(
            window=30, poll=0.2,
            sleep=lambda secs: (time.sleep(min(secs, 0.2)) or True),
            stopping=lambda: False, now=time.time)
        try:
            got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1",
                              home_url="https://dc/")
            check("unstick . a cleared page ended the hold . the file is "
                  "returned", got.name, "a.pdf")
        except M.NoFileAvailable as e:
            check_true("unstick . a cleared page ended the hold . the "
                       "file is returned", False,
                       "a retry after a cleared prompt was read as an "
                       "absence: {}".format(e))
        check("unstick . a cleared page ended the hold . the file was "
              "re-requested exactly once", len(drv.visited), 2)
        check("unstick . a cleared page ended the hold . exactly one "
              "prompt is recorded", len(v.prompts), 1)

    # And nothing is navigated when there was no prompt at all.
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, drv = _browser_fetcher(inc, [_Tab(drop=("b.pdf", b"%PDF-b"))])
        v = M.VerificationWatch(window=30, sleep=_Clock().sleep,
                                stopping=lambda: False, now=_Clock().now)
        M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=2",
                    home_url="https://dc/")
        check("unstick . an ordinary fetch navigates once and no more",
              len(drv.visited), 1)


def test_a_completed_navigation_that_changed_nothing_is_read_as_nothing():
    """The fast path's half of the same rule.

    A navigation can complete without rendering anything new — and the
    page then on screen belongs to the PREVIOUS request. Reading it is
    how a tab parked on the repository home page turns every failed
    fetch into "the server redirected me, so there is no file".

    Driven through the rendered branch specifically: the no-move test
    above stalls, which skips this branch entirely, so the guard here
    had no check of its own.
    """
    import tempfile
    base = "https://dc.example.edu"
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        # A navigation that completes and renders nothing (no title, no
        # source, no drop, no stall), with the tab already parked on the
        # home page from something earlier.
        f, drv = _browser_fetcher(inc, [_Tab()])
        f.start_within = 0.2
        f.settle_timeout = 0.5
        drv._title = "A Repository"
        drv._source = "<html>home</html>"
        drv.current_url = base + "/"
        try:
            f.fetch(base + "/cgi/viewcontent.cgi?article=1000")
            check_true("rendered-nomove . it raises", False)
        except M.NoFileAvailable as e:
            check_true("rendered-nomove . a page that did not change is "
                       "not this request's answer", False,
                       "called the file missing on the strength of the "
                       "previous request's page: {}".format(e))
        except M.BrowserFetchFailed:
            check_true("rendered-nomove . a page that did not change is "
                       "not this request's answer", True)


def test_a_resume_does_not_re_report_the_prompt_it_is_resuming_from():
    """Live proof, 2026-09-18, off campus: three prompts recorded where
    two had happened.

        1 | 10:51:53 | at request 12  | 11 files since | Cleared while
                                                          the run waited
        2 | 10:52:10 | at request 12  |  0 files since | Cleared without
                                                          stopping the run
        3 | 11:23:42 | at request 279 | 266 files since

    Row 2 is the same file request as row 1, seventeen seconds later.
    v1.33.5 taught `_attempt` to ignore a challenge page that was
    already on screen and did not teach `resume`, so the resume watched
    the frozen tab and filed what it saw — the thirteenth time a rule
    in this module reached one of a pair and not the other.

    The fixture is a transfer that BEGINS and completes seventeen
    ticks apart, because an instant download is found on settle's first
    look and the tab is never glanced at. That is why the suite could
    not reproduce this.
    """
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        first = _Tab(title="Just a moment...", stall=True)
        first.timed_transfer = (0.3, 1.2, "slow.pdf", b"%PDF-slow-thesis")
        f, drv = _browser_fetcher(inc, [first])
        f.settle_timeout = M.BROWSER_SETTLE_TIMEOUT  # as production (v1.34.1)
        f.start_within = 5.0
        v = M.VerificationWatch(
            window=30, poll=0.2,
            sleep=lambda secs: (time.sleep(min(secs, 0.2)) or True),
            stopping=lambda: False, now=time.time)
        got = M.fetch_one(f, v, "https://dc/viewcontent.cgi?article=1014",
                          home_url="https://dc/")
        check("resume-prompt . the slow transfer completes", got.name,
              "slow.pdf")
        check("resume-prompt . ONE prompt, not one per resume",
              len(v.prompts), 1)
        check_true("resume-prompt . and it is the one that stopped the run",
                   v.prompts and v.prompts[0]["paused"] is True,
                   str(v.prompts))
        check("resume-prompt . the URL was requested once", len(drv.visited),
              1)


# ---------------------------------------------------------------------------
# v1.33.9. THE BROWSER PATH'S ONE BLIND SPOT, AND ITS EXIT.
#
# Driving a real browser buys the ability to answer a verification prompt
# and pays for it in status codes. For most outcomes the tab substitutes.
# For a 404 it does not substitute at all: measured 2026-09-22, the
# navigation never commits, no page renders, the document identity is
# unchanged and nothing arrives — so the tab is not merely ambiguous, it
# is silent. Two off-campus runs wrote 51 such rows as `Failed:` while the
# server had answered HTTP 404 both times.
#
# THE FIXTURE MUST BEHAVE LIKE THE WORLD. Every case below starts the tab
# ALREADY on the home page and stalls the navigation, because that is what
# the probe measured — a fake that starts blank, or that renders something,
# reaches a different branch and proves nothing about this one.
# ---------------------------------------------------------------------------
class _NetResp:
    """A response the module's own urllib path can read."""

    def __init__(self, body, headers=None, landed="", status=200):
        self._body = body
        self.headers = dict(headers or {})
        self._landed = landed
        self.status = status

    def read(self, n=None):
        # Counted: v1.34 asks about natives with headers only, and the
        # only way to know a body was NOT read is to count the reads.
        self.reads = getattr(self, "reads", 0) + 1
        # A STREAM, as a real response is: each read continues where the
        # last stopped, and the end is b"". Returning the same prefix on
        # every call was harmless while the module read bodies in one
        # gulp; v1.34 streams them to disk a megabyte at a time, and a
        # fake that never reaches its end makes that loop forever.
        pos = getattr(self, "_pos", 0)
        chunk = (self._body[pos:] if n is None or n < 0
                 else self._body[pos:pos + n])
        self._pos = pos + len(chunk)
        return chunk

    def geturl(self):
        return self._landed

    def __enter__(self):
        # Each request reads the body from its start, even when a test
        # hands the same response to several requests.
        self._pos = 0
        return self

    def __exit__(self, *a):
        return False


def _challenge_error(code=403, header=True, body=True):
    """A refusal that says it is a challenge — by header, by body, or both.

    Two INDEPENDENT signals, so they get two independent cases. A single
    fixture carrying both is satisfied by either one working, which makes
    the other untested: planting the header check away left this suite
    green because the body markers covered for it.
    """
    return urllib.error.HTTPError(
        "http://x/y", code, "Forbidden",
        ({"cf-mitigated": "challenge", "server": "cloudflare"} if header
         else {"server": "cloudflare"}),
        io.BytesIO(b"<html><title>Just a moment...</title></html>" if body
                   else b"<html><body>Forbidden.</body></html>"))


def _silent_browser(folder, script=None):
    """The measured live shape: tab already parked, navigation stalls.

    Returns the fetcher, the driver, and the URL to ask for.

    SAID PLAINLY BECAUSE PLANTING SHOWED IT: parking the tab on the home
    page first is realism, not load-bearing. Replacing `at=` with a blank
    tab leaves every check below green, because the second opinion does
    not depend on where the tab was — that is the point of it. The
    positioning is kept so the fixture matches what the probe measured
    on 2026-09-22, and so a future reader is not misled about the shape
    of the live case; it is NOT evidence that anything here is tested.
    """
    f, drv = _browser_fetcher(
        folder, script if script is not None else [_Tab(stall=True)],
        at=("Example Commons", "https://dc.example.edu/",
            "<html><body>welcome</body></html>"))
    return f, drv, ("https://dc.example.edu/context/honors-theses/"
                    "article/1003/type/native/viewcontent?preview_mode=1")


def _asking(answer):
    """urlopen that answers file requests with `answer`, pages with a page.

    Scoped to viewcontent URLs so a worker's own admin traffic is not
    turned into the thing under test — a fixture that breaks everything
    proves nothing about the one path being measured.
    """
    asked = []

    def fake(req, timeout=None):
        url = getattr(req, "full_url", None) or str(req)
        if "viewcontent" not in url:
            return _NetResp(b"<html><body>admin</body></html>", landed=url)
        asked.append(url)
        if isinstance(answer, Exception):
            raise answer
        if callable(answer):
            return answer(url)
        return answer
    fake.asked = asked
    return fake


def test_a_404_the_browser_cannot_hear_is_still_an_answer():
    """The defect this pins: seven days of `Failed:` for a plain 404.

    Behavioral, through the real fetcher and the real fetch_file, with
    only the socket replaced.
    """
    import tempfile
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        f, _drv, url = _silent_browser(os.path.join(d, "_incoming"))
        f.use_cookie("c=1")
        M.reset_rate_state()
        try:
            urllib.request.urlopen = _asking(_http(404, "Not Found"))
            try:
                got = f.fetch(url)
            except M.NoFileAvailable as e:
                check_true("second-opinion . a 404 the tab cannot show is "
                           "still a definite absence", "404" in str(e),
                           str(e))
                check_true("second-opinion . and it says where the answer "
                           "came from", "network layer" in str(e), str(e))
            except Exception as e:
                check_true("second-opinion . a 404 the tab cannot show is "
                           "still a definite absence", False,
                           "got {}: {}".format(type(e).__name__, e))
            else:
                check_true("second-opinion . a 404 the tab cannot show is "
                           "still a definite absence", False,
                           "it returned {!r}".format(got.name))
            # 410 is the same answer in different clothing.
            f2, _d2, url2 = _silent_browser(os.path.join(d, "_inc2"))
            f2.use_cookie("c=1")
            urllib.request.urlopen = _asking(_http(410, "Gone"))
            try:
                f2.fetch(url2)
            except M.NoFileAvailable:
                check_true("second-opinion . 410 is an absence too", True)
            except Exception as e:
                check_true("second-opinion . 410 is an absence too", False,
                           "got {}: {}".format(type(e).__name__, e))
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_the_second_opinion_keeps_the_bytes_when_there_are_bytes():
    """A file urllib can reach must not be written down as anything else.

    Recording "it exists but was not fetched" would be a row every
    re-run reads and retries through the same blind spot — a loop, not
    a report.
    """
    import tempfile
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        f, _drv, url = _silent_browser(os.path.join(d, "_incoming"))
        f.use_cookie("c=1")
        M.reset_rate_state()
        try:
            urllib.request.urlopen = _asking(_NetResp(
                b"%PDF-1.7 real bytes",
                {"Content-Type": "application/pdf",
                 "Content-Disposition": 'inline; filename="thesis.pdf"'},
                landed=url))
            got = f.fetch(url)
            check("second-opinion . the file comes back named", got.name,
                  "thesis.pdf")
            check("second-opinion . with its bytes", got.nbytes, 19)
            dest = os.path.join(d, "saved.pdf")
            got.save_as(dest)
            check_true("second-opinion . and it saves like any other fetch",
                       os.path.exists(dest)
                       and open(dest, "rb").read().startswith(b"%PDF"))
        except Exception as e:
            check_true("second-opinion . a reachable file comes back as a "
                       "file", False, "got {}: {}".format(type(e).__name__, e))
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_a_refusal_that_names_a_challenge_holds_for_a_person():
    """A 403 is never an absence, and a challenge is never a failure.

    The standing rule: a 403 means *you may not have this*, never *this
    does not exist*. Asking the network layer must not become a new way
    to turn one into the other.
    """
    import tempfile
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        f, _drv, url = _silent_browser(os.path.join(d, "_incoming"))
        f.use_cookie("c=1")
        M.reset_rate_state()
        try:
            # Each signal on its own, in its own fetcher. The vendor's
            # own header is the stronger one and must stand alone; the
            # interstitial's words must stand alone too, for a refusal
            # that arrives without the header.
            #
            # Both are REPORTED, never held for — see
            # test_a_challenge_the_network_meets_is_not_a_hold below for
            # why that distinction cost a live run.
            for n, (hdr, body, what) in enumerate((
                    (True, False, "the header alone"),
                    (False, True, "the interstitial's words alone"))):
                fn, _dn, urln = _silent_browser(
                    os.path.join(d, "_inc_c{}".format(n)))
                fn.use_cookie("c=1")
                urllib.request.urlopen = _asking(
                    _challenge_error(header=hdr, body=body))
                try:
                    fn.fetch(urln)
                except M.BrowserFetchFailed as e:
                    check_true("second-opinion . a challenge refusal is "
                               "named as one — {}".format(what),
                               "verification challenge" in str(e), str(e))
                    check_true("second-opinion . and concludes nothing "
                               "about the file — {}".format(what),
                               "establishes whether the file exists"
                               in str(e), str(e))
                except Exception as e:
                    check_true("second-opinion . a challenge refusal is "
                               "named as one — {}".format(what), False,
                               "got {}: {}".format(type(e).__name__, e))
                else:
                    check_true("second-opinion . a challenge refusal is "
                               "named as one — {}".format(what), False,
                               "it returned a file")

            # A plain 403 carries neither signal. It stays what it was:
            # unexplained. Not an absence, not a challenge.
            f2, _d2, url2 = _silent_browser(os.path.join(d, "_inc2"))
            f2.use_cookie("c=1")
            urllib.request.urlopen = _asking(_http(403, "Forbidden"))
            try:
                f2.fetch(url2)
            except M.BrowserFetchFailed as e:
                check_true("second-opinion . a bare 403 stays unexplained",
                           "403" in str(e), str(e))
                check_true("second-opinion . and the message keeps BOTH "
                           "observations", "tab:" in str(e)
                           and "network layer" in str(e), str(e))
            except Exception as e:
                check_true("second-opinion . a bare 403 stays unexplained",
                           False,
                           "got {}: {}".format(type(e).__name__, e))
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_pushback_is_never_flattened_into_a_browser_failure():
    """A throttle says nothing about whether the file exists.

    This module's oldest regression, arriving by a new route: the second
    opinion must let RateLimited and LoginRequired through untouched,
    because the worker's lockout handling and its _PARTIAL path are
    reached by catching exactly those.
    """
    import tempfile
    real, real_sleep = urllib.request.urlopen, M._sleep
    M._sleep = lambda secs: True
    with tempfile.TemporaryDirectory() as d:
        f, _drv, url = _silent_browser(os.path.join(d, "_incoming"))
        f.use_cookie("c=1")
        M.reset_rate_state()
        try:
            urllib.request.urlopen = _asking(_http(429, "Too Many Requests"))
            try:
                f.fetch(url)
            except M.RateLimited:
                check_true("second-opinion . a throttle stays a throttle",
                           True)
            except Exception as e:
                check_true("second-opinion . a throttle stays a throttle",
                           False,
                           "got {}: {}".format(type(e).__name__, e))

            f2, _d2, url2 = _silent_browser(os.path.join(d, "_inc2"))
            f2.use_cookie("c=1")
            M.reset_rate_state()
            urllib.request.urlopen = _asking(_NetResp(
                b"<html><body><form><input name='login'>"
                b"<input name='password'></form></body></html>",
                {"Content-Type": "text/html"},
                landed="https://dc.example.edu/cgi/login.cgi?return_to=x"))
            try:
                f2.fetch(url2)
            except M.LoginRequired:
                check_true("second-opinion . a session expiry stays one",
                           True)
            except Exception as e:
                check_true("second-opinion . a session expiry stays one",
                           False,
                           "got {}: {}".format(type(e).__name__, e))
        finally:
            urllib.request.urlopen = real
            M._sleep = real_sleep
            M.reset_rate_state()


def test_no_cookie_degrades_loudly_and_asks_nothing():
    """A guard that quietly does nothing turns a bug into a mystery.

    The way this feature dies in production is a fetcher that was never
    given the session cookie. It must not ask anyway, and it must not
    fall silent about why it did not.
    """
    import tempfile
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        f, _drv, url = _silent_browser(os.path.join(d, "_incoming"))
        # deliberately no use_cookie()
        asked = _asking(_http(404, "Not Found"))
        M.reset_rate_state()
        try:
            urllib.request.urlopen = asked
            try:
                f.fetch(url)
            except M.BrowserFetchFailed as e:
                check_true("second-opinion . with no cookie it says so",
                           "no session cookie" in str(e), str(e))
                check_true("second-opinion . and does not claim an absence",
                           "404" not in str(e), str(e))
            except Exception as e:
                check_true("second-opinion . with no cookie it says so",
                           False,
                           "got {}: {}".format(type(e).__name__, e))
            check("second-opinion . and the network was never asked",
                  len(asked.asked), 0)
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_a_reached_file_is_never_asked_about():
    """The second opinion is bounded by failures, not by files.

    A fetch the browser completed must cost no extra request — this runs
    on every file of a collection pass.
    """
    import tempfile
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        f, _drv = _browser_fetcher(inc, [_Tab(drop=("a.pdf", b"%PDF-x"))])
        f.use_cookie("c=1")
        asked = _asking(_http(404, "Not Found"))
        M.reset_rate_state()
        try:
            urllib.request.urlopen = asked
            got = f.fetch("https://dc.example.edu/cgi/viewcontent.cgi?"
                          "article=1&context=etd")
            check("second-opinion . a browser fetch still returns the "
                  "browser's file", got.name, "a.pdf")
            check("second-opinion . and asked the network nothing",
                  len(asked.asked), 0)
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_an_absence_the_tab_states_is_not_re_asked():
    """The tab CAN say some things, and when it does it is believed.

    "No PDF has been provided" is the server answering. Asking again
    over another path would spend a request to be told the same thing.

    TWO BRANCHES SAY THIS, and the first version of this test drove only
    one. The absence check exists on the rendered fast path AND after the
    wait, and planting the second away left the suite green because the
    fixture rendered a page and never reached it. A check that verifies
    the adjacent branch is worse than no check.
    """
    import tempfile
    sorry = "<html><body>No PDF has been provided.</body></html>"
    url = ("https://dc.example.edu/cgi/viewcontent.cgi?"
           "article=1&context=etd")
    cases = (
        # The page renders on this navigation: the fast path reads it.
        # Planting the fast path's absence branch away does NOT turn this
        # case red — the post-wait branch catches the same page a second
        # later and still asks nobody. (An older check,
        # "fast-absence . a rendered absence is read immediately", is
        # what goes red, and correctly: what is lost is the speed, not
        # the answer.) Kept because the two branches backing each other
        # up is itself worth pinning.
        ("rendered now", [_Tab(title="sorry", source=sorry)], None),
        # The navigation times out and the absence page goes on screen a
        # moment later — a NEW document, served for this request. Read
        # after the wait, by the other branch.
        #
        # Until v1.34 this case parked the tab on the absence page BEFORE
        # the navigation (`at=`) and stalled, and asserted that the page
        # was believed. That is the stale-tab defect written down as an
        # expectation: a page left by an earlier request, trusted as the
        # answer to this one. See
        # test_a_page_left_by_an_earlier_request_is_not_an_absence.
        ("read after the wait", [_late_sorry(sorry)], None),
    )
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        for n, (what, script, at) in enumerate(cases):
            inc = os.path.join(d, "_inc{}".format(n))
            f, _drv = _browser_fetcher(inc, script, at=at)
            f.use_cookie("c=1")
            asked = _asking(_http(404, "Not Found"))
            M.reset_rate_state()
            try:
                urllib.request.urlopen = asked
                try:
                    f.fetch(url)
                except M.NoFileAvailable:
                    check_true("second-opinion . a stated absence is "
                               "believed — {}".format(what), True)
                except Exception as e:
                    check_true("second-opinion . a stated absence is "
                               "believed — {}".format(what), False,
                               "got {}: {}".format(type(e).__name__, e))
                check("second-opinion . and costs no second request — {}"
                      .format(what), len(asked.asked), 0)
            finally:
                urllib.request.urlopen = real
                M.reset_rate_state()


def _late_sorry(page, delay=0.05):
    tab = _Tab(stall=True)
    tab.late_render = (delay, "sorry", page)
    return tab


def test_a_page_left_by_an_earlier_request_is_not_an_absence():
    """The stale-tab absence — v1.34's first fix.

    Measured 2026-09-29: a native navigation usually never commits, so
    the tab keeps whatever the LAST rendered request put there. When
    that was an absence page, every later native that never committed
    was written "No native file" from it — 178 of 181 in one run. A
    probe then watched a real file arrive under exactly such a page.

    The rendered fast path already asked `served()`; the post-wait
    branch did not. Both cases below park the tab on an earlier
    request's absence page and stall the navigation, which is the live
    shape, and neither may conclude anything from that page.
    """
    import tempfile
    sorry = "<html><body>No PDF has been provided.</body></html>"
    earlier = ("sorry", "https://dc.example.edu/cgi/viewcontent.cgi?"
               "article=9&context=etd", sorry)
    url = ("https://dc.example.edu/context/etd/article/1/type/native/"
           "viewcontent?preview_mode=1")
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        # 1. Nothing arrives. The evidence must come from the network,
        #    and the message must say where it came from.
        f, _drv = _browser_fetcher(os.path.join(d, "_a"),
                                   [_Tab(stall=True)], at=earlier)
        f.use_cookie("c=1")
        asked = _asking(_http(404, "Not Found"))
        M.reset_rate_state()
        try:
            urllib.request.urlopen = asked
            try:
                f.fetch(url)
                check_true("stale-tab . an earlier page decides nothing",
                           False, "the fetch returned a file")
            except M.NoFileAvailable as e:
                check_true("stale-tab . the absence rests on the server's "
                           "answer, not the earlier page",
                           "HTTP 404" in str(e) and "page says" not in str(e),
                           str(e))
            except Exception as e:
                check_true("stale-tab . the absence rests on the server's "
                           "answer, not the earlier page", False,
                           "got {}: {}".format(type(e).__name__, e))
            check("stale-tab . so the network is asked once",
                  len(asked.asked), 1)
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()

        # 2. A REAL file under the stale page, slow to start. Before
        #    v1.34 this was written "No native file". The network holds
        #    the file, so the right answer is the file.
        tab = _Tab(stall=True)
        tab.timed_transfer = (0.4, 0.5, "late.pdf", b"%PDF-1.7 late")
        f, _drv = _browser_fetcher(os.path.join(d, "_b"), [tab], at=earlier)
        f.settle_timeout = M.BROWSER_SETTLE_TIMEOUT  # as production (v1.34.1)
        f.use_cookie("c=1")
        asked = _asking(_NetResp(
            b"%PDF-1.7 real bytes",
            {"Content-Type": "application/pdf",
             "Content-Disposition": 'inline; filename="thesis.pdf"'},
            landed=url))
        M.reset_rate_state()
        try:
            urllib.request.urlopen = asked
            try:
                got = f.fetch(url)
                check_true("stale-tab . a slow real file is not written "
                           "absent", bool(got and got.name), repr(got))
            except M.NoFileAvailable as e:
                check_true("stale-tab . a slow real file is not written "
                           "absent", False, "FALSE ABSENCE: {}".format(e))
            except Exception as e:
                check_true("stale-tab . a slow real file is not written "
                           "absent", False,
                           "got {}: {}".format(type(e).__name__, e))
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()

        # 3. Off campus the network cannot answer. Then NOTHING is
        #    established, and the message says the page was not this
        #    request's — rather than naming it as if it were.
        f, _drv = _browser_fetcher(os.path.join(d, "_c"),
                                   [_Tab(stall=True)], at=earlier)
        f.use_cookie("c=1")
        M.reset_rate_state()
        try:
            urllib.request.urlopen = _asking(_http(
                403, "Forbidden", {"cf-mitigated": "challenge"}))
            try:
                f.fetch(url)
                check_true("stale-tab . unanswerable stays unanswered",
                           False, "the fetch returned a file")
            except M.NoFileAvailable as e:
                check_true("stale-tab . unanswerable stays unanswered",
                           False, "FALSE ABSENCE: {}".format(e))
            except M.BrowserFetchFailed as e:
                check_true("stale-tab . unanswerable stays unanswered",
                           True)
                check_true("stale-tab . and the message says the page "
                           "was left by an earlier request",
                           "earlier request" in str(e), str(e))
            except Exception as e:
                check_true("stale-tab . unanswerable stays unanswered",
                           False, "got {}: {}".format(type(e).__name__, e))
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_a_native_is_asked_about_before_the_browser_is_sent():
    """v1.34: headers first for natives, Chrome only for a file that exists.

    Measured 2026-09-29: a native with no file costs about 44 s through
    the browser (a navigation that never commits, the wait for a file,
    then the second opinion) and about one second asked directly. The
    file itself, when there is one, still comes through Chrome.
    """
    import tempfile
    native = ("https://dc.example.edu/context/etd/article/1/type/native/"
              "viewcontent?preview_mode=1")
    derivative = ("https://dc.example.edu/cgi/viewcontent.cgi?"
                  "article=1&context=etd&preview_mode=1")
    check("status-first . on by default in production",
          M.BrowserFetcher("unused").status_first, True)
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        def fetcher(sub, script):
            f, drv = _browser_fetcher(os.path.join(d, sub), script)
            f.status_first = True
            f.use_cookie("c=1")
            return f, drv

        # 1. No file: answered at once, and the browser is never sent.
        f, drv = fetcher("_a", [_Tab(stall=True)])
        asked = _asking(_http(404, "Not Found"))
        M.reset_rate_state()
        try:
            urllib.request.urlopen = asked
            try:
                f.fetch(native)
                check_true("status-first . a 404 is an absence", False,
                           "a file came back")
            except M.NoFileAvailable as e:
                check_true("status-first . a 404 is an absence, and says "
                           "how it was learned",
                           "HTTP 404" in str(e) and "before sending the "
                           "browser" in str(e), str(e))
            check("status-first . the browser is never sent",
                  len(drv.visited), 0)
            check("status-first . one request", len(asked.asked), 1)
            check("status-first . and it is counted", f.status_asks, 1)
            check_true("status-first . and the count is reported when the "
                       "fetcher closes",
                       any("1 native file(s) were asked about" in n
                           for n in f.close()), "")
        finally:
            urllib.request.urlopen = real

        # 2. A file: headers only, then Chrome fetches it as always.
        f, drv = fetcher("_b", [_Tab(drop=("orig.docx", b"PK\x03\x04 doc"))])
        head = _NetResp(b"PK\x03\x04 the whole body",
                        {"Content-Type": "application/octet-stream"},
                        landed=native)
        asked = _asking(head)
        try:
            urllib.request.urlopen = asked
            got = f.fetch(native)
            check("status-first . a present native still comes through "
                  "Chrome", len(drv.visited), 1)
            check("status-first . and the file is Chrome's",
                  getattr(got, "name", ""), "orig.docx")
            check("status-first . the network's body is never read",
                  getattr(head, "reads", 0), 0)
            check("status-first . one request to ask", len(asked.asked), 1)
        except Exception as e:
            check_true("status-first . a present native still comes "
                       "through Chrome", False,
                       "got {}: {}".format(type(e).__name__, e))
        finally:
            urllib.request.urlopen = real

        # 3. A web page is not a file. The browser reads it.
        f, drv = fetcher("_c", [_Tab(stall=True)])
        page = _NetResp(b"<html>sorry</html>",
                        {"Content-Type": "text/html; charset=utf-8"},
                        landed=native)
        try:
            urllib.request.urlopen = _asking(page)
            try:
                f.fetch(native)
            except Exception:
                pass
            check("status-first . a web page sends the browser, not an "
                  "absence", len(drv.visited), 1)
        finally:
            urllib.request.urlopen = real

        # 4. Off campus: challenged. The browser path runs, and the network
        #    is NOT asked a second time about the same URL.
        f, drv = fetcher("_d", [_Tab(stall=True)])
        asked = _asking(_challenge_error())
        try:
            urllib.request.urlopen = asked
            try:
                f.fetch(native)
                check_true("status-first . challenged falls back", False,
                           "a file came back")
            except M.NoFileAvailable as e:
                check_true("status-first . challenged falls back", False,
                           "FALSE ABSENCE: {}".format(e))
            except M.BrowserFetchFailed as e:
                check_true("status-first . challenged falls back, and "
                           "says the network could not answer",
                           "asked first" in str(e)
                           and "verification challenge" in str(e), str(e))
            check("status-first . the browser was sent", len(drv.visited), 1)
            check("status-first . and the network was asked once, not "
                  "twice", len(asked.asked), 1)
        finally:
            urllib.request.urlopen = real

        # 5. Derivatives are not asked about.
        f, drv = fetcher("_e", [_Tab(drop=("thesis.pdf", b"%PDF-1.7 x"))])
        asked = _asking(_http(404, "Not Found"))
        try:
            urllib.request.urlopen = asked
            f.fetch(derivative)
            check("status-first . a derivative is not asked about",
                  len(asked.asked), 0)
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_a_row_that_cannot_be_confirmed_here_has_a_status_of_its_own():
    """v1.34, decision F: neither downloaded, nor absent, nor retried blindly.

    Measured 2026-09-29 off campus: 337 native rows where the network
    layer was challenged. Every one was a failure that a re-run from the
    same network would meet again, one request at a time.
    """
    import ast
    import tempfile
    check_true("unconfirmed . it is a failure, never an absence",
               issubclass(M.NotConfirmedHere, M.BrowserFetchFailed)
               and not M.is_definite_absence(M.NotConfirmedHere("x")))
    # The network layer's challenge produces it.
    real = urllib.request.urlopen
    try:
        urllib.request.urlopen = _asking(_challenge_error())
        try:
            M.ask_the_network("https://dc.example.edu/cgi/viewcontent.cgi?"
                              "article=1&context=etd", "c=1", "observed")
            check_true("unconfirmed . a challenged second opinion is "
                       "not-confirmed", False, "returned")
        except M.NotConfirmedHere:
            check_true("unconfirmed . a challenged second opinion is "
                       "not-confirmed", True)
        except Exception as e:
            check_true("unconfirmed . a challenged second opinion is "
                       "not-confirmed", False,
                       "got {}: {}".format(type(e).__name__, e))
    finally:
        urllib.request.urlopen = real
        M.reset_rate_state()

    # A re-run leaves them out unless asked, and says so once.
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?article=7&context=etd"
    with tempfile.TemporaryDirectory() as d:
        rows = [
            _row(1, "7", "native", "public", "current",
                 M.UNCONFIRMED + ": nothing began arriving within 8s",
                 base + "/context/etd/article/7/type/native/viewcontent", adm),
            _row(2, "8", "native", "public", "current",
                 M.UNCONFIRMED + ": nothing began arriving within 8s",
                 base + "/context/etd/article/8/type/native/viewcontent",
                 base + "/cgi/editor.cgi?article=8&context=etd"),
            _row(3, "9", "primary", "public", "current",
                 "Failed: HTTP 500 Internal Server Error",
                 base + "/cgi/viewcontent.cgi?article=9&context=etd",
                 base + "/cgi/editor.cgi?article=9&context=etd"),
        ]
        p = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        got, skipped, _a, _k = M.read_failed_rows(p)
        check("unconfirmed . a re-run leaves them out by default",
              sorted(r["article"] for r in got), ["9"])
        check_true("unconfirmed . and says how many, once",
                   sum("2 row(s) could not be confirmed" in x
                       for x in skipped) == 1, repr(skipped))
        got, skipped, _a, _k = M.read_failed_rows(p, include_unconfirmed=True)
        check("unconfirmed . and includes them when asked",
              sorted(r["article"] for r in got), ["7", "8", "9"])

    # BOTH workers give it its own status and count it — a rule that
    # reaches one of a pair and not the other is this module's commonest
    # defect. Structural, and said so: download_worker cannot be driven
    # by this suite at all.
    tree = ast.parse(inspect.getsource(M))
    for fn in ("download_worker", "retry_worker"):
        node = next(n for n in ast.walk(tree)
                    if isinstance(n, ast.FunctionDef) and n.name == fn)
        tests = [n for n in ast.walk(node) if isinstance(n, ast.Call)
                 and getattr(n.func, "id", "") == "isinstance"
                 and len(n.args) == 2
                 and getattr(n.args[1], "id", "") == "NotConfirmedHere"]
        lines = [n for n in ast.walk(node) if isinstance(n, ast.Call)
                 and getattr(n.func, "id", "") == "unconfirmed_line"]
        check_true("unconfirmed . {} gives it its own status".format(fn),
                   len(tests) == 1, "{} isinstance test(s)".format(len(tests)))
        check_true("unconfirmed . {} reports the count once per "
                   "structure".format(fn), len(lines) == 1,
                   "{} call(s)".format(len(lines)))

    # The page offers the choice, and the re-run endpoint passes it on.
    page = M.PAGE if isinstance(M.PAGE, str) else M.PAGE.decode()
    check_true("unconfirmed . the page has the checkbox",
               'id="retryunconfirmed"' in page)
    src = inspect.getsource(M.make_handler)
    check_true("unconfirmed . /api/retry passes it to the worker",
               '"retry_unconfirmed": unconfirmed_v' in src)


def test_a_large_download_is_waited_for_by_its_progress():
    """v1.34, items I and J: no total ceiling; a stall is the limit.

    Measured 2026-09-29: a 3.05 GB video took about three minutes, close
    to the old fixed 300 s ceiling, and the module now has to expect
    10 GB. A transfer that was still growing used to be abandoned at
    300 s and its partial thrown away, so the row failed on every re-run.
    """
    import tempfile
    import threading as _th
    check("large . no total ceiling in production",
          M.BROWSER_SETTLE_TIMEOUT, None)
    check("large . the stall limit is 120 s", M.BROWSER_STALL_LIMIT, 120.0)
    check("large . holds and transfers tick alike",
          (M.VERIFY_TICK, M.PROGRESS_TICK), (60.0, 60.0))

    def grow(folder, name, steps, every, then_finish=True):
        part = os.path.join(folder, name + ".crdownload")

        def run():
            for k in range(steps):
                try:
                    with open(part, "ab") as fh:
                        fh.write(b"x" * 4096)
                except OSError:
                    return
                time.sleep(every)
            if then_finish:
                try:
                    os.replace(part, os.path.join(folder, name))
                except OSError:
                    pass
        t = _th.Thread(target=run, daemon=True)
        t.start()
        return t

    with tempfile.TemporaryDirectory() as d:
        # 1. Keeps growing for far longer than the stall limit, then
        #    finishes. It must be waited for, and it must say so.
        w = M.DirectoryDownloadWatch(d)
        w.begin()
        said = []
        t = grow(d, "video.mp4", steps=12, every=0.1)
        name, note = w.settle(start_within=1.0, timeout=None,
                              stall_limit=0.5, progress=said.append,
                              tick=0.3)
        t.join(2)
        check("large . a transfer that keeps growing is waited for", name,
              "video.mp4")
        check_true("large . and reports its progress while it waits",
                   any(x.startswith("Still downloading video.mp4") for x in
                       said), repr(said))
        check_true("large . with observed sizes only — no total, no ETA",
                   all("%" not in x and "remaining" not in x.lower()
                       and "eta" not in x.lower() for x in said)
                   and any("total size not known" in x for x in said),
                   repr(said))

        # 2. Stops growing. Announced first, then given up.
        w.begin()
        said = []
        t = grow(d, "stuck.mp4", steps=2, every=0.05, then_finish=False)
        name, note = w.settle(start_within=1.0, timeout=None,
                              stall_limit=0.6, progress=said.append,
                              tick=10.0)
        t.join(2)
        check("large . a stalled transfer is given up", name, "")
        check_true("large . and the reason is the stall, with what arrived",
                   note.startswith("no growth for") and "stuck.mp4" in note,
                   note)
        check_true("large . the stall is announced before it fails",
                   any(x.startswith("No growth for") for x in said),
                   repr(said))

        # 3. Stop ends a long wait at once.
        w.begin()
        t = grow(d, "long.mp4", steps=40, every=0.05, then_finish=False)
        real_hook = M._STOP_HOOK[0]
        M.set_stop_hook(lambda: True)
        t0 = time.time()
        try:
            w.settle(start_within=1.0, timeout=None, stall_limit=30.0)
            check_true("large . Stop ends a long wait", False, "returned")
        except M.StopRequested:
            check_true("large . Stop ends a long wait",
                       time.time() - t0 < 2.0)
        finally:
            M.set_stop_hook(real_hook)
            t.join(3)

    # 4. fetch_one wires the reporter, so neither worker can forget to.
    with tempfile.TemporaryDirectory() as d:
        f, _drv = _browser_fetcher(os.path.join(d, "_w"),
                                   [_Tab(drop=("a.pdf", b"%PDF-1.7 a"))])
        v = M.VerificationWatch(window=0)
        M.fetch_one(f, v, "https://dc.example.edu/cgi/viewcontent.cgi?"
                    "article=1&context=etd")
        check_true("large . fetch_one gives the fetcher somewhere to report",
                   f.say is not None)

    # 5. The stall limit is a page control, validated by the one function
    #    both endpoints call, and handed to the fetcher by both workers.
    check_true("large . an out-of-range stall limit is refused",
               bool(M._fetch_option_error(
                   {"fetch_path": M.FETCH_BROWSER, "stall_limit": 0})))
    check("large . a sane one is accepted", M._fetch_option_error(
        {"fetch_path": M.FETCH_BROWSER, "stall_limit": 120.0}), "")
    page = M.PAGE if isinstance(M.PAGE, str) else M.PAGE.decode()
    check_true("large . the page has the control",
               'id="stalllimit"' in page)
    check("large . make_fetcher hands it over",
          M.make_fetcher(M.FETCH_BROWSER, "", "unused",
                         stall_limit=45).stall_limit, 45.0)
    import ast
    tree = ast.parse(inspect.getsource(M))
    for fn in ("download_worker", "retry_worker"):
        node = next(n for n in ast.walk(tree)
                    if isinstance(n, ast.FunctionDef) and n.name == fn)
        calls = [n for n in ast.walk(node) if isinstance(n, ast.Call)
                 and getattr(n.func, "id", "") == "make_fetcher"]
        check_true("large . {} passes the stall limit on".format(fn),
                   len(calls) == 1 and any(k.arg == "stall_limit"
                                           for k in calls[0].keywords))


def test_a_body_the_network_fetches_is_streamed_to_disk():
    """v1.34, item I: a 10 GB answer must not be read into memory."""
    import tempfile
    body = b"%PDF-1.7 " + b"y" * (3 * 1024 * 1024)
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        try:
            urllib.request.urlopen = _asking(_NetResp(
                body, {"Content-Type": "application/pdf",
                       "Content-Disposition": 'inline; filename="big.pdf"'},
                landed="https://dc.example.edu/x"))
            path, name, ctype = M.fetch_file_to(
                "https://dc.example.edu/cgi/viewcontent.cgi?article=1",
                "c=1", 30, d)
            check("stream . the whole body lands on disk",
                  os.path.getsize(path), len(body))
            check("stream . named from the response", name, "big.pdf")
            check_true("stream . hidden from the download watcher",
                       os.path.basename(path).startswith("."))
            check("stream . nothing half-written is left",
                  [n for n in os.listdir(d) if n.endswith(".part")], [])
            # A page is refused, and its half-file removed.
            urllib.request.urlopen = _asking(_NetResp(
                b"<html><body>no</body></html>",
                {"Content-Type": "text/html"}, landed="https://dc/x"))
            before = set(os.listdir(d))
            try:
                M.fetch_file_to("https://dc.example.edu/cgi/viewcontent.cgi?"
                                "article=2", "c=1", 30, d)
                check_true("stream . a page is not a file", False)
            except ValueError:
                check_true("stream . a page is not a file", True)
            check("stream . and leaves nothing behind",
                  set(os.listdir(d)) - before, set())
            # The second opinion uses it when it has somewhere to put it.
            urllib.request.urlopen = _asking(_NetResp(
                body, {"Content-Type": "application/pdf"},
                landed="https://dc.example.edu/x"))
            got = M.ask_the_network("https://dc.example.edu/cgi/viewcontent."
                                    "cgi?article=3", "c=1", "o", 30,
                                    folder=d)
            check_true("stream . the second opinion streams too",
                       got.data is None and got.nbytes == len(body))
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_one_file_that_cannot_be_saved_does_not_end_the_run():
    """v1.34, item C: the filename cap, and SaveFailed.

    Measured 2026-09-21: saved names reached 168 characters (214 with
    natives), and 2 of 1,133 would pass Windows' 260-character path limit
    at a plausible destination. `save_as` then raised OSError, the nearest
    handler in BOTH workers was the whole-run one, and one file ended the
    run — at the same row on every re-run.
    """
    import tempfile
    import importlib.util
    # 1. The cap keeps the tokens infer_meta reads, and the extension.
    long_orig = "A" * 300 + ".docx"
    name = M.build_filename("honors-theses", "1363", "native", "public",
                            "current", long_orig)
    check_true("filename-cap . a long original name is capped",
               len(name) <= M.FILENAME_MAX_CHARS, str(len(name)))
    check_true("filename-cap . the leading tokens survive",
               name.startswith("honors-theses_1363_native_public_current_"),
               name)
    check_true("filename-cap . and so does the extension",
               name.endswith(".docx"), name)
    short = M.build_filename("etd", "7", "primary", "public", "current",
                             "thesis.pdf")
    check("filename-cap . a short name is untouched", short,
          "etd_7_primary_public_current_thesis.pdf")
    here = os.path.dirname(os.path.abspath(M.__file__))
    idg = os.path.join(here, "dc_image_describer.py")
    if os.path.exists(idg):
        spec = importlib.util.spec_from_file_location("_idg", idg)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        with tempfile.TemporaryDirectory() as d:
            root = os.path.join(d, "run")
            fpath = os.path.join(root, "honors-theses", name)
            meta = mod.infer_meta(fpath, root)
            check("filename-cap . the Image Description Generator still "
                  "relinks it", (meta.get("ctx"), meta.get("article")),
                  ("honors-theses", "1363"))
    else:
        check_true("filename-cap . relink (image describer absent)", True,
                   "skipped")

    # 2. A save the operating system refuses is one failed file.
    class _Refused:
        def __init__(self, path):
            self.path = path

        def save_as(self, dest):
            raise OSError(36, "File name too long")

    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        os.makedirs(inc)
        left = os.path.join(inc, "x.pdf")
        open(left, "wb").write(b"%PDF")
        try:
            M.place_file(_Refused(left), d, "etd_1_primary_x.pdf")
            check_true("save-failed . a refused save raises SaveFailed",
                       False, "returned")
        except M.SaveFailed as e:
            check_true("save-failed . a refused save raises SaveFailed, "
                       "naming the path's length",
                       "characters" in str(e) and "too long" in str(e),
                       str(e))
        check_true("save-failed . and the incoming copy is not left behind",
                   not os.path.exists(left))

    # 3. Through the real retry_worker: the refused file fails, the run
    #    goes on, and the next file is saved.
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article={}&context=etd"
    rows = [_row(i, str(i), "primary", "public", "current",
                 "Failed: HTTP 429 Too Many Requests",
                 base + "/cgi/viewcontent.cgi?article={}&context=etd".format(i),
                 adm.format(i), access=M.ACCESS_OPEN) for i in (1, 2, 3)]

    class _FakeDriver:
        page_source = "<html><body>admin</body></html>"
        current_url = base + "/"
        # A real driver always has a title; v1.34.1 reads it to prove the
        # tab answers before the run begins (ready_the_tab).
        title = "Digital Commons"

        def get(self, url):
            pass

    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "fetch_record_file",
              "LISTING_SETTLE")}
    real_save = M.Fetched.save_as

    def refusing(self, dest):
        if "_2_" in os.path.basename(dest):
            raise OSError(36, "File name too long")
        return real_save(self, dest)

    real_log = M.log
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: _FakeDriver()
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            M.fetch_record_file = (
                lambda url, cookie, native_url=None, **kw:
                (b"%PDF-1.7 body", "f.pdf", "application/pdf", False))
            M.Fetched.save_as = refusing
            M.log = lambda msg: None
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}}, out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            "fetch_path": M.FETCH_NETWORK,
                            "max_files": 0, "max_mb": 0,
                            "reconcile": False}, report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            M.Fetched.save_as = real_save
            M.log = real_log
        got = _rows_by_article(out)
    check_true("save-failed . the refused file is a failed row",
               str(got.get("2", ("",))[0]).startswith("Failed: could not "
                                                    "be saved"),
               repr(got.get("2")))
    check("save-failed . and the run went on to the next file",
          str(got.get("3", ("",))[0]).startswith("Downloaded"), True)
    import ast
    tree = ast.parse(inspect.getsource(M))
    node = next(n for n in ast.walk(tree)
                if isinstance(n, ast.FunctionDef) and n.name ==
                "download_worker")
    calls = [n for n in ast.walk(node) if isinstance(n, ast.Call)
             and getattr(n.func, "id", "") == "place_file"]
    handlers = [h for h in ast.walk(node) if isinstance(h, ast.ExceptHandler)
                and getattr(h.type, "id", "") == "SaveFailed"]
    check_true("save-failed . download_worker saves through the same "
               "function and catches the same failure (structural: the "
               "suite cannot drive that worker)",
               len(calls) == 1 and len(handlers) == 1,
               "{} call(s), {} handler(s)".format(len(calls), len(handlers)))


def test_the_preflight_downloads_a_file_twice_before_the_run():
    """v1.34, item B (open item 13): a cheap direct test, not a theory.

    Operators' Chrome settings will differ on a managed workstation, and
    how a CDP download folder interacts with policy, Safe Browsing or
    antivirus is not known. So the check DOWNLOADS: a few bytes from this
    computer, twice, through the attached Chrome.
    """
    import tempfile
    # The server itself serves an attachment.
    srv = M._PreflightServer()
    try:
        with urllib.request.urlopen(srv.url(1), timeout=5) as r:
            body = r.read()
            cd = r.headers.get("Content-Disposition", "")
        check("preflight . the check file is served", body, M.PREFLIGHT_BYTES)
        check_true("preflight . as an attachment", "attachment" in cd, cd)
    finally:
        srv.stop()

    big = type("U", (), {"free": 500 * 1024 ** 3})
    with tempfile.TemporaryDirectory() as d:
        # 1. Downloads work: two files arrive, and they are cleaned up.
        f, drv = _browser_fetcher(os.path.join(d, "_ok"), [])
        notes = M.preflight_downloads(f, d, disk_usage=lambda p: big)
        check("preflight . it downloads twice", len(drv.preflights), 2)
        check_true("preflight . and says so", any("2 test file(s)" in n
                                                  for n in notes), notes)
        check("preflight . and leaves nothing behind",
              os.listdir(os.path.join(d, "_ok")), [])

        # 2. Downloads blocked: the navigation is accepted and nothing
        #    arrives. That is the failure; the run must not start.
        f, drv = _browser_fetcher(os.path.join(d, "_no"), [])
        drv.downloads_blocked = True
        real = M.PREFLIGHT_START_WITHIN
        M.PREFLIGHT_START_WITHIN = 0.3
        seen = []
        try:
            M.preflight_downloads(f, d, disk_usage=lambda p: big,
                                  checks=seen)
            check_true("preflight . blocked downloads stop the run", False,
                       "it passed")
        except M.NotConfigured as e:
            check_true("preflight . blocked downloads stop the run, saying "
                       "what was observed and not guessing why",
                       "did not arrive" in str(e)
                       and "cannot tell which" in str(e), str(e))
        finally:
            M.PREFLIGHT_START_WITHIN = real
        # v1.34.1: the refusal is recorded as a check before it is raised.
        check("preflight . a refused download test is recorded as one",
              [(r[0], r[1]) for r in seen][-1:], [("Download test", "refused")])

        # 3. Free space: refused when there is almost none, warned when a
        #    10 GB file may not fit.
        f, drv = _browser_fetcher(os.path.join(d, "_sp"), [])
        tiny = type("U", (), {"free": 200 * 1024 ** 2})
        some = type("U", (), {"free": 5 * 1024 ** 3})
        try:
            M.preflight_downloads(f, d, disk_usage=lambda p: tiny)
            check_true("preflight . almost no space is refused", False)
        except M.NotConfigured:
            check_true("preflight . almost no space is refused", True)
        notes = M.preflight_downloads(f, d, disk_usage=lambda p: some)
        check_true("preflight . 5 GB free is a warning about 10 GB files",
                   any("10 GB" in n for n in notes), notes)

        # 4. Depth: a destination too deep for Windows is named.
        deep = os.path.join(d, *(["folder-name-long-enough"] * 6))
        os.makedirs(deep)
        # On Windows the same destination is REFUSED, and that is the
        # module working. The first version of this case expected the
        # refusal and never caught it, so the one platform the check
        # exists for failed the suite (CI, 2026-09-29, both Windows jobs).
        try:
            notes = M.preflight_downloads(
                M.NetworkFetcher(), deep, disk_usage=lambda p: big)
            if os.name == "nt":
                check_true("preflight . too deep for Windows is refused",
                           False, "it passed: {}".format(notes))
            else:
                check_true("preflight . too deep for Windows is said",
                           any("too long for Windows" in n for n in notes),
                           notes)
        except M.NotConfigured as e:
            check_true("preflight . too deep for Windows is refused on "
                       "Windows, and only there",
                       os.name == "nt" and "260" in str(e), str(e))

    # 5. Both workers call it, once, and turn NotConfigured into a refusal.
    import ast
    tree = ast.parse(inspect.getsource(M))
    for fn in ("download_worker", "retry_worker"):
        node = next(n for n in ast.walk(tree)
                    if isinstance(n, ast.FunctionDef) and n.name == fn)
        calls = [n for n in ast.walk(node) if isinstance(n, ast.Call)
                 and getattr(n.func, "id", "") == "preflight_downloads"]
        handlers = [h for h in ast.walk(node)
                    if isinstance(h, ast.ExceptHandler)
                    and getattr(h.type, "id", "") == "NotConfigured"]
        check_true("preflight . {} runs it and refuses on it".format(fn),
                   len(calls) == 1 and len(handlers) == 1,
                   "{} call(s), {} handler(s)".format(len(calls),
                                                     len(handlers)))


def test_files_that_only_the_network_delivers_are_noticed():
    """v1.34, item D: a quiet switch to the network layer is said aloud.

    Chrome's permission to download automatically is per site, so the
    localhost preflight cannot see it being refused. What the run CAN
    see is files that Chrome never delivered arriving over the network
    layer's second opinion instead — correctly, and slowly.
    """
    import tempfile
    url = "https://dc.example.edu/cgi/viewcontent.cgi?article=1&context=etd"
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        f, _drv = _browser_fetcher(
            os.path.join(d, "_i"),
            [_Tab(stall=True), _Tab(stall=True), _Tab(stall=True),
             _Tab(drop=("a.pdf", b"%PDF-1.7 a")), _Tab(stall=True)])
        f.use_cookie("c=1")
        said = []
        f.say = said.append
        try:
            urllib.request.urlopen = _asking(_NetResp(
                b"%PDF-1.7 net", {"Content-Type": "application/pdf"},
                landed=url))
            for _ in range(5):
                M.reset_rate_state()
                f.fetch(url)
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()
        check("network-delivered . counted", f.network_deliveries, 4)
        check("network-delivered . said once, at three in a row",
              sum("in a row came over the network layer" in x
                  for x in said), 1)
        check("network-delivered . and a Chrome delivery resets the run of "
              "them", f._network_in_a_row, 1)
        check_true("network-delivered . the total is reported at close",
                   any("4 file(s) came over the network layer" in n
                       for n in f.close()))


def test_a_refused_start_names_the_field_and_is_logged():
    """v1.34, item G. On 2026-09-29 Start answered a bare "Bad request.".

    A number box the browser cannot read reports itself as EMPTY, so a
    field that looked filled in sent nothing, and nothing said which.
    """
    session = M.load_session(Path("/nonexistent/session.json"))
    session["hub_url"] = "http://127.0.0.1:8750"
    srv = ThreadingHTTPServer(
        ("127.0.0.1", 0), M.make_handler(session, M.build_page(session)))
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    logs = []
    real_log = M.log
    M.log = lambda m: logs.append(str(m))
    try:
        with tempfile.TemporaryDirectory() as d:
            base = {"out_dir": d, "mode": "all", "parents": [],
                    "primary": True, "supp": False, "native": False,
                    "public": True, "hidden": True, "versions": "current",
                    "published": True, "unpublished": False,
                    "cooldowns": "3", "delay": "5.0", "page_delay": "1.0",
                    "fetch_path": M.FETCH_BROWSER, "verify_window": 900}
            code, j = _post(port, "/api/start", dict(base, cooldowns=""))
            check_true("bad-field . a blank box is named",
                       code == 400 and "Wait out server cooldowns"
                       in j.get("error", ""), repr((code, j)))
            check_true("bad-field . and the refusal is in the log",
                       any(l.startswith("Start refused:") and "cooldowns"
                           in l for l in logs), repr(logs[-2:]))
            code, j = _post(port, "/api/start", dict(base, delay="5,0"))
            check_true("bad-field . an unreadable value is named, with what "
                       "it received", code == 400 and "Gap before a file "
                       "download" in j.get("error", "") and "5,0"
                       in j.get("error", ""), repr((code, j)))
            code, j = _post(port, "/api/start", dict(base, cooldowns="3.0"))
            check("bad-field . a whole number written 3.0 is accepted",
                  (code, j.get("error")),
                  (400, "Load or map a hierarchy first."))
            code, j = _post(port, "/api/start", dict(base, max_files=""))
            check("bad-field . a blank session limit still means none",
                  (code, j.get("error")),
                  (400, "Load or map a hierarchy first."))
            code, j = _post(port, "/api/retry",
                            {"out_dir": d, "report_path": d,
                             "page_delay": "x"})
            check_true("bad-field . the re-run endpoint names it the same "
                       "way", code == 400 and "Gap before an admin page"
                       in j.get("error", ""), repr((code, j)))
    finally:
        M.log = real_log
        srv.shutdown()
        srv.server_close()

    # The rest of G, read off the module.
    check("session-defaults . files per session is 0 (no limit)",
          M.SESSION_LIMIT_FILES, 0)
    check("session-defaults . megabytes per session is 0 (no limit)",
          M.SESSION_LIMIT_MB, 0)
    check("verify-window . the default is 15 minutes", M.VERIFY_WINDOW,
          900.0)
    page = M.PAGE if isinstance(M.PAGE, str) else M.PAGE.decode()
    # SOURCE check, said so: the page's progress text is formatted by
    # this expression, and a fraction must print as a percentage.
    check_true("progress-text . a fraction is shown as a percentage "
               "(source check)", "Math.round(100*p.current/p.total)" in page)
    # Heartbeat: a quiet log reads as a frozen module.
    said = []
    real_log = M.log
    real_last = M._LAST_LOG[0]
    try:
        M.log = lambda m: said.append(str(m))
        M._LAST_LOG[0] = time.time() - M.HEARTBEAT_SECONDS - 1
        M.set_progress(5, 10, "etd · item 5/10 · 3 file(s)")
        check_true("heartbeat . a long quiet spell is broken by one line",
                   any(x.startswith("Still working — etd · item 5/10")
                       for x in said), repr(said))
        said.clear()
        M._LAST_LOG[0] = time.time()
        M.set_progress(6, 10, "etd · item 6/10")
        check("heartbeat . and not while the log is busy", said, [])
    finally:
        M.log = real_log
        M._LAST_LOG[0] = real_last
    # The bridge line leads with the outcome and never says FAILED.
    src = inspect.getsource(M.attach_chrome)
    check_true("bridge-line . leads with what the connection reports",
               '"command timeout: connection reports {}' in src)
    check_true("bridge-line . a lever that is unavailable is not FAILED",
               'set_via = "FAILED' not in src)
    import ast
    tree = ast.parse(inspect.getsource(M))
    node = next(n for n in ast.walk(tree)
                if isinstance(n, ast.FunctionDef) and n.name == "retry_worker")
    shows = [n for n in ast.walk(node) if isinstance(n, ast.Call)
             and getattr(n.func, "id", "") == "show_progress"]
    check_true("progress-records . the re-run reports progress by records "
               "at every row", len(shows) >= 2, str(len(shows)))


def test_the_download_statement_comes_from_the_profile():
    """v1.34, item H: an institution's terms are shown, never compiled in.

    The page shows the profile's `branding.download_statement`, escaped,
    and an institution-neutral sentence when there is none. The words of
    a real arrangement must never reach this file — the wording checks
    elsewhere in this suite search the source for vendor and licence
    vocabulary, and the generic build refuses the statement's phrases.
    """
    s = M.load_session(Path("/nonexistent/session.json"))
    s.setdefault("branding", {})["download_statement"] = (
        "Scripts are fine <here> & there.")
    page = M.build_page(s).decode("utf-8")
    check_true("statement . the profile's words are shown, escaped",
               "Scripts are fine &lt;here&gt; &amp; there." in page)
    s["branding"]["download_statement"] = ""
    page = M.build_page(s).decode("utf-8")
    check_true("statement . with none, the neutral sentence",
               M.NEUTRAL_DOWNLOAD_STATEMENT.replace("'", "&#x27;") in page,
               "")
    check_true("statement . no token is left unreplaced",
               "__STATEMENT__" not in page)


def test_the_worker_gives_the_browser_fetcher_its_cookie():
    """End to end, and the check that keeps the feature alive.

    Everything above passes with a fetcher a test handed a cookie to. In
    production the WORKER hands it over, and until v1.33.9 that call was
    a no-op — so nothing anywhere would have noticed it going missing.
    This drives the real worker, over the real browser path, against a
    server that answers 404, and asks what the report says.
    """
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article=7&context=etd"
    rows = [_row(1, "7", "native", "public", "current",
                 "Failed: nothing began arriving within 8s",
                 base + "/context/etd/article/7/type/native/viewcontent",
                 adm)]
    # The measured live shape, at the worker level: the tab is parked on
    # the home page and the navigation never commits.
    drv = _FakeBrowser([_Tab(stall=True)], None, None, None,
                       ("Example Commons", base + "/",
                        "<html><body>welcome</body></html>"))
    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "LISTING_SETTLE",
              "BROWSER_START_WITHIN")}
    real_urlopen, real_log = urllib.request.urlopen, M.log
    asked = _asking(_http(404, "Not Found"))
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: drv
            M.cookie_header = lambda driver: "session=abc"
            M.LISTING_SETTLE = 0
            M.BROWSER_START_WITHIN = 0.3
            M.log = lambda msg: None
            urllib.request.urlopen = asked
            M.STOP_EVENT.clear()
            M.reset_rate_state()
            M.retry_worker({"base_url": base, "settings": {}}, out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            "reconcile": False,
                            "fetch_path": M.FETCH_BROWSER,
                            "verify_window": 300},
                           report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
            urllib.request.urlopen = real_urlopen
            M.log = real_log
            M.reset_rate_state()

        check_true("second-opinion . the worker's fetcher did ask the "
                   "network", len(asked.asked) >= 1,
                   "the worker never passed the cookie through, so the "
                   "whole feature is inert in production")
        from openpyxl import load_workbook
        reports = []
        for root, _dirs, files in os.walk(out):
            reports += [os.path.join(root, fn) for fn in files
                        if fn.endswith(".xlsx") and "MasterLog" not in fn]
        check("second-opinion . one structure report was written",
              len(reports), 1)
        wbk = load_workbook(reports[0])
        try:
            ws = wbk.active
            hdr = [c.value for c in ws[1]]
            col = hdr.index("Status")
            got = [r[col].value for r in ws.iter_rows(min_row=2) if r[col].value]
        finally:
            wbk.close()
        # Read against the module's own labels rather than a literal: the
        # word differs by file kind ("No native file" here), and a test
        # carrying its own copy of that vocabulary is a test that goes
        # green on the wrong word the day the vocabulary moves.
        check_true("second-opinion . and the row is an absence, not a "
                   "failure",
                   bool(got) and str(got[0]).startswith(M.ABSENT_PREFIXES),
                   "status: {!r}".format(got))
        check_true("second-opinion . the row names what the server said",
                   bool(got) and "404" in str(got[0]),
                   "status: {!r}".format(got))


def test_a_challenge_the_network_meets_is_not_a_hold():
    """The v1.33.9 defect, live within two hours of shipping.

    A hold exists so a PERSON can clear a prompt the browser is
    showing, and `hold()` watches the tab to know when they have. A
    challenge served to urllib appears nowhere on that tab, and nothing
    a person does in Chrome clears it — the only thing that would is
    forwarding Chrome's clearance cookie into urllib, which is the one
    disguise this project will not build.

    So raising VerificationRequired from a urllib 403 manufactured a
    hold against the wrong witness. Measured off campus, 2026-09-22:
    three prompts in three file requests, no files fetched, ~23s to
    raise each one and 2s to "clear" it, three holds on one file, run
    over as NotVerified. 886 rows never requested.

    DRIVEN THROUGH fetch_one, not through the fetcher. Every other
    check in this block asks what the fetcher raises; the damage was
    done one layer up, by the loop that holds and counts. A test that
    stops at the fetcher cannot see a hold happen at all.
    """
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        f, _drv, url = _silent_browser(os.path.join(d, "_incoming"))
        f.use_cookie("c=1")
        polls = []
        # The tab is FINE — it was never challenged, which is the live
        # shape and the whole point. What is asserted is that this is
        # never CALLED: a hold that polls a tab nobody challenged is the
        # defect. Said plainly because planting showed it — flipping the
        # return value to True changes no result here, since with the
        # fix in place the hold is never entered at all. The value
        # matters only with the defect restored, where False is what
        # reproduces the instant "cleared in 2s".
        f.is_challenge = lambda: (polls.append(1), False)[1]
        held = []
        v = M.VerificationWatch(
            window=30, poll=0.1,
            sleep=lambda secs: True,
            stopping=lambda: False, now=time.time,
            announce=lambda msg: held.append(str(msg)))
        M.reset_rate_state()
        try:
            urllib.request.urlopen = _asking(_challenge_error())
            try:
                M.fetch_one(f, v, url, home_url="https://dc.example.edu/")
            except M.NotVerified as e:
                check_true("network-challenge . a refusal urllib cannot "
                           "answer must not end the run as unverified",
                           False, str(e))
            except M.BrowserFetchFailed as e:
                check_true("network-challenge . it is reported, not held "
                           "for", "verification challenge" in str(e),
                           str(e))
            except Exception as e:
                check_true("network-challenge . it is reported, not held "
                           "for", False,
                           "got {}: {}".format(type(e).__name__, e))
            else:
                check_true("network-challenge . it is reported, not held "
                           "for", False, "it returned a file")
            check("network-challenge . nothing was recorded as a prompt",
                  len(v.prompts), 0)
            check("network-challenge . and no hold was entered",
                  len(held), 0)
            check("network-challenge . so the tab was never polled for a "
                  "clearance it could not show", len(polls), 0)
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()


def test_a_challenge_the_browser_shows_is_still_a_hold():
    """The other half of the pair, checked in the same breath.

    Narrowing what holds is exactly the change that quietly stops a
    REAL prompt from holding, and this module's signature defect is a
    rule reaching one half of a pair. So: a challenge on the tab still
    stops the run, still asks a person, still resumes when they answer.
    """
    real = urllib.request.urlopen
    with tempfile.TemporaryDirectory() as d:
        inc = os.path.join(d, "_incoming")
        # A challenge page renders; the person clears it and the file
        # arrives, as it does in the world.
        first = _Tab(title="Just a moment...", stall=True)
        first.timed_clear = (0.3, "thesis.pdf", b"%PDF-real")
        f, _drv = _browser_fetcher(inc, [first])
        f.settle_timeout = M.BROWSER_SETTLE_TIMEOUT  # as production (v1.34.1)
        f.use_cookie("c=1")
        v = M.VerificationWatch(
            window=30, poll=0.1,
            sleep=lambda secs: (time.sleep(min(secs, 0.1)) or True),
            stopping=lambda: False, now=time.time)
        M.reset_rate_state()
        try:
            # The network layer is challenged too — off campus it would
            # be. It must not matter: the tab is the one being answered.
            urllib.request.urlopen = _asking(_challenge_error())
            try:
                got = M.fetch_one(f, v, "https://dc.example.edu/cgi/"
                                        "viewcontent.cgi?article=1",
                                  home_url="https://dc.example.edu/")
                check("network-challenge . a prompt on the TAB still "
                      "delivers the file", got.name, "thesis.pdf")
            except Exception as e:
                check_true("network-challenge . a prompt on the TAB still "
                           "delivers the file", False,
                           "got {}: {}".format(type(e).__name__, e))
            check_true("network-challenge . and it is still recorded",
                       len(v.prompts) >= 1,
                       "prompts: {}".format(v.prompts))
        finally:
            urllib.request.urlopen = real
            M.reset_rate_state()



def test_one_run_per_chrome_across_processes():
    """v1.34.1: a second COPY of the module must not start a second run.

    2026-09-30: the shell started a second copy on the next port while the
    first was mid-run, and the in-process claim could not see it. Both
    copies drive the same Chrome and set its download folder.
    """
    from http.server import BaseHTTPRequestHandler
    saved = M.RUN_LOCK_DIR
    with tempfile.TemporaryDirectory() as d:
        M.RUN_LOCK_DIR = d
        try:
            addr = "127.0.0.1:9222"
            path = M.run_lock_path(addr)
            live, dead = (lambda i: True), (lambda i: False)

            ok, h = M.claim_run_lock(addr, 8815, token="A", holder_check=live)
            check("run-lock . a free Chrome is claimed", ok, True)
            check("run-lock . the lock names its holder's port",
                  M._read_run_lock(path).get("port"), 8815)
            ok, h = M.claim_run_lock(addr, 8816, token="B", holder_check=live)
            check("run-lock . a live holder refuses a second copy",
                  (ok, h.get("port")), (False, 8815))
            msg = M.run_lock_refusal(h)
            check_true("run-lock . the refusal says where the running copy is",
                       "http://127.0.0.1:8815" in msg, msg)
            M.release_run_lock(addr, token="B")
            check_true("run-lock . a copy that does not hold it cannot "
                       "release it", os.path.exists(path))
            ok, h = M.claim_run_lock(addr, 8816, token="B", holder_check=dead)
            check("run-lock . a stale lock is taken over",
                  (ok, M._read_run_lock(path).get("token")), (True, "B"))
            M.release_run_lock(addr, token="B")
            check_true("run-lock . the holder releases it",
                       not os.path.exists(path))

            # A lock another copy is still writing is unreadable for a
            # moment; that is someone starting, not something dead.
            open(path, "w").close()
            ok, _ = M.claim_run_lock(addr, 8816, token="B", holder_check=dead)
            check("run-lock . an unreadable lock seconds old is held", ok,
                  False)
            old = time.time() - 60
            os.utime(path, (old, old))
            ok, _ = M.claim_run_lock(addr, 8816, token="B", holder_check=dead)
            check("run-lock . an unreadable lock a minute old is stale", ok,
                  True)
            M.release_run_lock(addr, token="B")
            check_true("run-lock . another Chrome has a lock of its own",
                       M.run_lock_path("127.0.0.1:9223") != path)

            # The holder is judged by asking it — the module's REAL page.
            session = M.load_session(Path("/nonexistent/session.json"))
            srv = ThreadingHTTPServer(
                ("127.0.0.1", 0),
                M.make_handler(session, M.build_page(session)))
            port = srv.server_address[1]
            threading.Thread(target=srv.serve_forever, daemon=True).start()
            before = M.STATE["phase"]
            try:
                info = {"port": port, "token": M.RUN_LOCK_TOKEN}
                check("run-lock . an idle copy does not hold a run",
                      M.holder_is_running(info), False)
                with M.LOCK:
                    M.STATE["phase"] = "downloading"
                try:
                    check("run-lock . a running copy with the lock's token "
                          "holds it", M.holder_is_running(info), True)
                    check("run-lock . a running copy with another token "
                          "does not", M.holder_is_running(
                              dict(info, token="other")), False)
                finally:
                    with M.LOCK:
                        M.STATE["phase"] = before
                check("run-lock . nothing listening does not hold it",
                      M.holder_is_running({"port": 1, "token": "x"}), False)

                # The endpoint, end to end. Another copy — a separate
                # server answering as a running downloader with the
                # lock's token — holds this Chrome. This copy is idle and
                # its own claim is free; Start must still be refused.
                class Other(BaseHTTPRequestHandler):
                    def do_GET(self):
                        body = json.dumps({"run_lock_token": "HELD",
                                           "phase": "downloading"}).encode()
                        self.send_response(200)
                        self.send_header("Content-Length", str(len(body)))
                        self.end_headers()
                        self.wfile.write(body)

                    def log_message(self, *a):
                        pass
                other = ThreadingHTTPServer(("127.0.0.1", 0), Other)
                oport = other.server_address[1]
                threading.Thread(target=other.serve_forever,
                                 daemon=True).start()
                try:
                    with open(path, "w", encoding="utf-8") as fh:
                        json.dump({"port": oport, "token": "HELD",
                                   "started": "2026-09-30 09:08:56"}, fh)
                    code, j = _post(port, "/api/map", {"out_dir": d})
                    check("run-lock . Start in a second copy is refused",
                          code, 409)
                    check_true("run-lock . and it names the running copy",
                               "http://127.0.0.1:{}".format(oport)
                               in (j.get("error") or ""), str(j))
                    check("run-lock . the refused copy is not left claimed",
                          M.STATE["claimed"], False)
                    check("run-lock . the holder's lock is untouched",
                          M._read_run_lock(path).get("token"), "HELD")
                finally:
                    other.shutdown()
                    other.server_close()
            finally:
                srv.shutdown()
                srv.server_close()

            # A run holds the lock while it runs and gives it back when it
            # ends — including when the worker raises on its first line.
            import types
            H = M.make_handler(session, M.build_page(session))
            h = H.__new__(H)
            h.server = types.SimpleNamespace(
                server_address=("127.0.0.1", 5555))
            quiet = threading.excepthook       # the "raises" case is meant to
            threading.excepthook = lambda args: None   # raise; say nothing
            for how in ("returns", "raises"):
                go, seen = threading.Event(), {}

                def target(how=how, go=go, seen=seen):
                    seen["held"] = M._read_run_lock(path).get("token")
                    go.wait(5)
                    if how == "raises":
                        raise RuntimeError("worker failed at once")
                started = h._claim_and_start(target, ())
                deadline = time.time() + 5
                while "held" not in seen and time.time() < deadline:
                    time.sleep(0.01)
                check("run-lock . a worker that {} held the lock while "
                      "running".format(how),
                      (started, seen.get("held")), (True, M.RUN_LOCK_TOKEN))
                go.set()
                deadline = time.time() + 5
                while M.STATE["claimed"] and time.time() < deadline:
                    time.sleep(0.01)
                check("run-lock . and a worker that {} gave it back".format(
                      how), os.path.exists(path), False)
            threading.excepthook = quiet
        finally:
            M.RUN_LOCK_DIR = saved



def test_fetch_commentary_is_progress_in_both_workers():
    """v1.34.1: a long transfer's ticks are not warnings, in either worker.

    2026-09-30: download_worker handed fetch_one its `warn`, retry_worker
    plain `log`, so two healthy minute-ticks for a 3 GB video were that
    run's "2 warning(s)". One of a pair again. Asked with `ast` of the
    call itself — the announce argument, sixth positional — not a grep.
    """
    import ast
    import textwrap as _tw
    for worker in (M.download_worker, M.retry_worker):
        tree = ast.parse(_tw.dedent(inspect.getsource(worker)))
        calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
                 and getattr(n.func, "id", "") == "fetch_one"]
        check("fetch-says . {} fetches through fetch_one once".format(
              worker.__name__), len(calls), 1)
        arg = calls[0].args[5] if calls and len(calls[0].args) > 5 else None
        check_true("fetch-says . {} hands it fetch_says(...)".format(
                   worker.__name__),
                   isinstance(arg, ast.Call)
                   and getattr(arg.func, "id", "") == "fetch_says",
                   ast.dump(arg) if arg is not None else "no sixth argument")
    # The log is capped at 200 lines, so read the LAST line rather than
    # comparing lengths, which stop changing once the cap is reached.
    M.fetch_says("etd")("Still downloading v.mp4 — 1.00 GB so far")
    line = M.STATE["log"][-1] if M.STATE["log"] else ""
    check_true("fetch-says . a tick is logged, prefixed, and not a warning",
               "'etd' — Still downloading" in line and "WARNING" not in line,
               line)


def test_every_downloaded_row_says_how_it_was_fetched():
    """v1.34.1: the "Fetched via" column, in both workers.

    2026-09-30: one file came over the network layer because Chrome
    delivered nothing, and its row said only "Downloaded". download_worker
    cannot be driven end to end by this suite, so its half is asked of the
    code with `ast`: every row it builds that says "Downloaded" must carry
    fetched_via(...). retry_worker's half is driven, in the rerun-browser
    test.
    """
    import ast
    import textwrap as _tw

    class Br:
        path = M.FETCH_BROWSER

    class Net:
        path = M.FETCH_NETWORK

    class Got:
        pass
    plain, second = Got(), Got()
    second.via_network = True
    check("fetched-via . the browser path", M.fetched_via(plain, Br()),
          "Chrome")
    check("fetched-via . the network path", M.fetched_via(plain, Net()),
          "network layer")
    check("fetched-via . Chrome delivered nothing",
          M.fetched_via(second, Br()),
          "network layer — Chrome delivered nothing")
    check("fetched-via . the column is last, so older readers are safe",
          M.REPORT_HEADERS[-1], "Fetched via")
    check("fetched-via . widths match headers",
          len(M.REPORT_WIDTHS), len(M.REPORT_HEADERS))
    for worker in (M.download_worker, M.retry_worker):
        tree = ast.parse(_tw.dedent(inspect.getsource(worker)))
        done = [n for n in ast.walk(tree) if isinstance(n, ast.Tuple)
                and any(isinstance(e, ast.Constant) and e.value == "Downloaded"
                        for e in n.elts)]
        check_true("fetched-via . {} builds a Downloaded row".format(
                   worker.__name__), bool(done))
        for n in done:
            last = n.elts[-1]
            check_true("fetched-via . {} line {}: its last cell is "
                       "fetched_via(...)".format(worker.__name__, n.lineno),
                       isinstance(last, ast.Call)
                       and getattr(last.func, "id", "") == "fetched_via",
                       ast.dump(last))


def test_a_tab_that_does_not_answer_stops_the_run_before_anything():
    """v1.34.1: the 270-second hang of 2026-09-30, and what it left behind.

    The run's first command to Chrome waited 270 s on a tab that never
    answered and ended as "Unexpected error" with a page of chromedriver
    stack frames; and because open() recorded the download-folder change
    only on success, close() neither restored it nor said so.
    """
    import ast
    import textwrap as _tw

    class TimeoutException(Exception):
        """Named as selenium names it; the module matches on the name."""

    class Ok:
        title = "Digital Commons"

        def __init__(self):
            self.went = []

        def set_page_load_timeout(self, s):
            pass

        def get(self, url):
            self.went.append(url)

    class Slow(Ok):
        def get(self, url):
            time.sleep(2.0)

    class PageLoadTimeout(Ok):
        def get(self, url):
            self.went.append(url)
            raise TimeoutException("timeout: page load")

    class Renderer(Ok):
        def get(self, url):
            raise TimeoutException(
                "Message: timeout: Timed out receiving message from "
                "renderer: 270.000\n  (Session info: chrome=154.0)\n"
                "Stacktrace:\n0   chromedriver   0x000000010512ca84 "
                "chromedriver + 3508868")

    saved = (M.TAB_READY_SECONDS, M.TAB_READY_GRACE)
    M.TAB_READY_SECONDS, M.TAB_READY_GRACE = 0.3, 0.2
    try:
        d = Ok()
        check("tab-ready . an answering tab is moved to the site",
              (M.ready_the_tab(d, "https://dc/"), d.went),
              ("Digital Commons", ["https://dc/"]))
        d = PageLoadTimeout()
        check("tab-ready . a page that is slow to finish loading still "
              "answers", M.ready_the_tab(d, "https://dc/"), "Digital Commons")
        for name, drv in (("no answer at all", Slow()),
                          ("the renderer timing out", Renderer())):
            t0 = time.time()
            try:
                M.ready_the_tab(drv, "https://dc/")
                check_true("tab-ready . {} stops the run".format(name), False,
                           "it passed")
            except M.TabNotAnswering as e:
                msg = str(e)
                check_true("tab-ready . {} stops the run, saying nothing "
                           "was requested and what to do".format(name),
                           "did not answer" in msg
                           and "Nothing was requested" in msg
                           and "start it again" in msg, msg)
                check_true("tab-ready . {} never shows stack frames".format(
                           name), "chromedriver +" not in msg
                           and "Stacktrace" not in msg, msg)
            check_true("tab-ready . {} is bounded by the clock".format(name),
                       time.time() - t0 < 1.5,
                       "{:.1f}s".format(time.time() - t0))
    finally:
        M.TAB_READY_SECONDS, M.TAB_READY_GRACE = saved

    # Both workers ready the tab BEFORE the fetcher exists — so nothing
    # touches Chrome's download folder until the tab has answered.
    for worker in (M.download_worker, M.retry_worker):
        tree = ast.parse(_tw.dedent(inspect.getsource(worker)))
        calls = {}
        for n in ast.walk(tree):
            if isinstance(n, ast.Call) and isinstance(n.func, ast.Name):
                calls.setdefault(n.func.id, []).append(n.lineno)
        # Through tab_ready_check() since v1.34.2, which times it and
        # records it; that function is the one place ready_the_tab is
        # called from.
        ready = calls.get("tab_ready_check", [])
        built = calls.get("make_fetcher", [])
        check_true("tab-ready . {} readies the tab before it builds the "
                   "fetcher".format(worker.__name__),
                   len(ready) == 1 and built and ready[0] < min(built),
                   "tab_ready_check at {}, make_fetcher at {}".format(
                       ready, built))
        handlers = calls.get("unexpected_failure", [])
        check("tab-ready . {} words an unexpected failure through "
              "unexpected_failure()".format(worker.__name__),
              len(handlers), 1)

    # The restore is attempted even when setting the folder failed.
    class Refuses:
        def __init__(self):
            self.cdp = []

        class timeouts:
            page_load = 300

        def execute_cdp_cmd(self, cmd, params):
            self.cdp.append(params.get("behavior"))
            if params.get("behavior") == "allow":
                raise TimeoutException("timed out")
            return {}
    with tempfile.TemporaryDirectory() as d:
        f = M.BrowserFetcher(os.path.join(d, "_i"))
        drv = Refuses()
        try:
            f.open(drv)
        except TimeoutException:
            pass
        notes = f.close()
        check("tab-ready . a failed handover is still put back",
              drv.cdp, ["allow", "default"])
        check_true("tab-ready . and the log says so",
                   any("set back" in n for n in notes), notes)

    # The words a person sees.
    e = TimeoutException(
        "Message: timeout: Timed out receiving message from renderer: "
        "270.000\n  (Session info: chrome=154.0.8037.58)\nStacktrace:\n"
        "0   chromedriver   0x000000010512ca84 chromedriver + 3508868")
    said = M.unexpected_failure(e)
    check_true("tab-ready . a renderer timeout is named as the browser",
               said.startswith("Chrome did not answer"), said)
    check_true("tab-ready . with its first line and no stack frames",
               "receiving message from renderer: 270.000" in said
               and "chromedriver +" not in said
               and "Session info" not in said, said)


def test_a_rerun_on_a_tab_that_does_not_answer_touches_nothing():
    """v1.34.1, at the worker: refused before the download folder is set."""
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?display=submission&article=7&context=etd"
    rows = [_row(1, "7", "primary", "public", "current",
                 "Failed: HTTP 500 Server Error",
                 base + "/cgi/viewcontent.cgi?article=7&context=etd", adm)]

    class Frozen(_FakeBrowser):
        def get(self, url):
            time.sleep(2.0)

    drv = Frozen([_Tab(drop=("a.pdf", b"%PDF-one"))])
    saved = {k: getattr(M, k) for k in
             ("_attach_or_fail", "cookie_header", "LISTING_SETTLE",
              "TAB_READY_SECONDS", "TAB_READY_GRACE")}
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: drv
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            M.TAB_READY_SECONDS, M.TAB_READY_GRACE = 0.3, 0.2
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}}, out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            "reconcile": False,
                            "fetch_path": M.FETCH_BROWSER,
                            "verify_window": 300},
                           report_dir=out)
        finally:
            for k, v in saved.items():
                setattr(M, k, v)
        summary = M.STATE.get("summary") or ""
        check_true("frozen-tab . the re-run says why it did not start",
                   summary.startswith("Not started:")
                   and "did not answer" in summary, summary)
        check("frozen-tab . Chrome's download folder was never touched",
              drv.cdp, [])
        check("frozen-tab . no file was requested", drv.visited, [])

# ---------------------------------------------------------------------------
# v1.34.2 · the duplicate defect, driven through a re-run of a record row.
# ---------------------------------------------------------------------------
def _pickvers_real(rows):
    """A revisions table in the layout the real page uses.

    Captured 2026-09-08 (an ETD record) and re-checked 2026-10-01: each
    file link sits in its own <p class="file"> beside a radio of its own,
    named per TYPE — `text.native`, `text.pdf` — so the current native and
    the current PDF are chosen separately, and may sit in different
    revisions. The older fixture, _pickvers_html, has one radio per row;
    a fake that cannot select per type is how per-type selection went
    unseen.

    rows: newest first, each (date, [(href, text, radio_name|None,
    checked), ...]). radio_name None = a link with no radio (a cover
    letter).
    """
    out = []
    for date, links in rows:
        cells = []
        for href, text, radio, checked in links:
            if radio:
                cells.append(
                    "<td class='act'><p class='file'><label>Select</label>"
                    "<input type='radio' name='{r}' value='1'{c}>"
                    "<a href=\"{h}\">{t}</a></p></td>".format(
                        r=radio, c=" checked='checked'" if checked else "",
                        h=href, t=text))
            else:
                cells.append("<td class='act'><a href=\"{}\">{}</a></td>"
                             .format(href, text))
        out.append("<tr><td><strong>An Editor</strong></td>"
                   "<td><ul><li>Revision</li></ul></td><td>{}</td>{}"
                   "<td class='version'>&nbsp;</td></tr>".format(
                       date, "".join(cells)))
    return ("<html><body><table id='revision'><thead><tr>"
            "<th>User</th><th>Comment</th><th>Date</th><th>Native</th>"
            "<th>PDF</th><th colspan='2'>Previous Versions</th></tr></thead>"
            "<tbody>" + "".join(out) + "</tbody></table></body></html>")


def _report_rows(run_dir):
    """Every row of every structure report under run_dir, as dicts."""
    import glob as _g
    import openpyxl
    rows = []
    found = 0
    for f in _g.glob(os.path.join(run_dir, "**", "*.xlsx"), recursive=True):
        if "MasterLog" in os.path.basename(f):
            continue
        found += 1
        wbk = openpyxl.load_workbook(f, read_only=True)
        try:
            data = list(wbk.active.values)
        finally:
            wbk.close()
        hdr = [str(h) for h in data[0]]
        rows += [dict(zip(hdr, r)) for r in data[1:]]
    if not found:
        raise AssertionError("no structure report under " + run_dir)
    return rows


def _rerun_one_record(pickvers_html, answer, plan_opts, ctx="ev", rows=None):
    """Re-run one Not-attempted RECORD row through retry_worker.

    That is exactly what re-running the honors_symposium _PARTIAL does:
    the record is re-planned from the Plan the report recorded, then its
    files are fetched. `answer(url, native_url)` plays the server and
    returns a Fetched or raises; every call is recorded.
    Returns (calls, rows, saved filenames).
    """
    base = "https://dc.example.edu"
    adm = base + "/cgi/editor.cgi?window=abstract&article=1000&context=" + ctx
    if rows is None:
        rows = [_row(1, "1000", M.RECORD_KIND, "", "", M.NOT_ATTEMPTED, adm,
                     adm, plan=M.plan_spec(plan_opts))]
    drv = _FakeBrowser([])
    calls = []

    def _fetch_one(fetcher, verify, url, native_url="", *a, **kw):
        calls.append((url, native_url))
        return answer(url, native_url)

    keep = {k: getattr(M, k) for k in
            ("_attach_or_fail", "cookie_header", "LISTING_SETTLE",
             "fetch_html", "fetch_one", "_throttle")}
    real_log = M.log
    with tempfile.TemporaryDirectory() as d:
        path = _fake_report(d, "DC_FileDownload_Report_{}_S.xlsx".format(ctx),
                            rows)
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: drv
            M.cookie_header = lambda driver: "c=1"
            M.LISTING_SETTLE = 0
            M._throttle = lambda *a, **k: None
            M.fetch_html = (lambda url, cookie, timeout=0:
                            pickvers_html if "pickvers" in url
                            else "<html><body></body></html>")
            M.fetch_one = _fetch_one
            M.log = lambda msg: msg
            M.STOP_EVENT.clear()
            M.retry_worker({"base_url": base, "settings": {}}, out, path,
                           {"delay": 0, "page_delay": 0, "cooldowns": 0,
                            "reconcile": False,
                            "fetch_path": M.FETCH_BROWSER,
                            "verify_window": 0},
                           report_dir=out)
        finally:
            for k, v in keep.items():
                setattr(M, k, v)
            M.log = real_log
        got = _report_rows(out)
        files = sorted(n for _r, _d, fs in os.walk(out) for n in fs
                       if not n.endswith(".xlsx")
                       and not n.startswith("DC_FileDownload_Log_"))
        logs = sorted(n for _r, _d, fs in os.walk(out) for n in fs
                      if n.startswith("DC_FileDownload_Log_"))
    return calls, got, files, logs


def _start_one_record(pickvers_html, answer, opts, ctx="ev",
                      stop_after=None):
    """Drive download_worker — a fresh Start — over one record.

    download_worker was long described as undrivable ("it needs Chrome and
    a loaded hierarchy"); what it needs is a handful of seams, all stubbed
    here, and the duplicate defect was a FRESH-START defect on 2026-10-01,
    so a check that only drove the re-run worker would have been the
    one-of-a-pair failure in the test suite itself.

    Stubbed: the Chrome attach, the tab check, the fetcher and the
    preflight, the per-structure session check, the listing (one Posted
    record), the admin pages (fetch_html), and fetch_one, which plays the
    server through `answer(url, native_url)`. Everything between — the
    planner, the dedupe, SavedFiles, the rows and the report — is real.
    `stop_after` = raise StopRequested on that fetch call (1-based).
    Returns (calls, rows, saved filenames, log lines).
    """
    base = "https://dc.example.edu"
    calls, lines = [], []

    class _Fetcher:
        def open(self, driver):
            pass

        def close(self):
            return []

        def use_cookie(self, cookie):
            pass

    def _fetch_one(fetcher, verify, url, native_url="", *a, **kw):
        calls.append((url, native_url))
        if stop_after is not None and len(calls) == stop_after:
            raise M.StopRequested("stop pressed")
        return answer(url, native_url)

    item = {"article": "1000", "title": "A talk", "release": "",
            "access": "", "af": 0}
    keep = {k: getattr(M, k) for k in
            ("_attach_or_fail", "cookie_header", "ready_the_tab",
             "make_fetcher", "preflight_downloads", "session_still_open",
             "hold_if_challenged", "crawl_listing", "fetch_html",
             "fetch_one", "_throttle", "log", "MODEL")}
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "out")
        os.makedirs(out, exist_ok=True)
        try:
            M._attach_or_fail = lambda session: _FakeBrowser([])
            M.cookie_header = lambda driver: "c=1"
            M.ready_the_tab = lambda *a, **k: None
            M.make_fetcher = lambda *a, **k: _Fetcher()
            M.preflight_downloads = lambda *a, **k: []
            M.session_still_open = lambda *a, **k: ("", False)
            M.hold_if_challenged = lambda *a, **k: None
            M.crawl_listing = lambda *a, **k: ([dict(item)], False)
            M._throttle = lambda *a, **k: None
            M.fetch_html = (lambda url, cookie, timeout=0:
                            pickvers_html if "pickvers" in url
                            else "<html><body></body></html>")
            M.fetch_one = _fetch_one
            M.log = lambda msg: lines.append(msg)
            with M.LOCK:
                M.MODEL = {"levels": {ctx: 1}, "types": {ctx: "Event"},
                           "chains": [["root", ctx]]}
            M.STOP_EVENT.clear()
            run = dict({"delay": 0, "page_delay": 0, "cooldowns": 0,
                        "fetch_path": M.FETCH_BROWSER, "verify_window": 0},
                       **opts)
            M.download_worker({"base_url": base, "settings": {}}, out,
                              "all", [], run, report_dir=out)
        finally:
            for k, v in keep.items():
                if k == "MODEL":
                    with M.LOCK:
                        M.MODEL = v
                else:
                    setattr(M, k, v)
        got = _report_rows(out)
        files = sorted(n for _r, _d, fs in os.walk(out) for n in fs
                       if not n.endswith(".xlsx")
                       and not n.startswith("DC_FileDownload_Log_"))
        logs = sorted(n for _r, _d, fs in os.walk(out) for n in fs
                      if n.startswith("DC_FileDownload_Log_"))
    return calls, got, files, lines, logs


_ALL_NATIVES = {"primary": True, "supp": False, "native": True,
                "public": True, "hidden": True, "versions": "all",
                "published": True, "unpublished": False}


def test_a_video_is_fetched_once_in_all_versions_mode():
    """THE PLANT for the duplicate defect, v1.34.2.

    honors_symposium, 2026-10-01, All versions with natives on: each video
    record came down three times — X.mp4, X_2.mp4, X_3.mp4, identical
    sizes — because three jobs named one file at three addresses:

      1. the primary, which has no derivative and falls back to the native;
      2. the natives job, at the short path-form URL;
      3. the current revision's native, at its dated CGI URL.

    The record here is that one: a current revision whose native is
    selected and an original revision with a native of its own. The
    server answers the way it did: the derivative is absent so the
    primary arrives through its native; any address for the current
    native gives the current bytes.

    Must hold: ONE current copy and ONE original copy on disk, and the
    current native requested once.
    """
    cur = ("viewcontent.cgi?type=native&article=1000&unstamped=yes"
           "&date=1592928048&preview_mode=1&context=ev&/1592928048-text.native")
    old = ("viewcontent.cgi?type=native&article=1000&unstamped=yes"
           "&date=1592926724&preview_mode=1&context=ev&/1592926724-text.native")
    html = _pickvers_real([
        ("Tue Jun 23 09:00:00 2020", [(cur, "MPEG-4", "text.native", True)]),
        ("Tue Jun 23 08:38:00 2020", [(old, "MPEG-4", "text.native", False)]),
    ])
    CURRENT, ORIGINAL = b"\x00\x00\x00\x18ftypmp42-current", \
                        b"\x00\x00\x00\x18ftypmp42-original"

    def answer(url, native_url):
        if "date=1592926724" in url:
            return M.Fetched(data=ORIGINAL, name="talk.mp4")
        if "type/native" in url or "type=native" in url:
            return M.Fetched(data=CURRENT, name="talk.mp4")
        # The derivative: no PDF behind it, so the native answers.
        if native_url:
            return M.Fetched(data=CURRENT, name="talk.mp4", used_native=True)
        raise M.NoFileAvailable("No PDF has been provided")

    # BOTH workers. honors_symposium was a fresh Start (download_worker);
    # re-running its _PARTIAL is retry_worker. A check on one of them is
    # the one-of-a-pair failure, in the suite.
    for who, drive in (("start", _start_one_record),
                       ("re-run", _rerun_one_record)):
        got = drive(html, answer, _ALL_NATIVES)
        calls, rows, files = got[0], got[1], got[2]
        done = [r for r in rows if r["Status"] == "Downloaded"]
        current = [r for r in done if r["Version"] == "current"]
        check("dup-video {} . one current copy in the report".format(who),
              len(current), 1)
        check("dup-video {} . one original copy in the report".format(who),
              len([r for r in done if r["Version"] == "original"]), 1)
        check("dup-video {} . two files on disk, not four".format(who),
              len(files), 2)
        check_true("dup-video {} . no _2 or _3 copy on disk".format(who),
                   not any(re.search(r"_\d\.mp4$", f) for f in files),
                   ", ".join(files))
        check("dup-video {} . two requests: the primary and the "
              "original".format(who), len(calls), 2)
        check_true("dup-video {} . the saved current copy is labelled "
                   "native".format(who),
                   current and current[0]["File Kind"] == "native",
                   str(current[0]["File Kind"] if current else None))
        check("dup-video {} . nothing else in the report".format(who),
              len(rows), 2)
        # And the run's whole log lands beside its reports, from BOTH
        # workers — the pair rule, for save_run_log.
        check("dup-video {} . the run's log is saved as a .txt".format(who),
              len(got[-1]), 1)


def test_a_pdf_record_keeps_stamped_and_unstamped_named_apart():
    """The sspeach shape, v1.34.2.

    A Word upload in a peer-reviewed journal, All versions with natives on.
    The editor selected the current PDF from the NEWER revision and the
    current native from the OLDER one — the real page selects per type.
    Article 1004 on 2026-10-01 was read as two "current" revisions, and its
    older PDF was saved as current.

    Must hold:
      - the stamped default stream is kept, as `current`;
      - the revisions table's unstamped copy of the current PDF is kept
        too — different bytes — and labelled `current-unstamped`;
      - the current native comes down ONCE, though two addresses name it;
      - the older revision's PDF is `original`, not `current`;
      - the cover letter beside the selected PDF is the current one.
    """
    pdf_new = ("viewcontent.cgi?type=pdf&article=1000&unstamped=yes"
               "&date=1701352305&preview_mode=1&context=jj&/1701352305-text.pdf")
    pdf_old = ("viewcontent.cgi?type=pdf&article=1000&unstamped=yes"
               "&date=1691017077&preview_mode=1&context=jj&/1691017077-text.pdf")
    nat_old = ("viewcontent.cgi?type=native&article=1000&unstamped=yes"
               "&date=1691017070&preview_mode=1&context=jj"
               "&/1691017070-text.native")
    cover = "viewcoverletter.cgi?article=1000&context=jj&date=1701352305"
    html = _pickvers_real([
        ("Thu Nov 30 05:51:00 2023", [(pdf_new, "PDF", "text.pdf", True),
                                      (cover, "Cover letter", None, False)]),
        ("Wed Aug  2 15:57:00 2023", [(nat_old, "MS Word", "text.native",
                                       True),
                                      (pdf_old, "PDF", "text.pdf", False)]),
    ])

    def answer(url, native_url):
        if "viewcoverletter" in url:
            return M.Fetched(data=b"%PDF-cover", name="cover.pdf")
        if "type/native" in url or "type=native" in url:
            return M.Fetched(data=b"PK\x03\x04docx", name="paper.docx")
        if "date=1701352305" in url:
            return M.Fetched(data=b"%PDF-unstamped-new", name="paper.pdf")
        if "date=1691017077" in url:
            return M.Fetched(data=b"%PDF-unstamped-old", name="paper.pdf")
        return M.Fetched(data=b"%PDF-stamped-with-cover-page", name="p.pdf")

    for who, drive in (("start", _start_one_record),
                       ("re-run", _rerun_one_record)):
        got = drive(html, answer, _ALL_NATIVES, ctx="jj")
        calls, rows, files = got[0], got[1], got[2]
        done = {(r["File Kind"], r["Version"]): r for r in rows
                if r["Status"] == "Downloaded"}
        check("pdf-shape {} . one row per distinct file".format(who),
              sorted(done), sorted([
                  ("primary", "current"),
                  ("primary", M.VERSION_CURRENT_UNSTAMPED),
                  ("native", "current"), ("primary", "original"),
                  ("coverletter", "current")]))
        check("pdf-shape {} . and no row beyond them".format(who),
              len(rows), 5)
        check("pdf-shape {} . five files on disk".format(who), len(files), 5)
        check_true("pdf-shape {} . no _2 copy on disk".format(who),
                   not any(re.search(r"_\d\.\w+$", f) for f in files),
                   ", ".join(files))
        check("pdf-shape {} . the current native was asked for "
              "once".format(who), sum(1 for u, _n in calls if "native" in u), 1)
        check_true("pdf-shape {} . the unstamped copy says so in its "
                   "filename".format(who),
                   any(M.VERSION_CURRENT_UNSTAMPED in f for f in files),
                   ", ".join(files))
        old_pdf = done.get(("primary", "original"))
        check_true("pdf-shape {} . the older PDF is the original, at its "
                   "own URL".format(who),
                   old_pdf is not None and "1691017077" in old_pdf["File URL"],
                   str(old_pdf and old_pdf["File URL"]))


def test_a_rerun_stopped_mid_record_keeps_the_rest_of_the_record():
    """Eighteenth one-of-a-pair, found reading for v1.34.2.

    A re-run that stops part-way through a RE-PLANNED record wrote the
    files it had fetched and dropped the record row, because file rows
    existed for it — so the record's remaining files were in no row at
    all, and the next re-run could not find them. download_worker has
    written them since v1.33 (unreached_job_rows); retry_worker never
    did. Reproduced by driving the worker before the fix: one row,
    `native current Downloaded`, and the original revision nowhere.

    Re-running the honors_symposium _PARTIAL is exactly this case: 344
    record rows, each re-planned in All-versions mode, any of which a
    prompt or a Stop can interrupt.
    """
    cur = ("viewcontent.cgi?type=native&article=1000&unstamped=yes"
           "&date=1592928048&preview_mode=1&context=ev&/1592928048-text.native")
    old = ("viewcontent.cgi?type=native&article=1000&unstamped=yes"
           "&date=1592926724&preview_mode=1&context=ev&/1592926724-text.native")
    html = _pickvers_real([
        ("Tue Jun 23 09:00:00 2020", [(cur, "MPEG-4", "text.native", True)]),
        ("Tue Jun 23 08:38:00 2020", [(old, "MPEG-4", "text.native", False)]),
    ])

    def answer(url, native_url):
        M.STOP_EVENT.set()          # pressed while this file arrives
        return M.Fetched(data=b"\x00\x00\x00\x18ftypmp42",
                         name="talk.mp4", used_native=bool(native_url))

    try:
        calls, rows, files, _logs = _rerun_one_record(html, answer,
                                                      _ALL_NATIVES)
    finally:
        M.STOP_EVENT.clear()
    got = sorted((r["File Kind"], r["Version"], r["Status"]) for r in rows)
    check("rerun-stop . the rest of the record is recorded, and only that",
          got, sorted([("native", "current", "Downloaded"),
                       ("native", "original", M.NOT_ATTEMPTED)]))
    saved_row = [r for r in rows if r["Status"] == "Downloaded"]
    check_true("rerun-stop . the fallback row points at the native it got",
               saved_row and "type/native" in saved_row[0]["File URL"],
               str(saved_row and saved_row[0]["File URL"]))
    check("rerun-stop . no record row left beside the file rows",
          [r for r in rows if r["File Kind"] == M.RECORD_KIND], [])


def test_a_stop_after_the_fallback_leaves_no_duplicate_row():
    """Driven: a fresh Start stopped after the primary fell back.

    The record has a current native (reached by the fallback) and an
    original revision. Stop arrives as the original is asked for. The
    _PARTIAL must hold the current native as Downloaded and the original
    as Not attempted — and NOT the natives job, which a re-run would
    otherwise fetch as a second copy.
    """
    cur = ("viewcontent.cgi?type=native&article=1000&unstamped=yes"
           "&date=1592928048&preview_mode=1&context=ev&/1592928048-text.native")
    old = ("viewcontent.cgi?type=native&article=1000&unstamped=yes"
           "&date=1592926724&preview_mode=1&context=ev&/1592926724-text.native")
    html = _pickvers_real([
        ("Tue Jun 23 09:00:00 2020", [(cur, "MPEG-4", "text.native", True)]),
        ("Tue Jun 23 08:38:00 2020", [(old, "MPEG-4", "text.native", False)]),
    ])

    def answer(url, native_url):
        if native_url:
            return M.Fetched(data=b"\x00\x00\x00\x18ftypmp42",
                             name="talk.mp4", used_native=True)
        return M.Fetched(data=b"\x00\x00\x00\x18ftypmp42-o",
                         name="talk.mp4")

    want = sorted([("native", "current", "Downloaded"),
                   ("native", "original", M.NOT_ATTEMPTED)])
    # Two ways to stop, and they leave different jobs unreached. (a) the
    # Stop arrives as the original is asked for: the natives job was
    # already skipped. (b) Stop is pressed WHILE the primary downloads, so
    # it is seen at the top of the very next job — the natives job — and
    # that job is in the unreached list. Only (b) reaches the writer's
    # filter; (a) alone left a planted removal of it green.
    calls, rows, files, _lines, _logs = _start_one_record(
        html, answer, _ALL_NATIVES, stop_after=2)
    check("stop-dup start (a) . rows: the saved native and the unreached "
          "original", sorted((r["File Kind"], r["Version"], r["Status"])
                             for r in rows), want)
    check("stop-dup start (a) . one file on disk", len(files), 1)

    def answer_then_stop(url, native_url):
        M.STOP_EVENT.set()          # pressed while this file arrives
        return answer(url, native_url)

    calls, rows, files, _lines, _logs = _start_one_record(
        html, answer_then_stop, _ALL_NATIVES)
    M.STOP_EVENT.clear()
    check("stop-dup start (b) . one request before the stop", len(calls), 1)
    check("stop-dup start (b) . rows: the saved native and the unreached "
          "original — not the natives job",
          sorted((r["File Kind"], r["Version"], r["Status"]) for r in rows),
          want)


def test_a_stop_between_fallback_and_native_leaves_no_duplicate_to_rerun():
    """The unreached-rows writer asks SavedFiles too (v1.34.2).

    A run stopped after the primary fell back to the native, and before
    the natives job: written as Not attempted, the natives job would be
    fetched by the re-run — the duplicate the run itself declined.
    """
    item = {"article": "1000", "title": "T"}
    jobs = [
        {"kind": "native", "visibility": "public", "version": "current",
         "version_date": "", "orig_hint": "", "access": "", "release": "",
         "embargo": "",
         "url": "https://dc/context/ev/article/1000/type/native/viewcontent"},
        {"kind": "supp1", "visibility": "public", "version": "current",
         "version_date": "", "orig_hint": "x.pdf", "access": "",
         "release": "", "embargo": "",
         "url": "https://dc/cgi/viewcontent.cgi?filename=0&article=1000"},
    ]
    saved = M.SavedFiles()
    saved.saved("ev", "1000", {"kind": "native", "version": "current",
                               "url": "https://dc/other"}, "X.mp4")
    rows = M.unreached_job_rows(
        jobs, 0, 5, item, "Posted", "pub", "adm",
        already=lambda job: saved.already("ev", "1000", job))
    check("stop-dup . only the file not yet saved is left to re-run",
          [r[4] for r in rows], ["supp1"])
    check("stop-dup . its order follows on without a gap",
          [r[0] for r in rows], [6])
    check("stop-dup . the summary says what was skipped, once asked",
          M.SavedFiles().summary(), "")


# ---------------------------------------------------------------------------
# v1.34.2 · a clean start, the whole log, and a form that survives a reopen
# ---------------------------------------------------------------------------
def _get(port, path):
    with urllib.request.urlopen("http://127.0.0.1:{}{}".format(port, path),
                                timeout=10) as r:
        return r.status, r.read().decode("utf-8")


def test_a_clean_start_keeps_the_settings_and_the_hierarchy():
    """Jeff, 2026-10-01: a way to clear the module and its log for a clean
    start, only when no run is active. Driven through the endpoints."""
    session = M.load_session(Path("/nonexistent/session.json"))
    session["hub_url"] = "http://127.0.0.1:8750"
    srv = ThreadingHTTPServer(
        ("127.0.0.1", 0), M.make_handler(session, M.build_page(session)))
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    real_log_print = M.log
    keep = {k: (dict(v) if isinstance(v, dict) else
                list(v) if isinstance(v, list) else v)
            for k, v in M.STATE.items()}
    try:
        M.begin_run_log()
        for n in range(250):
            M.log("line {}".format(n))
        with M.LOCK:
            M.STATE["phase"] = "downloading"
            M.STATE["summary"] = "the last run"
            M.STATE["last_file"] = "/tmp/master.xlsx"
            M.STATE["hier"] = dict(M.STATE["hier"], loaded=True, count=3)
            M.STATE["form"] = {"fields": {"rptdir": "/Volumes/Reports"},
                               "parents": ["etd"]}
        code, txt = _get(port, "/api/log")
        check("clean . the whole run is served, past the page's 200 lines",
              len(txt.strip().split("\n")), 250)
        check("clean . while the page itself keeps 200",
              len(M.STATE["log"]), 200)
        code, j = _post(port, "/api/clear", {})
        check("clean . refused while a run is active", code, 409)
        check("clean . and nothing was cleared",
              (len(M.STATE["log"]), M.STATE["summary"]), (200, "the last run"))
        with M.LOCK:
            M.STATE["phase"] = "done"
            M.STATE["claimed"] = True
        code, j = _post(port, "/api/clear", {})
        check("clean . refused while a run still holds the claim", code, 409)
        with M.LOCK:
            M.STATE["claimed"] = False
        before = M.STATE["cleared"]
        code, j = _post(port, "/api/clear", {})
        check("clean . allowed once no run is active", code, 200)
        check("clean . the log, summary and links are gone",
              (M.STATE["log"], M.STATE["summary"], M.STATE["last_file"],
               M.STATE["phase"]), ([], "", None, "idle"))
        check("clean . the run's whole log too", _get(port, "/api/log")[1], "")
        check("clean . the loaded hierarchy stays",
              M.STATE["hier"]["loaded"], True)
        check("clean . the form stays, destinations and all",
              ((M.STATE["form"] or {}).get("fields") or {}).get("rptdir"),
              "/Volumes/Reports")
        check("clean . every open page is told", M.STATE["cleared"],
              before + 1)
    finally:
        srv.shutdown()
        with M.LOCK:
            M.STATE.clear()
            M.STATE.update(keep)
        M.log = real_log_print


def test_the_whole_log_starts_with_the_run_and_is_saved_beside_it():
    """Copy log = exactly the current run. Driven: a line logged before
    the run is not in it, a line logged by the run is, and the .txt the
    workers write beside the master log holds the same text."""
    session = M.load_session(Path("/nonexistent/session.json"))
    session["hub_url"] = "http://127.0.0.1:8750"
    srv = ThreadingHTTPServer(
        ("127.0.0.1", 0), M.make_handler(session, M.build_page(session)))
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    real_worker = M.retry_worker
    real_lock = (M.claim_run_lock, M.release_run_lock)
    ran = threading.Event()
    try:
        M.log("a line from the PREVIOUS run")

        def _fake_worker(*a, **k):
            M.log("a line from THIS run")
            ran.set()

        M.retry_worker = _fake_worker
        M.claim_run_lock = lambda addr, port: (True, None)
        M.release_run_lock = lambda addr: None
        with tempfile.TemporaryDirectory() as d:
            rpt = _fake_report(d, "DC_FileDownload_Report_etd_S.xlsx", [
                _row(1, "7", "primary", "public", "current",
                     M.NOT_ATTEMPTED, "https://dc/cgi/viewcontent.cgi?a=7",
                     "https://dc/cgi/editor.cgi?article=7")])
            code, j = _post(port, "/api/retry", {
                "out_dir": d, "report_dir": "", "report_path": rpt,
                "fetch_path": M.FETCH_BROWSER, "verify_window": 0})
            check("whole-log . the re-run was accepted", code, 200)
            ran.wait(5)
            time.sleep(0.2)
            txt = _get(port, "/api/log")[1]
            check_true("whole-log . the run's own line is in it",
                       "a line from THIS run" in txt, txt[-200:])
            check_true("whole-log . the previous run's is not",
                       "PREVIOUS run" not in txt, txt[:200])
            path = M.save_run_log(d, "20261001_120000")
            saved = Path(path).read_text(encoding="utf-8") if path else ""
            check_true("whole-log . a .txt is written beside the reports",
                       os.path.basename(path or "")
                       == "DC_FileDownload_Log_20261001_120000.txt", path)
            check_true("whole-log . and holds the run, saying where it went",
                       "a line from THIS run" in saved
                       and "saved as DC_FileDownload_Log_" in saved, saved)
            check("whole-log . no folder, no file, no error",
                  M.save_run_log(os.path.join(d, "missing"), "x"), "")
    finally:
        M.retry_worker = real_worker
        M.claim_run_lock, M.release_run_lock = real_lock
        srv.shutdown()


def test_the_form_is_kept_for_a_reopened_page():
    """2026-10-01: a reopened tab came back with "Save reports to" blank
    and the mode on Start. The module now keeps the form; this drives the
    endpoint and then the page's own restore code in node."""
    good = {"fields": {"rptdir": "/Volumes/Reports", "jobresume": True,
                       "jobnew": False, "verall": True},
            "parents": ["etd", "honors_symposium"]}
    check("form . a good form is accepted",
          M.validate_form_state(good)["fields"]["rptdir"], "/Volumes/Reports")
    for bad, why in ((dict(good, fields={"evil": "x"}), "an unknown field"),
                     (dict(good, fields={"rptdir": ["x"]}), "a list value"),
                     (dict(good, parents="etd"), "parents not a list"),
                     ("nope", "not an object")):
        try:
            M.validate_form_state(bad)
            check_true("form . refuses {}".format(why), False)
        except ValueError:
            check_true("form . refuses {}".format(why), True)
    page = M.build_page({"base_url": "https://dc", "settings": {}}).decode()
    missing = [f for f in M.FORM_FIELDS
               if 'id="{}"'.format(f) not in page]
    check("form . every kept field is a field on the page", missing, [])

    import shutil as _shutil
    import subprocess as _sp
    node = _shutil.which("node")
    if not node:
        return
    start = page.find("const FORM_IDS=")
    end = page.find("FORM_IDS.forEach(id=>{\n  const el=document."
                    "getElementById(id); if(!el) return;\n  el.addEvent")
    check_true("form . the restore code was found", start > 0 and end > start)
    js = page[start:end]
    harness = (
        "const els={};\n"
        "M_FIELDS.forEach(id=>{els[id]={id:id,type:(id==='rptdir'||"
        "id==='dldir'||id==='retryrpt'||id==='repdir')?'text':'checkbox',"
        "value:'',checked:false};});\n"
        "els.jobnew.checked=true;\n"
        "const document={getElementById:id=>els[id]||null};\n"
        "let checked=new Set(); let hierLoaded=false, nodes=[];\n"
        "let modeShown=0; function applyJobMode(){modeShown++;}\n"
        "function syncScope(){}\n"
        "function api(){}\n" + js +
        "\nrestoreForm(" + json.dumps(good) + ");\n"
        "console.log(JSON.stringify([els.rptdir.value, els.jobresume.checked,"
        " els.jobnew.checked, els.verall.checked, pendingParents, modeShown,"
        " formRestored]));\n").replace(
            "M_FIELDS", json.dumps(list(M.FORM_FIELDS)))
    r = _run_node(node, harness)
    check("form . the page's restore ran", r.returncode, 0)
    if r.returncode == 0:
        got = json.loads(r.stdout.strip())
        check("form . a reopened page gets its destination back",
              got[0], "/Volumes/Reports")
        check("form . and its mode", got[1:3], [True, False])
        check("form . and its options", got[3], True)
        check("form . and the parents it had checked, for the hierarchy",
              got[4], ["etd", "honors_symposium"])
        check("form . and shows the steps for that mode", got[5] >= 1, True)
    else:
        check_true("form . node said", False, r.stderr[-300:])


def test_the_idle_line_says_what_to_do_next():
    """It said "load or map a hierarchy first" after one was loaded, and
    in re-run mode, which needs none. Run in node against the page."""
    import shutil as _shutil
    import subprocess as _sp
    node = _shutil.which("node")
    if not node:
        return
    page = M.build_page({"base_url": "https://dc", "settings": {}}).decode()
    start = page.find("function idleStatus(){")
    end = page.find("\n}\n", start) + 3
    js = page[start:end]
    out = []
    for loaded, resume in ((False, False), (True, False), (False, True)):
        harness = ("const status={className:'',textContent:'Idle.'};\n"
                   "const jobResume={checked:%s}; let hierLoaded=%s;\n"
                   % ("true" if resume else "false",
                      "true" if loaded else "false")
                   + js + "\nidleStatus(); console.log(status.textContent);")
        r = _run_node(node, harness)
        out.append(r.stdout.strip() if r.returncode == 0 else r.stderr[-200:])
    check("idle . nothing loaded asks for a hierarchy",
          out[0], "Idle — load or map a hierarchy first.")
    check("idle . a loaded hierarchy moves on to the options",
          out[1], "Idle — choose the scope and options, then Start downloads.")
    check("idle . re-run mode asks for the report, not a hierarchy",
          out[2], "Idle — name the previous report, then Re-run from report.")
    check_true("idle . and the poller repaints it every time it is idle",
               "}else if(s.phase==='idle'){\n    idleStatus();" in page,
               "the old guard repainted only an empty line")


def test_the_messages_say_only_what_is_true():
    """v1.34.2, the untrue and awkward messages from 2026-10-01's runs."""
    check("msg . the first prompt has no 'previous prompt'",
          M.prompt_announcement(1, 0, first=True),
          "Verification asked for at file request 1 of this run — the "
          "first prompt of the run, after 0 file(s).")
    check_true("msg . a later one still counts from the previous",
               "after the previous prompt" in
               M.prompt_announcement(151, 149, first=False),
               M.prompt_announcement(151, 149, first=False))
    stopped = [(1, "honors_symposium", "Event", 391, 47, 2, 0,
                "STOPPED mid-structure (487 row(s): 142 downloaded, 345 "
                "never reached)", "17:37", "x.xlsx")]
    check("msg . a structure stopped part-way is not 'done'",
          M.structures_done_text(stopped, 1),
          "0 of 1 structure(s) finished, 1 stopped part-way")
    done_and_stopped = [(1, "a", "Series", 3, 3, 0, 0, "OK", "", ""),
                        stopped[0]]
    check("msg . finished ones are counted apart from the stopped one",
          M.structures_done_text(done_and_stopped, 5),
          "1 of 5 structure(s) finished, 1 stopped part-way")
    check("msg . and nothing is said about stopping when nothing stopped",
          M.structures_done_text(done_and_stopped[:1], 1),
          "1 of 1 structure(s) finished")

    # The Verifications sheet's column says what its first row means.
    check_true("msg . the sheet's count says it may run from the start",
               "Files since previous prompt or run start"
               in M.VERIFY_HEADERS, str(M.VERIFY_HEADERS))

    # "1 page requests".
    check("msg . one saved page request is singular",
          M.showall_note(M.LISTING_PAGE_SIZE + 5).endswith(
              "saving about 1 page request"), True)
    check("msg . two are plural",
          M.showall_note(2 * M.LISTING_PAGE_SIZE + 5).endswith(
              "saving about 2 page requests"), True)


def test_ready_the_tab_says_what_it_did_and_puts_the_timeout_back():
    """v1.34.2: it was silent, and it left Chrome's page-load timeout at
    its own 20 s — which the fetcher then 'restored' as if it were the
    bridge's setting."""
    class Tab:
        title = "Digital Commons"
        current_url = ""

        class timeouts:
            page_load = 300000          # ms, as Selenium reports it

        def __init__(self):
            self.calls = []

        def set_page_load_timeout(self, s):
            self.calls.append(s)

        def get(self, url):
            self.current_url = url

    d = Tab()
    checks, said = [], []
    title = M.tab_ready_check(d, "https://dc/", checks, said.append)
    check("tab-check . returns the title", title, "Digital Commons")
    check("tab-check . the page-load timeout is put back as it was",
          d.calls[-1], 300.0)
    check("tab-check . and was bounded during the move",
          d.calls[0], float(M.TAB_READY_SECONDS))
    check("tab-check . a Checks row, passed", (checks[0][0], checks[0][1]),
          ("Tab ready", "passed"))
    check_true("tab-check . which says where and how long",
               "moved to https://dc/ and answered in " in checks[0][2],
               checks[0][2])
    check_true("tab-check . and the log says it too",
               said and said[0].startswith("Chrome's tab is ready: "),
               str(said))


def test_the_first_prompt_is_announced_as_the_first():
    """Driven through fetch_one(): the first prompt of a run said "0
    file(s) after the previous prompt" on 2026-10-01, when there had been
    no previous prompt."""
    lines = []

    class Asks:
        """Asks once per round of three fetches, then delivers."""
        def __init__(self):
            self.n = 0

        def fetch(self, url, native_url=""):
            self.n += 1
            if self.n in (1, 4):
                raise M.VerificationRequired("Just a moment...")
            return M.Fetched(data=b"%PDF-x", name="a.pdf")

        def is_challenge(self):
            return False

    verify = M.VerificationWatch(window=5, announce=lines.append)
    f = Asks()
    M.fetch_one(f, verify, "https://dc/1", "", 1, announce=lines.append)
    M.fetch_one(f, verify, "https://dc/2", "", 2, announce=lines.append)
    M.fetch_one(f, verify, "https://dc/3", "", 3, announce=lines.append)
    asked = [l for l in lines if l.startswith("Verification asked for")]
    check("first-prompt . two prompts were announced", len(asked), 2)
    check_true("first-prompt . the first says it is the first",
               asked and "the first prompt of the run" in asked[0]
               and "previous prompt" not in asked[0], str(asked[:1]))
    check_true("first-prompt . the second counts from the first",
               len(asked) > 1 and "after the previous prompt" in asked[1],
               str(asked[1:]))


def test_a_rerun_takes_a_records_primary_before_its_native():
    """Live, 2026-10-02: article 1061 was stopped mid-transfer, leaving
    its primary AND its native as Not attempted rows. The re-run sorted a
    record's rows by kind name, so `native` came first; the primary then
    fell back to that same native and saved it again. SavedFiles cannot
    see a fallback before the fetch, so the order has to be right.

    The rows are written native first, as a report may hold them.
    """
    base = "https://dc.example.edu"
    ctx = "ev"
    adm = base + "/cgi/editor.cgi?window=abstract&article=1061&context=" + ctx
    nat = base + "/context/ev/article/1061/type/native/viewcontent?preview_mode=1"
    pri = base + "/cgi/viewcontent.cgi?article=1061&context=ev&preview_mode=1"
    rows = [_row(1, "1061", "native", "public", "current", M.NOT_ATTEMPTED,
                 nat, adm),
            _row(2, "1061", "primary", "public", "current", M.NOT_ATTEMPTED,
                 pri, adm)]

    def answer(url, native_url):
        # No derivative: the primary arrives through its native.
        return M.Fetched(data=b"\x00\x00\x00\x18ftypmp42", name="talk.mp4",
                         used_native=bool(native_url))

    calls, got, files, _logs = _rerun_one_record("", answer, _ALL_NATIVES,
                                                 ctx=ctx, rows=rows)
    check("rerun-order . one copy of the video, not two", len(files), 1)
    check("rerun-order . one request", len(calls), 1)
    check_true("rerun-order . and it was the primary's, with its fallback",
               calls and calls[0][0] == pri and calls[0][1], str(calls))
    check("rerun-order . the merge puts a record's primary first",
          [r["kind"] for r in M.merge_rows([("s", [
              {"context": "c", "article": "1", "kind": k, "url": k,
               "visibility": "public", "version": "current",
               "version_date": ""} for k in ("coverletter", "native",
                                             "primary", "supp1")])])],
          ["primary", "coverletter", "native", "supp1"])


def main():
    for fn in (test_a_rerun_on_a_tab_that_does_not_answer_touches_nothing,
               test_a_tab_that_does_not_answer_stops_the_run_before_anything,
               test_every_downloaded_row_says_how_it_was_fetched,
               test_fetch_commentary_is_progress_in_both_workers,
               test_one_run_per_chrome_across_processes,
               test_content_link_type, test_native_variant,
               test_parse_pickvers, test_parse_additional_files,
               test_fetch_record_file, test_sniff_extension,
               test_build_filename, test_dedupe_jobs, test_plan_item_jobs,
               test_report_shape, test_endpoints, test_native_file_url,
               test_listing_urls, test_parse_edikit_listing,
               test_parse_gallery_listing,
               test_flavor_detection_is_one_directional,
               test_listing_row_hazards, test_crawl_listing,
               test_supplemental_gate, test_describe_error,
               test_no_file_available_names_both_attempts,
               test_retry_after_parsing, test_open_retries_only_rate_limits,
               test_rate_limit_never_becomes_no_file, test_adaptive_delay,
               test_pushback_classification,
               test_definite_absence_classification,
               test_redirect_loop_is_retried_like_a_throttle,
               test_absence_requires_two_real_answers,
               test_a_wait_is_announced, test_stop_writes_what_it_has,
               test_retry_after_is_honored_not_clamped,
               test_a_long_wait_stays_stoppable,
               test_report_destinations,
               test_reports_are_written_where_the_split_says,
               test_a_cooldown_is_waited_out_within_budget,
               test_a_cooldown_wait_is_stoppable,
               test_the_retry_loop_terminates_whatever_the_budget,
               # Windows CI, 2026-09-21
               test_every_workbook_this_module_opens_can_be_closed,
               test_read_failed_rows,
               test_retry_native_fallback_is_offered_only_where_it_is_valid,
               test_the_retry_endpoint_is_wired,
               test_the_run_sets_its_own_pace,
               test_the_paged_fallback_still_works_if_showall_is_ignored,
               test_progress_and_eta_measure_records_not_structures,
               test_a_refused_showall_falls_back_to_paging,
               test_one_bad_state_does_not_discard_the_others,
               test_revision_links_resolve_against_their_page,
               test_a_regeneration_link_is_never_harvested,
               test_the_family_is_not_re_guessed_once_known,
               test_a_long_eta_is_readable,
               test_an_empty_record_state_is_not_a_failure,
               test_access_state_is_read_and_never_guessed,
               # open item 8 (2026-09-21): one marker list, from the profile
               test_the_identity_gate_and_this_suite_read_one_list,
               test_the_listing_carries_the_release_option,
               test_a_type_column_is_not_always_a_release_column,
               test_the_embargo_parser_harvests_no_links,
               test_an_absence_is_reported_as_an_absence,
               test_a_definite_absence_is_not_retried,
               test_the_plan_round_trips,
               test_unreached_records_are_written_and_resumable,
               test_an_unreached_record_is_replanned_not_refetched,
               test_the_early_exits_record_what_they_never_reached,
               test_a_failed_one_page_listing_falls_back_to_paging,
               test_a_run_of_403s_is_a_block_not_eighty_problems,
               test_both_workers_watch_for_the_block,
               test_a_blocked_rerun_stops_and_says_so_once,
               test_a_session_is_capped_in_both_units,
               test_a_session_re_reads_the_listing_before_it_fetches,
               test_a_session_acts_on_what_the_listing_says_now,
               test_a_stopped_rerun_records_what_it_never_reached,
               test_the_page_does_not_shout_its_evidence_at_staff,
               test_a_full_session_is_not_reported_as_an_interruption,
               # v1.33.8 — an expired session skipped a real record
               test_a_bounce_is_recognized_by_where_the_bytes_came_from,
               test_an_empty_listing_is_not_every_record_going_away,
               test_the_browser_never_loads_the_big_listing,
               test_the_rerun_summary_counts_what_happened,
               test_pages_and_downloads_are_paced_apart,
               test_the_plan_workbook_carries_the_job_that_made_it,
               test_a_session_takes_its_share_and_says_what_is_left,
               test_sessions_do_not_redo_each_others_work,
               test_a_blocked_run_stops_instead_of_grinding,
               test_inventory_mode_records_without_fetching,
               test_no_name_is_used_without_being_bound,
               test_a_stopped_rerun_writes_what_it_had,
               test_the_rerun_plans_as_it_goes,
               test_two_simultaneous_starts_produce_one_run,
               test_the_rerun_control_is_findable_and_honestly_named,
               test_a_stopped_run_accounts_for_every_row,
               test_a_row_of_the_wrong_width_is_refused,
               # v1.33 — the browser fetch path and the verification prompt
               test_a_download_is_a_transition_not_an_appearance,
               test_the_browser_path_tells_the_three_outcomes_apart,
               test_a_challenge_does_not_spend_a_second_request_on_the_native,
               test_the_browser_gets_its_download_folder_back,
               test_a_fetched_file_is_moved_not_copied_and_never_empty,
               test_the_hold_waits_passively_and_says_how_long_is_left,
               test_stop_works_while_the_run_is_holding,
               test_the_prompt_record_is_the_audit_trail,
               test_a_prompt_is_retried_properly_and_bounded,
               test_nobody_confirming_is_a_stop_reason_not_a_failure,
               test_a_partial_is_cleared_before_the_retry,
               test_the_in_flight_records_remaining_files_are_not_lost,
               test_both_workers_fetch_the_same_way,
               test_a_rerun_through_the_browser_lands_files_and_an_audit_trail,
               test_a_rerun_nobody_verifies_stops_without_calling_it_a_failure,
               test_the_page_offers_the_fetch_path_and_the_window,
               test_the_modules_wording_names_no_vendor_and_no_agreement,
               test_both_endpoints_refuse_a_fetch_path_they_cannot_take,
               # v1.33.1 — the two ways a prompt went unrecorded
               test_a_prompt_is_seen_while_the_run_waits_for_the_file,
               test_a_prompt_cleared_in_the_gap_is_still_recorded,
               test_the_outcome_says_what_was_observed_not_who_acted,
               test_the_page_between_structures_is_checked_too,
               test_a_rendered_page_is_read_at_once_not_after_the_file_wait,
               test_a_prompt_on_the_derivative_survives_the_native_fallback,
               test_both_workers_check_the_page_between_structures,
               # v1.33.2 — what the second live run found
               test_a_record_whose_planning_fails_still_gets_a_row,
               test_a_stopped_run_still_reports_what_the_monitoring_saw,
               test_a_record_that_will_not_replan_is_still_reported,
               test_a_video_is_fetched_once_in_all_versions_mode,
               test_a_pdf_record_keeps_stamped_and_unstamped_named_apart,
               test_a_stop_between_fallback_and_native_leaves_no_duplicate_to_rerun,
               test_a_stop_after_the_fallback_leaves_no_duplicate_row,
               test_a_rerun_stopped_mid_record_keeps_the_rest_of_the_record,
               test_a_rerun_takes_a_records_primary_before_its_native,
               test_a_missing_native_is_given_time_to_say_so,
               test_a_clean_start_keeps_the_settings_and_the_hierarchy,
               test_the_whole_log_starts_with_the_run_and_is_saved_beside_it,
               test_the_form_is_kept_for_a_reopened_page,
               test_the_idle_line_says_what_to_do_next,
               test_the_messages_say_only_what_is_true,
               test_the_first_prompt_is_announced_as_the_first,
               test_ready_the_tab_says_what_it_did_and_puts_the_timeout_back,
               # v1.33.3 — the third way the server says "no file"
               test_a_redirect_to_the_site_root_is_noticed_not_concluded,
               # v1.33.7 — a redirect from the page it redirects to
               test_a_redirect_from_the_page_it_redirects_to_is_still_a_redirect,
               test_a_browser_that_will_not_say_falls_back_rather_than_claims,
               test_a_download_leaves_the_document_alone,
               # v1.33.4 — a prompt that interrupts a transfer
               test_a_prompt_that_interrupts_a_transfer_does_not_cost_the_transfer,
               test_a_hold_that_ends_in_bytes_waits_instead_of_asking_again,
               test_a_redirect_found_after_the_wait_is_noticed_too,
               test_arriving_means_since_this_fetch_began,
               test_whatever_ended_the_hold_a_transfer_in_progress_is_waited_for,
               # v1.33.5 — the tab is only a witness when it moves
               test_a_frozen_challenge_page_is_not_a_new_prompt_every_time,
               test_a_real_prompt_is_still_caught_while_the_tab_is_frozen,
               test_a_tab_that_never_moved_is_not_a_redirect,
               test_a_completed_navigation_that_changed_nothing_is_read_as_nothing,
               test_the_tab_is_put_back_after_a_prompt_so_it_can_speak_again,
               test_a_prompt_cleared_while_waiting_gives_the_file_another_chance,
               test_both_workers_put_the_tab_back_after_a_prompt,
               test_a_resume_does_not_re_report_the_prompt_it_is_resuming_from,
               # v1.33.9 — the one answer the tab cannot express
               test_a_404_the_browser_cannot_hear_is_still_an_answer,
               test_the_second_opinion_keeps_the_bytes_when_there_are_bytes,
               test_a_refusal_that_names_a_challenge_holds_for_a_person,
               test_pushback_is_never_flattened_into_a_browser_failure,
               test_no_cookie_degrades_loudly_and_asks_nothing,
               test_a_reached_file_is_never_asked_about,
               test_an_absence_the_tab_states_is_not_re_asked,
               # v1.34 — the stale-tab absence
               test_a_page_left_by_an_earlier_request_is_not_an_absence,
               test_a_native_is_asked_about_before_the_browser_is_sent,
               test_a_redirect_is_judged_by_the_page_it_lands_on,
               test_a_row_that_cannot_be_confirmed_here_has_a_status_of_its_own,
               test_a_large_download_is_waited_for_by_its_progress,
               test_a_body_the_network_fetches_is_streamed_to_disk,
               test_one_file_that_cannot_be_saved_does_not_end_the_run,
               test_the_preflight_downloads_a_file_twice_before_the_run,
               test_files_that_only_the_network_delivers_are_noticed,
               test_a_refused_start_names_the_field_and_is_logged,
               test_the_download_statement_comes_from_the_profile,
               test_the_worker_gives_the_browser_fetcher_its_cookie,
               # v1.33.10 — a challenge urllib meets is nobody's to clear
               test_a_challenge_the_network_meets_is_not_a_hold,
               test_a_challenge_the_browser_shows_is_still_a_hold):
        try:
            fn()
        except Exception as e:
            # WHERE, not just what. A FAIL that names an exception and no
            # line is a diagnosis nobody can act on — and on 2026-09-21
            # nine Windows failures said `NotADirectoryError` against a
            # path that three plausible call sites could have produced,
            # none of which turned out to be right. The frames cost two
            # lines of output and save a round trip to a runner nobody
            # here can reproduce locally.
            import traceback
            where = traceback.extract_tb(e.__traceback__)
            trail = " <- ".join(
                "{}:{} in {}".format(os.path.basename(f.filename),
                                     f.lineno, f.name)
                for f in where[-3:])
            FAIL.append("{} raised {}: {}\n           at {}".format(
                fn.__name__, e.__class__.__name__, e, trail))

    print("Batch File Downloader v{} — verification".format(
        M.MANIFEST["version"]))
    print("  passed: {}".format(len(PASS)))
    print("  failed: {}".format(len(FAIL)))
    if XFAIL:
        print("  expected failures: {} (suspected defects, not breakage)"
              .format(len(XFAIL)))
    if XPASS:
        print("  UNEXPECTED PASSES: {} — resolve these".format(len(XPASS)))
    for f in FAIL:
        print("  FAIL  " + f)
    for x in XFAIL:
        print("  XFAIL " + x)
    for x in XPASS:
        print("  XPASS " + x)
    # Only a real FAIL fails the run. An XFAIL is a hypothesis awaiting real
    # HTML; an XPASS is a note to a human, not a broken build.
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
